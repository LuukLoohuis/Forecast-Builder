"""Tab 'CF kosten' (sinds v19): kosten per post en per kwartaal, zoals het FO-tabblad 'CF - kosten' (blok 'Onderdelen').

Per post: naam, 'schuift mee met' (start bouw / verkoop / vast), het gerealiseerde bedrag (t/m het actuals-kwartaal, één bedrag)
en de prognose per periode (kolom = periode van het Model, dus dezelfde volgorde als de rijen op tab Invoer). 'Ophalen uit FO'
vult de tab; zelf plakken of aanpassen mag ook.

Wat het werkboek ermee doet:
  * verkoopscenario's (tab Scenario's): de Basis blijft je eigen cashflow (tab Invoer). In een scenario schuiven alleen de
    prognosekosten van de posten op 'start bouw' mee met de start bouw van dat scenario en die op 'verkoop' met 'verkoop vanaf';
    de rest blijft staan. Model: cumulatieve prognosekosten per soort (layout.VS_G cksb / ckv), scenario = basis - eigen
    curve + verschoven curve (wat vóór het eerste prognosekwartaal zou vallen komt in dat kwartaal, wat na de laatste periode zou
    vallen in de laatste periode). Zonder posten op deze tab schuiven, zoals voorheen, alle kosten vanaf de geplande start bouw mee.
  * grootste lasten (dia 2): de drie posten met de hoogste kosten t/m het dieptepunt van de Basis (gerealiseerd + prognose),
    als je op het Dashboard geen lasten invult.
  * controle: per prognosekwartaal het totaal van de posten tegen de kosten op tab Invoer.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import layout as LY
from .layout import FIX, H, ROW1, ROWN, dabs, h, hm
from .data import standaard_mee
from .sheets import eur_m
from .styles import *  # noqa: F401,F403

NR, KWT, KOS, CUM, STAND = (FIX[k] for k in ("nr", "kwartaal", "kosten", "cum_opbr", "stand"))
N_PER = ROWN - ROW1 + 1
C = LY.KF_C


def mee_code(ref):
    """Formulefragment: tekst 'schuift mee met' -> 1 (start bouw), 2 (verkoop), 0 (vast of leeg)."""
    return (f'IF(ISNUMBER(SEARCH("bouw",{ref}&"")),1,IF(ISNUMBER(SEARCH("verkoop",{ref}&"")),2,0))')


# =============================================================================
#  TABBLAD
# =============================================================================
def bouw_tab(wb, data):
    ws = wb.create_sheet(LY.KF_TAB)
    ws.sheet_properties.codeName = "shCFKosten"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.freeze_panes = f"{L(LY.KF_COL1)}{LY.KF_R1}"
    for col, w in (("A", 2.5), ("B", 34), ("C", 13), ("D", 11), ("E", 12), ("F", 12), ("G", 4), ("H", 4)):
        ws.column_dimensions[col].width = w
    for col in (C["code"], C["sleutel"]):
        ws.column_dimensions[col].hidden = True
    for i in range(N_PER):
        ws.column_dimensions[L(LY.KF_COL1 + i)].width = 9
    n, pa = hm("n"), hm("pos_actuals")
    dal_pos = hm("vs0_dal_pos")
    posten = list(getattr(data, "kosten_posten", None) or [])
    R1, RN = LY.KF_R1, LY.KF_RN
    c_last = L(LY.KF_COL1 + N_PER - 1)
    pos_rij = f"${L(LY.KF_COL1)}${LY.KF_R_POS}:${c_last}${LY.KF_R_POS}"
    GRIJS_FL = fill("F2F3F5")
    BD_RIJ = Border(bottom=side("thin", LINE2))
    BD_TOP = Border(top=side("thin", LINE))

    # ---- kop -------------------------------------------------------------------------------------------------------------------
    put(ws, "B1", f'="CF KOSTEN · KOSTEN PER POST EN PER KWARTAAL · STAND PER "&UPPER({hm("kw_actuals")})', f=F_NOTE)
    ws.row_dimensions[2].height = 24
    put(ws, "B2", f'=IF({hm("kf_aan")}=0,"Nog geen kostenposten: gebruik Ophalen uit FO of plak het blok \'Onderdelen\' uit het FO",'
                  f'"Kosten "&{eur_m("SUM(" + LY.kf_rng("totaal") + ")")}&" in "&COUNTIF({LY.kf_rng("naam")},"?*")&" posten · "'
                  f'&{eur_m(hm("kf_sb_tot"))}&" schuift mee met de start bouw, "&{eur_m(hm("kf_v_tot"))}&" met de verkoop")',
        f=font(14, True, TXT), al=AL_VCENTER)
    put(ws, "B3", "Zoals het FO-tabblad 'CF - kosten' (blok 'Onderdelen'): één rij per post, één kolom per kwartaal, bedragen in euro's. "
                  "'Schuift mee met' bepaalt wat er in een verkoopscenario met een andere start bouw of verkoop verschuift; de Basis "
                  "blijft je eigen cashflow (tab Invoer).", f=font(9, False, TXT2))
    put(ws, "B4", "start bouw = bouwkosten die met de bouw meeschuiven (aanneemsom, meerwerk, aansluitkosten, reserve na oplevering) · "
                  "verkoop = schuift mee met de start van de verkoop (promotie, verkoopkosten) · vast = blijft staan (grond, honoraria, "
                  "leges, rente, AK)", f=F_NOTE8)

    # ---- 1 · posten ------------------------------------------------------------------------------------------------------------
    put(ws, f"B{LY.KF_R_POS - 1}", "1 · KOSTEN PER POST · Ophalen uit FO vult dit blok; zelf aanpassen mag · grijze kolommen = gerealiseerd "
                                    "(tellen niet mee: het gerealiseerde bedrag staat in kolom E)", f=F_SECTION)
    ws.row_dimensions[LY.KF_R_POS].hidden = True
    for i in range(N_PER):
        put(ws, f"{L(LY.KF_COL1 + i)}{LY.KF_R_POS}", i + 1, f=F_NOTE8)
    rk = LY.KF_R_KOP
    koppen = [(C["naam"], LY.KF_LABEL_KOP, AL_LEFT), (C["mee"], "Schuift mee met", AL_CENTER), (C["totaal"], "Totaal", AL_RIGHT),
              (C["real"], f'="Gerealiseerd t/m "&{hm("kw_actuals")}', AL_RIGHT_WRAP),
              (C["dal"], f'={hm("kf_dal_lbl")}', AL_RIGHT_WRAP), (C["code"], "code", AL_RIGHT),
              (C["sleutel"], "sleutel", AL_RIGHT)]
    for col, t, al in koppen:
        put(ws, f"{col}{rk}", t, f=F_HDR8, fl=FL_HDR, al=al)
    ws.row_dimensions[rk].height = 26
    for i in range(N_PER):
        cel = f"{L(LY.KF_COL1 + i)}{rk}"
        put(ws, cel, f'=IF({i + 1}>{n},"",INDEX(Model!${KWT}${ROW1}:${KWT}${ROWN},{i + 1}))', f=F_HDR8, fl=FL_HDR, al=AL_RIGHT)
    ws.conditional_formatting.add(f"{L(LY.KF_COL1)}{rk}:{c_last}{rk}",
                                  FormulaRule(formula=[f"{L(LY.KF_COL1)}${LY.KF_R_POS}<={pa}"], fill=PatternFill(fill_type="solid", bgColor=GREY_HDR)))

    dv = DataValidation(type="list", formula1='"' + ",".join(LY.KF_MEE) + '"', allow_blank=True)
    dv.promptTitle, dv.prompt = "Schuift mee met", "start bouw, verkoop of vast (blijft staan)"
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(f"{C['mee']}{R1}:{C['mee']}{RN}")
    for j in range(LY.N_KF):
        r = R1 + j
        post = posten[j] if j < len(posten) else {}
        bedragen = list(post.get("bedragen") or [])
        put(ws, f"{C['naam']}{r}", post.get("naam") or None, f=F_INPUT, fl=FL_INPUT, al=AL_LEFT, bd=BD_RIJ)
        put(ws, f"{C['mee']}{r}", (post.get("mee") or standaard_mee(post.get("naam"))) if post.get("naam") else None,
            f=F_INPUT, fl=FL_INPUT, al=AL_CENTER, bd=BD_RIJ)
        rij = f"{L(LY.KF_COL1)}{r}:{c_last}{r}"
        put(ws, f"{C['totaal']}{r}", f'=IF(TRIM({C["naam"]}{r}&"")="","",N({C["real"]}{r})+SUMIFS({rij},{pos_rij},">"&{pa},{pos_rij},"<="&{n}))',
            f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT, bd=BD_RIJ)
        put(ws, f"{C['real']}{r}", post.get("gerealiseerd") or None, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT, al=AL_RIGHT, bd=BD_RIJ)
        put(ws, f"{C['dal']}{r}", f'=IF(TRIM({C["naam"]}{r}&"")="","",N({C["real"]}{r})+IF({dal_pos}<={pa},0,'
                                  f'SUMIFS({rij},{pos_rij},">"&{pa},{pos_rij},"<="&{dal_pos})))', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT, bd=BD_RIJ)
        put(ws, f"{C['code']}{r}", f'=IF(TRIM({C["naam"]}{r}&"")="",0,{mee_code(C["mee"] + str(r))})', f=F_NOTE8)
        # sleutel voor de grootste lasten: kosten t/m dieptepunt, gelijke bedragen op volgorde van de rijen
        put(ws, f"{C['sleutel']}{r}", f'=IF(OR(TRIM({C["naam"]}{r}&"")="",N({C["dal"]}{r})<=0),"",{C["dal"]}{r}-ROW()/1000000)', f=F_NOTE8)
        for i in range(N_PER):
            v = bedragen[i] if i < len(bedragen) else None
            put(ws, f"{L(LY.KF_COL1 + i)}{r}", v if v not in (None, 0) else None, f=F_INPUT, nf=FMT_INT, al=AL_RIGHT, bd=BD_RIJ)
    blok = f"{L(LY.KF_COL1)}{R1}:{c_last}{RN}"
    ws.conditional_formatting.add(blok, FormulaRule(formula=[f"{L(LY.KF_COL1)}${LY.KF_R_POS}<={pa}"], fill=GRIJS_FL, font=font(9, False, GREY_TXT)))
    ws.conditional_formatting.add(blok, FormulaRule(formula=[f"{L(LY.KF_COL1)}${LY.KF_R_POS}>{n}"], fill=GRIJS_FL))

    # ---- totalen en controle ---------------------------------------------------------------------------------------------------
    rt, ri, rv, rc = LY.KF_R_TOT, LY.KF_R_INV, LY.KF_R_VERSCHIL, LY.KF_R_CONTROLE
    put(ws, f"B{rt}", "Totaal posten", f=F_BOLD9, al=AL_LEFT, bd=BD_TOP)
    for col in (C["totaal"], C["real"], C["dal"]):
        put(ws, f"{col}{rt}", f"=SUM({col}{R1}:{col}{RN})", f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT, bd=BD_TOP)
    put(ws, f"C{rt}", None, bd=BD_TOP)
    put(ws, f"B{ri}", "Kosten op tab Invoer (je eigen cashflow)", f=F_CALC9, al=AL_LEFT)
    put(ws, f"D{ri}", f"=SUM(Model!${KOS}${ROW1}:${KOS}${ROWN})", f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
    put(ws, f"E{ri}", f'=SUMIF(Model!${NR}${ROW1}:${NR}${ROWN},"<="&{pa},Model!${KOS}${ROW1}:${KOS}${ROWN})', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
    put(ws, f"B{rv}", "Verschil posten − tab Invoer (prognose)", f=F_CALC9, al=AL_LEFT)
    for i in range(N_PER):
        col = L(LY.KF_COL1 + i)
        put(ws, f"{col}{rt}", f'=IF({i + 1}>{n},"",SUM({col}{R1}:{col}{RN}))', f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT, bd=BD_TOP)
        put(ws, f"{col}{ri}", f'=IF({i + 1}>{n},"",INDEX(Model!${KOS}${ROW1}:${KOS}${ROWN},{i + 1}))', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{col}{rv}", f'=IF(OR({i + 1}>{n},{i + 1}<={pa}),"",{col}{rt}-{col}{ri})', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
    rng_v = f"{L(LY.KF_COL1)}{rv}:{c_last}{rv}"
    ws.conditional_formatting.add(rng_v, FormulaRule(formula=[f'AND(ISNUMBER({L(LY.KF_COL1)}{rv}),ABS({L(LY.KF_COL1)}{rv})>=1000)'],
                                                     font=font(9, True, RED)))
    put(ws, f"B{rc}", f"={hm('kf_controle')}", f=font(9, True, GREEN), al=AL_LEFT)
    ws.conditional_formatting.add(f"B{rc}", FormulaRule(formula=[f"OR({hm('kf_verschil')}>=1000,{hm('kf_buiten')}>=1)"], font=font(9, True, RED)))

    # ---- 2 · waar bestaat de financieringsbehoefte uit -------------------------------------------------------------------------
    r2 = rc + 3
    # ligt het dieptepunt in de realisatie, dan kan het niet per post: het FO geeft de gerealiseerde kosten als één bedrag per post
    put(ws, f"B{r2}", f'="2 · WAAR BESTAAT DE MAX. FINANCIERINGSBEHOEFTE UIT · kosten "&{hm("kf_dal_lbl")}&IF({hm("kf_dal_real")}=1,'
                      f'" (het dieptepunt "&{hm("dal_basis_kw")}&" ligt in de realisatie; het FO geeft de gerealiseerde kosten als één bedrag per post)",'
                      f'" van de Basis")', f=F_SECTION)
    for col, t, al in (("B", "Kostenpost", AL_LEFT), ("C", "Schuift mee met", AL_CENTER),
                       ("D", f'=IF({hm("kf_dal_real")}=1,"gerealiseerd","t/m dieptepunt")', AL_RIGHT),
                       ("E", "aandeel", AL_RIGHT), ("F", "dia 2", AL_CENTER)):
        put(ws, f"{col}{r2 + 1}", t, f=F_HDR8, fl=FL_HDR, al=al)
    sleutels = LY.kf_rng("sleutel")
    TOP_FL = fill("FFF2CC")
    for k in range(1, 6):
        r = r2 + 1 + k
        m = f"MATCH(LARGE({sleutels},{k}),{sleutels},0)"
        put(ws, f"B{r}", f'=IFERROR(INDEX({LY.kf_rng("naam")},{m}),"")', f=F_CALC9, al=AL_LEFT, bd=BD_RIJ)
        put(ws, f"C{r}", f'=IFERROR(INDEX({LY.kf_rng("mee")},{m})&"","")', f=F_CALC9, al=AL_CENTER, bd=BD_RIJ)
        put(ws, f"D{r}", f'=IFERROR(INDEX({LY.kf_rng("dal")},{m}),"")', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT, bd=BD_RIJ)
        put(ws, f"E{r}", f'=IF(OR(D{r}="",{hm("kf_dal_tot")}<=0),"",D{r}/{hm("kf_dal_tot")})', f=F_CALC9, nf=FMT_PCT0, al=AL_RIGHT, bd=BD_RIJ)
        put(ws, f"F{r}", f'=IF(AND({k}<=3,{hm("kl_bron")}=1,B{r}<>""),"top {k}","")', f=font(8, True, INPUT_YELLOW_FONT), al=AL_CENTER, bd=BD_RIJ)
        if k <= 3:
            ws.conditional_formatting.add(f"B{r}:F{r}", FormulaRule(formula=[f"AND({hm('kl_bron')}=1,$B{r}<>\"\")"], fill=TOP_FL))
    r = r2 + 7
    put(ws, f"B{r}", f'="Kosten "&{hm("kf_dal_lbl")}&" (alle posten)"', f=F_BOLD9, al=AL_LEFT, bd=BD_TOP)
    put(ws, f"D{r}", f"={hm('kf_dal_tot')}", f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT, bd=BD_TOP)
    for col in ("C", "E", "F"):
        put(ws, f"{col}{r}", None, bd=BD_TOP)
    put(ws, f"B{r + 1}", "Max. financieringsbehoefte (Basis)", f=F_CALC9, al=AL_LEFT)
    put(ws, f"D{r + 1}", f"={hm('dal_basis')}", f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
    put(ws, f"B{r + 2}", f'=IF({hm("kl_bron")}=1,"Geel = de drie grootste posten op dia 2 (kaart grootste lasten), want op het Dashboard zijn '
                         f'geen lasten ingevuld.","Dia 2 toont de grootste lasten die je op het Dashboard hebt ingevuld; maak die leeg om deze '
                         f'top 3 te gebruiken.")', f=F_NOTE8)

    ws.print_area = f"A1:{L(LY.KF_COL1 + 23)}{r + 2}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# =============================================================================
#  MODEL
# =============================================================================
def bouw_model(ws):
    """Kolommen cksb / ckv (cumulatieve prognosekosten van de posten die met de start bouw / de verkoop meeschuiven) en de
    hulpcellen kf_* en kl_* (grootste lasten voor dia 2)."""
    G = LY.VS_G
    n, pa = h("n"), h("pos_actuals")
    codes, blok = LY.kf_rng("code"), LY.kf_blok()
    for col, t in ((G["cksb"], "Kosten mee met start bouw cum"), (G["ckv"], "Kosten mee met verkoop cum")):
        put(ws, f"{col}7", t, f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
        ws.column_dimensions[col].width = 11
    for r in range(ROW1, ROWN + 1):
        p = r - 1
        leeg = f'${NR}{r}=""'
        for col, code in ((G["cksb"], 1), (G["ckv"], 2)):
            put(ws, f"{col}{r}", f'=IF({leeg},"",N({col}{p})+IF(${NR}{r}<={pa},0,SUMIF({codes},{code},INDEX({blok},0,${NR}{r}))))',
                f=F_CALC, nf=FMT_INT)
    rng = lambda col: f"${col}${ROW1}:${col}${ROWN}"  # noqa: E731
    sb_mee, v_mee = h("vs_sb_mee"), h("vs_v_mee")
    sb_tot, v_tot = h("kf_sb_tot"), h("kf_v_tot")
    deel_sb = f'AND({sb_mee}=1,ABS({sb_tot})>=1)'
    deel_v = f'AND({v_mee}=1,ABS({v_tot})>=1)'
    oud = f'IF({sb_mee}=1,"kosten vanaf start bouw schuiven mee met de start bouw","kosten gelijk in alle scenario\'s")'
    handmatig = ",".join(f'TRIM({dabs(f"last{i}_naam")}&"")<>"",N({dabs(f"last{i}_bedrag")})<>0' for i in range(1, LY.N_LASTEN + 1))
    sleutels = LY.kf_rng("sleutel")
    kfb = f"{LY.KF_BLAD}!"
    C = LY.KF_C
    hulp = {
        "kf_aan": ("CF kosten: posten ingevuld (1/0)", f'=IF(COUNTIF({LY.kf_rng("naam")},"?*")>0,1,0)'),
        "kf_sb_tot": ("CF kosten: prognose die meeschuift met de start bouw (€)", f'=IF({n}=0,0,N(INDEX({rng(G["cksb"])},{n})))'),
        "kf_v_tot": ("CF kosten: prognose die meeschuift met de verkoop (€)", f'=IF({n}=0,0,N(INDEX({rng(G["ckv"])},{n})))'),
        "vs_v_mee": ("Scenario's: een scenario verschuift de verkoop (1/0)",
                     "=IF(OR(" + ",".join(f"AND({h(f'vs{s}_aan')}=1,{h(f'vs{s}_d_v')}<>0)" for s in range(1, LY.N_VS + 1)) + "),1,0)"),
        "kf_tekst": ("Scenario's: regel over de kosten (tab)",
                     f'=IF({h("kf_aan")}=0,{oud},IF(AND(NOT({deel_sb}),NOT({deel_v})),"kosten gelijk in alle scenario\'s",'
                     f'"kosten per post (tab CF kosten): "&IF({deel_sb},{eur_m(sb_tot)}&" schuift mee met de start bouw","")'
                     f'&IF(AND({deel_sb},{deel_v}),", ","")&IF({deel_v},{eur_m(v_tot)}&" met de verkoop","")&", de rest blijft staan"))'),
        "kf_tekst_kort": ("Scenario's: regel over de kosten (dia 4)",
                          f'=IF({h("kf_aan")}=0,{oud},IF(AND(NOT({deel_sb}),NOT({deel_v})),"kosten gelijk in alle scenario\'s",'
                          f'IF({deel_sb},"bouwkosten "&{eur_m(sb_tot)}&" schuiven mee met de start bouw","")'
                          f'&IF(AND({deel_sb},{deel_v}),", ","")&IF({deel_v},IF({deel_sb},"","kosten ")&{eur_m(v_tot)}&" met de verkoop","")))'),
        "kf_dal_real": ("CF kosten: dieptepunt van de Basis ligt in de realisatie (1/0)", f'=IF({h("vs0_dal_pos")}<={h("pos_actuals")},1,0)'),
        "kf_dal_lbl": ("CF kosten: kosten t/m ... (label)",
                       f'=IF({h("kf_dal_real")}=1,"gerealiseerd t/m "&{h("kw_actuals")},"t/m dieptepunt "&{h("dal_basis_kw")})'),
        "kf_dal_tot": ("CF kosten: kosten t/m het dieptepunt van de Basis (€; in de realisatie: gerealiseerd t/m actuals)",
                       f'=SUM({LY.kf_rng("dal")})'),
        "kf_real_gap": ("CF kosten: verschil gerealiseerd (kolom E) met tab Invoer t/m actuals (€)",
                        f'=IF({h("kf_aan")}=0,0,ABS(N({kfb}${C["real"]}${LY.KF_R_TOT})-N({kfb}${C["real"]}${LY.KF_R_INV})))'),
        "kf_buiten": ("CF kosten: bedragen in kolommen na de laatste periode (€, tellen niet mee)",
                      f'=IF({h("kf_aan")}=0,0,ABS(SUM({blok})-SUM({kfb}${L(LY.KF_COL1)}${LY.KF_R_TOT}:${L(LY.KF_COL1 + ROWN - ROW1)}${LY.KF_R_TOT})))'),
        "kf_verschil": ("CF kosten: grootste verschil met tab Invoer: gerealiseerd of een prognosekwartaal (€)",
                        f'=IF({h("kf_aan")}=0,0,MAX(0,{h("kf_real_gap")},'
                        + f'MAX({kfb}${L(LY.KF_COL1)}${LY.KF_R_VERSCHIL}:${L(LY.KF_COL1 + ROWN - ROW1)}${LY.KF_R_VERSCHIL}),'
                        + f'-MIN({kfb}${L(LY.KF_COL1)}${LY.KF_R_VERSCHIL}:${L(LY.KF_COL1 + ROWN - ROW1)}${LY.KF_R_VERSCHIL})))'),
        "kf_controle": ("CF kosten: controletekst",
                        f'=IF({h("kf_aan")}=0,"Leeg: in een scenario met een andere start bouw schuiven alle kosten vanaf de geplande start bouw mee.",'
                        f'IF({h("kf_buiten")}>=1,"Let op: er staan bedragen ("&{eur_m(h("kf_buiten"))}&") in kolommen na de laatste periode van tab '
                        f'Invoer; die tellen niet mee. Zet ze in de laatste periode.",'
                        f'IF({h("kf_verschil")}<1000,"✓ Gerealiseerd en elk prognosekwartaal sluiten aan op de kosten op tab Invoer.",'
                        f'IF({h("kf_real_gap")}>=1000,"Let op: gerealiseerd t/m "&{h("kw_actuals")}&" wijkt €"&FIXED({h("kf_real_gap")},0)&'
                        f'" af van tab Invoer (actuals verschoven? werk kolom E bij of haal opnieuw op uit het FO). De Basis volgt tab Invoer.",'
                        f'"Let op: een prognosekwartaal wijkt tot €"&FIXED({h("kf_verschil")},0)&" af van tab Invoer. De Basis volgt tab Invoer; '
                        f'de posten bepalen alleen wat er in een scenario verschuift."))))'),
        "kl_bron": ("Grootste lasten dia 2: 1 = automatisch uit tab CF kosten, 0 = Dashboard (ingevuld gaat voor)",
                    f"=IF(OR({handmatig}),0,IF({h('kf_aan')}=1,1,0))"),
    }
    for i in range(1, LY.N_LASTEN + 1):
        m = f"MATCH(LARGE({sleutels},{i}),{sleutels},0)"
        hulp[f"kl{i}_naam"] = (f"Grootste lasten dia 2: naam {i}",
                               f'=IF({h("kl_bron")}=1,IFERROR(INDEX({LY.kf_rng("naam")},{m})&"",""),TRIM({dabs(f"last{i}_naam")}&""))')
        hulp[f"kl{i}_bedrag"] = (f"Grootste lasten dia 2: bedrag {i} (€)",
                                 f'=IF({h("kl_bron")}=1,IFERROR(N(INDEX({LY.kf_rng("dal")},{m})),0),N({dabs(f"last{i}_bedrag")}))')
    for naam, (label, formule) in hulp.items():
        r = H[naam]
        put(ws, f"{LY.M_HULP_LABEL}{r}", label, f=F_NOTE8)
        put(ws, f"{LY.M_HULP}{r}", formule, f=F_CALC, al=AL_LEFT_TOP)
    for naam in ("kf_sb_tot", "kf_v_tot", "kf_dal_tot", "kf_verschil", "kf_real_gap", "kf_buiten") + tuple(f"kl{i}_bedrag" for i in range(1, LY.N_LASTEN + 1)):
        ws[f"{LY.M_HULP}{H[naam]}"].number_format = FMT_INT
