"""Invoer uit een FO-werkboek (Financieel Overzicht van Heijmans Vastgoed) halen.

Het FO heeft per project dezelfde opbouw (tabbladen '1. Cashflow', 'CF - opbrengsten', 'FO - actuals'); alles wordt op
labels gezocht, niet op vaste celadressen, zodat kleine verschillen tussen projecten geen kwaad kunnen:
  * '1. Cashflow': de kwartaalcashflow (koprij met 'Jaar' in kolom B; daaronder jaar, Q, kosten, % kosten, omzet, % omzet,
    CF, CF x 1000, CF vorig kwartaal) -> tab Invoer B:J. Rijen zonder Q (jaartotalen) blijven jaarrijen.
  * 'FO - actuals': laatste boekjaar/kwartaal met boekingen -> Dashboard 'Actuals t/m'.
  * 'CF - opbrengsten': per woningtype een blok 'Termijnen k - naam' (bouwplanning: per termijn het aantal woningen per
    kwartaal), een blok 'Omzet HVG Termijnen' (koopsom en termijnen in %, € en € per woning; kopersmeerwerk, kadastrale
    kosten, overige) en het verkooptempo (verkocht per kwartaal; 'actuals' = al verkocht). Daaruit: naam, aantal, koopsom
    per woning (grond + termijnen, excl. btw), grondtermijn %, start bouw, bouwtermijnen (naam, %, bouwkwartaal; een termijn
    die over meer kwartalen is verdeeld wordt gesplitst), extra's per woning en verkocht/transport per kwartaal.
  * 'CF - kosten', blok 'Onderdelen' (label in kolom B, kwartalen Q1..Q4 met de jaartallen in de rij erboven): per kostenpost
    de naam, het gerealiseerde bedrag (de kolom(men) onder 'Actuals', anders de kolom vóór het eerste kwartaal; kwartalen t/m
    actuals tellen ook als gerealiseerd) en de prognose per kwartaal -> tab CF kosten, met 'schuift mee met' op de naam
    (data.standaard_mee). Rijen tot de eerste rij zonder label of met 'Totaal'.
  * DAEB-fees: een rij 'AK fee' of 'Bijkomende kosten' met alleen een totaal in € (geen %, geen kwartalen) is een
    fee-component; elke rij daarna met een % of € én een kwartaal of 'actuals' is een termijn van die component, wat het
    label ook is ('na akkoord SO', 'bij start bouw', ...). 'Onvoorzien' wordt overgeslagen. Zo'n type wordt DAEB: alleen
    de fees, geen koopsom, bouwtermijnen of extra's. Of een los DAEB-blok onderaan 'CF - opbrengsten' (bv. 'Opbrengsten DAEB',
    de componenten met hun totaal boven de koprij): dat wordt een eigen DAEB-type (_lees_daeb), met het aantal woningen en de
    verkoop van tabblad 'Verkooptempo'. Bedragen per termijn uit de cellen, per kwartaal gesplitst (_fee_delen); bij DAEB is
    transport = verkoop.
Meldingen (ontbrekend tabblad, samengevoegde termijnen, overgeslagen rijen) staan in p.fo['waarschuwingen'].
De macro 'Ophalen uit FO' in het werkboek doet hetzelfde (zie vba/CashflowNaarPowerPoint_v10.bas, UitFOOphalen).
"""
import os
import re
from datetime import datetime

from openpyxl import load_workbook

from . import layout as LY
from .data import PARAM_STANDAARD, ProjectData, TypeData, standaard_mee

KW = {"Q1": 1, "Q2": 2, "Q3": 3, "Q4": 4}


def _txt(v):
    return str(v).strip() if v is not None else ""


def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _jaar(v):
    """Jaartal uit een getal of tekst ('2026'); None als het geen jaar is."""
    try:
        j = int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None
    return j if 1990 <= j <= 2100 else None


def _kw(v):
    t = _txt(v).upper().replace(" ", "")
    if t in KW:
        return KW[t]
    if t in ("1", "2", "3", "4"):
        return int(t)
    return None


def _zoek_blad(wb, *namen):
    """Werkblad op naam (hoofdletterongevoelig, spaties genegeerd); None als het ontbreekt."""
    norm = lambda s: re.sub(r"\s+", "", s.lower())  # noqa: E731
    for ws in wb.worksheets:
        if norm(ws.title) in [norm(n) for n in namen]:
            return ws
    return None


def _kolomkaart(ws, rij_q):
    """{kolom: (jaar, kw)} uit een koprij met Q1..Q4 en de jaartallen in de rij erboven (één jaartal per vier kwartalen);
    ook de kolom met 'actuals' (of None)."""
    kaart, actuals = {}, None
    jaar = None
    for c in range(1, ws.max_column + 1):
        j = _jaar(ws.cell(rij_q - 1, c).value)
        if j is not None:
            jaar = j
        v = ws.cell(rij_q, c).value
        if _txt(v).lower().startswith("actual"):
            actuals = c
        q = _kw(v)
        if q is not None and jaar is not None:
            kaart[c] = (jaar, q)
    return kaart, actuals


def _tellingen(ws, r, kaart, c_act=None):
    """[(jaar, kw, aantal)] van de kwartaalcellen in rij r (aantal <> 0); de kolom 'actuals' als (None, None, aantal) vooraan."""
    uit = []
    if c_act:
        v = _num(ws.cell(r, c_act).value)
        if v:
            uit.append((None, None, v))
    for c, (j, q) in sorted(kaart.items()):
        v = _num(ws.cell(r, c).value)
        if v:
            uit.append((j, q, v))
    return uit


def _typenaam(ws, r, k):
    """Naam van het type: uit de bloktitel 'Termijnen k - naam', anders de cel naast de titel (bv. '<300K VON'), anders de
    tekst na 'Termijnen ' als dat geen nummer is, anders 'Type k'."""
    t = _txt(ws.cell(r, 2).value)
    if " - " in t:
        return t.split(" - ", 1)[1].strip() or f"Type {k}"
    naast = ws.cell(r, 3).value
    if _txt(naast) and _num(naast) is None:
        return _txt(naast)
    rest = t[10:].strip() if t.lower().startswith("termijnen ") else ""
    if rest and not rest.replace(".", "").isdigit():
        return rest
    return f"Type {k}"


def voeg_vorige_prognose_toe(project, pad_vorig):
    """Vorige prognose uit een ouder FO: per type de start bouw (blok 4) en per fee-termijn het kwartaal (blok 6). Typen op naam, anders op
    positie; fee-termijnen op mijlpaal, anders op positie. Meldingen in project.fo['waarschuwingen']."""
    vorig = lees_fo(pad_vorig)
    meldingen = project.fo.setdefault("waarschuwingen", [])
    op_naam = {t.naam.strip().lower(): t for t in vorig.types}
    for k, t in enumerate(project.types):
        tv = op_naam.get(t.naam.strip().lower()) or (vorig.types[k] if k < len(vorig.types) else None)
        if tv is None:
            meldingen.append(f"{t.naam}: niet gevonden in het vorige FO ({os.path.basename(pad_vorig)}); vorige prognose leeg gelaten.")
            continue
        if tv.naam.strip().lower() != t.naam.strip().lower():
            meldingen.append(f"{t.naam}: op naam niet gevonden in het vorige FO; het type op dezelfde positie ({tv.naam}) is gebruikt.")
        t.vorig_start_jaar, t.vorig_start_kw = tv.start_jaar, tv.start_kw
        if any(lab for lab, _, _ in t.fee_termijnen):
            kw_vorig = {str(lab).strip().lower(): kw for lab, _, kw in tv.fee_termijnen if lab}
            t.fee_vorig = []
            for i, (lab, _, _) in enumerate(t.fee_termijnen):
                kw = kw_vorig.get(str(lab).strip().lower()) if lab else None
                if kw is None and lab and i < len(tv.fee_termijnen) and tv.fee_termijnen[i][0]:
                    kw = tv.fee_termijnen[i][2]
                t.fee_vorig.append(kw)
    project.fo["vorig_bestand"] = os.path.basename(pad_vorig)
    return project


def _is_fee_label(ll):
    return bool(re.search(r"\bfees?\b", ll)) or ll.startswith("ak-fee") or ll.startswith("bijkomende kosten") or ll.startswith("onvoorzien")


def _binnen_limiet(rijen, n_max, naam, wat, waarschuwingen, afkappen=False):
    """[(label, waarde, bouwkwartaal)] terugbrengen tot n_max rijen: gesplitste delen (label 'x (i/n)') van dezelfde termijn worden
    samengevoegd bij het grootste deel (kleinste deel eerst), met een melding. Zonder splitsingen om samen te voegen: foutmelding,
    of met afkappen=True (zoals de macro) vallen de laatste rijen weg, met een melding."""
    rijen = list(rijen)
    samengevoegd = []
    while len(rijen) > n_max:
        basis = lambda lab: re.sub(r" \(\d+/\d+\)$", "", lab)
        groepen = {}
        for i, (lab, w, kw) in enumerate(rijen):
            if re.search(r" \(\d+/\d+\)$", lab):
                groepen.setdefault(basis(lab), []).append(i)
        groepen = {b: ix for b, ix in groepen.items() if len(ix) > 1}
        if not groepen:
            if not afkappen:
                raise ValueError(f"{naam}: {len(rijen)} {wat} in het FO, het werkboek heeft ruimte voor {n_max}.")
            waarschuwingen.append(f"{naam}: {len(rijen)} {wat} in het FO, ruimte voor {n_max}; de laatste zijn niet overgenomen.")
            rijen = rijen[:n_max]
            break
        b, ix = max(groepen.items(), key=lambda kv: len(kv[1]))
        klein = min(ix, key=lambda i: abs(rijen[i][1] or 0))
        groot = max((i for i in ix if i != klein), key=lambda i: abs(rijen[i][1] or 0))
        lab_g, w_g, kw_g = rijen[groot]
        rijen[groot] = (lab_g, round((w_g or 0) + (rijen[klein][1] or 0), 6), kw_g)
        del rijen[klein]
        samengevoegd.append(b)
        n = len([i for i, (lab, _, _) in enumerate(rijen) if basis(lab) == b and re.search(r" \(\d+/\d+\)$", lab)])
        k = 0
        for i, (lab, w, kw) in enumerate(rijen):
            if basis(lab) == b and re.search(r" \(\d+/\d+\)$", lab):
                k += 1
                rijen[i] = (b if n == 1 else f"{b} ({k}/{n})", w, kw)
    if samengevoegd:
        waarschuwingen.append(f"{naam}: meer dan {n_max} {wat} na het splitsen over kwartalen; kleinste delen samengevoegd bij het "
                              f"grootste deel van {', '.join(sorted(set(samengevoegd)))}.")
    return rijen


def _idx(j, q):
    return j * 4 + q


def _kw_tekst(idx):
    """kwartaalindex -> 'jaar Qk' (de schrijfwijze van blok 5 en 6; zoals KwTekst in de macro)"""
    return f"{(idx - 1) // 4} Q{idx - ((idx - 1) // 4) * 4}"


def _fee_delen(lab, pct, eur, tel, act):
    """Eén fee-termijn -> [(mijlpaal, waarde, kwartaal)] voor blok 5. De bedragen komen uit de cellen: de kolom 'actuals' is één deel
    'actuals' ('x (ontvangen)' als er ook kwartalen zijn), elk kwartaal een eigen deel met zijn kalenderkwartaal ('x (1/2)' ...; ook
    een kwartaal t/m actuals houdt zijn eigen kwartaal). Tellen de cellen samen op tot 1 of minder (alleen een teken), dan het % of €
    van de rij, naar rato. Zonder bedrag in een kwartaal of in 'actuals': []. Zelfde regels als FeeDelen in de macro."""
    ontv = act or 0.0
    later = [(_idx(j, q), v) for j, q, v in tel]
    if not ontv and not later:
        return []
    som = ontv + sum(v for _, v in later)
    if abs(som) > 1 or som == 0 or not (pct or eur):
        waarde = lambda v: round(v, 2)  # noqa: E731
    else:
        waarde = lambda v: round((pct or eur) * v / som, 6)  # noqa: E731
    uit = []
    if ontv:
        uit.append((lab + (" (ontvangen)" if later else ""), waarde(ontv), "actuals"))
    for i, (idx, v) in enumerate(later):
        uit.append((lab + (f" ({i + 1}/{len(later)})" if len(later) > 1 else ""), waarde(v), _kw_tekst(idx)))
    return uit


def _fee_rijen(fees, naam, waarschuwingen):
    """[(component (naam, totaal), [(mijlpaal, waarde, kwartaal)])] -> (fee_comp, fee_termijnen) voor blok 5: maximaal N_FEE_COMP
    componenten met elk N_FEE_PER_COMP rijen (meer delen: kleinste gesplitste delen samengevoegd, daarna afgekapt)."""
    if len(fees) > LY.N_FEE_COMP:
        waarschuwingen.append(f"{naam}: {len(fees)} fee-componenten in het FO, alleen de eerste {LY.N_FEE_COMP} zijn overgenomen "
                              f"({', '.join(c[0] for c, _ in fees[LY.N_FEE_COMP:])} niet).")
    rijen = []
    for c, termijnen_c in fees[:LY.N_FEE_COMP]:
        rijen_c = _binnen_limiet(termijnen_c, LY.N_FEE_PER_COMP, naam, f"fee-termijnen van '{c[0]}'", waarschuwingen, afkappen=True)
        rijen += rijen_c + [("", None, None)] * (LY.N_FEE_PER_COMP - len(rijen_c))
    return [(c[0], c[1]) for c, _ in fees][:LY.N_FEE_COMP], rijen


def _zonder_voorvoegsel(lab, comp_naam):
    """'AK fee na akkoord SO' bij component 'AK' -> 'na akkoord SO': de woorden vooraan die ook in de componentnaam staan (of
    'fee'/'fees') vallen weg; blijft er niets over (of valt er niets weg), dan het label zoals het is."""
    namen = {w.lower() for w in comp_naam.split()} | {"fee", "fees"}
    w = lab.split()
    i = 0
    while i < len(w) and w[i].lower() in namen:
        i += 1
    return " ".join(w[i:]) if 0 < i < len(w) else lab


def _lees_daeb(wb, wo, namen, waarschuwingen):
    """Los DAEB-blok op 'CF - opbrengsten' (bv. 'Opbrengsten DAEB': koprij met 'actuals' en Q1..Q4, de jaartallen erboven). Boven de
    koprij per fee-component de naam (kolom C) en het totaal (kolom E, anders D); daaronder per termijn een rij (label, %, de
    bedragen in 'actuals' en de kwartalen), groepen gescheiden door een lege rij, tot 'Totaal' of drie lege rijen. Een termijn hoort
    bij de component met hetzelfde eerste woord ('AK fee na akkoord SO' -> 'AK'), anders bij de component van zijn groep; de
    woorden die al in de componentnaam staan vallen weg uit de mijlpaal. Start bouw: 'Startbouw' boven een kwartaalkolom. Aantal
    woningen en verkoop: tabblad 'Verkooptempo', de rij die bij geen woningtype hoort (de corporatie); transport = verkoop.
    -> TypeData (soort DAEB) of None. Zelfde regels als FODaeb in de macro."""
    r0 = None
    for r in range(2, wo.max_row + 1):
        b = _txt(wo.cell(r, 2).value).lower()
        if "daeb" in b and not b.startswith("termijnen ") and not b.startswith("omzet") and len(_kolomkaart(wo, r)[0]) >= 4:
            r0 = r
            break
    if r0 is None:
        return None
    titel = _txt(wo.cell(r0, 2).value)
    kaart, c_act = _kolomkaart(wo, r0)
    comps = []
    for rr in range(max(1, r0 - 3), r0):
        lab = _txt(wo.cell(rr, 3).value)
        tot = _num(wo.cell(rr, 5).value)
        if tot is None:
            tot = _num(wo.cell(rr, 4).value)
        if lab and _num(wo.cell(rr, 3).value) is None and tot and not lab.lower().startswith("onvoorzien"):
            comps.append((lab, round(tot, 2)))
    if not comps:
        waarschuwingen.append(f"'{wo.title}': blok '{titel}' (rij {r0}) zonder fee-componenten (naam in kolom C en totaal in kolom E "
                              f"boven de koprij): DAEB-fees niet overgenomen.")
        return None
    fees = [(c, []) for c in comps]
    groep, leeg, in_groep = 1, 0, False
    r = r0 + 1
    while r <= r0 + 40:
        lab = _txt(wo.cell(r, 2).value)
        if not lab:
            leeg += 1
            if leeg >= 3:
                break
            if in_groep:
                groep, in_groep = groep + 1, False
            r += 1
            continue
        leeg = 0
        if lab.lower().startswith("totaal"):
            break
        in_groep = True
        eerste = lab.split()[0].lower()
        ci = next((i for i, (cn, _) in enumerate(comps) if cn.split()[0].lower() == eerste), min(groep, len(comps)) - 1)
        pct, eur = _num(wo.cell(r, 3).value), _num(wo.cell(r, 4).value)
        act = _num(wo.cell(r, c_act).value) if c_act else None
        delen = _fee_delen(_zonder_voorvoegsel(lab, comps[ci][0]), pct, eur, _tellingen(wo, r, kaart), act)
        if delen:
            fees[ci][1].extend(delen)
        elif pct or eur:
            waarschuwingen.append(f"'{wo.title}': rij '{lab}' in het blok '{titel}' heeft geen bedrag in een kwartaal of in 'actuals' "
                                  f"en is overgeslagen.")
        r += 1
    naam = " ".join(w for w in titel.split() if w.lower() not in ("opbrengsten", "omzet", "fee", "fees")) or "DAEB"
    t = TypeData(naam, None, None, None, None, None, [])
    t.soort = "DAEB"
    for rr in range(max(1, r0 - 2), r0):
        c = next((c for c in sorted(kaart) if _txt(wo.cell(rr, c).value).lower().replace(" ", "").startswith("startbouw")), None)
        if c is not None:
            t.start_jaar, t.start_kw = kaart[c]
            break
    t.fee_comp, t.fee_termijnen = _fee_rijen(fees, naam, waarschuwingen)
    # aantal en verkoop: tabblad 'Verkooptempo', de rij die bij geen woningtype hoort
    t._verkoop = []
    wv = _zoek_blad(wb, "Verkooptempo")
    kop = None
    if wv is not None:
        kop = next(((rr, c) for rr in range(1, min(wv.max_row, 30) + 1) for c in range(1, 7)
                    if _txt(wv.cell(rr, c).value).lower() == "verkooptempo"), None)
    kandidaten = []
    if kop is not None:
        rk, cn = kop
        kaart_v, c_act_v = _kolomkaart(wv, rk)
        c_aantal = next((c for c in range(1, wv.max_column + 1) if _txt(wv.cell(rk, c).value).lower() == "aantal"), None)
        rr = rk + 1
        while _txt(wv.cell(rr, cn).value):
            nm = _txt(wv.cell(rr, cn).value)
            if nm.lower() not in namen:
                tel = _tellingen(wv, rr, kaart_v, c_act_v)
                aantal = _num(wv.cell(rr, c_aantal).value) if c_aantal else None
                if not aantal:
                    aantal = sum(v for _, _, v in tel) or None
                if aantal and aantal > 0:
                    kandidaten.append((nm, aantal, tel))
            rr += 1
    if kandidaten:
        nm, aantal, tel = next((k for k in kandidaten if "daeb" in k[0].lower()), kandidaten[0])
        t.aantal = int(aantal) if aantal == int(aantal) else aantal
        t._verkoop = tel
        if len(kandidaten) > 1:
            waarschuwingen.append(f"{naam}: op 'Verkooptempo' horen {len(kandidaten)} rijen bij geen woningtype "
                                  f"({', '.join(k[0] for k in kandidaten)}); '{nm}' is gebruikt voor de DAEB-woningen.")
        waarschuwingen.append(f"{naam}: {t.aantal} woningen en de verkoop uit 'Verkooptempo' (rij '{nm}'); transport = verkoop "
                              f"(het FO heeft geen transport voor DAEB).")
    else:
        waarschuwingen.append(f"{naam}: aantal woningen niet gevonden (tabblad 'Verkooptempo', een rij die bij geen woningtype hoort); "
                              f"vul het in op tab Woningtypes: zonder aantal tellen de fees niet mee.")
    t._transport = list(t._verkoop)
    return t


def lees_fo(pad, actuals=None):
    """actuals = (jaar, kwartaal) als terugval wanneer 'FO - actuals' geen kwartaal geeft (build.py --fo --from: die van het
    werkboek, zoals de macro de Dashboard-cellen houdt); bepaalt ook de grens gerealiseerd/prognose op tab CF kosten."""
    wb = load_workbook(pad, data_only=True)
    p = ProjectData()
    p.params = dict(PARAM_STANDAARD)
    if actuals and actuals[0] and actuals[1]:
        p.params["actuals_jaar"], p.params["actuals_kw"] = int(actuals[0]), int(actuals[1])
    p.fo = {"bestand": os.path.basename(pad), "datum": datetime.now().strftime("%d-%m-%Y %H:%M"), "waarschuwingen": []}
    waarschuwingen = p.fo["waarschuwingen"]

    # ---- 1. Cashflow -> periodes ----------------------------------------------------------------------------------
    ws = _zoek_blad(wb, "1. Cashflow", "Cashflow")
    if ws is None:
        raise ValueError("Tabblad '1. Cashflow' niet gevonden in het FO.")
    # koprij 'Jaar' ergens in A1:F30; de kolommen erachter: Q, kosten, % kosten, omzet, % omzet, CF, CF x 1000, CF vorig
    kop = next(((r, c) for r in range(1, 31) for c in range(1, 7) if _txt(ws.cell(r, c).value).lower() == "jaar"), None)
    if kop is None:
        raise ValueError(f"Koprij 'Jaar' niet gevonden op '{ws.title}' (gezocht in A1:F30).")
    r, cj = kop[0] + 1, kop[1]
    while _jaar(ws.cell(r, cj).value) is not None:
        q = _kw(ws.cell(r, cj + 1).value)
        v = [_num(ws.cell(r, cj + k).value) for k in range(2, 9)]
        p.periodes.append({"jaar": _jaar(ws.cell(r, cj).value), "kw": f"Q{q}" if q else None,
                           "kosten": v[0], "pct_kosten": v[1], "omzet": v[2], "pct_omzet": v[3], "cf": v[4], "cf1000": v[5], "cf_vorig": v[6]})
        r += 1
    for rr in range(r, r + 6):          # totalen onder de tabel ('Totaal:' in de kolom van Q)
        if _txt(ws.cell(rr, cj + 1).value).lower().startswith("totaal"):
            p.fo["kosten"] = _num(ws.cell(rr, cj + 2).value)
            p.fo["opbrengsten"] = _num(ws.cell(rr, cj + 4).value)
    n_per = len(p.periodes)
    periode_van = {}
    for i, per in enumerate(p.periodes):
        if per["kw"]:
            periode_van[_idx(per["jaar"], int(per["kw"][1]))] = i

    # ---- FO - actuals -> actuals t/m -------------------------------------------------------------------------------
    wa = _zoek_blad(wb, "FO - actuals", "actuals")
    if wa is not None:
        koppen = {_txt(wa.cell(1, c).value).lower(): c for c in range(1, wa.max_column + 1)}
        cj, cq, cw = koppen.get("boekjaar"), koppen.get("q"), koppen.get("waarde/trval")
        laatste = None
        if cj and cq:
            for r in range(2, wa.max_row + 1):
                j, q = _jaar(wa.cell(r, cj).value), _kw(wa.cell(r, cq).value)
                if j and q and (cw is None or wa.cell(r, cw).value is not None):
                    if laatste is None or _idx(j, q) > _idx(*laatste):
                        laatste = (j, q)
        if laatste:
            p.params["actuals_jaar"], p.params["actuals_kw"] = laatste
            p.fo["actuals"] = f"Q{laatste[1]} {laatste[0]}"
        else:
            waarschuwingen.append("'FO - actuals' heeft geen rij met boekjaar en kwartaal: 'Actuals t/m' is niet uit het FO gehaald.")
    else:
        waarschuwingen.append("tabblad 'FO - actuals' ontbreekt: 'Actuals t/m' is niet uit het FO gehaald.")
    idx_act = _idx(p.params["actuals_jaar"], p.params["actuals_kw"])

    # ---- CF - opbrengsten -> woningtypes ----------------------------------------------------------------------------
    wo = _zoek_blad(wb, "CF - opbrengsten", "CF-opbrengsten", "Opbrengsten")
    if wo is None:
        raise ValueError("Tabblad 'CF - opbrengsten' niet gevonden in het FO.")
    blokken, omzet, verkoop = [], [], None
    for r in range(1, wo.max_row + 1):
        b = _txt(wo.cell(r, 2).value)
        lb = b.lower()
        if lb.startswith("termijnen ") and " in €" not in lb and _txt(wo.cell(r - 1, 6).value).lower().startswith("# won"):
            blokken.append(r)
        elif lb.startswith("omzet hvg") or lb.startswith("omzet termijnen"):
            omzet.append(r)
        elif lb == "verkooptempo" and verkoop is None:
            verkoop = r
    p.types = []
    for k, r0 in enumerate(blokken):
        naam = _typenaam(wo, r0, k + 1)
        aantal = _num(wo.cell(r0 - 1, 7).value) or 0
        kaart, c_act = _kolomkaart(wo, r0)
        # bouwplanning: per termijn het aantal woningen per kwartaal
        planning = []
        r = r0 + 1
        while _txt(wo.cell(r, 2).value):
            planning.append((_txt(wo.cell(r, 2).value), _tellingen(wo, r, kaart), _num(wo.cell(r, c_act).value) if c_act else None))
            r += 1
        start = None
        for tn, tel, _ in planning:
            if tel:
                start = (tel[0][0], tel[0][1])
                break
        if start and start[0] is None:
            start = None
        t = TypeData(naam, int(aantal) if aantal == int(aantal) else aantal, None, start[0] if start else None, start[1] if start else None, None, [])
        if k < len(omzet):
            ro = omzet[k]
            kaart_o, c_act_o = _kolomkaart(wo, ro)
            koop_pw, koop_tot, transport = 0.0, 0.0, []
            termijnen, extras, fees = [], [], []
            r = ro + 1
            comp = None
            while _txt(wo.cell(r, 2).value):
                lab = _txt(wo.cell(r, 2).value)
                ll = lab.lower()
                pct, eur, pw = _num(wo.cell(r, 3).value), _num(wo.cell(r, 4).value), _num(wo.cell(r, 5).value)
                tel = _tellingen(wo, r, kaart_o)
                act = _num(wo.cell(r, c_act_o).value) if c_act_o else None
                comp_rij = bool(eur) and not pct and not tel and not act      # alleen een totaal in €: fee-component
                if ll == "koopsom":
                    koop_pw, koop_tot = pw or 0.0, eur or 0.0
                    transport = _tellingen(wo, r, kaart_o, c_act_o)
                elif (comp_rij and _is_fee_label(ll)) or (comp is not None and comp_rij):
                    # DAEB: componentrij (totaal in €, geen % en geen kwartalen); 'onvoorzien' doet niet mee
                    comp = (lab, eur)
                    if ll.startswith("onvoorzien"):
                        comp = (None, None)
                    else:
                        fees.append((comp, []))
                elif comp is not None and (pct or eur) and (tel or act):
                    # termijn van de laatste component (label vrij: 'na akkoord SO', 'bij start bouw', ...); per kwartaal gesplitst
                    if comp[0] is not None:
                        fees[-1][1].extend(_fee_delen(lab, pct, eur, tel, act))
                elif comp is not None and _is_fee_label(ll):
                    waarschuwingen.append(f"{naam}: rij '{lab}' in het fee-blok heeft geen bedrag met kwartaal en is overgeslagen.")
                elif pct is not None and pct != 0:
                    termijnen.append((lab, pct, pw or 0.0, tel))
                elif pw and abs(pw) >= 0.005:                 # extra per woning (minder dan een cent telt niet)
                    extras.append((lab, pw, tel, act))
                r += 1
            # koopsom per woning = grond (rij Koopsom) + termijnen per woning
            koopsom = koop_pw + sum(pw for _, _, pw, _ in termijnen)
            t.koopsom = round(koopsom, 2) if koopsom else None
            t.grond_pct = round(koop_pw / koopsom, 6) if koopsom else None
            idx_start = _idx(*start) if start else None
            for lab, pct, pw, tel in termijnen:
                # kwartalen uit de bouwplanning (zelfde naam), anders uit het €-blok; verdeeld over meer kwartalen -> gesplitst
                tel_p = next((tp for tn, tp, _ in planning if tn.lower() == lab.lower() and tp), None) or tel
                if tel_p and idx_start:
                    som = sum(v for _, _, v in tel_p)
                    delen = [(j, q, v / som) for j, q, v in tel_p] if som else []
                    for i, (j, q, deel) in enumerate(delen):
                        suffix = f" ({i + 1}/{len(delen)})" if len(delen) > 1 else ""
                        t.termijnen.append((lab + suffix, round(pct * deel, 6), _idx(j, q) - idx_start + 1))
                else:
                    t.termijnen.append((lab, pct, None))
            for lab, pw, tel, act in extras:
                tel_p = next((tp for tn, tp, _ in planning if tn.lower() == lab.lower() and tp), None) or tel
                if tel_p and idx_start and not act:
                    som = sum(v for _, _, v in tel_p)
                    delen = [(j, q, v / som) for j, q, v in tel_p] if som else []
                    for i, (j, q, deel) in enumerate(delen):
                        suffix = f" ({i + 1}/{len(delen)})" if len(delen) > 1 else ""
                        t.extras.append((lab + suffix, round(pw * deel, 2), _idx(j, q) - idx_start + 1))
                else:
                    t.extras.append((lab, round(pw, 2), None))        # bij transport
            t.termijnen = _binnen_limiet(t.termijnen, LY.N_TERMIJNEN, naam, "bouwtermijnen", waarschuwingen)
            t.extras = _binnen_limiet(t.extras, LY.N_EXTRA, naam, "extra's", waarschuwingen)
            if fees:
                t.soort = "DAEB"
                t.koopsom, t.grond_pct, t.termijnen, t.extras = None, None, [], []
                t.fee_comp, t.fee_termijnen = _fee_rijen(fees, naam, waarschuwingen)
            t._transport = transport
        t._blok = k                                   # positie van het blok in het FO (ook lege blokken tellen mee in het verkooptempo)
        t._titel = _txt(wo.cell(r0, 2).value)         # bloktitel: zelfde titel op 'CF - kosten' (termijnblok van de bouwkosten)
        t._titel_nr = sum(1 for rr in blokken[:k] if _txt(wo.cell(rr, 2).value).lower() == t._titel.lower())   # dubbele titels: n-de blok
        p.types.append(t)
    namen = {t.naam.strip().lower() for t in p.types}
    p.types = [t for t in p.types if (t.aantal or 0) > 0]
    daeb = _lees_daeb(wb, wo, namen, waarschuwingen)     # los DAEB-blok (fees zonder bouwplanning) als laatste type
    if daeb is not None:
        if len(p.types) < LY.N_TYPES:
            p.types.append(daeb)
        else:
            waarschuwingen.append(f"{daeb.naam}: het werkboek heeft ruimte voor {LY.N_TYPES} woningtypes; het DAEB-blok is niet overgenomen.")
    # verkoop per kwartaal per type (blok Verkooptempo: rijen in dezelfde volgorde als de typen) en transport
    verkocht = {}
    if verkoop is not None:
        kaart_v, c_act_v = _kolomkaart(wo, verkoop)
        r = verkoop + 1
        i = 0
        while _txt(wo.cell(r, 2).value):
            verkocht[i] = _tellingen(wo, r, kaart_v, c_act_v)   # kwartaalcellen tellen op tot het aantal; 'actuals' = al verkocht
            r += 1
            i += 1
    _lees_bouwkosten(wb, p.types, idx_act, waarschuwingen)
    typen_namen = [t.naam for t in p.types]
    nT = len(p.types)
    p.verkocht = [[None] * nT for _ in range(n_per)]
    p.transport = [[None] * nT for _ in range(n_per)]
    pos_act = periode_van.get(idx_act)
    if pos_act is None and periode_van:
        pos_act = min(periode_van.values())

    def zet(matrix, k, i, v):
        if v:
            matrix[i][k] = (matrix[i][k] or 0) + int(round(v))

    for k, t in enumerate(p.types):
        v_t = t._verkoop if hasattr(t, "_verkoop") else verkocht.get(getattr(t, "_blok", k), [])
        if t.soort == "DAEB" and not getattr(t, "_transport", None):
            t._transport = v_t                          # DAEB: het FO heeft geen transport (geen koopsom); transport = verkoop
        for jj, q, v in v_t:
            i = pos_act if jj is None else periode_van.get(_idx(jj, q))
            if i is not None:
                zet(p.verkocht, k, i, v)
        for jj, q, v in getattr(t, "_transport", []):
            i = pos_act if jj is None else periode_van.get(_idx(jj, q))
            if i is not None:
                zet(p.transport, k, i, v)
        for attr in ("_transport", "_verkoop", "_blok", "_titel", "_titel_nr"):
            if hasattr(t, attr):
                delattr(t, attr)
    p.fo["typen"] = typen_namen
    p.kosten_posten = _lees_kosten(wb, periode_van, idx_act, n_per, waarschuwingen)
    return p


def _lees_bouwkosten(wb, types, idx_act, waarschuwingen):
    """FO-tabblad 'CF - kosten': per type het termijnblok met dezelfde titel als op 'CF - opbrengsten' ('Termijnen ...', met
    '# won' in de rij erboven): per termijn de naam en het bedrag per woning (kolom E) en het aantal woningen per kwartaal (de
    bouwplanning) -> TypeData.bouwkosten [(termijn, bedrag €, bouwkwartaal t.o.v. start bouw)]. Wat t/m actuals ligt (kolom 'actuals'
    en kwartalen t/m het actuals-kwartaal) wordt één deel in het actuals-kwartaal ('x (gerealiseerd)'), elk later kwartaal een eigen
    deel ('(1/3)' ...); meer dan tien regels: kleinste latere delen samengevoegd (_binnen_limiet; het gerealiseerde deel doet niet
    mee, zodat er geen prognose in de realisatie belandt). Dubbele titels: het n-de type met die titel krijgt het n-de blok. Rijen tot
    'Aanneemsom'/'Meerwerk' of de eerste lege rij; termijnen zonder bedrag per woning (kopersmeerwerk) tellen niet mee."""
    wk = _zoek_blad(wb, "CF - kosten", "CF-kosten", "CF kosten")
    if wk is None:
        return
    blokken = {}
    for r in range(2, wk.max_row + 1):
        b = _txt(wk.cell(r, 2).value).lower()
        if b.startswith("termijnen ") and any(_txt(wk.cell(r - 1, c).value).lower().startswith("# won") for c in range(2, 13)):
            blokken.setdefault(b, []).append(r)
    for t in types:
        titel = getattr(t, "_titel", "") or ""
        if not titel:                                  # los DAEB-blok: geen termijnblok van de aannemer
            continue
        rr = blokken.get(titel.lower(), [])
        nr = getattr(t, "_titel_nr", 0)
        r0 = rr[nr] if nr < len(rr) else None
        if r0 is None:
            if blokken:
                waarschuwingen.append(f"'CF - kosten': geen termijnblok '{titel}'{f' (nr. {nr + 1})' if nr else ''} voor {t.naam}: "
                                      f"geen bouwkosten per termijn.")
            continue
        kaart, c_act = _kolomkaart(wk, r0)
        idx_start = _idx(t.start_jaar, t.start_kw) if t.start_jaar and t.start_kw else None
        rijen = []
        r = r0 + 1
        while _txt(wk.cell(r, 2).value):
            lab = _txt(wk.cell(r, 2).value)
            if lab.lower().startswith("aanneemsom") or lab.lower().startswith("meerwerk"):
                break
            pw = _num(wk.cell(r, 5).value)
            if pw:
                real = (_num(wk.cell(r, c_act).value) or 0.0) if c_act else 0.0
                later = []
                for c, (j, q) in sorted(kaart.items()):
                    v = _num(wk.cell(r, c).value)
                    if not v:
                        continue
                    if _idx(j, q) <= idx_act:
                        real += v
                    else:
                        later.append((_idx(j, q), v))
                if not real and not later:
                    rijen.append((lab, round(pw * (t.aantal or 0), 2), None))
                if real:
                    rijen.append((lab + (" (gerealiseerd)" if later else ""), round(pw * real, 2),
                                  idx_act - idx_start + 1 if idx_start else None))
                for i, (idx, v) in enumerate(later):
                    suffix = f" ({i + 1}/{len(later)})" if len(later) > 1 else ""
                    rijen.append((lab + suffix, round(pw * v, 2), idx - idx_start + 1 if idx_start else None))
            r += 1
        t.bouwkosten = _binnen_limiet(rijen, LY.N_TERMIJNEN, t.naam, "bouwkosten-termijnen", waarschuwingen, afkappen=True)


def _lees_kosten(wb, periode_van, idx_act, n_per, waarschuwingen):
    """FO-tabblad 'CF - kosten', blok 'Onderdelen' -> [{naam, mee, gerealiseerd, bedragen per periode}] (zie de moduletekst)."""
    wk = _zoek_blad(wb, "CF - kosten", "CF-kosten", "CF kosten")
    if wk is None:
        waarschuwingen.append("tabblad 'CF - kosten' ontbreekt: tab CF kosten blijft leeg (in een scenario met een andere start bouw "
                              "schuiven dan alle kosten vanaf de geplande start bouw mee).")
        return []
    r_q = next((r for r in range(2, wk.max_row + 1) if _txt(wk.cell(r, 2).value).lower() == "onderdelen"), None)
    kaart, _ = _kolomkaart(wk, r_q) if r_q else ({}, None)
    if not kaart:
        waarschuwingen.append("'CF - kosten': blok 'Onderdelen' met kwartalen niet gevonden: tab CF kosten blijft leeg.")
        return []
    c1 = min(kaart)
    act = set()
    for rr in (r_q - 1, r_q):
        for c in range(3, c1):
            if _txt(wk.cell(rr, c).value).lower().startswith("actual"):
                act.add(c)
                for mr in wk.merged_cells.ranges:
                    if mr.min_row <= rr <= mr.max_row and mr.min_col <= c <= mr.max_col:
                        act.update(range(mr.min_col, mr.max_col + 1))
    if not act:
        act = {c1 - 1}
    laatste = max(periode_van.values()) if periode_van else None
    posten, buiten = [], []
    r = r_q + 1
    while _txt(wk.cell(r, 2).value) and not _txt(wk.cell(r, 2).value).lower().startswith("totaal"):
        naam = _txt(wk.cell(r, 2).value)
        real = sum(_num(wk.cell(r, c).value) or 0.0 for c in act)
        bedragen = [None] * n_per
        for c, (j, q) in sorted(kaart.items()):
            v = _num(wk.cell(r, c).value)
            if not v:
                continue
            i = periode_van.get(_idx(j, q))
            if _idx(j, q) <= idx_act:
                real += v                           # kwartaal t/m actuals: gerealiseerd
            elif i is None and laatste is not None:
                buiten.append(naam)                 # na de laatste periode van '1. Cashflow': in de laatste periode
                bedragen[laatste] = (bedragen[laatste] or 0.0) + v
            elif i is not None:
                bedragen[i] = (bedragen[i] or 0.0) + v
        posten.append({"naam": naam, "mee": standaard_mee(naam), "gerealiseerd": round(real, 2) if real else None,
                       "bedragen": [round(b, 2) if b else None for b in bedragen]})
        r += 1
    if buiten:
        waarschuwingen.append(f"'CF - kosten': bedragen na de laatste periode van '1. Cashflow' staan in de laatste periode "
                              f"({', '.join(sorted(set(buiten)))}).")
    if len(posten) > LY.N_KF:
        waarschuwingen.append(f"'CF - kosten': {len(posten)} posten, alleen de eerste {LY.N_KF} zijn overgenomen "
                              f"({', '.join(p_['naam'] for p_ in posten[LY.N_KF:])} niet).")
        posten = posten[:LY.N_KF]
    return posten
