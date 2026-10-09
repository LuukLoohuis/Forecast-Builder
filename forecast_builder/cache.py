"""Cachewaarden in het gebouwde werkboek zetten.

openpyxl schrijft formules zonder uitkomst. Excel rekent bij het openen alles opnieuw uit (fullCalcOnLoad), maar in
de beveiligde weergave, in voorbeeldweergaven en in viewers die niet rekenen blijven de cellen en grafieken dan leeg.
Deze module rekent het werkboek door met LibreOffice (Nederlandse landinstelling, dus FIXED geeft een komma), leest de
uitkomsten en schrijft ze als cachewaarde in de cellen en in de grafiekreeksen. Zonder LibreOffice wordt deze stap
overgeslagen; het werkboek werkt dan nog steeds, alleen zonder vooraf ingevulde waarden.
"""
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from . import layout as LY

MACRO = """<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE script:module PUBLIC "-//OpenOffice.org//DTD OfficeDocument 1.0//EN" "module.dtd">
<script:module xmlns:script="http://openoffice.org/2000/script" script:name="Module1" script:language="StarBasic">
    Sub RecalculateAndSave()
      ThisComponent.calculateAll()
      ThisComponent.store()
      ThisComponent.close(True)
    End Sub
</script:module>"""
LOCALE_ITEMS = """<item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale" oor:op="fuse"><value>{loc}</value></prop></item>
<item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooLocale" oor:op="fuse"><value>{loc}</value></prop></item>
"""
EXCEL_FOUTEN = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NUM!", "#NAME?", "#NULL!"}


def soffice_pad():
    return shutil.which("soffice") or shutil.which("libreoffice")


def herbereken(pad_xlsx, werkmap, locale="nl-NL", timeout=300):
    """Rekent een kopie van het werkboek door met LibreOffice; geeft het pad van die kopie (of None)."""
    soffice = soffice_pad()
    if not soffice:
        return None
    werkmap = Path(werkmap)
    profiel = werkmap / "lo_profiel"
    kopie = werkmap / ("herberekend_" + os.path.basename(pad_xlsx))
    shutil.copyfile(pad_xlsx, kopie)
    url = profiel.as_uri()
    subprocess.run([soffice, "--headless", "--terminate_after_init", f"-env:UserInstallation={url}"],
                   capture_output=True, timeout=timeout)
    macro_map = profiel / "user" / "basic" / "Standard"
    if not macro_map.exists():
        return None
    (macro_map / "Module1.xba").write_text(MACRO)
    reg = profiel / "user" / "registrymodifications.xcu"
    if reg.exists():
        tekst = reg.read_text(encoding="utf-8")
        if "</oor:items>" in tekst:
            reg.write_text(tekst.replace("</oor:items>", LOCALE_ITEMS.format(loc=locale) + "</oor:items>"), encoding="utf-8")
    r = subprocess.run([soffice, "--headless", f"-env:UserInstallation={url}",
                        "vnd.sun.star.script:Standard.Module1.RecalculateAndSave?language=Basic&location=application", str(kopie)],
                       capture_output=True, timeout=timeout)
    if r.returncode != 0 or not kopie.exists():
        return None
    return str(kopie)


_LEEG_STR = re.compile(r'<c r="([A-Z]+[0-9]+)"[^>]*\st="str"[^>]*>(?:<f[^>]*>.*?</f>|<f[^>]*/>)?(?:<v></v>|<v/>)?</c>', re.S)


def lees_waarden(pad):
    """Uitkomsten per blad; een formule met uitkomst "" wordt als "" opgenomen (openpyxl geeft daar None voor)."""
    wb = load_workbook(pad, data_only=True)
    waarden = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wb.worksheets}
    with zipfile.ZipFile(pad) as z:
        for bestand, blad in _bladen(z).items():
            xml = z.read(bestand).decode("utf-8")
            for m in _LEEG_STR.finditer(xml):
                waarden.setdefault(blad, {}).setdefault(m.group(1), "")
    return waarden


def _v_xml(waarde):
    """(extra attributen, <v>-inhoud) voor een cachewaarde."""
    if isinstance(waarde, bool):
        return ' t="b"', "1" if waarde else "0"
    if isinstance(waarde, (int, float)):
        return "", repr(waarde) if isinstance(waarde, float) else str(waarde)
    s = str(waarde)
    if s in EXCEL_FOUTEN:
        return ' t="e"', s
    return ' t="str"', escape(s)


_CEL = re.compile(r'<c r="([A-Z]+[0-9]+)"((?:\s+[a-z]+="[^"]*")*)>(<f(?:\s[^>]*)?>.*?</f>)<v></v></c>', re.S)


def _patch_sheet(xml, waarden):
    def vervang(m):
        ref, attrs, f = m.group(1), m.group(2), m.group(3)
        w = waarden.get(ref)
        if w is None:
            return m.group(0)
        if w == "":
            attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
            return f'<c r="{ref}"{attrs} t="str">{f}<v></v></c>'
        extra, v = _v_xml(w)
        attrs = re.sub(r'\s+t="[^"]*"', "", attrs)
        return f'<c r="{ref}"{attrs}{extra}>{f}<v>{v}</v></c>'
    return _CEL.sub(vervang, xml)


def _str_cache(teksten):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{escape(str(t))}</c:v></c:pt>' for i, t in enumerate(teksten) if t not in (None, ""))
    return f'<c:strCache><c:ptCount val="{len(teksten)}"/>{pts}</c:strCache>'


def _num_cache(getallen):
    pts = "".join(f'<c:pt idx="{i}"><c:v>{repr(float(g))}</c:v></c:pt>' for i, g in enumerate(getallen)
                  if isinstance(g, (int, float)) and not isinstance(g, bool))
    return f'<c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="{len(getallen)}"/>{pts}</c:numCache>'


def _patch_chart(xml, model, n):
    """Zet str-/numCache in elke reeks: categorieën (g_kwartaal), waarden (g_*) en reeksnamen (Model!$BY$..)."""
    kolom = {"g_kwartaal": LY.FIX["kwartaal"]}
    kolom.update({naam: col for naam, col in LY.M.items() if naam.startswith("g_")})
    for k in range(1, LY.N_TYPES + 1):
        kolom[f"g_v{k}"] = LY.m_col("gv", k)
        kolom[f"g_t{k}"] = LY.m_col("gt", k)
    lengte = {}
    bt_n = model.get(f"{LY.M_HULP}{LY.H['bt_n']}")
    bt_n = int(bt_n) if isinstance(bt_n, (int, float)) else 0
    pt_start = model.get(f"{LY.M_HULP}{LY.H['pt_start']}")
    pt_n = model.get(f"{LY.M_HULP}{LY.H['pt_n']}")
    pt_start = int(pt_start) if isinstance(pt_start, (int, float)) and pt_start >= 1 else 1
    pt_n = int(pt_n) if isinstance(pt_n, (int, float)) and pt_n >= 1 else n
    start = {}
    for key, col in [(k, LY.PT[k]) for k in LY.PT_REEKSEN] + [("kwartaal", LY.FIX["kwartaal"]), ("realisatie", LY.M["g_realisatie"])]:
        kolom[f"g_pt_{key}"] = col
        lengte[f"g_pt_{key}"] = pt_n
        start[f"g_pt_{key}"] = LY.ROW1 + pt_start - 1
    for key in ["label"] + LY.BT_REEKSEN:
        kolom[f"g_bt_{key}"] = LY.BT[key]
        lengte[f"g_bt_{key}"] = max(1, bt_n)

    def reeks(naam):
        col = kolom.get(naam)
        if col is None:
            return None
        r0 = start.get(naam, LY.ROW1)
        return [model.get(f"{col}{r}") for r in range(r0, r0 + lengte.get(naam, n))]

    def cat(m):
        naam = m.group(1)
        waarden = reeks(naam)
        return m.group(0) if waarden is None else f"<c:f>[0]!{naam}</c:f>{_str_cache(waarden)}</c:strRef>"

    def val(m):
        naam = m.group(1)
        waarden = reeks(naam)
        return m.group(0) if waarden is None else f"<c:f>[0]!{naam}</c:f>{_num_cache(waarden)}</c:numRef>"

    def tx(m):
        ref = m.group(1)
        w = model.get(ref)
        return m.group(0) if w is None else f"<c:f>Model!{m.group(2)}</c:f>{_str_cache([w])}</c:strRef>"

    xml = re.sub(r"<c:f>\[0\]!(g_[a-z_0-9]+)</c:f></c:strRef>", cat, xml)
    xml = re.sub(r"<c:f>\[0\]!(g_[a-z_0-9]+)</c:f></c:numRef>", val, xml)
    xml = re.sub(r"<c:f>Model!(\$([A-Z]+\$[0-9]+))</c:f></c:strRef>", lambda m: tx(_M(m)), xml)
    if "g_bt_" in xml:
        # bouwtermijnengrafiek: de vaste ondergrens van beide tijd-assen op t_first (Model!BY97), zoals het Model die uitrekent
        # (twee jaar vóór de eerste termijn); de macro zet de PowerPoint-grafiek op dezelfde cel
        t_first = model.get(f"{LY.M_HULP}{LY.H['t_first']}")
        t_end = model.get(f"{LY.M_HULP}{LY.H['t_end']}")
        if isinstance(t_first, (int, float)) and t_first > 0:
            xml = re.sub(r'<c:min val="[^"]*"/>', f'<c:min val="{int(t_first)}"/>', xml)
        if isinstance(t_end, (int, float)) and t_end > 0:
            xml = re.sub(r'<c:max val="[^"]*"/>', f'<c:max val="{int(t_end)}"/>', xml)
    return xml


class _M:
    """Kleine adapter: group(1) = celverwijzing zonder $, group(2) = originele tekst."""
    def __init__(self, m):
        self._m = m

    def group(self, i):
        if i == 1:
            return self._m.group(2).replace("$", "")
        if i == 2:
            return self._m.group(1)
        return self._m.group(0)


def _bladen(z):
    """Bestandsnaam in het zip -> bladnaam, via workbook.xml en de rels."""
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    def norm(t):
        t = t.lstrip("/")
        return t if t.startswith("xl/") else "xl/" + t          # LibreOffice schrijft relatieve doelen (worksheets/sheet1.xml)
    doel = {}
    for rel in re.findall(r"<Relationship\b[^>]*>", rels):
        if "worksheet" not in rel:
            continue
        t = re.search(r'Target="([^"]+)"', rel)
        i = re.search(r'Id="([^"]+)"', rel)
        if t and i:
            doel[i.group(1)] = norm(t.group(1))
    uit = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb):
        if m.group(2) in doel:
            uit[doel[m.group(2)]] = m.group(1)
    return uit


def voeg_caches_toe(pad_in, pad_uit, waarden):
    """Schrijft cachewaarden in cellen en grafieken van pad_in en bewaart als pad_uit."""
    model = waarden.get("Model", {})
    n = model.get(LY.h("n").replace("$", ""))
    n = int(n) if isinstance(n, (int, float)) else 0
    with zipfile.ZipFile(pad_in) as zin:
        bladen = _bladen(zin)
        with zipfile.ZipFile(pad_uit, "w", zipfile.ZIP_DEFLATED) as zout:
            for naam in zin.namelist():
                data = zin.read(naam)
                if naam in bladen and bladen[naam] in waarden:
                    data = _patch_sheet(data.decode("utf-8"), waarden[bladen[naam]]).encode("utf-8")
                elif re.fullmatch(r"xl/charts/chart\d+\.xml", naam) and n > 0:
                    data = _patch_chart(data.decode("utf-8"), model, n).encode("utf-8")
                zout.writestr(naam, data)


def caches_vullen(paden, werkmap=None):
    """Rekent het eerste .xlsx uit `paden` door en zet de uitkomsten als cache in alle paden. Geeft True als gelukt."""
    xlsx = next((p for p in paden if p.lower().endswith(".xlsx")), None)
    if xlsx is None:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        werk = werkmap or tmp
        kopie = herbereken(xlsx, werk)
        if kopie is None:
            return False
        waarden = lees_waarden(kopie)
        for pad in paden:
            tussen = os.path.join(werk, "cache_" + os.path.basename(pad))
            voeg_caches_toe(pad, tussen, waarden)
            shutil.move(tussen, pad)
    return True
