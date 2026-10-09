"""Bouwt de vijf tabbladen. Alle formules staan hier."""
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from . import chartxml as CX, layout as LY
from .layout import FIX, H, M, N_TYPES, ROW1, ROWN, h, hm, dabs, m_col, m_range_abs
from .styles import *  # noqa: F401,F403

MINUS = "−"       # typografisch minteken, zoals in het bestaande werkboek
AMBER_TXT = "B87400"   # tekstkleur van de scenariolijn (donkerder dan de lijn zelf, leesbaar op wit)
VERKOOP_TITEL = ("Verkoop en transport per woningtype en per kwartaal · aantal woningen · cijfer boven de stapel = totaal per kwartaal · "
                 "grijs vlak = gerealiseerde kwartalen")
APOS = "’"
AUTEUR = "Claude"


def _comment(text):
    c = Comment(text, AUTEUR)
    c.width, c.height = 320, 110
    return c


def eur_m(ref, dec=1):
    """Formulefragment: bedrag als '−€1,2M' (M = mln)."""
    return f'IF({ref}<0,"{MINUS}","")&"€"&FIXED(ABS({ref})/1000000,{dec})&"M"'


def eur_mln(ref, dec=1):
    return f'IF({ref}<0,"{MINUS}","")&"€"&FIXED(ABS({ref})/1000000,{dec})&" mln"'


# =============================================================================
#  WONINGTYPES
# =============================================================================
def bouw_woningtypes(wb, data):
    ws = wb.create_sheet("Woningtypes")
    ws.sheet_properties.codeName = "shWoningtypes"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"{L(LY.WT_COL1)}{LY.WT_R_NAAM}"
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 36
    for k in range(1, N_TYPES + 1):
        ws.column_dimensions[L(LY.wt_col(k, 0))].width = 27
        ws.column_dimensions[L(LY.wt_col(k, 1))].width = 9
        ws.column_dimensions[L(LY.wt_col(k, 2))].width = 7
    ws.row_dimensions[LY.WT_R_TERMIJNKOP].height = 18
    ws.row_dimensions[4].height = 22      # ruimte voor de knoppen Type invoegen / Type verwijderen (.xlsm)

    put(ws, "B1", "Woningtypes — per type één blok van drie kolommen", f=F_TITLE)
    put(ws, "B2", "Blauw op lichtblauw = invoer. Naam en aantal altijd; koopsom, start bouw en termijnen alleen als je het "
                  "verkooptempo-model gebruikt (schakelaar op tab Dashboard). Elk type heeft zijn eigen termijnen.", f=F_NOTE)
    put(ws, "B3", f"Type toevoegen: vul het eerstvolgende lege blok, of voeg op de gewenste plek drie kolommen in (en twee kolommen op tab "
                  f"Invoer). Type verwijderen: haal hier drie kolommen weg en op tab Invoer twee. Alles schuift mee; maximaal {N_TYPES} types. "
                  "Met macro's: knoppen 'Type invoegen' en 'Type verwijderen' (alleen in de .xlsm).", f=F_NOTE)
    put(ws, f"B{LY.WT_R_HEADER - 1}", "1 · WONINGTYPES · naam, aantal, koopsom en start bouw", f=F_SECTION)
    labels = {
        LY.WT_R_NAAM: "Type (naam)",
        LY.WT_R_AANTAL: "Aantal woningen",
        LY.WT_R_KOOPSOM: "Koopsom per woning (€)",
        LY.WT_R_STARTJAAR: "Start bouw · jaar",
        LY.WT_R_STARTKW: "Start bouw · kwartaal (1-4)",
        LY.WT_R_STARTTEKST: "Start bouw (controle)",
        LY.WT_R_GROND: "Grondtermijn (bij notarieel transport)",
        LY.WT_R_TOTAAL: "Totaal bouwtermijnen (100%, óf 100% − grond) · kw = bouwtijd",
    }
    for r, t in labels.items():
        put(ws, f"B{r}", t, f=F_NOTE, al=AL_VCENTER)
    for i in range(LY.N_TERMIJNEN):
        put(ws, f"B{LY.WT_R_T1 + i}", f"Bouwtermijn {i + 1}", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_X_TITEL}", "3 · EXTRA OPBRENGSTEN PER WONING · € per woning, bovenop de koopsom (kopersmeerwerk, kadastrale kosten, "
                                   "overige) · kw = bouwkwartaal waarin het binnenkomt · leeg of 0 = bij notarieel transport", f=F_SECTION)
    ws[f"B{LY.WT_R_X_KOP}"].comment = _comment(
        "Extra opbrengsten per woning in euro's, naast de koopsom: kopersmeerwerk, kadastrale kosten/rentes, overige opbrengsten, "
        "verschillen. Per regel een bedrag per woning en het bouwkwartaal waarin het binnenkomt (zoals bij de termijnen); leeg of 0 = "
        "in het kwartaal van notarieel transport. Het model telt ze mee in de opbrengst per woning; de koopsomknop op het Dashboard "
        "werkt er niet op.")
    for i in range(LY.N_EXTRA):
        put(ws, f"B{LY.WT_R_X1 + i}", f"Extra opbrengst {i + 1}", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_X_TOTAAL}", "Totaal extra per woning (€)", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_X_TOTAAL2}", "Koopsom + extra per woning (€)", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_V_TITEL}", "4 · VORIGE PROGNOSE · start bouw volgens de vorige prognose: de lichtrode blokjes in de bouwtermijnengrafiek "
                                   "(Dashboard en dia 5) · leeg = geen vorige prognose", f=F_SECTION)
    put(ws, f"B{LY.WT_R_VSTARTJAAR}", "Start bouw vorige prognose · jaar", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_VSTARTKW}", "Start bouw vorige prognose · kwartaal (1-4)", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_VSTARTTEKST}", "Vorige prognose (controle)", f=F_NOTE, al=AL_VCENTER)
    ws[f"B{LY.WT_R_VSTARTJAAR}"].comment = _comment("Start bouw zoals die in de vorige prognose stond. De bouwtermijnen schuiven dan in de grafiek "
                                                     "van lichtrood (vorige prognose) naar de typekleur (huidige planning). Leeg laten als er geen vorige "
                                                     "prognose is: dan staan er alleen blokjes in de typekleur.")
    put(ws, f"B{LY.WT_R_TERMIJNTITEL}", "2 · TERMIJNEN PER TYPE · % van de koopsom · kw = bouwkwartaal waarin de termijn vervalt "
                                        "(1 = kwartaal van start bouw van dat type) · naam vrij", f=F_SECTION)
    ws[f"B{LY.WT_R_TERMIJNKOP}"].comment = _comment(
        "Bouwtermijnen: per type een eigen lijst. Naam vrij (bijvoorbeeld 'Casco gereed'), percentage van de koopsom en het "
        "bouwkwartaal waarin de termijn vervalt, geteld vanaf de start bouw van dat type. Niet van toepassing: leeg laten.")
    ws[f"B{LY.WT_R_GROND}"].comment = _comment(
        "Eerste termijn van de koopsom: de grondtermijn. Komt binnen in de periode van notarieel transport en schuift dus mee "
        "als je met het verkooptempo speelt.")

    # blok 5: soort en fees (DAEB)
    put(ws, f"B{LY.WT_R_D_TITEL}", "5 · SOORT EN FEES · niet-DAEB = koopsom en bouwtermijnen (blok 1 en 2) · DAEB = alleen fees, "
                                   "bedragen voor het hele type; koopsom, grondtermijn, bouwtermijnen en extra's tellen dan niet mee", f=F_SECTION)
    put(ws, f"B{LY.WT_R_SOORT}", "Soort (niet-DAEB / DAEB)", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_D_NOTE}", "Per termijn een bedrag in € of een percentage van het componenttotaal (bv. 30%) · kwartaal = wanneer het binnenkomt: "
                                  "schrijf het kalenderkwartaal ('2026 Q4'), of een getal t.o.v. start bouw van dit type (1 = kwartaal van start "
                                  "bouw, 2 = het kwartaal erna, 0 = het kwartaal vóór start bouw, -1 = twee ervoor), of 'actuals' = al ontvangen", f=F_NOTE8)
    put(ws, f"B{LY.WT_R_F_KOP}", "Component · totaal (€) · gepland", f=F_NOTE, al=AL_VCENTER)
    for i in range(LY.N_FEE_COMP):
        put(ws, f"B{LY.WT_R_FC1 + i}", f"Fee-component {i + 1}", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_FT_KOP}", "Termijnen: mijlpaal · € of % · kwartaal", f=F_NOTE, al=AL_VCENTER)
    for i in range(LY.N_TERMIJNEN):
        c = i // LY.N_FEE_PER_COMP + 1
        put(ws, f"B{LY.WT_R_FT1 + i}", f"Component {c} · termijn {i % LY.N_FEE_PER_COMP + 1}", f=F_NOTE, al=AL_VCENTER)
    put(ws, f"B{LY.WT_R_F_TOTAAL}", "Totaal fees gepland (€)", f=F_NOTE, al=AL_VCENTER)
    # blok 6: vorige prognose van de fee-kwartalen
    put(ws, f"B{LY.WT_R_VF_TITEL}", "6 · VORIGE PROGNOSE FEES (DAEB) · per mijlpaal het kwartaal volgens de vorige prognose, zelfde schrijfwijze als "
                                    "blok 5 ('2026 Q4', bouwkwartaal of 'actuals') · leeg = zoals nu (een bouwkwartaal schuift mee met de start bouw van "
                                    "blok 4) · lichtrood in de bouwtermijnengrafiek", f=F_SECTION)
    put(ws, f"B{LY.WT_R_VF_KOP}", "Mijlpaal · kwartaal nu · kwartaal vorige prognose", f=F_NOTE, al=AL_VCENTER)
    for i in range(LY.N_TERMIJNEN):
        put(ws, f"B{LY.WT_R_VF1 + i}", f"Fee-termijn {i + 1}", f=F_NOTE, al=AL_VCENTER)
    ws[f"B{LY.WT_R_VF_KOP}"].comment = _comment(
        "De knop 'Ophalen uit FO' bewaart (na een vraag) de planning die vóór het ophalen in dit werkboek stond als vorige prognose: start "
        "bouw per type in blok 4 en deze fee-kwartalen in blok 6. De knop 'Vorige prognose uit FO' haalt beide uit een ouder FO. "
        "Handmatig invullen mag ook.")
    ws[f"B{LY.WT_R_SOORT}"].comment = _comment(
        "DAEB: sociale huurwoningen met gescheiden koop- en aannemingsovereenkomst. Vastgoed krijgt dan geen koopsom "
        "en geen bouwtermijnen, maar fees: vul hieronder per component het totaalbedrag en de termijnen in. Aantal, start bouw en "
        "verkocht/transport op tab Invoer blijven nodig (voor de verkoopgrafiek en de tabel): zet bij een DAEB-type het hele aantal in het "
        "kwartaal van tekenen (verkocht) en van levering (transport). Een project mag DAEB en niet-DAEB mengen.")
    ws[f"B{LY.WT_R_FT_KOP}"].comment = _comment(
        "Per termijn: de mijlpaal (naam vrij), het bedrag in euro's of een percentage van het componenttotaal (een getal van 1 of kleiner "
        "telt als percentage), en het kwartaal: '2026 Q4' of 'Q4 2026', een bouwkwartaal als getal (1 = kwartaal van start bouw, 0 = het "
        "kwartaal ervoor, 7 = zes kwartalen na start bouw) of 'actuals' voor termijnen die al in je eigen cashflow zitten. Percentages "
        "worden niet opgeschaald: 'gepland' laat zien hoeveel van het totaal is ingepland; het restant telt niet mee.")

    dv_kw4 = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True, error="Kwartaal 1 t/m 4.")
    dv_kw = DataValidation(type="whole", operator="between", formula1="1", formula2="60", allow_blank=True, error="Bouwkwartaal: heel getal 1 t/m 60.")
    dv_kwx = DataValidation(type="whole", operator="between", formula1="0", formula2="60", allow_blank=True,
                            error="Bouwkwartaal: heel getal 1 t/m 60; leeg of 0 = bij notarieel transport.")
    dv_pct = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True, error="Percentage tussen 0% en 100%.")
    dv_soort = DataValidation(type="list", formula1='"niet-DAEB,DAEB"', allow_blank=True, error="Kies niet-DAEB of DAEB (leeg = niet-DAEB).")
    dv_eur = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True, error="Bedrag in euro's (of een percentage van 1 of kleiner).")
    # invoerbericht bij het kwartaal van een fee-termijn (verschijnt als de cel geselecteerd is)
    dv_fkw = DataValidation(type=None, allow_blank=True)
    dv_fkw.showInputMessage = True
    dv_fkw.promptTitle = "Kwartaal van deze termijn"
    dv_fkw.prompt = ("Kalenderkwartaal: '2026 Q4' of 'Q4 2026'. Of een getal t.o.v. start bouw van dit type: 1 = kwartaal van start bouw, "
                     "2 = het kwartaal erna, 0 = kwartaal ervoor, -1 = twee ervoor. Of 'actuals' = al ontvangen.")
    dv_vkw = DataValidation(type=None, allow_blank=True)
    dv_vkw.showInputMessage = True
    dv_vkw.promptTitle = "Kwartaal vorige prognose"
    dv_vkw.prompt = ("Kwartaal van deze termijn volgens de vorige prognose, zelfde schrijfwijze als in blok 5 ('2026 Q4', bouwkwartaal of "
                     "'actuals'). Leeg = zoals nu; een bouwkwartaal schuift mee met de start bouw van blok 4.")
    for dv in (dv_kw4, dv_kw, dv_kwx, dv_pct, dv_soort, dv_eur):
        dv.showErrorMessage = True
        dv.errorTitle = "Ongeldige invoer"
    for dv, titel, tekst in ((dv_kw, "Bouwkwartaal", "Kwartaal waarin de termijn vervalt, geteld vanaf start bouw van dit type: 1 = kwartaal van start bouw, "
                                                      "2 = het kwartaal erna, enz. Verschuift mee als start bouw verschuift."),
                             (dv_kwx, "Bouwkwartaal", "Kwartaal waarin het binnenkomt, geteld vanaf start bouw van dit type (1 = kwartaal van start bouw). "
                                                       "Leeg of 0 = bij notarieel transport.")):
        dv.showInputMessage = True
        dv.promptTitle = titel
        dv.prompt = tekst
    for dv in (dv_kw4, dv_kw, dv_kwx, dv_pct, dv_soort, dv_eur, dv_fkw, dv_vkw):
        ws.add_data_validation(dv)

    totaal_cellen = []
    for k in range(1, N_TYPES + 1):
        c0, c1, c2 = (L(LY.wt_col(k, o)) for o in range(3))
        t = data.type_(k)
        # kop: positioneel ("TYPE 1" volgt de plek van het blok, ook na invoegen of verwijderen van kolommen)
        put(ws, f"{c0}{LY.WT_R_HEADER}", f'="TYPE "&COLUMN()/{LY.WT_W}', f=F_HDR, fl=FL_HDR, al=AL_CENTER_CONT)
        for c in (c1, c2):
            style(ws, f"{c}{LY.WT_R_HEADER}", f=F_HDR, fl=FL_HDR, al=AL_CENTER_CONT)
        put(ws, f"{c0}{LY.WT_R_NAAM}", t.naam or None, f=F_INPUT, fl=FL_INPUT, al=AL_LEFT_TOP)
        put(ws, f"{c0}{LY.WT_R_AANTAL}", t.aantal, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT, al=AL_LEFT_TOP)
        put(ws, f"{c0}{LY.WT_R_KOOPSOM}", t.koopsom, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT, al=AL_LEFT_TOP)
        put(ws, f"{c0}{LY.WT_R_STARTJAAR}", t.start_jaar, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_LEFT_TOP)
        put(ws, f"{c0}{LY.WT_R_STARTKW}", t.start_kw, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_LEFT_TOP)
        dv_kw4.add(f"{c0}{LY.WT_R_STARTKW}")
        put(ws, f"{c0}{LY.WT_R_STARTTEKST}",
            f'=IF(OR(N({c0}{LY.WT_R_AANTAL})=0,N({c0}{LY.WT_R_STARTJAAR})=0),"–","Q"&N({c0}{LY.WT_R_STARTKW})&" "&N({c0}{LY.WT_R_STARTJAAR}))',
            f=F_NOTE, al=AL_LEFT_TOP)
        # termijnen
        put(ws, f"{c0}{LY.WT_R_TERMIJNKOP}", "termijn", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
        put(ws, f"{c1}{LY.WT_R_TERMIJNKOP}", "%", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c2}{LY.WT_R_TERMIJNKOP}", "kw", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c0}{LY.WT_R_GROND}", "Grondtermijn (bij transport)", f=F_NOTE)
        put(ws, f"{c1}{LY.WT_R_GROND}", t.grond_pct, f=F_INPUT, fl=FL_INPUT, nf=FMT_PCT)
        put(ws, f"{c2}{LY.WT_R_GROND}", "transport", f=F_NOTE8, al=Alignment(horizontal="right"))
        dv_pct.add(f"{c1}{LY.WT_R_GROND}")
        for i in range(LY.N_TERMIJNEN):
            r = LY.WT_R_T1 + i
            if i < len(t.termijnen):
                naam, pct, kw = t.termijnen[i]
            else:
                naam, pct, kw = (data.termijn_namen[i] if i < len(data.termijn_namen) else None), None, None
            put(ws, f"{c0}{r}", naam, f=F_INPUT, fl=FL_INPUT)
            put(ws, f"{c1}{r}", pct, f=F_INPUT, fl=FL_INPUT, nf=FMT_PCT)
            put(ws, f"{c2}{r}", kw, f=F_INPUT, fl=FL_INPUT, nf="0")
            dv_pct.add(f"{c1}{r}")
            dv_kw.add(f"{c2}{r}")
        put(ws, f"{c1}{LY.WT_R_TOTAAL}", f"=SUM({c1}{LY.WT_R_T1}:{c1}{LY.WT_R_TN})", f=F_BOLD9, nf=FMT_PCT, al=AL_RIGHT)
        # blok 3: extra opbrengsten per woning (€)
        put(ws, f"{c0}{LY.WT_R_X_KOP}", "omschrijving", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
        put(ws, f"{c1}{LY.WT_R_X_KOP}", "€/won", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c2}{LY.WT_R_X_KOP}", "kw", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        for i in range(LY.N_EXTRA):
            r = LY.WT_R_X1 + i
            if i < len(t.extras):
                xn, eur, kw = t.extras[i]
            else:
                xn, eur, kw = (data.extra_namen[i] if i < len(data.extra_namen) else None) or None, None, None
            put(ws, f"{c0}{r}", xn, f=F_INPUT, fl=FL_INPUT)
            put(ws, f"{c1}{r}", eur, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT)
            put(ws, f"{c2}{r}", kw, f=F_INPUT, fl=FL_INPUT, nf="0")
            dv_kwx.add(f"{c2}{r}")
        put(ws, f"{c1}{LY.WT_R_X_TOTAAL}", f"=SUM({c1}{LY.WT_R_X1}:{c1}{LY.WT_R_XN})", f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{c1}{LY.WT_R_X_TOTAAL2}", f"=N({c0}{LY.WT_R_KOOPSOM})+{c1}{LY.WT_R_X_TOTAAL}", f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT)
        # blok 4: vorige prognose
        put(ws, f"{c0}{LY.WT_R_VSTARTJAAR}", t.vorig_start_jaar, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_LEFT_TOP)
        put(ws, f"{c0}{LY.WT_R_VSTARTKW}", t.vorig_start_kw, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_LEFT_TOP)
        dv_kw4.add(f"{c0}{LY.WT_R_VSTARTKW}")
        put(ws, f"{c0}{LY.WT_R_VSTARTTEKST}",
            f'=IF(OR(N({c0}{LY.WT_R_VSTARTJAAR})=0,N({c0}{LY.WT_R_VSTARTKW})=0),"–","Q"&N({c0}{LY.WT_R_VSTARTKW})&" "&N({c0}{LY.WT_R_VSTARTJAAR}))',
            f=F_NOTE, al=AL_LEFT_TOP)
        put(ws, f"{c2}{LY.WT_R_TOTAAL}", f'=IF(COUNT({c2}{LY.WT_R_T1}:{c2}{LY.WT_R_TN})=0,"",MAX({c2}{LY.WT_R_T1}:{c2}{LY.WT_R_TN}))',
            f=F_BOLD9, nf="0", al=AL_RIGHT)
        totaal_cellen.append(f"{c1}{LY.WT_R_TOTAAL}")
        # blok 5: soort en fees
        put(ws, f"{c0}{LY.WT_R_SOORT}", t.soort or None, f=F_INPUT, fl=FL_INPUT, al=AL_LEFT_TOP)
        dv_soort.add(f"{c0}{LY.WT_R_SOORT}")
        put(ws, f"{c0}{LY.WT_R_F_KOP}", "component", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
        put(ws, f"{c1}{LY.WT_R_F_KOP}", "totaal €", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c2}{LY.WT_R_F_KOP}", "gepland", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        gepland_cellen = []
        for i in range(LY.N_FEE_COMP):
            r = LY.WT_R_FC1 + i
            fn, tot = t.fee_comp[i] if i < len(t.fee_comp) else (None, None)
            put(ws, f"{c0}{r}", fn or data.fee_comp_namen[i] or None, f=F_INPUT, fl=FL_INPUT)
            put(ws, f"{c1}{r}", tot, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT)
            dv_eur.add(f"{c1}{r}")
            r1, rN = LY.WT_R_FT1 + i * LY.N_FEE_PER_COMP, LY.WT_R_FT1 + (i + 1) * LY.N_FEE_PER_COMP - 1
            w = f"{c1}{r1}:{c1}{rN}"
            # gepland = som van de bedragen (> 1) plus de percentages (<= 1) x totaal; als % van het totaal, of in € zonder totaal
            gepland = f'(SUMPRODUCT(--ISNUMBER({w}),--({w}>1),{w})+SUMPRODUCT(--ISNUMBER({w}),--({w}<=1),{w})*N({c1}{r}))'
            put(ws, f"{c2}{r}", f'=IF(COUNT({w})=0,IF(N({c1}{r})>0,0,""),IF(N({c1}{r})>0,{gepland}/{c1}{r},'
                                f'IF(SUMPRODUCT(--ISNUMBER({w}),--({w}<=1),--({w}>0))>0,"% zonder totaal",{gepland})))', f=F_BOLD9, al=AL_RIGHT,
                nf='[<=9]0%;#,##0')
            gepland_cellen.append(f"{c2}{r}")
        put(ws, f"{c0}{LY.WT_R_FT_KOP}", "mijlpaal", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
        put(ws, f"{c1}{LY.WT_R_FT_KOP}", "€ of %", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c2}{LY.WT_R_FT_KOP}", "kwartaal", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        for i in range(LY.N_TERMIJNEN):
            r = LY.WT_R_FT1 + i
            mn, w_, kw = t.fee_termijnen[i] if i < len(t.fee_termijnen) else (None, None, None)
            put(ws, f"{c0}{r}", mn or data.fee_mijlpalen[i] or None, f=F_INPUT, fl=FL_INPUT)
            put(ws, f"{c1}{r}", w_, f=F_INPUT, fl=FL_INPUT, nf='[<=1]0%;#,##0', al=AL_RIGHT)
            put(ws, f"{c2}{r}", kw, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_RIGHT)
            dv_eur.add(f"{c1}{r}")
            dv_fkw.add(f"{c2}{r}")
        wt_ = f"{c1}{LY.WT_R_FT1}:{c1}{LY.WT_R_FTN}"
        put(ws, f"{c1}{LY.WT_R_F_TOTAAL}", f"=SUM(Model!{m_col('fee', k)}$6)", f=F_BOLD9, nf=FMT_INT, al=AL_RIGHT)
        # rood als een component met totaal niet (bijna) volledig is ingepland
        ws.conditional_formatting.add(" ".join(gepland_cellen),
                                      FormulaRule(formula=[f'AND(UPPER(TRIM(${c0}${LY.WT_R_SOORT}&""))="DAEB",N({c1}{LY.WT_R_FC1})>0,'
                                                           f'ISNUMBER({c2}{LY.WT_R_FC1}),ABS({c2}{LY.WT_R_FC1}-1)>=0.005)'],
                                                  font=Font(name=ARIAL, bold=True, color=RED)))
        # bij DAEB: koopsom, grondtermijn, bouwtermijnen en extra's grijs (tellen niet mee)
        daeb = f'UPPER(TRIM(${c0}${LY.WT_R_SOORT}&""))="DAEB"'
        ws.conditional_formatting.add(f"{c0}{LY.WT_R_KOOPSOM} {c0}{LY.WT_R_GROND}:{c2}{LY.WT_R_TOTAAL} {c0}{LY.WT_R_X1}:{c2}{LY.WT_R_X_TOTAAL2}",
                                      FormulaRule(formula=[daeb], font=Font(name=ARIAL, color="B0B6BE"),
                                                  fill=PatternFill(fill_type="solid", bgColor="F2F3F5", fgColor="F2F3F5")))
        # blok 6: vorige prognose van de fee-kwartalen (mijlpaal en huidig kwartaal uit blok 5, vorig kwartaal invoer); grijs als niet-DAEB
        put(ws, f"{c0}{LY.WT_R_VF_KOP}", "mijlpaal", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
        put(ws, f"{c1}{LY.WT_R_VF_KOP}", "nu", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        put(ws, f"{c2}{LY.WT_R_VF_KOP}", "vorig", f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
        for i in range(LY.N_TERMIJNEN):
            r, rf = LY.WT_R_VF1 + i, LY.WT_R_FT1 + i
            put(ws, f"{c0}{r}", f'=IF({c0}{rf}="","",{c0}{rf})', f=F_CALC9, al=AL_LEFT)
            put(ws, f"{c1}{r}", f'=IF({c2}{rf}="","",{c2}{rf})', f=F_CALC9, nf="0", al=AL_RIGHT)
            put(ws, f"{c2}{r}", t.fee_vorig[i] if i < len(t.fee_vorig) else None, f=F_INPUT, fl=FL_INPUT, nf="0", al=AL_RIGHT)
            dv_vkw.add(f"{c2}{r}")
        ws.conditional_formatting.add(f"{c0}{LY.WT_R_VF1}:{c2}{LY.WT_R_VFN}",
                                      FormulaRule(formula=[f"NOT({daeb})"], font=Font(name=ARIAL, color="B0B6BE"),
                                                  fill=PatternFill(fill_type="solid", bgColor="F2F3F5", fgColor="F2F3F5")))
    # rood als de bouwtermijnen niet optellen tot 100% (van de aanneemsom) of tot 100% − grondtermijn (van de koopsom); per blok, relatief
    c0, c1 = L(LY.wt_col(1, 0)), L(LY.wt_col(1, 1))
    ws.conditional_formatting.add(
        " ".join(totaal_cellen),
        FormulaRule(formula=[f"AND(N({c0}{LY.WT_R_AANTAL})>0,UPPER(TRIM({c0}{LY.WT_R_SOORT}&\"\"))<>\"DAEB\",{c1}{LY.WT_R_TOTAAL}<>0,"
                             f"ABS({c1}{LY.WT_R_TOTAAL}-1)>=0.00005,ABS({c1}{LY.WT_R_TOTAAL}-(1-N({c1}{LY.WT_R_GROND})))>=0.00005)"],
                    font=Font(name=ARIAL, bold=True, color=RED)))
    put(ws, f"B{LY.WT_R_TOTAAL + 2}", "Bouwtermijnen vul je in zoals ze bij jullie staan: als % van de aanneemsom (samen 100%) of als % van de koopsom "
                                      "(samen 100% − grondtermijn). Het model schaalt ze zelf naar 100% − grondtermijn, zodat grondtermijn + bouwtermijnen "
                                      "altijd de hele koopsom vormen; terugrekenen hoeft niet.", f=F_NOTE8)
    put(ws, f"B{LY.WT_R_TOTAAL + 3}", "Termijnen die bij transport al vervallen zijn, komen in het transportkwartaal in één keer binnen, "
                                      "bovenop de grondtermijn. Een type zonder start bouw telt alleen de grondtermijn mee.", f=F_NOTE8)
    return ws


# =============================================================================
#  INVOER
# =============================================================================
def bouw_invoer(wb, data):
    ws = wb.create_sheet("Invoer")
    ws.sheet_properties.codeName = "shInvoer"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "D8"
    for col, w in zip("ABCDEFGHIJK", (2, 7, 6, 13, 9, 13, 9, 11, 11, 12, 3)):
        ws.column_dimensions[col].width = w
    for k in range(1, N_TYPES + 1):
        ws.column_dimensions[L(LY.in_col(k, 0))].width = 11
        ws.column_dimensions[L(LY.in_col(k, 1))].width = 11
    ws.row_dimensions[7].height = 30

    put(ws, "B1", "Invoer — cashflow uit je eigen bestand, verkoop en transport per periode", f=F_TITLE)
    put(ws, "B2", "Blauw op lichtblauw = invoer. Grijs = wordt niet gebruikt, mag gewoon meegeplakt worden. "
                  "Woningtypes, koopsommen en termijnen staan op tab Woningtypes.", f=F_NOTE)
    put(ws, "B3", "Kopieer in tab '4. Cashflow' het blok B4:J(laatste rij) en plak hier in B8 met Plakken speciaal → Waarden. "
                  "Kosten en omzet in euro's, 'CF vorig kwartaal' in x € 1.000.", f=F_NOTE)
    put(ws, "B4", "Je eigen cashflowbestand blijft ongewijzigd: je kopieert er alleen cijfers uit. Rijen (kwartalen) mag je hier gewoon "
                  "verwijderen of invoegen; lege rijen tellen niet mee. Maximaal 60 kwartalen. Sneller: de knop 'Ophalen uit FO' (.xlsm, rechtsboven) "
                  "leest '1. Cashflow', 'CF - opbrengsten' en 'FO - actuals' uit het FO-werkboek en vult deze tab en tab Woningtypes.", f=F_NOTE)
    put(ws, "B5", "1 · CASHFLOW UIT JE EIGEN BESTAND", f=F_SECTION)
    # titel in de smalle kolom K: loopt over de paren heen en schuift niet mee als er kolommen bij L worden ingevoegd
    put(ws, f"{L(LY.IN_COL1 - 1)}5", "2 · VERKOCHT EN GETRANSPORTEERD · aantal woningen per periode, per woningtype · "
                                     "verkocht naast notarieel transport · leeg = 0", f=F_SECTION)
    ws[f"{L(LY.IN_COL1 - 1)}5"].comment = _comment(
        "Per type twee kolommen naast elkaar: verkocht en notarieel getransporteerd. Bij transport komt de grondtermijn binnen. "
        "Type toevoegen of verwijderen: twee kolommen invoegen of verwijderen (zie tab Woningtypes).")

    koppen = ["Jaar", "Q", "Kosten", "% Kosten", "Omzet", "% Omzet", "CF", "CF (x 1000)", "CF vorig kwartaal"]
    grijs = {"E", "G", "H", "I"}
    for i, kop in enumerate(koppen):
        col = L(2 + i)
        put(ws, f"{col}7", kop, f=F_HDR, fl=FL_HDR_GREY if col in grijs else FL_HDR,
            al=AL_LEFT_WRAP if col in ("B", "C") else AL_RIGHT_WRAP)
    ws["C7"].comment = _comment("Q1 t/m Q4 (of 1 t/m 4). Staat hier een jaartal of iets anders, dan telt de regel als een heel jaar.")
    ws["J7"].comment = _comment("Cumulatieve cashflow volgens de vorige prognose, in x € 1.000 (kolom J in je eigen bestand). "
                                "Alles nul of leeg = geen vergelijking.")
    for k in range(1, N_TYPES + 1):
        cv, ct = L(LY.in_col(k, 0)), L(LY.in_col(k, 1))
        # typenaam boven het paar: positioneel, zodat de kop meeschuift met invoegen/verwijderen van kolommen
        kk = f"((COLUMN()-{LY.IN_COL1})/{LY.IN_W}+1)"
        naam = f"INDEX(Woningtypes!$A${LY.WT_R_NAAM}:${LY.WT_RANGE_END}${LY.WT_R_NAAM},1,{LY.WT_W}*{kk})"
        put(ws, f"{cv}{LY.IN_R_TYPE}", f'=IF({naam}="","type "&{kk},{naam})', f=F_BOLD9, al=AL_CENTER_CONT)
        style(ws, f"{ct}{LY.IN_R_TYPE}", f=F_BOLD9, al=AL_CENTER_CONT)
        put(ws, f"{cv}{LY.IN_R_KOP}", "Verkocht", f=F_HDR, fl=FL_HDR, al=AL_RIGHT_WRAP)
        put(ws, f"{ct}{LY.IN_R_KOP}", "Transport", f=F_HDR, fl=FL_HDR, al=AL_RIGHT_WRAP)

    for i in range(LY.N_PERIODS):
        r = ROW1 + i
        rij = data.periodes[i] if i < len(data.periodes) else {}
        for j, naam in enumerate(["jaar", "kw", "kosten", "pct_kosten", "omzet", "pct_omzet", "cf", "cf1000", "cf_vorig"]):
            col = L(2 + j)
            v = rij.get(naam)
            if col in ("B", "C"):
                put(ws, f"{col}{r}", v, f=F_INPUT, fl=FL_INPUT, al=AL_LEFT_TOP)
            elif col in grijs:
                put(ws, f"{col}{r}", v, f=F_GREY, nf=FMT_INT)
            else:
                put(ws, f"{col}{r}", v, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT)
        for k in range(1, N_TYPES + 1):
            vv = data.verkocht[i][k - 1] if i < len(data.verkocht) and k - 1 < len(data.verkocht[i]) else None
            tv = data.transport[i][k - 1] if i < len(data.transport) and k - 1 < len(data.transport[i]) else None
            put(ws, f"{L(LY.in_col(k, 0))}{r}", vv, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT)
            put(ws, f"{L(LY.in_col(k, 1))}{r}", tv, f=F_INPUT, fl=FL_INPUT, nf=FMT_INT)
    return ws


# =============================================================================
#  MODEL
# =============================================================================
def bouw_model(wb):
    ws = wb.create_sheet("Model")
    ws.sheet_properties.codeName = "shModel"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"{FIX['idx']}{ROW1}"
    laatste = LY.M_BLOK1 + len(LY.BLOKKEN) * N_TYPES - 1
    for c in range(1, laatste + 1):
        ws.column_dimensions[L(c)].width = 11
    ws.column_dimensions[LY.M_SPACER].width = 3
    ws.column_dimensions[LY.M_HULP_LABEL].width = 36
    ws.column_dimensions[LY.M_HULP].width = 30
    ws.column_dimensions[L(LY.M_BLOK1 - 1)].width = 3
    ws.row_dimensions[7].height = 34
    NR, BRON, KW, KWT, KWV, IDX, FASE, KOS, OPB, STAND, CUM, VORIG = (FIX[k] for k in
        ("nr", "bron", "kw", "kwartaal", "kwartaal_vol", "idx", "fase", "kosten", "opbr", "stand", "cum_opbr", "vorig"))
    inv_b = LY.in_rng("B")
    inv_j = LY.in_rng("J")

    put(ws, "A1", "Model — rekenblad, niets invullen", f=F_TITLE)
    put(ws, "A2", "Rij i van dit blad is de i-de gevulde rij op tab Invoer (kolom Bronrij): lege rijen tellen niet mee en rijen invoegen of "
                  "verwijderen op Invoer kan gewoon. Elke kolom rekent met de invoer, de typeblokken (vanaf kolom CA), kolommen links ervan of de "
                  "vorige rij: geen kringverwijzingen. #N/B in grafiekkolommen is normaal. Basis = je eigen cashflow (kolom J); scenario's = basis + "
                  "effect van de knoppen; rente = negatieve stand van het vorige kwartaal × jaarrente / 4.", f=F_NOTE)
    put(ws, f"{M['model_basis']}4", "modelkolommen: opbrengsten volgens woningtypes en termijnen, scenario's, rente en de 'getoonde' standen (met of zonder rente in de basis)", f=F_NOTE8)
    put(ws, f"{M['g_voorfinanciering']}4", "grafiekkolommen (x € 1 mln)", f=F_NOTE8)
    bloklabels = {
        "vcum": "rij 5: verkocht vóór start bouw · rij 6: aantal woningen",
        "tcum": "rij 5: aantal (types met start bouw) · rij 6: koopsom per woning",
        "tup": "rij 5: schaal bouwtermijnen = (100% − grond) / som bouwtermijnen · rij 6: grondtermijn %",
        "tdown": "rij 6: getransporteerd t/m actuals",
        "verv": "rij 6: index start bouw (jaar × 4 + kwartaal; 99999 = geen)",
        "tscn": "transport cumulatief in de scenariolijn (verschuiving uit de derde knoppenkolom)",
        "gv": "grafiek: verkocht per kwartaal per type (rij 7 = typenaam)",
        "gt": "grafiek: getransporteerd per kwartaal per type",
        "vx": "rij 6: extra opbrengsten per woning bij transport (€) · rijen: extra's vervallen per bouwkwartaal (€ per woning)",
        "fee": "rij 5: DAEB (1/0) · rij 6: totaal fees gepland (€) · rijen: fees vervallen t/m de periode (€, cumulatief, hele type)",
    }
    for blok, t in bloklabels.items():
        put(ws, f"{m_col(blok, 1)}4", t, f=font(8, True, TXT2))

    koppen = {
        NR: "Nr", BRON: "Bronrij (Invoer)", KW: "Kw", KWT: "Kwartaal", KWV: "Kwartaal volledig", IDX: "Idx", FASE: "Fase",
        KOS: "Kosten", OPB: "Opbrengsten", STAND: "Stand", CUM: "Cum opbrengsten", VORIG: "Vorige prognose",
        M["model_basis"]: "Model basis", M["model_up"]: "Model upside", M["model_down"]: "Model downside",
        M["cum_opbr_up"]: "Cum opbr upside", M["cum_opbr_down"]: "Cum opbr downside",
        M["opbr_up"]: "Opbrengsten upside", M["opbr_down"]: "Opbrengsten downside",
        M["kosten_up"]: "Kosten upside", M["kosten_down"]: "Kosten downside",
        M["rente_basis"]: "Rente basis (berekend)", M["stand_basis_rente"]: "Stand basis incl. rente",
        M["rente_up"]: "Rente upside", M["stand_up_raw"]: "Stand upside (berekend)",
        M["rente_down"]: "Rente downside", M["stand_down_raw"]: "Stand downside (berekend)",
        M["stand_basis"]: "Stand basis (getoond)", M["kosten_basis"]: "Kosten basis (getoond)",
        M["stand_up"]: "Stand upside (getoond)", M["stand_down"]: "Stand downside (getoond)",
        M["pos_basis"]: "Positief", M["pos_up"]: "Positief upside", M["pos_down"]: "Positief downside", M["pos_vorig"]: "Positief vorig",
        M["model_scn"]: "Model scenariolijn", M["cum_opbr_scn"]: "Cum opbr scenariolijn", M["opbr_scn"]: "Opbrengsten scenariolijn",
        M["kosten_scn"]: "Kosten scenariolijn", M["rente_scn"]: "Rente scenariolijn", M["stand_scn_raw"]: "Stand scenariolijn (berekend)",
        M["stand_scn"]: "Stand scenariolijn (getoond)", M["pos_scn"]: "Positief scenariolijn",
        M["g_voorfinanciering"]: "Voorfinanciering", M["g_positief_saldo"]: "Positief saldo", M["g_opbrengsten"]: "Opbrengsten",
        M["g_kosten"]: "Kosten", M["g_stand"]: "Stand", M["g_vorige"]: "Vorige", M["g_band_onder"]: "Band onder",
        M["g_bandbreedte"]: "Bandbreedte", M["g_downside"]: "Downside", M["g_upside"]: "Upside", M["g_punt_nu"]: "Punt nu",
        M["g_punt_dal"]: "Punt dal", M["g_eind_upside"]: "Eind upside", M["g_eind_downside"]: "Eind downside",
        M["g_eind_basis"]: "Eind basis", M["g_verkocht"]: "Verkocht", M["g_getransporteerd"]: "Getransporteerd",
        M["g_verkocht_cum"]: "Verkocht cum", M["g_getransporteerd_cum"]: "Getransporteerd cum",
        M["g_verkocht_pct"]: "Verkocht pct", M["g_getransporteerd_pct"]: "Getransporteerd pct",
        M["g_opbrengsten_pct"]: "Opbrengsten pct", M["g_kosten_pct"]: "Kosten pct",
        M["g_scenario"]: "Scenariolijn", M["g_eind_scenario"]: "Eind scenariolijn", M["g_realisatie"]: "Realisatie (waas)",
        M["g_punt_uitverkocht"]: "Punt uitverkocht", M["g_punt_alles_transport"]: "Punt laatste transport",
        M["stap_ok"]: "Stap ok (controle)",
    }
    for k in range(1, N_TYPES + 1):
        koppen[m_col("vcum", k)] = f"Verkocht cum {k}"
        koppen[m_col("tcum", k)] = f"Transport cum {k}"
        koppen[m_col("tup", k)] = f"Transport cum upside {k}"
        koppen[m_col("tdown", k)] = f"Transport cum downside {k}"
        koppen[m_col("verv", k)] = f"Bouwtermijnen vervallen {k}"
        koppen[m_col("tscn", k)] = f"Transport cum scenariolijn {k}"
        koppen[m_col("vx", k)] = f"Extra's vervallen {k} (€/won)"
        koppen[m_col("fee", k)] = f"Fees vervallen {k} (€)"
    for col, t in koppen.items():
        put(ws, f"{col}7", t, f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    for k in range(1, N_TYPES + 1):   # typenaam als reeksnaam van de verkoopgrafiek (vaste cel, schuift niet mee met Woningtypes)
        naam = LY.wt_ref(LY.WT_R_NAAM, k)
        for blok in ("gv", "gt"):
            put(ws, f"{m_col(blok, k)}7", f'=IF({naam}="","Type {k}",{naam})', f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    put(ws, f"{LY.M_HULP_LABEL}7", "Hulpcellen", f=F_BOLD9)

    # ---- rij 5 en 6: per type ------------------------------------------------
    idx_rng = f"${IDX}${ROW1}:${IDX}${ROWN}"
    BT = LY.BT
    rng_bt = lambda naam: f"${BT[naam]}${LY.BT_ROW1}:${BT[naam]}${LY.BT_ROWN}"  # noqa: E731
    for k in range(1, N_TYPES + 1):
        vc, tc, tu, td, bv, ts, gv, gt, vx, fe = (m_col(b, k) for b in LY.BLOKKEN)
        aantal = f"N({LY.wt_ref(LY.WT_R_AANTAL, k)})"
        # DAEB (blok 5): rij 5 = vlag, rij 6 = totaal fees (gepland en geldig) uit de bouwtermijnentabel
        put(ws, f"{fe}5", f'=IF(UPPER(TRIM({LY.wt_ref(LY.WT_R_SOORT, k)}&""))="DAEB",1,0)', f=F_CALC, nf="0", al=AL_RIGHT)
        put(ws, f"{fe}6", f'=SUMIFS({rng_bt("bedrag")},{rng_bt("k")},{k},{rng_bt("gebruikt")},1)', f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        # extra opbrengsten per woning die bij transport binnenkomen (kw leeg of 0)
        x_eur = LY.wt_ref(0, k, 1, rows=(LY.WT_R_X1, LY.WT_R_XN))
        x_kw = LY.wt_ref(0, k, 2, rows=(LY.WT_R_X1, LY.WT_R_XN))
        put(ws, f"{vx}6", f'=SUMPRODUCT((N(+{x_kw})<1)*N(+{x_eur}))', f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{vc}5", f'=IF(OR({bv}$6>=99999,{fe}$5=1,COUNTIF({idx_rng},"<"&{bv}$6)=0),0,'
                          f'N(INDEX({vc}${ROW1}:{vc}${ROWN},COUNTIF({idx_rng},"<"&{bv}$6))))', f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{vc}6", f"={aantal}", f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{tc}5", f"=IF(OR({bv}$6>=99999,{fe}$5=1),0,{vc}$6)", f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{tc}6", f"=N({LY.wt_ref(LY.WT_R_KOOPSOM, k)})", f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        put(ws, f"{tu}6", f"=N({LY.wt_ref(LY.WT_R_GROND, k, 1)})", f=F_CALC, nf=FMT_PCT, al=AL_RIGHT)
        pct_k = LY.wt_ref(0, k, 1, rows=(LY.WT_R_T1, LY.WT_R_TN))
        put(ws, f"{tu}5", f'=IF(SUM({pct_k})=0,0,MAX(0,1-{tu}$6)/SUM({pct_k}))', f=F_CALC, nf="0.0000", al=AL_RIGHT)
        put(ws, f"{td}6", f"=IF({h('pos_actuals')}=0,0,N(INDEX({tc}${ROW1}:{tc}${ROWN},{h('pos_actuals')})))", f=F_CALC, nf=FMT_INT, al=AL_RIGHT)
        sj, sk = LY.wt_ref(LY.WT_R_STARTJAAR, k), LY.wt_ref(LY.WT_R_STARTKW, k)
        put(ws, f"{bv}6", f"=IF(AND({aantal}>0,N({sj})>0),N({sj})*4+N({sk}),99999)", f=F_CALC, nf="0", al=AL_RIGHT)

    # ---- periode-rijen ------------------------------------------------------
    up_shift, down_shift, scn_shift = dabs("shift_up"), dabs("shift_down"), dabs("shift_scn")
    for r in range(ROW1, ROWN + 1):
        p = r - 1
        i = r - ROW1 + 1
        leeg = f'${NR}{r}=""'
        jaar = LY.in_col_ref("B", r)
        kw = LY.in_col_ref("C", r)
        f = {}
        # bronrij: de i-de gevulde rij op Invoer (matrixformule); daarna is alles positioneel
        # koprij 'Jaar' is altijd de eerste gevulde cel van het bereik, dus periode i = de (i+1)-de gevulde cel;
        # geen verwijzing naar één cel (die zou #REF! worden als de gebruiker die rij verwijdert)
        f[BRON] = ArrayFormula(f"{BRON}{r}", f'=IFERROR(SMALL(IF({inv_b}<>"",ROW({inv_b})-MIN(ROW({inv_b}))+1),{i + 1}),"")')
        f[NR] = f'=IF(${BRON}{r}="","",{i})'
        kwv = f'VALUE({kw}&"")'
        f[KW] = (f'=IF({leeg},"",IF(LEFT({kw}&"",1)="Q",IFERROR(VALUE(MID({kw},2,1)),0),'
                 f'IFERROR(IF(AND({kwv}>=1,{kwv}<=4),{kwv},0),0)))')
        f[KWT] = f'=IF({leeg},"",IF({KW}{r}=0,{jaar}&"",IF({KW}{r}=1,"Q1 {APOS}"&RIGHT({jaar},2),"Q"&{KW}{r})))'
        f[KWV] = f'=IF({leeg},"",IF({KW}{r}=0,{jaar}&"","Q"&{KW}{r}&" {APOS}"&RIGHT({jaar},2)))'
        f[IDX] = f'=IF({leeg},"",IFERROR(VALUE(TRIM({jaar}&"")),0)*4+IF({KW}{r}=0,4,{KW}{r}))'
        f[FASE] = f'=IF({leeg},"",IF({IDX}{r}<={h("idx_actuals")},"Realisatie","Prognose"))'
        f[KOS] = f'=IF({leeg},"",IFERROR(N({LY.in_col_ref("D", r)}),0))'
        f[OPB] = f'=IF({leeg},"",IFERROR(N({LY.in_col_ref("F", r)}),0))'
        f[STAND] = f'=IF({leeg},"",N({STAND}{p})+{OPB}{r}-{KOS}{r})'
        f[CUM] = f'=IF({leeg},"",N({CUM}{p})+{OPB}{r})'
        vorig = LY.in_col_ref("J", r)
        f[VORIG] = f'=IF(OR({leeg},{h("vorig_aanwezig")}=0),NA(),IF(ISNUMBER({vorig}),{vorig}*1000,NA()))'
        # typeblokken
        for k in range(1, N_TYPES + 1):
            vc, tc, tu, td, bv, ts, gv, gt, vx, fe = (m_col(b, k) for b in LY.BLOKKEN)
            # fees (DAEB) die t/m deze periode vervallen: cumulatief, hele type, uit de bouwtermijnentabel (geldige banen)
            f[fe] = (f'=IF({leeg},"",SUMIFS({rng_bt("bedrag")},{rng_bt("k")},{k},{rng_bt("gebruikt")},1,{rng_bt("idxb")},"<="&{IDX}{r},'
                     f'{rng_bt("idxb")},">="&${IDX}${ROW1}))')
            f[vc] = f'=IF({leeg},"",N({vc}{p})+IFERROR(N({LY.in_ref(r, k, 0)}),0))'
            f[tc] = f'=IF({leeg},"",N({tc}{p})+IFERROR(N({LY.in_ref(r, k, 1)}),0))'
            f[gv] = f'=IF(OR({leeg},{vc}$6=0),NA(),{vc}{r}-N({vc}{p}))'
            f[gt] = f'=IF(OR({leeg},{vc}$6=0),NA(),{tc}{r}-N({tc}{p}))'
            for col, shift in ((tu, up_shift), (td, down_shift), (ts, scn_shift)):
                f[col] = (f'=IF({leeg},"",IF(OR({FASE}{r}="Realisatie",${NR}{r}={h("n")}),{tc}{r},'
                          f'MAX({td}$6,N(INDEX({tc}${ROW1}:{tc}${ROWN},MAX(1,MIN({h("n")},${NR}{r}-N({shift}))))))))')
            kw_rng = LY.wt_ref(0, k, 2, rows=(LY.WT_R_T1, LY.WT_R_TN))
            pct_rng = LY.wt_ref(0, k, 1, rows=(LY.WT_R_T1, LY.WT_R_TN))
            # bouwtermijnen geschaald naar 100% − grondtermijn (rij 5 van blok tup), zodat termijnen als % van de aanneemsom mogen
            f[bv] = f'=IF({leeg},"",IF({IDX}{r}<{bv}$6,0,SUMIF({kw_rng},"<="&({IDX}{r}-{bv}$6+1),{pct_rng})*{tu}$5))'
            # extra's (€ per woning) met een bouwkwartaal >= 1: vervallen zoals termijnen; bij transport (rij 6) telt het model apart
            x_kw = LY.wt_ref(0, k, 2, rows=(LY.WT_R_X1, LY.WT_R_XN))
            x_eur = LY.wt_ref(0, k, 1, rows=(LY.WT_R_X1, LY.WT_R_XN))
            f[vx] = f'=IF({leeg},"",IF({IDX}{r}<{bv}$6,0,SUMIFS({x_eur},{x_kw},">=1",{x_kw},"<="&({IDX}{r}-{bv}$6+1))))'
        tc1, tcN = m_col("tcum", 1), m_col("tcum", N_TYPES)
        koopsom = m_range_abs("tcum", 6)
        grond = m_range_abs("tup", 6)
        verv = f"{m_col('verv', 1)}{r}:{m_col('verv', N_TYPES)}{r}"
        x_transport = m_range_abs("vx", 6)
        x_verv = f"{m_col('vx', 1)}{r}:{m_col('vx', N_TYPES)}{r}"
        # opbrengst per getransporteerde woning = koopsom × (grond + vervallen termijnen) [× koopsomknop per scenario]
        #                                        + extra's bij transport + extra's vervallen per bouwkwartaal
        # DAEB-typen (vlag in rij 5 van blok fee) tellen geen koopsom of extra's, maar de fees die t/m de periode vervallen (in alle
        # vier de modelkolommen gelijk: de knoppen raken de fees niet)
        nodaeb = f"(1-{m_range_abs('fee', 5)})"
        fee_r = f"{m_col('fee', 1)}{r}:{m_col('fee', N_TYPES)}{r}"
        for naam, blok, ks in (("model_basis", "tcum", None), ("model_up", "tup", "koopsom_up"), ("model_down", "tdown", "koopsom_down"),
                               ("model_scn", "tscn", "koopsom_scn")):
            rng = f"{m_col(blok, 1)}{r}:{m_col(blok, N_TYPES)}{r}"
            factor = f"*(1+{h(ks)})" if ks else ""
            f[M[naam]] = f'=IF({leeg},"",SUMPRODUCT({rng}*({koopsom}*({grond}+{verv}){factor}+{x_transport}+{x_verv})*{nodaeb})+SUM({fee_r}))'
        cum_rng = f"${CUM}${ROW1}:${CUM}${ROWN}"
        for naam, model, shift in (("cum_opbr_up", "model_up", up_shift), ("cum_opbr_down", "model_down", down_shift),
                                   ("cum_opbr_scn", "model_scn", scn_shift)):
            f[M[naam]] = (f'=IF({leeg},"",IF({h("model_aan")}=1,{CUM}{r}+{M[model]}{r}-{M["model_basis"]}{r},'
                          f'IF(OR({FASE}{r}="Realisatie",${NR}{r}={h("n")}),{CUM}{r},MAX({h("cum_opbr_actuals")},'
                          f'N(INDEX({cum_rng},MAX(1,MIN({h("n")},${NR}{r}-N({shift})))))))))')
        f[M["opbr_up"]] = f'=IF({leeg},"",IF({FASE}{r}="Realisatie",{OPB}{r},({M["cum_opbr_up"]}{r}-N({M["cum_opbr_up"]}{p}))*(1+N({dabs("opbr_up")}))))'
        f[M["opbr_down"]] = f'=IF({leeg},"",IF({FASE}{r}="Realisatie",{OPB}{r},({M["cum_opbr_down"]}{r}-N({M["cum_opbr_down"]}{p}))*(1+N({dabs("opbr_down")}))))'
        f[M["kosten_up"]] = f'=IF({leeg},"",IF({FASE}{r}="Prognose",{KOS}{r}*(1+N({dabs("kosten_up")})),{KOS}{r}))'
        f[M["kosten_down"]] = f'=IF({leeg},"",IF({FASE}{r}="Prognose",{KOS}{r}*(1+N({dabs("kosten_down")})),{KOS}{r}))'
        f[M["opbr_scn"]] = f'=IF({leeg},"",IF({FASE}{r}="Realisatie",{OPB}{r},({M["cum_opbr_scn"]}{r}-N({M["cum_opbr_scn"]}{p}))*(1+N({dabs("opbr_scn")}))))'
        f[M["kosten_scn"]] = f'=IF({leeg},"",IF({FASE}{r}="Prognose",{KOS}{r}*(1+N({dabs("kosten_scn")})),{KOS}{r}))'

        # rente: over de negatieve stand van het vorige kwartaal, alleen in prognosekwartalen, t/m de rente-eindindex
        def rente(stand_col, eind, pct):
            return (f'=IF({leeg},"",IF(AND({FASE}{r}="Prognose",{pct}>0,{IDX}{r}<{eind}),'
                    f'MAX(0,-N({stand_col}{p}))*{pct}/4,0))')
        f[M["rente_basis"]] = rente(M["stand_basis_rente"], h("rente_eind_basis"), h("rente"))
        f[M["stand_basis_rente"]] = f'=IF({leeg},"",N({M["stand_basis_rente"]}{p})+{OPB}{r}-{KOS}{r}-{M["rente_basis"]}{r})'
        f[M["rente_up"]] = rente(M["stand_up_raw"], h("rente_eind_up"), h("rente_pct_up"))
        f[M["stand_up_raw"]] = f'=IF({leeg},"",N({M["stand_up_raw"]}{p})+{M["opbr_up"]}{r}-{M["kosten_up"]}{r}-{M["rente_up"]}{r})'
        f[M["rente_down"]] = rente(M["stand_down_raw"], h("rente_eind_down"), h("rente_pct_down"))
        f[M["stand_down_raw"]] = f'=IF({leeg},"",N({M["stand_down_raw"]}{p})+{M["opbr_down"]}{r}-{M["kosten_down"]}{r}-{M["rente_down"]}{r})'
        f[M["rente_scn"]] = rente(M["stand_scn_raw"], h("rente_eind_scn"), h("rente_pct_scn"))
        f[M["stand_scn_raw"]] = f'=IF({leeg},"",N({M["stand_scn_raw"]}{p})+{M["opbr_scn"]}{r}-{M["kosten_scn"]}{r}-{M["rente_scn"]}{r})'
        # getoond: met rente in de basis (ja) of alleen het verschil in rente in de scenario's (nee)
        rb = h("rente_in_basis")
        f[M["stand_basis"]] = f'=IF({leeg},"",IF({rb}=1,{M["stand_basis_rente"]}{r},{STAND}{r}))'
        f[M["kosten_basis"]] = f'=IF({leeg},"",{KOS}{r}+IF({rb}=1,{M["rente_basis"]}{r},0))'
        f[M["stand_up"]] = f'=IF({leeg},"",IF({rb}=1,{M["stand_up_raw"]}{r},{M["stand_up_raw"]}{r}+({STAND}{r}-{M["stand_basis_rente"]}{r})))'
        f[M["stand_down"]] = f'=IF({leeg},"",IF({rb}=1,{M["stand_down_raw"]}{r},{M["stand_down_raw"]}{r}+({STAND}{r}-{M["stand_basis_rente"]}{r})))'
        f[M["stand_scn"]] = f'=IF({leeg},"",IF({rb}=1,{M["stand_scn_raw"]}{r},{M["stand_scn_raw"]}{r}+({STAND}{r}-{M["stand_basis_rente"]}{r})))'
        for naam, bron in (("pos_basis", M["stand_basis"]), ("pos_up", M["stand_up"]), ("pos_down", M["stand_down"]), ("pos_vorig", VORIG),
                           ("pos_scn", M["stand_scn"])):
            f[M[naam]] = f'=IF(ISNUMBER({bron}{r}),IF({bron}{r}>0,1,0),0)'
        sb, su, sd, kb = M["stand_basis"], M["stand_up"], M["stand_down"], M["kosten_basis"]
        na = f'=IF({leeg},NA(),'
        f[M["g_voorfinanciering"]] = f'{na}MIN({sb}{r}/1000000,0))'
        f[M["g_positief_saldo"]] = f'{na}MAX({sb}{r}/1000000,0))'
        f[M["g_opbrengsten"]] = f'{na}{OPB}{r}/1000000)'
        f[M["g_kosten"]] = f'{na}-{kb}{r}/1000000)'
        f[M["g_stand"]] = f'{na}{sb}{r}/1000000)'
        f[M["g_vorige"]] = f'{na}{VORIG}{r}/1000000)'
        f[M["g_band_onder"]] = f'{na}MIN({su}{r},{sd}{r})/1000000)'
        f[M["g_bandbreedte"]] = f'{na}ABS({su}{r}-{sd}{r})/1000000)'
        f[M["g_downside"]] = f'{na}{sd}{r}/1000000)'
        f[M["g_upside"]] = f'{na}{su}{r}/1000000)'
        f[M["g_punt_nu"]] = f'{na}IF(${NR}{r}={h("pos_actuals")},{sb}{r}/1000000,NA()))'
        f[M["g_punt_dal"]] = f'{na}IF({sb}{r}={h("dal_basis")},{sb}{r}/1000000,NA()))'
        f[M["g_eind_upside"]] = f'{na}IF(${NR}{r}={h("n")},{su}{r}/1000000,NA()))'
        f[M["g_eind_downside"]] = f'{na}IF(${NR}{r}={h("n")},{sd}{r}/1000000,NA()))'
        f[M["g_eind_basis"]] = f'{na}IF(${NR}{r}={h("n")},{sb}{r}/1000000,NA()))'
        # scenariolijn: alleen als de knop 'tonen' op ja staat (anders overal #N/B: de lijn verdwijnt)
        ss = M["stand_scn"]
        # de lijn begint bij het punt 'nu' (laatste gerealiseerde kwartaal): daarvóór valt hij samen met de basis
        f[M["g_scenario"]] = f'{na}IF(AND({h("scn_aan")}=1,OR({FASE}{r}="Prognose",${NR}{r}={h("pos_actuals")})),{ss}{r}/1000000,NA()))'
        f[M["g_eind_scenario"]] = f'{na}IF(AND({h("scn_aan")}=1,${NR}{r}={h("n")}),{ss}{r}/1000000,NA()))'
        f[M["g_realisatie"]] = f'{na}IF({FASE}{r}="Realisatie",1,NA()))'
        vcum, tcum = M["g_verkocht_cum"], M["g_getransporteerd_cum"]
        f[M["g_verkocht"]] = f'{na}{vcum}{r}-N({vcum}{p}))'
        f[M["g_getransporteerd"]] = f'{na}{tcum}{r}-N({tcum}{p}))'
        f[vcum] = f'{na}SUM({m_col("vcum", 1)}{r}:{m_col("vcum", N_TYPES)}{r}))'
        f[tcum] = f'{na}SUM({tc1}{r}:{tcN}{r}))'
        f[M["g_verkocht_pct"]] = f'{na}IF({h("woningen")}=0,0,{vcum}{r}/{h("woningen")}))'
        f[M["g_getransporteerd_pct"]] = f'{na}IF({h("woningen")}=0,0,{tcum}{r}/{h("woningen")}))'
        f[M["g_opbrengsten_pct"]] = f'{na}IF({h("opbr_totaal")}=0,0,{CUM}{r}/{h("opbr_totaal")}))'
        f[M["g_kosten_pct"]] = f'{na}IF({h("kosten_totaal")}=0,0,({CUM}{r}-{sb}{r})/{h("kosten_totaal")}))'
        f[M["g_punt_uitverkocht"]] = f'{na}IF(${NR}{r}={h("uitverkocht_pos")},{M["g_verkocht"]}{r},NA()))'
        f[M["g_punt_alles_transport"]] = f'{na}IF(${NR}{r}={h("alles_transport_pos")},{M["g_getransporteerd"]}{r},NA()))'
        if r == ROW1:
            f[M["stap_ok"]] = f'=IF({leeg},1,1)'
        else:
            f[M["stap_ok"]] = f'=IF(OR(${NR}{r}="",${NR}{p}=""),1,IF({IDX}{r}-{IDX}{p}=1+3*({KW}{r}=0),1,0))'
        for col, formule in f.items():
            put(ws, f"{col}{r}", formule, f=F_CALC, nf=FMT_INT)
        for col in (KWT, KWV, FASE):
            ws[f"{col}{r}"].alignment = AL_LEFT_TOP
        for col in range(LY.M_BLOK1 + 4 * N_TYPES, LY.M_BLOK1 + 5 * N_TYPES):
            ws[f"{L(col)}{r}"].number_format = FMT_PCT
        for naam in ("g_voorfinanciering", "g_positief_saldo", "g_opbrengsten", "g_kosten", "g_stand", "g_vorige", "g_band_onder",
                     "g_bandbreedte", "g_downside", "g_upside", "g_punt_nu", "g_punt_dal", "g_eind_upside", "g_eind_downside", "g_eind_basis",
                     "g_scenario", "g_eind_scenario"):
            ws[f"{M[naam]}{r}"].number_format = "0.00"
        for naam in ("g_verkocht_pct", "g_getransporteerd_pct", "g_opbrengsten_pct", "g_kosten_pct"):
            ws[f"{M[naam]}{r}"].number_format = FMT_PCT

    # ---- bouwtermijnentabel (tijdlijn): raster van alle type × termijn-combinaties, daarna de compacte lijst voor de grafiek:
    #      twee koprijen (jaartallen, kwartaalnummers), per bouwtermijn een baan per woningtype dat hem heeft, een lege rij
    #      tussen de termijnen. Bij een DAEB-type (blok 5) zijn de tien banen de fee-termijnen, gegroepeerd per component.
    #      Tijd = kwartaalindex (jaar × 4 + kwartaal): één eenheid per kwartaal. ------------------------------------------------
    put(ws, f"{BT['j']}4", "bouwtermijnentabel (grafiek): raster van alle type × termijn-combinaties (bij DAEB: fee-termijnen) met groepsnaam, "
                           "€ (fees), kwartaalindex (uit bouwkwartaal, 'jaar Qk' of 'actuals'), rang en positie; vanaf 'nr' de compacte lijst: "
                           "koprij jaartallen, koprij kwartaalnummers, per groep een baan per type (lichtrood = vorige prognose, typekleur = "
                           "huidige planning, grijs = gerealiseerd) en een lege rij tussen de groepen", f=F_NOTE8)
    bt_koppen = {"j": "j", "k": "type", "i": "termijn", "gebruikt": "Gebruikt", "knaam": "Typenaam", "tnaam": "Groep (termijn / component)",
                 "mnaam": "Baan (termijn / mijlpaal)", "comp": "Fee-component", "bedrag": "Bedrag fee (€)", "waarde": "Invoer % of €",
                 "kw": "Kwartaalcel", "kwr": "Kwartaal vorig (blok 6)", "fout": "Fout (fee)", "idxb": "Idx huidig", "idxr": "Idx vorig", "uniek": "Uniek", "rang": "Rang groep",
                 "pos": "Positie", "eerste": "Eerste baan", "nr": "Nr (uniek)", "bron": "Bronrij", "soort": "Soort", "label": "Termijn",
                 "kb": "Type", "ib": "Idx huidig", "ir": "Idx vorig"}
    bt_koppen.update(LY.BT_SEG_NAMEN)
    for naam, kop in bt_koppen.items():
        put(ws, f"{BT[naam]}7", kop, f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    # reeksnamen van de typereeksen, jaarreeksen en kwartaalreeksen (de Dashboard-grafiek en het PowerPoint-blok lezen deze cellen;
    # '(uit)' = de macro haalt de reeks uit de grafiek in de presentatie)
    k_rng, gebruikt_rng = rng_bt("k"), rng_bt("gebruikt")
    for k in range(1, N_TYPES + 1):
        naam_k = LY.wt_ref(LY.WT_R_NAAM, k)
        knaam = f'IF({naam_k}="","Type {k}",{naam_k}&"")'
        aan = f"COUNTIFS({k_rng},{k},{gebruikt_rng},1)>0"
        put(ws, f"{BT[f'cur{k}']}7", f'=IF({aan},{knaam},"(uit)")', f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
        put(ws, f"{BT[f'prev{k}']}7", f'=IF({aan},{knaam}&" · vorige prognose","(uit)")', f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    for y in range(1, LY.BT_NJ + 1):
        put(ws, f"{BT[f'jr{y}']}7", f'=IF({y}<={h("bt_jaren")},TEXT({h("jaar_first")}+{y - 1},"0"),"(uit)")', f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    for q_ in range(1, LY.BT_NK + 1):
        put(ws, f"{BT[f'kw{q_}']}7", str((q_ - 1) % 4 + 1), f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    ws.column_dimensions[L(LY.BT_COL1 - 1)].width = 3
    ws.column_dimensions[BT["label"]].width = 30
    for naam in BT:
        if naam != "label":
            ws.column_dimensions[BT[naam]].width = 11 if not naam.startswith(("jr", "kw")) else 5
    uniek_rng, tnaam_rng, mnaam_rng, idxb_rng, idxr_rng = rng_bt("uniek"), rng_bt("tnaam"), rng_bt("mnaam"), rng_bt("idxb"), rng_bt("idxr")
    j_rng, i_rng, rang_rng, pos_rng, nr_rng, eerste_rng, comp_rng = (rng_bt(x) for x in ("j", "i", "rang", "pos", "nr", "eerste", "comp"))
    tf, te, tn = h("t_first"), h("t_end"), h("t_nu_end")
    idx_first = f"${IDX}${ROW1}"
    for r in range(LY.BT_ROW1, LY.BT_ROWN + 1):
        j = r - LY.BT_ROW1 + 1
        f = {}
        f["j"] = j
        # ---- raster (alleen de eerste N_TYPES × N_TERMIJNEN rijen) ----
        if j <= LY.BT_LANES:
            k = (j - 1) // LY.N_TERMIJNEN + 1
            i = (j - 1) % LY.N_TERMIJNEN + 1
            c = (i - 1) // LY.N_FEE_PER_COMP + 1                      # fee-component van deze baan bij DAEB
            kw_bt = LY.wt_ref(LY.WT_R_T1 + i - 1, k, 2)
            pct_bt = LY.wt_ref(LY.WT_R_T1 + i - 1, k, 1)
            tn_bt = LY.wt_ref(LY.WT_R_T1 + i - 1, k, 0)
            kw_f = LY.wt_ref(LY.WT_R_FT1 + i - 1, k, 2)
            kw_vr = LY.wt_ref(LY.WT_R_VF1 + i - 1, k, 2)                  # kwartaal volgens de vorige prognose (blok 6)
            w_f = LY.wt_ref(LY.WT_R_FT1 + i - 1, k, 1)
            mn_f = LY.wt_ref(LY.WT_R_FT1 + i - 1, k, 0)
            cn_f = LY.wt_ref(LY.WT_R_FC1 + c - 1, k, 0)
            tot_f = LY.wt_ref(LY.WT_R_FC1 + c - 1, k, 1)
            naam_k = LY.wt_ref(LY.WT_R_NAAM, k)
            vj_k, vk_k = LY.wt_ref(LY.WT_R_VSTARTJAAR, k), LY.wt_ref(LY.WT_R_VSTARTKW, k)
            daeb = f"{m_col('fee', k)}$5"
            start = f"{m_col('verv', k)}$6"
            aantal = f"N({LY.wt_ref(LY.WT_R_AANTAL, k)})"
            gebruikt, waarde, bedrag, kwv, kwr, idxb, idxr_ref, rang = (f"${BT[x]}{r}" for x in ("gebruikt", "waarde", "bedrag", "kw", "kwr", "idxb", "idxr", "rang"))
            f["k"], f["i"] = k, i
            f["knaam"] = f'=IF({naam_k}="","Type {k}",{naam_k}&"")'
            f["comp"] = f'=IF({daeb}=1,{c},0)'
            f["mnaam"] = f'=IF({daeb}=1,IF({mn_f}="","termijn {i}",{mn_f}&""),IF({tn_bt}="","termijn {i}",{tn_bt}&""))'
            f["tnaam"] = f'=IF({daeb}=1,IF({cn_f}="","Component {c}",{cn_f}&""),${BT["mnaam"]}{r})'
            f["waarde"] = f'=IF({daeb}=1,IF(ISNUMBER({w_f}),{w_f},0),IF(ISNUMBER({pct_bt}),{pct_bt},0))'
            f["bedrag"] = f'=IF({daeb}=1,IF(AND({waarde}>0,{waarde}<=1),{waarde}*N({tot_f}),{waarde}),0)'
            f["kw"] = f'=IF({daeb}=1,IF(ISBLANK({kw_f}),"",{kw_f}),IF(ISBLANK({kw_bt}),"",{kw_bt}))'
            f["kwr"] = f'=IF({daeb}=1,IF(ISBLANK({kw_vr}),"",{kw_vr}),"")'
            # kwartaalindex: bouwkwartaal (getal) t.o.v. start bouw; bij fees ook 'jaar Qk' / 'Qk jaar' (tekst) of 'actuals'
            t = f'UPPER(TRIM({kwv}&""))'
            jaar = f'IFERROR(VALUE(LEFT({t},4)),IFERROR(VALUE(RIGHT({t},4)),0))'
            kwt = f'IFERROR(VALUE(MID({t},FIND("Q",{t})+1,1)),0)'
            tekst = f'IF(AND({jaar}>=1990,{jaar}<=2100,{kwt}>=1,{kwt}<=4),{jaar}*4+{kwt},99999)'
            actuals = f'IF({h("pos_actuals")}=0,N({idx_first}),{h("idx_actuals")})'
            rel = f'IF(OR({start}>=99999,{kwv}<-12,{kwv}>60),99999,{start}+{kwv}-1)'
            rel_t = f'IF(OR({start}>=99999,VALUE({t})<-12,VALUE({t})>60),99999,{start}+VALUE({t})-1)'
            f["idxb"] = (f'=IF({daeb}=1,IF(ISNUMBER({kwv}),{rel},IF({t}="ACTUALS",{actuals},IF({kwv}="",99999,'
                         f'IF(ISNUMBER(IFERROR(VALUE({t}),"")),{rel_t},{tekst})))),IF(AND(ISNUMBER({kwv}),N({kwv})>=1),{rel},99999))')
            f["gebruikt"] = f'=IF(AND({aantal}>0,{idxb}<99999,IF({daeb}=1,{bedrag}<>0,{waarde}<>0)),1,0)'
            f["fout"] = (f'=IF({daeb}=1,IF(OR({waarde}<0,AND({gebruikt}=0,OR({waarde}<>0,{bedrag}<>0,{kwv}<>"")),'
                        f'AND({gebruikt}=1,{kwr}<>"",{idxr_ref}>=99999)),1,0),IF(AND(ISNUMBER({w_f}),{w_f}<>0),1,0))')
            # kwartaalindex vorige prognose: bij fees het kwartaal uit blok 6 (zelfde schrijfwijze, bouwkwartaal t.o.v. start bouw van
            # blok 4, anders t.o.v. de huidige start); zonder blok 6: een bouwkwartaal schuift mee met blok 4, de rest staat waar het nu staat
            startv = f'IF(AND(N({vj_k})>0,N({vk_k})>=1),N({vj_k})*4+N({vk_k}),{start})'
            tr = f'UPPER(TRIM({kwr}&""))'
            jaar_r = f'IFERROR(VALUE(LEFT({tr},4)),IFERROR(VALUE(RIGHT({tr},4)),0))'
            kwt_r = f'IFERROR(VALUE(MID({tr},FIND("Q",{tr})+1,1)),0)'
            tekst_r = f'IF(AND({jaar_r}>=1990,{jaar_r}<=2100,{kwt_r}>=1,{kwt_r}<=4),{jaar_r}*4+{kwt_r},99999)'
            rel_r = f'IF(OR({startv}>=99999,{kwr}<-12,{kwr}>60),99999,{startv}+{kwr}-1)'
            rel_tr = f'IF(OR({startv}>=99999,VALUE({tr})<-12,VALUE({tr})>60),99999,{startv}+VALUE({tr})-1)'
            idx_vorig = f'IF(ISNUMBER({kwr}),{rel_r},IF({tr}="ACTUALS",{actuals},IF(ISNUMBER(IFERROR(VALUE({tr}),"")),{rel_tr},{tekst_r})))'
            f["idxr"] = (f'=IF({gebruikt}=0,99999,IF(AND({daeb}=1,{kwr}<>""),{idx_vorig},'
                         f'IF(AND(ISNUMBER({kwv}),N({vj_k})>0,N({vk_k})>=1),N({vj_k})*4+N({vk_k})+{kwv}-1,{idxb})))')
            if r == LY.BT_ROW1:
                f["uniek"] = f'={gebruikt}'
            else:
                eerder_t = f"${BT['tnaam']}${LY.BT_ROW1}:${BT['tnaam']}${r - 1}"
                eerder_g = f"${BT['gebruikt']}${LY.BT_ROW1}:${BT['gebruikt']}${r - 1}"
                f["uniek"] = f'=IF(AND({gebruikt}=1,SUMPRODUCT(--({eerder_t}=${BT["tnaam"]}{r}),--({eerder_g}=1))=0),1,0)'
            # rang van de groepsnaam = positie (in 'nr') van de eerste gebruikte rij met deze naam
            zelfde = f"({tnaam_rng}=${BT['tnaam']}{r})*({gebruikt_rng}=1)"
            f["rang"] = ArrayFormula(f"{BT['rang']}{r}", f'=IF({gebruikt}=1,MATCH(MIN(IF({zelfde},{j_rng})),{nr_rng},0),"")')
            # positie in de compacte lijst: banen van lagere rangen + scheidingsrijen + eerdere banen in deze groep (typevolgorde, dan
            # termijnvolgorde binnen het type)
            eerder_in_groep = f'({gebruikt_rng}=1)*({rang_rng}={rang})*(({k_rng}<{k})+({k_rng}={k})*({i_rng}<{i}))'
            f["pos"] = f'=IF({gebruikt}=1,SUMPRODUCT(({gebruikt_rng}=1)*({rang_rng}<{rang}))+{rang}-1+SUMPRODUCT({eerder_in_groep})+1,"")'
            f["eerste"] = f'=IF({gebruikt}=1,IF(SUMPRODUCT({eerder_in_groep})=0,1,0),"")'
        else:
            for naam_ in ("k", "i", "gebruikt", "knaam", "tnaam", "mnaam", "comp", "bedrag", "waarde", "kw", "kwr", "fout", "idxb", "idxr",
                          "uniek", "rang", "pos", "eerste"):
                f[naam_] = None
        # ---- compacte lijst ----
        f["nr"] = ArrayFormula(f"{BT['nr']}{r}", f'=IFERROR(SMALL(IF({uniek_rng}=1,{j_rng}),{j}),"")')
        bron, soort = f"${BT['bron']}{r}", f"${BT['soort']}{r}"
        if j <= LY.BT_KOP:
            f["bron"] = None
            f["soort"] = j
            f["label"] = " "
            f["kb"] = f["ib"] = f["ir"] = None
        else:
            f["bron"] = f'=IFERROR(MATCH({j - LY.BT_KOP},{pos_rng},0),"")'
            f["soort"] = f'=IF({bron}<>"",3,IF({j - LY.BT_KOP}<={h("bt_lanes")}+{h("bt_uniek")}-1,4,0))'
            f["label"] = (f'=IF({soort}=3,IF(INDEX({comp_rng},{bron})>0,INDEX({tnaam_rng},{bron})&" · "&INDEX({mnaam_rng},{bron}),'
                          f'IF(INDEX({eerste_rng},{bron})=1,INDEX({tnaam_rng},{bron})," ")),IF({soort}=4," ",""))')
            f["kb"] = f'=IF({soort}=3,INDEX({k_rng},{bron}),"")'
            f["ib"] = f'=IF({soort}=3,INDEX({idxb_rng},{bron}),"")'
            f["ir"] = f'=IF({soort}=3,INDEX({idxr_rng},{bron}),"")'
        kb, ib, ir = f"${BT['kb']}{r}", f"${BT['ib']}{r}", f"${BT['ir']}{r}"
        baan = f"{soort}=3"
        # blokjes van één kwartaal, afgekapt op de as [t_first, t_end]; buiten de as: breedte 0
        ibc, irc = f"MIN(MAX({ib},{tf}),{te})", f"MIN(MAX({ir},{tf}),{te})"
        blauw, rood = f"MAX(0,MIN({ib}+1,{te})-{ibc})", f"MAX(0,MIN({ir}+1,{te})-{irc})"
        s1, s2, s4 = (f"${BT[x]}{r}" for x in ("s1", "s2", "s4"))
        # stapel 1 (eerste as): onzichtbaar tot de as-ondergrens, grijs, onzichtbaar, vorige prognose (per type), grijs, onzichtbaar; jaarblokjes in koprij 1
        f["s0"] = f'=IF(OR({baan},{soort}=1),{tf},NA())'
        f["s1"] = f'=IF({baan},MAX(0,MIN({irc},{tn})-{tf}),NA())'
        f["s2"] = f'=IF({baan},MAX(0,{irc}-{tn}),NA())'
        for k in range(1, N_TYPES + 1):
            f[f"prev{k}"] = f'=IF(AND({baan},{kb}={k}),{rood},NA())'
        f["s4"] = f'=IF({baan},MAX(0,{tn}-({irc}+{rood})),NA())'
        f["s5"] = f'=IF({baan},MAX(0,{te}-({tf}+{s1}+{s2}+{rood}+{s4})),NA())'
        # stapel 2 (tweede as, bovenop): onzichtbaar, huidige planning (per type), onzichtbaar; kwartaalblokjes in koprij 2
        f["s6"] = f'=IF({baan},{ibc},IF({soort}=2,{tf},NA()))'
        for k in range(1, N_TYPES + 1):
            f[f"cur{k}"] = f'=IF(AND({baan},{kb}={k}),{blauw},NA())'
        f["s8"] = f'=IF({baan},MAX(0,{te}-{ibc}-{blauw}),NA())'
        for y in range(1, LY.BT_NJ + 1):
            f[f"jr{y}"] = f'=IF({y}<={h("bt_jaren")},4,NA())' if j == 1 else "#N/A"
        for q_ in range(1, LY.BT_NK + 1):
            f[f"kw{q_}"] = f'=IF({q_}<=4*{h("bt_jaren")},1,NA())' if j == 2 else "#N/A"
        for naam_, formule in f.items():
            if formule is None:
                continue
            nf = "0.00" if naam_ in ("s1", "s2", "s4", "s5", "s8") or naam_.startswith(("prev", "cur")) else ("0" if naam_ != "bedrag" else FMT_INT)
            put(ws, f"{BT[naam_]}{r}", formule, f=F_CALC, nf=nf, al=AL_LEFT_TOP if naam_ in ("label", "knaam", "tnaam", "mnaam", "kw") else AL_RIGHT)
            if formule == "#N/A":
                ws[f"{BT[naam_]}{r}"].data_type = "e"

    # ---- grafiek per woningtype: selectieblok (type uit de keuzecel), € mln per kwartaal ------------------------------
    PT = LY.PT
    put(ws, f"{PT['grond']}4", "grafiek per woningtype: het gekozen type (Dashboard-keuzecel, hulpcel pt_keuze) · € mln per kwartaal: grondtermijn, "
                               "bouwtermijnen, extra's, fees per component, cumulatief", f=F_NOTE8)
    for naam, kop in (("grond", "Grondtermijn"), ("bouw", "Bouwtermijnen"), ("extra", "Extra's"), ("fee1", "Fee 1"), ("fee2", "Fee 2"), ("cum", "Cumulatief"),
                      ("totaal", "Totaal kwartaal"), ("eind", "Eindpunt")):
        put(ws, f"{PT[naam]}7", kop, f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
        ws.column_dimensions[PT[naam]].width = 11
    ws.column_dimensions[L(LY.PT_COL1 - 1)].width = 3
    kz = h("pt_keuze")
    sel = lambda blok, row: f"INDEX({m_range_abs(blok, row)},{kz})"  # noqa: E731
    for r in range(ROW1, ROWN + 1):
        p = r - 1
        leeg = f'${NR}{r}=""'
        tc_r, tc_p = f"N({sel('tcum', r)})", f"N({sel('tcum', p)})"
        koopsom = f"(N({sel('tcum', 6)})*(1-N({sel('fee', 5)})))"
        grond = f"N({sel('tup', 6)})"
        verv_r, verv_p = f"N({sel('verv', r)})", f"N({sel('verv', p)})"
        vx6, vx_r, vx_p = f"N({sel('vx', 6)})", f"N({sel('vx', r)})", f"N({sel('vx', p)})"
        nod = f"(1-N({sel('fee', 5)}))"
        fee = lambda c: (f'SUMIFS({rng_bt("bedrag")},{rng_bt("k")},{kz},{rng_bt("comp")},{c},{rng_bt("gebruikt")},1,'  # noqa: E731
                         f'{rng_bt("idxb")},${IDX}{r})')
        f = {
            "grond": f'=IF({leeg},NA(),({tc_r}-{tc_p})*{koopsom}*{grond}/1000000)',
            "bouw": f'=IF({leeg},NA(),{koopsom}*({tc_r}*{verv_r}-{tc_p}*{verv_p})/1000000)',
            "extra": f'=IF({leeg},NA(),(({tc_r}-{tc_p})*{vx6}+({tc_r}*{vx_r}-{tc_p}*{vx_p}))*{nod}/1000000)',
            "fee1": f'=IF({leeg},NA(),{fee(1)}/1000000)',
            "fee2": f'=IF({leeg},NA(),{fee(2)}/1000000)',
        }
        f["cum"] = (f'=IF({leeg},NA(),N({PT["cum"]}{p})+{PT["grond"]}{r}+{PT["bouw"]}{r}+{PT["extra"]}{r}+{PT["fee1"]}{r}+{PT["fee2"]}{r})')
        f["totaal"] = f'=IF({leeg},NA(),{PT["grond"]}{r}+{PT["bouw"]}{r}+{PT["extra"]}{r}+{PT["fee1"]}{r}+{PT["fee2"]}{r})'
        f["eind"] = f'=IF({leeg},NA(),IF({r - ROW1 + 1}={h("n")},{PT["cum"]}{r},NA()))'
        for naam, formule in f.items():
            put(ws, f"{PT[naam]}{r}", formule, f=F_CALC, nf="0.00", al=AL_RIGHT)

    # ---- hulpcellen -----------------------------------------------------------
    sb, su, sd, ss = M["stand_basis"], M["stand_up"], M["stand_down"], M["stand_scn"]
    pb, pu, pd, pv, ps = M["pos_basis"], M["pos_up"], M["pos_down"], M["pos_vorig"], M["pos_scn"]
    rng = lambda col: f"${col}${ROW1}:${col}${ROWN}"  # noqa: E731
    D_ = rng(KWV)
    bron_van = lambda pos: f"INDEX({rng(BRON)},{pos})"  # noqa: E731  bronrij (Invoer) van modelpositie pos
    hulp = {
        "n": ("Aantal periodes", f"=COUNT({rng(NR)})"),
        "idx_actuals": ("Index actuals t/m", f"={dabs('actuals_jaar')}*4+{dabs('actuals_kw')}"),
        "pos_actuals": ("Positie actuals", f'=COUNTIF({rng(IDX)},"<="&{h("idx_actuals")})'),
        "kw_actuals": ("Kwartaal actuals", f'=IF({h("pos_actuals")}=0,"–",INDEX({D_},{h("pos_actuals")}))'),
        "stand_nu": ("Stand nu", f'=IF({h("pos_actuals")}=0,0,INDEX({rng(sb)},{h("pos_actuals")}))'),
        "model_aan": ("Verkooptempo-model aan (1/0)", f'=IF(LOWER({dabs("model_aan")}&"")="ja",1,0)'),
        "vorig_aanwezig": ("Vorige prognose aanwezig (1/0)", f'=IF(SUMIF({inv_j},">0")-SUMIF({inv_j},"<0")>0,1,0)'),
        "opbr_totaal": ("Opbrengsten totaal", f"=SUM({rng(OPB)})"),
        "kosten_totaal": ("Kosten totaal (incl. rente als die in de basis zit)", f"=SUM({rng(M['kosten_basis'])})"),
        "cum_opbr_actuals": ("Cum. opbrengsten t/m actuals", f'=IF({h("pos_actuals")}=0,0,INDEX({rng(CUM)},{h("pos_actuals")}))'),
        "dal_basis": ("Dieptepunt basis", f"=MIN({rng(sb)})"),
        "dal_basis_kw": ("Dieptepunt basis kwartaal", f'=IFERROR(INDEX({D_},MATCH({h("dal_basis")},{rng(sb)},0)),"–")'),
        "eind_basis": ("Eindsaldo basis", f'=IF({h("n")}=0,0,INDEX({rng(sb)},{h("n")}))'),
        "be_basis": ("Break-even basis", f'=IFERROR(INDEX({D_},MATCH(1,{rng(pb)},0)),"")'),
        "be_basis_pos": ("Positie break-even basis", f'=IFERROR(MATCH(1,{rng(pb)},0),0)'),
        "dal_up": ("Dieptepunt upside", f"=MIN({rng(su)})"),
        "dal_up_kw": ("Dieptepunt upside kwartaal", f'=IFERROR(INDEX({D_},MATCH({h("dal_up")},{rng(su)},0)),"–")'),
        "eind_up": ("Eindsaldo upside", f'=IF({h("n")}=0,0,INDEX({rng(su)},{h("n")}))'),
        "be_up": ("Break-even upside", f'=IFERROR(INDEX({D_},MATCH(1,{rng(pu)},0)),"")'),
        "be_up_pos": ("Positie break-even upside", f'=IFERROR(MATCH(1,{rng(pu)},0),0)'),
        "dal_down": ("Dieptepunt downside", f"=MIN({rng(sd)})"),
        "dal_down_kw": ("Dieptepunt downside kwartaal", f'=IFERROR(INDEX({D_},MATCH({h("dal_down")},{rng(sd)},0)),"–")'),
        "eind_down": ("Eindsaldo downside", f'=IF({h("n")}=0,0,INDEX({rng(sd)},{h("n")}))'),
        "be_down": ("Break-even downside", f'=IFERROR(INDEX({D_},MATCH(1,{rng(pd)},0)),"")'),
        "be_down_pos": ("Positie break-even downside", f'=IFERROR(MATCH(1,{rng(pd)},0),0)'),
        "laatste_kw": ("Laatste kwartaal", f'=IF({h("n")}=0,"–",INDEX({D_},{h("n")}))'),
        "dal_vorig": ("Dieptepunt vorige prognose (€)", f'=IF({h("vorig_aanwezig")}=0,0,MIN({inv_j})*1000)'),
        "eind_vorig": ("Eindsaldo vorige prognose (€)", f'=IF(OR({h("vorig_aanwezig")}=0,{h("n")}=0),0,N(INDEX({inv_j},{bron_van(h("n"))}))*1000)'),
        "vorig_nu": ("Vorige prognose nu (€)", f'=IF(OR({h("vorig_aanwezig")}=0,{h("pos_actuals")}=0),0,N(INDEX({inv_j},{bron_van(h("pos_actuals"))}))*1000)'),
        "be_vorig": ("Break-even vorige prognose", f'=IFERROR(INDEX({D_},MATCH(1,{rng(pv)},0)),"")'),
        "be_vorig_pos": ("Positie break-even vorige prognose", f'=IFERROR(MATCH(1,{rng(pv)},0),0)'),
        "woningen": ("Woningen totaal", f"=SUM({m_range_abs('vcum', 6)})"),
        "verkocht_plan": ("Verkocht totaal (planning)", f'=IF({h("n")}=0,0,INDEX({rng(M["g_verkocht_cum"])},{h("n")}))'),
        "transport_plan": ("Getransporteerd totaal (planning)", f'=IF({h("n")}=0,0,INDEX({rng(M["g_getransporteerd_cum"])},{h("n")}))'),
        "verkocht_actuals": ("Verkocht t/m actuals", f'=IF({h("pos_actuals")}=0,0,INDEX({rng(M["g_verkocht_cum"])},{h("pos_actuals")}))'),
        "transport_actuals": ("Getransporteerd t/m actuals", f'=IF({h("pos_actuals")}=0,0,INDEX({rng(M["g_getransporteerd_cum"])},{h("pos_actuals")}))'),
        "model_opbr": ("Model-opbrengst totaal", f'=IF({h("n")}=0,0,INDEX({rng(M["model_basis"])},{h("n")}))'),
        "verk_voor_start": ("Verkocht vóór start bouw (aantal)", f"=SUM({m_range_abs('vcum', 5)})"),
        "won_met_start": ("Woningen in types met start bouw", f"=SUM({m_range_abs('tcum', 5)})"),
        "verk_voor_start_pct": ("Verkocht vóór start bouw (%)", f'=IFERROR({h("verk_voor_start")}/{h("won_met_start")},0)'),
        "uitverkocht": ("Uitverkocht in",
                        f'=IF(OR({h("woningen")}=0,(COUNTIF({rng(M["g_verkocht_pct"])},"<"&0.99999)+1)>{h("n")}),"niet binnen de looptijd",'
                        f'INDEX({D_},(COUNTIF({rng(M["g_verkocht_pct"])},"<"&0.99999)+1)))'),
        "alles_transport": ("Alles getransporteerd in",
                            f'=IF(OR({h("woningen")}=0,(COUNTIF({rng(M["g_getransporteerd_pct"])},"<"&0.99999)+1)>{h("n")}),"niet binnen de looptijd",'
                            f'INDEX({D_},(COUNTIF({rng(M["g_getransporteerd_pct"])},"<"&0.99999)+1)))'),
        "lbl_nu": ("Label punt nu", f'="nu · "&{h("kw_actuals")}'),
        "lbl_dal": ("Label dieptepunt", f'="Dieptepunt "&{h("dal_basis_kw")}&"  "&{eur_m(h("dal_basis"))}'),
        "lbl_eind_up": ("Label eind upside", f'="Upside "&{eur_m(h("eind_up"))}'),
        "lbl_eind_down": ("Label eind downside", f'="Downside "&{eur_m(h("eind_down"))}'),
        "lbl_eind_basis": ("Label eind basis", f'="Basis "&{eur_m(h("eind_basis"))}'),
        "rente": ("Jaarrente", f"=N({dabs('rente')})"),
        "rente_tm_start": ("Rente t/m start bouw (1/0)", f'=IF(LOWER({dabs("rente_tm")}&"")="start bouw",1,0)'),
        "rente_in_basis": ("Rente ook in de basis (1/0)", f'=IF(LOWER({dabs("rente_basis")}&"")="ja",1,0)'),
        "start_idx": ("Index start bouw project (vroegste type)", f"=IF(MIN({m_range_abs('verv', 6)})>=99999,99999,MIN({m_range_abs('verv', 6)}))"),
        "start_tekst": ("Start bouw project", f'=IF({h("start_idx")}>=99999,"–","Q"&(MOD({h("start_idx")}-1,4)+1)&" "&INT(({h("start_idx")}-1)/4))'),
        "rente_eind_basis": ("Rente-eindindex basis (rente zolang idx < deze)", f'=IF({h("rente_tm_start")}=1,{h("start_idx")},99999)'),
        "rente_eind_up": ("Rente-eindindex upside", f'=IF(OR({h("rente_tm_start")}=0,{h("start_idx")}>=99999),99999,{h("start_idx")}+N({dabs("uitstel_up")}))'),
        "rente_eind_down": ("Rente-eindindex downside", f'=IF(OR({h("rente_tm_start")}=0,{h("start_idx")}>=99999),99999,{h("start_idx")}+N({dabs("uitstel_down")}))'),
        "rente_basis_tot": ("Rente basis totaal (berekend)", f"=SUM({rng(M['rente_basis'])})"),
        "rente_up_tot": ("Rente upside totaal", f"=SUM({rng(M['rente_up'])})"),
        "rente_down_tot": ("Rente downside totaal", f"=SUM({rng(M['rente_down'])})"),
        "rente_up_extra": ("Extra rente upside t.o.v. basis", f'={h("rente_up_tot")}-{h("rente_basis_tot")}'),
        "rente_down_extra": ("Extra rente downside t.o.v. basis", f'={h("rente_down_tot")}-{h("rente_basis_tot")}'),
        "rente_tekst": ("Rente-tekst", f'=IF({h("rente")}=0,"geen rente","rente "&FIXED({h("rente")}*100,1)&"%"&IF({h("rente_tm_start")}=1," t/m start bouw",'
                                       f'" over de hele looptijd")&IF({h("rente_in_basis")}=1,", ook in de basis",""))'
                                       f'&IF(OR({h("rente_pct_down")}<>{h("rente")},{h("rente_pct_up")}<>{h("rente")},{h("rente_pct_scn")}<>{h("rente")}),'
                                       f'" · per scenario "&FIXED({h("rente_pct_down")}*100,1)&"% / "&FIXED({h("rente_pct_up")}*100,1)&"% / "'
                                       f'&FIXED({h("rente_pct_scn")}*100,1)&"%","")'),
        "aantal_types": ("Aantal woningtypes (met aantal)", f'=COUNTIF({m_range_abs("vcum", 6)},">0")'),
        # scenariolijn (gele lijn in de cashflowgrafiek)
        "scn_aan": ("Scenariolijn tonen (1/0)", f'=IF(LOWER({dabs("scn_aan")}&"")="nee",0,1)'),
        "scn_naam": ("Naam scenariolijn", f'=IF(TRIM({dabs("scn_naam")}&"")="","Scenario",TRIM({dabs("scn_naam")}&""))'),
        "eind_scn": ("Eindsaldo scenariolijn", f'=IF({h("n")}=0,0,INDEX({rng(ss)},{h("n")}))'),
        "dal_scn": ("Dieptepunt scenariolijn", f"=MIN({rng(ss)})"),
        "dal_scn_kw": ("Dieptepunt scenariolijn kwartaal", f'=IFERROR(INDEX({D_},MATCH({h("dal_scn")},{rng(ss)},0)),"–")'),
        "be_scn": ("Break-even scenariolijn", f'=IFERROR(INDEX({D_},MATCH(1,{rng(ps)},0)),"")'),
        "be_scn_pos": ("Positie break-even scenariolijn", f'=IFERROR(MATCH(1,{rng(ps)},0),0)'),
        "rente_eind_scn": ("Rente-eindindex scenariolijn", f'=IF(OR({h("rente_tm_start")}=0,{h("start_idx")}>=99999),99999,{h("start_idx")}+N({dabs("uitstel_scn")}))'),
        "rente_scn_tot": ("Rente scenariolijn totaal", f"=SUM({rng(M['rente_scn'])})"),
        "rente_scn_extra": ("Extra rente scenariolijn t.o.v. basis", f'={h("rente_scn_tot")}-{h("rente_basis_tot")}'),
        "lbl_scenario": ("Legenda scenariolijn", f'={h("scn_naam")}&IF({h("scn_aan")}=1,""," (uit)")'),
        "lbl_eind_scn": ("Label eind scenariolijn", f'={h("scn_naam")}&"  "&{eur_m(h("eind_scn"))}'),
        "lbl_nu_cf": ("Label stand nu (cashflowgrafiek)", f'="Stand "&{h("kw_actuals")}&"  "&{eur_m(h("stand_nu"))}'),
        "lbl_eind_basis_cf": ("Label eindsaldo (cashflowgrafiek)", f'="Eindsaldo "&{h("laatste_kw")}&"  "&{eur_m(h("eind_basis"))}'),
        # verkoopgrafiek
        "uitverkocht_pos": ("Positie uitverkocht (0 = niet binnen de looptijd)",
                            f'=IF(OR({h("woningen")}=0,(COUNTIF({rng(M["g_verkocht_pct"])},"<"&0.99999)+1)>{h("n")}),0,COUNTIF({rng(M["g_verkocht_pct"])},"<"&0.99999)+1)'),
        "alles_transport_pos": ("Positie alles getransporteerd (0 = niet binnen de looptijd)",
                                f'=IF(OR({h("woningen")}=0,(COUNTIF({rng(M["g_getransporteerd_pct"])},"<"&0.99999)+1)>{h("n")}),0,COUNTIF({rng(M["g_getransporteerd_pct"])},"<"&0.99999)+1)'),
        "lbl_uitverkocht": ("Label uitverkocht", f'="Uitverkocht "&{h("uitverkocht")}'),
        "lbl_alles_transport": ("Label laatste transport", f'="Laatste transport "&{h("alles_transport")}'),
        "nog_verkopen": ("Nog te verkopen (prognose)", f'=MAX(0,{h("woningen")}-{h("verkocht_actuals")})'),
        "nog_transport": ("Nog te transporteren (prognose)", f'=MAX(0,{h("woningen")}-{h("transport_actuals")})'),
        # eigen jaarrente per scenario (lege knop = de algemene jaarrente) en koopsomknop (VON-prijs) per scenario
        "rente_pct_down": ("Jaarrente downside", f'=IF({dabs("rente_pct_down")}="",{h("rente")},N({dabs("rente_pct_down")}))'),
        "rente_pct_up": ("Jaarrente upside", f'=IF({dabs("rente_pct_up")}="",{h("rente")},N({dabs("rente_pct_up")}))'),
        "rente_pct_scn": ("Jaarrente scenariolijn", f'=IF({dabs("rente_pct_scn")}="",{h("rente")},N({dabs("rente_pct_scn")}))'),
        "koopsom_down": ("Koopsom downside (%)", f'=N({dabs("koopsom_down")})'),
        "koopsom_up": ("Koopsom upside (%)", f'=N({dabs("koopsom_up")})'),
        "koopsom_scn": ("Koopsom scenariolijn (%)", f'=N({dabs("koopsom_scn")})'),
        # bouwtermijnengrafiek: aantal rijen en de tijd-as in kwartaalindex (jaar × 4 + kwartaal): van 1 januari van het jaar twee
        # jaar vóór de eerste termijn (huidig of vorig; niet vóór de eerste periode) tot het einde van het laatste jaar van de
        # periodes; de VBA leest bt_n (rijen), t_first en t_end (as-grenzen), cache.py zet de Excel-grafiek op dezelfde grenzen
        "bt_min": ("Tijd-as: eerste termijn (kwartaalindex, huidig of vorige prognose)",
                   ArrayFormula(f"{LY.M_HULP}{LY.H['bt_min']}",
                                f"=IF(SUM(${LY.BT['gebruikt']}${LY.BT_ROW1}:${LY.BT['gebruikt']}${LY.BT_ROWN})=0,99999,"
                                f"MIN(IF(${LY.BT['gebruikt']}${LY.BT_ROW1}:${LY.BT['gebruikt']}${LY.BT_ROWN}=1,"
                                f"IF(${LY.BT['idxr']}${LY.BT_ROW1}:${LY.BT['idxr']}${LY.BT_ROWN}<${LY.BT['idxb']}${LY.BT_ROW1}:${LY.BT['idxb']}${LY.BT_ROWN},"
                                f"${LY.BT['idxr']}${LY.BT_ROW1}:${LY.BT['idxr']}${LY.BT_ROWN},${LY.BT['idxb']}${LY.BT_ROW1}:${LY.BT['idxb']}${LY.BT_ROWN}),99999)))")),
        "bt_lanes": ("Bouwtermijnentabel: aantal banen (type × termijn)", f"=SUM(${LY.BT['gebruikt']}${LY.BT_ROW1}:${LY.BT['gebruikt']}${LY.BT_ROWN})"),
        "bt_uniek": ("Bouwtermijnentabel: aantal unieke termijnnamen", f"=SUM(${LY.BT['uniek']}${LY.BT_ROW1}:${LY.BT['uniek']}${LY.BT_ROWN})"),
        "bt_n": ("Bouwtermijnentabel: aantal rijen (koprijen, banen, scheidingsrijen)",
                 f'=IF({h("bt_lanes")}=0,0,MIN({LY.BT_N},{LY.BT_KOP}+{h("bt_lanes")}+{h("bt_uniek")}-1))'),
        "jaar_first": ("Tijd-as: eerste jaar (twee jaar vóór de eerste termijn, niet vóór de eerste periode)",
                       f'=IF({h("n")}=0,0,MAX(INT(({IDX}{ROW1}-1)/4),IF({h("bt_min")}>=99999,0,INT(({h("bt_min")}-1)/4)-2)))'),
        "jaar_last": ("Tijd-as: laatste jaar", f'=IF({h("n")}=0,0,INT((INDEX({rng(IDX)},{h("n")})-1)/4))'),
        "bt_jaren": ("Tijd-as: aantal jaren (koprij)", f'=MAX(0,MIN({LY.BT_NJ},{h("jaar_last")}-{h("jaar_first")}+1))'),
        "t_first": ("Tijd-as: ondergrens (kwartaalindex 1 jan eerste jaar)", f'={h("jaar_first")}*4+1'),
        "t_end": ("Tijd-as: bovengrens (kwartaalindex na het laatste jaar)", f'=({h("jaar_last")}+1)*4+1'),
        "t_nu_end": ("Tijd einde laatste gerealiseerde kwartaal (kwartaalindex)",
                     f'=IF({h("pos_actuals")}=0,{h("t_first")},MIN(MAX({h("idx_actuals")}+1,{h("t_first")}),{h("t_end")}))'),
        "lbl_realisatie": ("Legenda gerealiseerd", f'="Gerealiseerd t/m "&{h("kw_actuals")}'),
        "bt_types": ("Aantal typen in de bouwtermijnengrafiek", f'=COUNTIF(${LY.BT["cur1"]}$7:${LY.BT[f"cur{N_TYPES}"]}$7,"<>(uit)")'),
        # fees (DAEB)
        "fee_fout_gepland": ("Fee-componenten met totaal die niet (bijna) volledig zijn gepland",
                             "=" + "+".join(f'IF(AND(N({LY.wt_ref(LY.WT_R_FC1 + c, k, 1)})>0,N({m_col("fee", k)}$5)=1,ISNUMBER({LY.wt_ref(LY.WT_R_FC1 + c, k, 2)}),'
                                            f'ABS(N({LY.wt_ref(LY.WT_R_FC1 + c, k, 2)})-1)>=0.005),1,0)'
                                            for k in range(1, N_TYPES + 1) for c in range(LY.N_FEE_COMP))),
        "fee_fout_termijn": ("Fee-termijnen zonder geldig kwartaal of bedrag", f"=SUM({rng_bt('fout')})"),
        "daeb_types": ("Aantal DAEB-typen", f"=SUM({m_range_abs('fee', 5)})"),
        # grafiek per woningtype (keuzecel op het Dashboard; naam -> typenummer; onbekend = 1)
        "pt_keuze": ("Grafiek per type: gekozen typenummer (getal = typenummer, anders naam)",
                     f'=IF(ISNUMBER({dabs("pt_keuze")}),MIN({N_TYPES},MAX(1,INT({dabs("pt_keuze")}))),'
                     f'IFERROR(MATCH({dabs("pt_keuze")},${m_col("gv", 1)}$7:${m_col("gv", N_TYPES)}$7,0),1))'),
        "pt_naam": ("Grafiek per type: naam", f'=INDEX(${m_col("gv", 1)}$7:${m_col("gv", N_TYPES)}$7,{h("pt_keuze")})'),
        "pt_daeb": ("Grafiek per type: DAEB (1/0)", f'=N(INDEX({m_range_abs("fee", 5)},{h("pt_keuze")}))'),
        "pt_aantal": ("Grafiek per type: aantal woningen", f'=N(INDEX({m_range_abs("vcum", 6)},{h("pt_keuze")}))'),
        "pt_totaal": ("Grafiek per type: totaal volgens planning (€ mln)", f'=IF({h("n")}=0,0,INDEX({rng(LY.PT["cum"])},{h("n")}))'),
        "pt_nu": ("Grafiek per type: t/m actuals volgens planning (€ mln)", f'=IF({h("pos_actuals")}=0,0,INDEX({rng(LY.PT["cum"])},{h("pos_actuals")}))'),
        "pt_titel": ("Grafiek per type: titel", f'={h("pt_naam")}&" · "&{h("pt_aantal")}&" woningen · "&IF({h("pt_daeb")}=1,"DAEB (fees)","koop")'
                                                f'&" · totaal €"&FIXED({h("pt_totaal")},1)&" mln volgens planning"'),
        "pt_subtitel": ("Grafiek per type: ondertitel",
                        f'="Per kwartaal · t/m "&{h("kw_actuals")}&" €"&FIXED({h("pt_nu")},1)&" mln, daarna €"&FIXED({h("pt_totaal")}-{h("pt_nu")},1)'
                        f'&" mln · grijs = gerealiseerde kwartalen (bedragen volgens planning; werkelijke ontvangsten staan in je eigen cashflow)"'
                        f'&IF(AND({h("pt_daeb")}=1,{h("fee_fout_gepland")}+{h("fee_fout_termijn")}>0)," · let op: fees niet volledig ingepland","")'),
        "pt_lbl_grond": ("Grafiek per type: reeksnaam grondtermijn", f'=IF(COUNTIF({rng(LY.PT["grond"])},">0")>0,"Grondtermijn","")'),
        "pt_lbl_bouw": ("Grafiek per type: reeksnaam bouwtermijnen", f'=IF(COUNTIF({rng(LY.PT["bouw"])},">0")>0,"Bouwtermijnen","")'),
        "pt_lbl_extra": ("Grafiek per type: reeksnaam extra's", f'=IF(COUNTIF({rng(LY.PT["extra"])},">0")>0,"Extra opbrengsten","")'),
        "pt_lbl_fee1": ("Grafiek per type: reeksnaam fee 1", f'=IF(COUNTIF({rng(LY.PT["fee1"])},">0")>0,'
                                                             f'INDEX(Woningtypes!$A${LY.WT_R_FC1}:${LY.WT_RANGE_END}${LY.WT_R_FC1},1,{LY.WT_COL1}+({h("pt_keuze")}-1)*{LY.WT_W})&"","")'),
        "pt_lbl_fee2": ("Grafiek per type: reeksnaam fee 2", f'=IF(COUNTIF({rng(LY.PT["fee2"])},">0")>0,'
                                                             f'INDEX(Woningtypes!$A${LY.WT_R_FC1 + 1}:${LY.WT_RANGE_END}${LY.WT_R_FC1 + 1},1,{LY.WT_COL1}+({h("pt_keuze")}-1)*{LY.WT_W})&"","")'),
        "pt_lbl_cum": ("Grafiek per type: reeksnaam cumulatief", '="Cumulatief"'),
        "pt_lbl_totaal": ("Grafiek per type: reeksnaam totaal per kwartaal", '="Totaal per kwartaal"'),
        "pt_lbl_eind": ("Grafiek per type: label eindpunt", f'="totaal €"&FIXED({h("pt_totaal")},1)&" mln"'),
        # tijdvak van de grafiek per type: vanaf 1 januari van het jaar twee jaar vóór de eerste opbrengst van het gekozen type
        # (niet vóór de eerste periode); pt_start = modelpositie van de eerste getoonde periode, pt_n = aantal getoonde periodes
        "pt_first": ("Grafiek per type: eerste kwartaal met opbrengst (kwartaalindex)",
                     ArrayFormula(f"{LY.M_HULP}{LY.H['pt_first']}",
                                  f'=IF(SUMPRODUCT(--(IFERROR({rng(LY.PT["totaal"])},0)>0.000001))=0,99999,'
                                  f'MIN(IF(IFERROR({rng(LY.PT["totaal"])},0)>0.000001,{rng(IDX)},99999)))')),
        "pt_jaar_first": ("Grafiek per type: eerste jaar van het tijdvak",
                          f'=IF({h("n")}=0,0,MAX(INT(({IDX}{ROW1}-1)/4),IF({h("pt_first")}>=99999,0,INT(({h("pt_first")}-1)/4)-2)))'),
        "pt_start": ("Grafiek per type: positie van de eerste getoonde periode",
                     ArrayFormula(f"{LY.M_HULP}{LY.H['pt_start']}",
                                  f'=IF({h("n")}=0,1,MIN({h("n")},IFERROR(MATCH(TRUE,{rng(IDX)}>={h("pt_jaar_first")}*4+1,0),1)))')),
        "pt_n": ("Grafiek per type: aantal getoonde periodes", f'=MAX(1,{h("n")}-{h("pt_start")}+1)'),
    }
    for naam, (label, formule) in hulp.items():
        r = H[naam]
        put(ws, f"{LY.M_HULP_LABEL}{r}", label, f=F_NOTE8)
        put(ws, f"{LY.M_HULP}{r}", formule, f=F_CALC, al=AL_LEFT_TOP)
    for naam in ("verk_voor_start_pct", "rente", "rente_pct_down", "rente_pct_up", "rente_pct_scn", "koopsom_down", "koopsom_up", "koopsom_scn"):
        ws[f"{LY.M_HULP}{H[naam]}"].number_format = FMT_PCT
    for naam in ("pt_totaal", "pt_nu"):
        ws[f"{LY.M_HULP}{H[naam]}"].number_format = "0.00"
    for naam in ("stand_nu", "opbr_totaal", "kosten_totaal", "cum_opbr_actuals", "dal_basis", "eind_basis", "dal_up", "eind_up",
                 "dal_down", "eind_down", "dal_vorig", "eind_vorig", "vorig_nu", "model_opbr", "rente_basis_tot", "rente_up_tot",
                 "rente_down_tot", "rente_up_extra", "rente_down_extra", "eind_scn", "dal_scn", "rente_scn_tot", "rente_scn_extra"):
        ws[f"{LY.M_HULP}{H[naam]}"].number_format = FMT_INT
    return ws


# =============================================================================
#  DASHBOARD
# =============================================================================
def bouw_dashboard(wb, data):
    ws = wb.create_sheet("Dashboard", 0)
    ws.sheet_properties.codeName = "shDashboard"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    for col, w in (("A", 2), ("B", 9), ("F", 2), ("G", 9), ("K", 2), ("L", 9), ("P", 2), ("Q", 9), ("V", 36), ("W", 12), ("X", 12), ("Y", 12)):
        ws.column_dimensions[col].width = w
    for r, hgt in ((1, 15.6), (2, 25.95), (3, 20), (5, 16.05), (6, 28.05), (7, 13.95), (8, 15), (10, 16.05), (11, 24), (12, 13.95), (13, 15)):
        ws.row_dimensions[r].height = hgt
    P = data.params
    rb = hm("rente_in_basis")

    # ---- kopregels ---------------------------------------------------------------
    put(ws, "B1", f'="CASHFLOW · STAND PER "&UPPER({hm("kw_actuals")})', f=F_NOTE)
    ws.merge_cells("B2:T2")
    put(ws, "B2", f'="Financiering piekt op €"&FIXED(ABS({hm("dal_basis")})/1000000,1)&" mln in "&{hm("dal_basis_kw")}'
                  f'&IF({hm("be_basis")}="",", geen break-even",", break-even in "&{hm("be_basis")})', f=font(14, True, TXT), al=AL_VCENTER)
    ws.merge_cells("B3:T3")
    put(ws, "B3", f'="Projectresultaat "&{eur_mln(hm("eind_basis"))}&" · marge "&IFERROR(FIXED({hm("eind_basis")}/{hm("opbr_totaal")}*100,1),"0")'
                  f'&"% · opbrengsten €"&FIXED({hm("opbr_totaal")}/1000000,1)&" mln · kosten €"&FIXED({hm("kosten_totaal")}/1000000,1)&" mln · basis = je eigen cashflow"'
                  f'&IF({rb}=1," + rente ("&FIXED({hm("rente_basis_tot")}/1000000,1)&" mln)","")', f=font(10, False, TXT2))

    # ---- KPI-kaarten -------------------------------------------------------------
    def kaart(col0, r0, titel, waarde, sub, verschil, nf=FMT_MLN, kleur=TXT, size=18):
        c0, c3 = L(col0), L(col0 + 3)
        for r in (r0, r0 + 1, r0 + 2, r0 + 3):
            ws.merge_cells(f"{c0}{r}:{c3}{r}")
        put(ws, f"{c0}{r0}", titel, f=F_NOTE, fl=FL_CARD, al=AL_LEFT,
            bd=Border(left=side(), right=side(), top=side()))
        for c in range(col0 + 1, col0 + 3):
            style(ws, f"{L(c)}{r0}", bd=Border(top=side()))
        style(ws, f"{c3}{r0}", bd=Border(top=side(), right=side()))
        put(ws, f"{c0}{r0 + 1}", waarde, f=font(size, True, kleur), fl=FL_CARD, nf=nf, al=AL_LEFT, bd=Border(left=side(), right=side()))
        style(ws, f"{c3}{r0 + 1}", bd=Border(right=side()))
        put(ws, f"{c0}{r0 + 2}", sub, f=font(9, False, TXT2), fl=FL_CARD, al=AL_LEFT, bd=Border(left=side(), right=side()))
        style(ws, f"{c3}{r0 + 2}", bd=Border(right=side()))
        put(ws, f"{c0}{r0 + 3}", verschil, f=F_NOTE, fl=FL_CARD, al=AL_LEFT,
            bd=Border(left=side(), right=side(), bottom=side("thin", LINE2)))
        for c in range(col0 + 1, col0 + 3):
            style(ws, f"{L(c)}{r0 + 3}", bd=Border(bottom=side("thin", LINE2)))
        style(ws, f"{c3}{r0 + 3}", bd=Border(right=side(), bottom=side("thin", LINE2)))

    dal, dal_v, eind, eind_v = hm("dal_basis"), hm("dal_vorig"), hm("eind_basis"), hm("eind_vorig")
    kaart(2, 5, "MAX. FINANCIERINGSBEHOEFTE", f"={dal}/1000000", f'="dieptepunt in "&{hm("dal_basis_kw")}',
          f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose","was "&{eur_m(dal_v)}&IF(ABS(({dal}-{dal_v}))<50000," · gelijk",'
          f'" · €"&FIXED(ABS(({dal}-{dal_v}))/1000000,1)&"M "&IF(({dal}-{dal_v})<0,"dieper","minder diep")))', kleur=RED)
    kaart(7, 5, "BREAK-EVEN", f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', "eerste kwartaal met positief saldo",
          f'=IF(OR({hm("be_basis_pos")}=0,{hm("be_vorig_pos")}=0),"geen vergelijking","was "&{hm("be_vorig")}&IF(({hm("be_basis_pos")}-{hm("be_vorig_pos")})=0," · gelijk",'
          f'" · "&ABS(({hm("be_basis_pos")}-{hm("be_vorig_pos")}))&" kw "&IF(({hm("be_basis_pos")}-{hm("be_vorig_pos")})>0,"later","eerder")))', nf="General")
    kaart(12, 5, "PROJECTRESULTAAT", f"={eind}/1000000",
          f'="eindsaldo in "&{hm("laatste_kw")}&" · marge "&IFERROR(FIXED({eind}/{hm("opbr_totaal")}*100,1),"0")&"%"',
          f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose","was "&{eur_m(eind_v)}&IF(ABS(({eind}-{eind_v}))<50000," · gelijk",'
          f'" · €"&FIXED(ABS(({eind}-{eind_v}))/1000000,1)&"M "&IF(({eind}-{eind_v})>0,"hoger","lager")))', kleur=GREEN)
    kaart(17, 5, "CUMULATIEVE CASHFLOW", f'={hm("stand_nu")}/1000000', f'="stand per "&{hm("kw_actuals")}',
          f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose","vorige prognose "&{eur_m(hm("vorig_nu"))})')
    kaart(2, 10, "DOWNSIDE · EINDSALDO", f'={hm("eind_down")}/1000000',
          f'="diepste dal "&{eur_m(hm("dal_down"))}&" in "&{hm("dal_down_kw")}',
          f'="break-even "&IF({hm("be_down")}="","niet bereikt",{hm("be_down")})', kleur=RED, size=16)
    kaart(7, 10, "BASIS · EINDSALDO", f"={eind}/1000000", f'="diepste dal "&{eur_m(dal)}&" in "&{hm("dal_basis_kw")}',
          f'="break-even "&IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', kleur=BLUE_DARK, size=16)
    kaart(12, 10, "UPSIDE · EINDSALDO", f'={hm("eind_up")}/1000000',
          f'="diepste dal "&{eur_m(hm("dal_up"))}&" in "&{hm("dal_up_kw")}',
          f'="break-even "&IF({hm("be_up")}="","niet bereikt",{hm("be_up")})', kleur=GREEN, size=16)
    kaart(17, 10, "BANDBREEDTE EINDSALDO", f'=ABS({hm("eind_up")}-{hm("eind_down")})/1000000',
          f'="van "&{eur_m(hm("eind_down"))}&" tot "&{eur_m(hm("eind_up"))}',
          f'="diepste dal van "&{eur_m("MIN(" + hm("dal_up") + "," + hm("dal_down") + ")")}&" tot "&{eur_m("MAX(" + hm("dal_up") + "," + hm("dal_down") + ")")}', size=16)

    put(ws, "B14", "Scenario's · cumulatieve cashflow: basis (donkerblauw) met de band tussen downside en upside · grijs vlak = gerealiseerde kwartalen · € mln", f=F_SECTION)
    put(ws, "B43", "Cashflow per kwartaal · balken: opbrengsten en kosten · lijn: cumulatieve cashflow · stippellijn: vorige prognose · "
                   "gele lijn: scenariolijn (knoppen rechts, kolom Scenario) · grijs vlak = gerealiseerd · € mln", f=F_SECTION)
    put(ws, "B69", VERKOOP_TITEL, f=F_SECTION)
    put(ws, f"B{LY.D_ROW_BT}", "Bouwtermijnen per woningtype · kwartaal waarin de termijn vervalt (tab Woningtypes) · per termijn een baan per type "
                               "in de typekleur · lichtrood = vorige prognose (blok 4, fees: blok 6) · grijs = gerealiseerd · DAEB: fee-termijnen per component · "
                               "tijd-as vanaf twee jaar vóór de eerste termijn", f=F_SECTION)

    # cel-legenda: één gekleurde chip per typeblok (zelfde kleur als in de grafiek), leeg en wit als het blok niet meedoet
    def chips(r, meedoen, tekst):
        ws.row_dimensions[r].height = 16
        for k, col in enumerate(LY.D_LEGENDA_CELLEN, start=1):
            naam = LY.wt_ref(LY.WT_R_NAAM, k)
            cel = f"{col}{r}"
            put(ws, cel, f'=IF({meedoen(k)},IF({naam}="","Type {k}",{naam}),"")', f=font(8, True, "FFFFFF"),
                fl=fill(CX.TYPEKLEUREN[k - 1]), al=Alignment(horizontal="center", vertical="center", shrink_to_fit=True))
            ws.conditional_formatting.add(cel, FormulaRule(formula=[f'{cel}=""'], fill=PatternFill(fill_type="solid", bgColor="FFFFFF", fgColor="FFFFFF")))
        ws.merge_cells(f"O{r}:T{r}")
        put(ws, f"O{r}", tekst, f=F_NOTE, al=AL_RIGHT)

    chips(LY.D_ROW_LEGENDA, lambda k: f'N(Model!{m_col("vcum", k)}$6)>0', "zelfde kleuren in beide panelen · ● = mijlpaal")
    chips(LY.D_ROW_LEGENDA_BT, lambda k: f'Model!${LY.BT[f"cur{k}"]}$7<>"(uit)"',
          f'="lichtrood = vorige prognose · grijs = gerealiseerd t/m "&{hm("kw_actuals")}')

    # grafiek per woningtype: keuzecel (naam uit de lijst op Model rij 7), chips per onderdeel, ondertitel
    put(ws, f"B{LY.D_ROW_PT}", "Opbrengsten per woningtype · kies het type in de blauwe cel · per kwartaal: grondtermijn, bouwtermijnen en extra's "
                               "(koop) of fees per component (DAEB) · lijn = cumulatief · grijs = gerealiseerd · € mln · tijdvak vanaf twee jaar vóór de eerste opbrengst", f=F_SECTION)
    r = LY.D_ROW_LEGENDA_PT
    ws.row_dimensions[r].height = 16
    kz = LY.D["pt_keuze"]
    ws.merge_cells(f"{kz}:C{r}")
    put(ws, kz, data.type_(1).naam or "Type 1", f=F_INPUT, fl=FL_INPUT, nf="@", al=Alignment(horizontal="left", vertical="center"))
    dv_pt = DataValidation(type="list", formula1=f"Model!${m_col('gv', 1)}$7:${m_col('gv', N_TYPES)}$7", allow_blank=True,
                           error="Kies een woningtype uit de lijst.")
    dv_pt.showErrorMessage = True
    dv_pt.errorTitle = "Ongeldige invoer"
    ws.add_data_validation(dv_pt)
    dv_pt.add(kz)
    put(ws, f"D{r}", "← kies het woningtype", f=F_NOTE8, al=AL_VCENTER)
    for naam, col in zip([x for x in LY.PT_REEKSEN if x != "cum"], LY.D_LEGENDA_PT_CELLEN):
        cel = f"{col}{r}"
        lbl = hm("pt_lbl_" + naam)
        put(ws, cel, f'=IF(LEN({lbl})>16,LEFT({lbl},15)&"…",{lbl})', f=font(8, True, "FFFFFF"), fl=fill(LY.PT_KLEUREN[naam]),
            al=Alignment(horizontal="center", vertical="center", shrink_to_fit=True))
        ws.conditional_formatting.add(cel, FormulaRule(formula=[f'{cel}=""'], fill=PatternFill(fill_type="solid", bgColor="FFFFFF", fgColor="FFFFFF")))
    ws.merge_cells(f"M{r}:T{r}")
    put(ws, f"M{r}", f'={hm("pt_naam")}&" · "&{hm("pt_aantal")}&" woningen · "&IF({hm("pt_daeb")}=1,"DAEB (fees)","koop")&" · totaal €"&FIXED({hm("pt_totaal")},1)&" mln"',
        f=F_NOTE, al=AL_RIGHT)

    # ---- parameters (kolom V/W/X/Y, toelichting in Z) --------------------------------
    put(ws, "V1", "KNOPPEN", f=F_TITLE)
    put(ws, "V2", "Geel = invoer. Cijfers vul je in op tab Invoer, woningtypes en termijnen op tab Woningtypes.", f=F_NOTE)

    def sectie(r, t):
        put(ws, f"V{r}", t, f=F_SECTION, bd=BD_SECTION)
        for c in "WXYZ":
            style(ws, f"{c}{r}", bd=BD_SECTION)

    def label(r, t, note=None):
        put(ws, f"V{r}", t, f=F_NOTE, al=AL_VCENTER)
        if note:
            put(ws, f"Z{r}", note, f=F_NOTE8, al=AL_VCENTER)

    def geel(ref, v, nf="General"):
        put(ws, ref, v, f=F_YELLOW, fl=FL_YELLOW, nf=nf, al=AL_CENTER)

    def info(ref, v, nf="General", kleur=TXT):
        put(ws, ref, v, f=font(10, True, kleur), nf=nf, al=AL_CENTER)

    sectie(4, "ALGEMEEN")
    label(5, LY.D_LABELS["actuals_jaar"], "t/m dit kwartaal veranderen de scenario's niets")
    geel("W5", P["actuals_jaar"], "0")
    geel("X5", P["actuals_kw"], "0")
    label(6, LY.D_LABELS["rente"], "over de negatieve stand van het vorige kwartaal, alleen in prognosekwartalen · 0% = geen rente")
    geel("W6", P["rente"], FMT_PCT)
    ws["W6"].comment = _comment("Rente per jaar over de negatieve stand (voorfinanciering) van het vorige kwartaal, gedeeld door 4. "
                                "Alleen in prognosekwartalen. 0% = geen rente.")
    label(7, LY.D_LABELS["rente_tm"], "start bouw = t/m het kwartaal vóór start bouw van het project (+ uitstel) · hele looptijd = alle prognosekwartalen")
    geel("W7", P["rente_tm"])
    label(8, LY.D_LABELS["rente_basis"], "nee = je eigen cashflow bevat al rente; scenario's tellen alleen de extra rente · ja = basis en scenario's krijgen rente")
    geel("W8", P["rente_basis"])
    ws["W8"].comment = _comment("nee: de basis blijft je eigen cashflow; downside en upside krijgen alleen het verschil in rente ten opzichte "
                                "van de basis (door verschuiven, uitstel of andere kosten). ja: ook de basis krijgt rente over de negatieve stand; "
                                "de kaarten en dia's rekenen daar dan mee.")
    label(9, "Start bouw project", "vroegste start bouw van de woningtypes (tab Woningtypes)")
    info("W9", f"={hm('start_tekst')}")

    sectie(11, "VERKOOPTEMPO-MODEL")
    label(12, LY.D_LABELS["model_aan"], "ja = grondtermijn schuift mee met transport · nee = alle prognose-opbrengsten schuiven")
    geel("W12", P["model_aan"])
    ws["W12"].comment = _comment("ja: het model gebruikt je woningtypes, grondtermijn en bouwtermijnen om alleen het verschil door een ander "
                                 "verkooptempo uit te rekenen. nee: de verschuiving verplaatst al je prognose-opbrengsten.")
    label(13, "Model dekt van je opbrengsten", "ter controle: aantal × koopsom volgens je planning, gedeeld door je eigen opbrengsten")
    info("W13", f'=IFERROR({hm("model_opbr")}/{hm("opbr_totaal")},0)', FMT_PCT0)

    r = LY.D_ROW_SCENARIO
    sectie(r, "SCENARIO'S")
    put(ws, f"W{r + 1}", "Downside", f=font(9, True, RED), al=AL_CENTER)
    put(ws, f"X{r + 1}", "Upside", f=font(9, True, GREEN), al=AL_CENTER)
    put(ws, f"Y{r + 1}", "Scenario", f=font(9, True, AMBER_TXT), al=AL_CENTER)   # data.py herkent deze kop
    put(ws, f"Z{r + 1}", "Scenario = de gele lijn in de cashflowgrafiek, met eigen knoppen; downside en upside vormen de band in de scenariografiek",
        f=F_NOTE8, al=AL_VCENTER)
    label(17, LY.D_LABELS["shift_down"], "+ = later, − = eerder")
    geel("W17", P["shift_down"], FMT_KW)
    geel("X17", P["shift_up"], FMT_KW)
    geel("Y17", P["shift_scn"], FMT_KW)
    ws["W17"].comment = _comment("Aantal kwartalen dat verkoop en notarieel transport later (+) of eerder (−) vallen dan in de basis. "
                                 "Geldt alleen voor prognosekwartalen.")
    label(18, LY.D_LABELS["uitstel_down"], "verlengt alleen de renteperiode (bij rente t/m start bouw); kosten en termijnen schuiven niet")
    geel("W18", P["uitstel_down"], FMT_KW)
    geel("X18", P["uitstel_up"], FMT_KW)
    geel("Y18", P["uitstel_scn"], FMT_KW)
    ws["W18"].comment = _comment("Start bouw zoveel kwartalen later dan gepland. Net als in het oude template verlengt dit alleen de periode "
                                 "waarover rente loopt: de kosten in je eigen cashflow en de bouwtermijnen blijven staan.")
    label(19, LY.D_LABELS["opbr_down"], "op alle prognose-opbrengsten (indexatie), ook buiten het verkooptempo-model")
    geel("W19", P["opbr_down"], FMT_PCT_SIGN)
    geel("X19", P["opbr_up"], FMT_PCT_SIGN)
    geel("Y19", P["opbr_scn"], FMT_PCT_SIGN)
    label(20, LY.D_LABELS["kosten_down"], "op alle prognosekosten (kostenindexatie)")
    geel("W20", P["kosten_down"], FMT_PCT_SIGN)
    geel("X20", P["kosten_up"], FMT_PCT_SIGN)
    geel("Y20", P["kosten_scn"], FMT_PCT_SIGN)
    label(21, LY.D_LABELS["rente_pct_down"], "leeg = de algemene jaarrente hierboven · zelfde regels: negatieve stand × rente / 4, alleen prognose")
    geel("W21", P["rente_pct_down"], FMT_PCT)
    geel("X21", P["rente_pct_up"], FMT_PCT)
    geel("Y21", P["rente_pct_scn"], FMT_PCT)
    ws["W21"].comment = _comment("Eigen jaarrente voor dit scenario. Leeg = de algemene jaarrente (knop Jaarrente bovenaan). "
                                 "De renteperiode (t/m start bouw + uitstel, of hele looptijd) blijft gelijk.")
    label(22, LY.D_LABELS["koopsom_down"], "hogere of lagere VON-prijs, dus grondtermijn én bouwtermijnen · werkt via het verkooptempo-model (knop 'ja')")
    geel("W22", P["koopsom_down"], FMT_PCT_SIGN)
    geel("X22", P["koopsom_up"], FMT_PCT_SIGN)
    geel("Y22", P["koopsom_scn"], FMT_PCT_SIGN)
    ws["W22"].comment = _comment("Koopsom (VON-prijs) van alle woningen zoveel procent hoger (+) of lager (−) dan op tab Woningtypes. "
                                 "Grondtermijn en bouwtermijnen schalen mee, op de momenten van het verkooptempo-model. Alleen met "
                                 "'Woningtypes en termijnen gebruiken = ja'; de knop Opbrengsten hieronder werkt op alle prognose-opbrengsten.")
    label(23, "Rente in de scenario's (totaal)")
    info("W23", f'={hm("rente_down_tot")}/1000000', FMT_MLN, RED)
    info("X23", f'={hm("rente_up_tot")}/1000000', FMT_MLN, GREEN)
    info("Y23", f'={hm("rente_scn_tot")}/1000000', FMT_MLN, AMBER_TXT)
    put(ws, "Z23", f'=IF(AND({hm("rente")}=0,{hm("rente_pct_down")}=0,{hm("rente_pct_up")}=0,{hm("rente_pct_scn")}=0),"zet een jaarrente om rente mee te rekenen",'
                   f'"rente over de basis zou €"&FIXED({hm("rente_basis_tot")}/1000000,2)&" mln zijn"&'
                   f'IF({rb}=1," (zit in de basis)"," (niet in de basis: scenario\'s tellen alleen het verschil)"))', f=F_NOTE8, al=AL_VCENTER)
    label(24, "Eindsaldo")
    info("W24", f'={hm("eind_down")}/1000000', FMT_MLN, RED)
    info("X24", f'={hm("eind_up")}/1000000', FMT_MLN, GREEN)
    info("Y24", f'={hm("eind_scn")}/1000000', FMT_MLN, AMBER_TXT)
    put(ws, "Z24", f'="basis "&{eur_m(hm("eind_basis"))}&" · scenario: dieptepunt "&{eur_m(hm("dal_scn"))}&" in "&{hm("dal_scn_kw")}'
                   f'&", break-even "&IF({hm("be_scn")}="","niet bereikt",{hm("be_scn")})', f=F_NOTE8, al=AL_VCENTER)
    label(25, LY.D_LABELS["scn_aan"], "nee = de gele lijn verdwijnt uit de grafiek (de knoppen blijven staan)")
    geel(LY.D["scn_aan"], P["scn_aan"])
    label(26, LY.D_LABELS["scn_naam"], "staat in de legenda en bij het eindpunt, bv. 'Versnelde verkoop' · leeg = Scenario")
    geel(LY.D["scn_naam"], P["scn_naam"])

    r = LY.D_ROW_VERKOOP
    sectie(r, "VERKOOP")
    label(r + 1, LY.D_LABELS["norm"])
    geel(LY.D["norm"], P["norm"], FMT_PCT0)
    label(r + 2, "Verkocht vóór start bouw")
    info(f"W{r + 2}", f'={hm("verk_voor_start_pct")}', FMT_PCT0)
    put(ws, f"Z{r + 2}", f'={hm("verk_voor_start")}&" van "&{hm("won_met_start")}&" woningen (types met start bouw)"', f=F_NOTE8, al=AL_VCENTER)
    label(r + 3, "Verkocht t/m nu")
    info(f"W{r + 3}", f'={hm("verkocht_actuals")}&" van "&{hm("woningen")}')
    put(ws, f"Z{r + 3}", f'="nog te verkopen in de prognose: "&{hm("nog_verkopen")}&" · uitverkocht in "&{hm("uitverkocht")}', f=F_NOTE8, al=AL_VCENTER)
    label(r + 4, "Getransporteerd t/m nu")
    info(f"W{r + 4}", f'={hm("transport_actuals")}&" van "&{hm("woningen")}')
    put(ws, f"Z{r + 4}", f'="nog te transporteren in de prognose: "&{hm("nog_transport")}&" · alles getransporteerd in "&{hm("alles_transport")}', f=F_NOTE8, al=AL_VCENTER)

    r = LY.D_ROW_POWERPOINT
    sectie(r, "POWERPOINT")
    label(r + 1, "Dia's maken met één klik")
    put(ws, f"V{r + 4}", "De knop staat hier in de .xlsm (macro's inschakelen). Zonder macro: blokken plakken vanaf tab PowerPoint.", f=F_NOTE8, al=AL_VCENTER)

    r = LY.D_ROW_CONTROLES
    sectie(r, "CONTROLES")
    vcum1, vcumN = m_col("vcum", 1), m_col("vcum", N_TYPES)
    tcum1, tcumN = m_col("tcum", 1), m_col("tcum", N_TYPES)
    b_rng = LY.in_rng("B")
    checks = [
        ("Periodes oplopend, zonder gaten",
         f'=IF(COUNTIF(Model!${M["stap_ok"]}${ROW1}:${M["stap_ok"]}${ROWN},0)=0,"OK","LET OP: jaar/kwartaal loopt niet op of er ontbreekt een kwartaal")'),
        ("Niet meer dan 60 periodes",
         f'=IF(COUNTA({b_rng})-1<={LY.N_PERIODS},"OK","LET OP: "&(COUNTA({b_rng})-1)&" gevulde rijen, alleen de eerste {LY.N_PERIODS} tellen mee")'),
        ("Elke gevulde rij heeft een jaar",
         f'=IF(SUMPRODUCT(({b_rng}="")*({LY.in_rng("C", LY.WT_RANGE_END)}<>""))=0,"OK","LET OP: rij met cijfers maar zonder jaar telt niet mee")'),
        ("Jaar is een getal tussen 1990 en 2100",
         f'=IF(SUMPRODUCT(({b_rng}<>"")*ISTEXT({b_rng}))-COUNTIF({b_rng},"Jaar")+COUNTIF({b_rng},"<1990")+COUNTIF({b_rng},">2100")=0,"OK",'
         f'"LET OP: jaar als tekst, datum of buiten 1990-2100 (tekst werkt, maar maak er een getal van)")'),
        ("Koprij 'Jaar' staat boven de periodes",
         f'=IF(IFERROR(INDEX({b_rng},MATCH(TRUE,INDEX({b_rng}<>"",0),0)),"")="Jaar","OK",'
         f'"LET OP: koprij \'Jaar\' op tab Invoer ontbreekt; de eerste gevulde rij telt dan niet mee")'),
        ("Verkocht = aantal woningen",
         f'=IF({hm("verkocht_plan")}={hm("woningen")},"OK","LET OP: "&{hm("verkocht_plan")}&" verkocht tegen "&{hm("woningen")}&" woningen")'),
        ("Getransporteerd = aantal woningen",
         f'=IF({hm("transport_plan")}={hm("woningen")},"OK","LET OP: "&{hm("transport_plan")}&" getransporteerd tegen "&{hm("woningen")}&" woningen")'),
        ("Geen transport vóór verkoop",
         f'=IF(SUMPRODUCT(--(Model!${tcum1}${ROW1}:${tcumN}${ROWN}>Model!${vcum1}${ROW1}:${vcumN}${ROWN}))=0,"OK",'
         f'"LET OP: in minstens één periode meer getransporteerd dan verkocht")'),
    ]
    tot = []
    kwc = []
    for k in range(1, N_TYPES + 1):
        aantal = f"N(Model!{m_col('vcum', k)}$6)"
        totaal = LY.wt_ref(LY.WT_R_TOTAAL, k, 1)
        grond_k = LY.wt_ref(LY.WT_R_GROND, k, 1)
        daeb_k = f"N(Model!{m_col('fee', k)}$5)=1"
        tot.append(f"OR({aantal}=0,{daeb_k},N({totaal})=0,ABS(N({totaal})-1)<0.00005,ABS(N({totaal})-(1-N({grond_k})))<0.00005)")
        kw_rng = LY.wt_ref(0, k, 2, rows=(LY.WT_R_T1, LY.WT_R_TN))
        pct_rng = LY.wt_ref(0, k, 1, rows=(LY.WT_R_T1, LY.WT_R_TN))
        kwc.append(f'OR({daeb_k},ABS(SUMIF({kw_rng},">=1",{pct_rng})-SUM({pct_rng}))<0.00005)')
    checks.append(("Bouwtermijnen per type 100% (of 100% − grond)",
                   f'=IF({hm("model_aan")}=0,"n.v.t. (model staat uit)",IF(AND({",".join(tot)}),"OK",'
                   f'"LET OP: bouwtermijnen tellen bij minstens één type niet op tot 100% of tot 100% − grondtermijn"))'))
    checks.append(("Elke bouwtermijn heeft een bouwkwartaal",
                   f'=IF({hm("model_aan")}=0,"n.v.t. (model staat uit)",IF(AND({",".join(kwc)}),"OK","LET OP: percentage zonder bouwkwartaal"))'))
    bt_ib = f"Model!${LY.BT['idxb']}${LY.BT_ROW1}:${LY.BT['idxb']}${LY.BT_ROWN}"
    bt_g = f"Model!${LY.BT['gebruikt']}${LY.BT_ROW1}:${LY.BT['gebruikt']}${LY.BT_ROWN}"
    idx_first = f"Model!${FIX['idx']}${ROW1}"
    idx_last = f"INDEX(Model!${FIX['idx']}${ROW1}:${FIX['idx']}${ROWN},MAX(1,{hm('n')}))"
    buiten = f'(COUNTIFS({bt_g},1,{bt_ib},"<"&{idx_first})+COUNTIFS({bt_g},1,{bt_ib},">"&{idx_last}))'
    checks.append(("Bouwtermijnen en fees vallen binnen de periodes",
                   f'=IF({buiten}=0,"OK","LET OP: "&{buiten}&" termijn(en) vallen buiten de periodes op tab Invoer: ze tellen niet mee in het model "'
                   f'&"en staan niet in de grafieken (al ontvangen fees: kwartaal actuals)")'))
    checks.append(("DAEB: fees per component volledig gepland",
                   f'=IF({hm("daeb_types")}=0,"n.v.t. (geen DAEB-type)",IF({hm("fee_fout_gepland")}=0,"OK","LET OP: bij "&{hm("fee_fout_gepland")}'
                   f'&" component(en) wijkt het geplande bedrag af van het totaal (tab Woningtypes, blok 5, kolom gepland)"))'))
    checks.append(("DAEB: elke fee-termijn heeft een geldig kwartaal en bedrag",
                   f'=IF({hm("daeb_types")}=0,"n.v.t. (geen DAEB-type)",IF({hm("fee_fout_termijn")}=0,"OK","LET OP: "&{hm("fee_fout_termijn")}'
                   f'&" fee-termijn(en) tellen niet mee: kwartaal niet herkend (2026 Q4, bouwkwartaal of actuals), percentage zonder totaal, "'
                   f'&"negatief bedrag, geen aantal of geen start bouw, of fees bij een type dat niet op DAEB staat"))'))
    kos_rng = f"Model!${FIX['kosten']}${ROW1}:${FIX['kosten']}${ROWN}"
    checks.append(("Invoer sluit aan op het FO (kosten en opbrengsten)",
                   f'=IF({dabs("fo_opbr")}="","n.v.t. (geen FO gekoppeld: knop \'Ophalen uit FO\' op tab Invoer)",'
                   f'IF(AND(ABS({hm("opbr_totaal")}-N({dabs("fo_opbr")}))<1,ABS(SUM({kos_rng})-N({dabs("fo_kosten")}))<1),"OK",'
                   f'"LET OP: tab Invoer wijkt af van het FO (opbrengsten "&{eur_mln(hm("opbr_totaal"))}&" tegen "&{eur_mln("N(" + dabs("fo_opbr") + ")")}'
                   f'&", kosten "&{eur_mln("SUM(" + kos_rng + ")")}&" tegen "&{eur_mln("N(" + dabs("fo_kosten") + ")")}&"); haal het FO opnieuw op"))'))
    checks.append(("Koopsomknop: verkooptempo-model op ja",
                   f'=IF(AND({hm("model_aan")}=0,OR({hm("koopsom_down")}<>0,{hm("koopsom_up")}<>0,{hm("koopsom_scn")}<>0)),'
                   f'"LET OP: zet \'Woningtypes en termijnen gebruiken\' op ja, anders doet de koopsomknop niets","OK")'))
    for i, (t, formule) in enumerate(checks):
        label(r + 1 + i, t)
        put(ws, f"W{r + 1 + i}", formule, f=font(9, True, GREEN), al=AL_LEFT)
    ws.conditional_formatting.add(f"W{r + 1}:W{r + len(checks)}",
                                  FormulaRule(formula=[f'LEFT($W{r + 1},6)="LET OP"'], font=Font(name=ARIAL, bold=True, color=RED)))

    r = LY.D_ROW_TYPES
    sectie(r, "WONINGTYPES (tab Woningtypes)")
    put(ws, f"W{r}", "aantal", f=font(8, True, TXT2), al=AL_CENTER, bd=BD_SECTION)
    put(ws, f"X{r}", "start bouw", f=font(8, True, TXT2), al=AL_CENTER, bd=BD_SECTION)
    put(ws, f"Y{r}", "soort", f=font(8, True, TXT2), al=AL_CENTER, bd=BD_SECTION)
    for k in range(1, N_TYPES + 1):
        naam = LY.wt_ref(LY.WT_R_NAAM, k)
        bv = f"Model!{m_col('verv', k)}$6"
        put(ws, f"V{r + k}", f'=IF({naam}="","–",{naam})', f=F_NOTE, al=AL_VCENTER)
        put(ws, f"W{r + k}", f'=IF(Model!{m_col("vcum", k)}$6=0,"",Model!{m_col("vcum", k)}$6)', f=F_CALC9, nf=FMT_INT, al=AL_CENTER)
        put(ws, f"X{r + k}", f'=IF({bv}>=99999,"–","Q"&(MOD({bv}-1,4)+1)&" {APOS}"&RIGHT(INT(({bv}-1)/4),2))', f=F_CALC9, al=AL_CENTER)
        put(ws, f"Y{r + k}", f'=IF(Model!{m_col("vcum", k)}$6=0,"",IF(Model!{m_col("fee", k)}$5=1,"DAEB","koop"))', f=F_CALC9, al=AL_CENTER)

    # ---- FO-koppeling ----------------------------------------------------------------------
    r = LY.D_ROW_FO
    sectie(r, "FO-KOPPELING (knop 'Ophalen uit FO' op tab Invoer)")
    fo = data.fo or {}
    for naam, waarde, nf in (("fo_bestand", fo.get("bestand"), None), ("fo_actuals", fo.get("actuals"), None),
                             ("fo_opbr", fo.get("opbrengsten"), FMT_INT), ("fo_kosten", fo.get("kosten"), FMT_INT), ("fo_datum", fo.get("datum"), None)):
        rij = int(LY.D[naam][1:])
        label(rij, LY.D_LABELS[naam])
        put(ws, LY.D[naam], waarde, f=F_CALC9, nf=nf, al=AL_LEFT)
    put(ws, f"Z{LY.D_ROW_FO + 1}", "de macro vult deze cellen; de controle hierboven vergelijkt tab Invoer met deze totalen", f=F_NOTE8)

    # ---- validaties -------------------------------------------------------------------
    for dv, cells in (
        (DataValidation(type="list", formula1='"ja,nee"', allow_blank=False, error="Kies ja of nee."), ["W12", "W8", LY.D["scn_aan"]]),
        (DataValidation(type="list", formula1='"start bouw,hele looptijd"', allow_blank=False, error="Kies 'start bouw' of 'hele looptijd'."), ["W7"]),
        (DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True, error="Kwartaal 1 t/m 4."), ["X5"]),
        (DataValidation(type="whole", operator="between", formula1="-20", formula2="20", allow_blank=True, error="Heel aantal kwartalen tussen -20 en 20."), ["W17:Y17"]),
        (DataValidation(type="whole", operator="between", formula1="0", formula2="40", allow_blank=True, error="Heel aantal kwartalen tussen 0 en 40."), ["W18:Y18"]),
        (DataValidation(type="decimal", operator="between", formula1="0", formula2="0.5", allow_blank=True, error="Jaarrente tussen 0% en 50%."), ["W6", "W21:Y21"]),
        (DataValidation(type="decimal", operator="between", formula1="-0.5", formula2="0.5", allow_blank=True, error="Koopsom tussen −50% en +50%."), ["W22:Y22"]),
    ):
        dv.showErrorMessage = True
        dv.errorTitle = "Ongeldige invoer"
        ws.add_data_validation(dv)
        for c in cells:
            dv.add(c)

    # ---- uitleg ----------------------------------------------------------------------------
    r = LY.D_ROW_UITLEG
    put(ws, f"B{r}", "Zo werkt het", f=font(11, True, TXT))
    uitleg = [
        ("Werkwijze", "1. tab Invoer: plak kosten en omzet uit je eigen bestand en vul verkoop en transport in · 2. tab Woningtypes: types, koopsom, "
                      "start bouw en termijnen · 3. zet hier rechts de knoppen · 4. tab PowerPoint: blokken naar de dia's"),
        ("Basis", "het basispad is altijd je eigen cashflow; dit bestand rekent je cijfers niet opnieuw uit (behalve rente, als je die ook in de basis zet)"),
        ("Scenario's", "upside en downside = basis + effect van de knoppen: verkoop/transport eerder of later, uitstel start bouw, opbrengsten %, kosten %, rente "
                       "(alleen op prognosekwartalen)"),
        ("Scenariolijn", "de gele lijn in de cashflowgrafiek is een derde scenario met eigen knoppen (kolom Scenario rechts); zet hem aan of uit en geef hem een naam; "
                         "hij staat los van de band in de scenariografiek"),
        ("Verkooptempo-model", "ja: bij verschuiven schuift de grondtermijn mee met het transport en blijven de bouwtermijnen bij de bouw; alleen dat verschil "
                               "komt bovenop je eigen opbrengsten"),
        ("", "nee: verschuiven verplaatst al je prognose-opbrengsten; woningtypes en termijnen zijn dan niet nodig"),
        ("Rente", "rente = negatieve stand van het vorige kwartaal × jaarrente / 4, t/m het kwartaal vóór start bouw van het project (of de hele looptijd). "
                  "'Rente ook in de basis' = nee: de scenario's tellen alleen de extra rente ten opzichte van de basis"),
        ("Uitstel start bouw", "verlengt alleen de renteperiode van dat scenario; kosten en bouwtermijnen blijven staan (zoals in het oude template)"),
        ("Jaarrente per scenario", "elk scenario kan een eigen jaarrente krijgen; leeg = de algemene jaarrente"),
        ("Koopsom", "per scenario de VON-prijs hoger of lager: grondtermijn en bouwtermijnen schalen mee via het verkooptempo-model (alleen met 'ja')"),
        ("Woningtypes", f"maximaal {N_TYPES} types; elk type heeft op tab Woningtypes een eigen blok van drie kolommen met koopsom, start bouw en eigen termijnen, "
                        "en op tab Invoer twee kolommen (verkocht | transport)"),
        ("Bouwtermijnen", "invullen als % van de aanneemsom (samen 100%) of van de koopsom (samen 100% − grondtermijn); het model schaalt ze naar "
                          "100% − grondtermijn, dus grondtermijn + bouwtermijnen = de koopsom"),
        ("Extra opbrengsten", "per type in € per woning (kopersmeerwerk, kadastrale kosten, overige) met een bouwkwartaal, of leeg = bij transport; "
                              "tellen mee in de modelopbrengst, niet in de koopsomknop"),
        ("Bouwtermijnengrafiek", "per bouwtermijn een baan per woningtype dat hem heeft (zelfde naam bij meerdere typen = één termijn); het blokje "
                                 "staat in het kwartaal waarin de termijn vervalt (start bouw + bouwkwartaal − 1) in de typekleur; lichtrood = "
                                 "hetzelfde met de start bouw uit de vorige prognose (tab Woningtypes blok 4); grijs = gerealiseerde kwartalen; "
                                 "bovenaan de jaartallen en de kwartaalnummers"),
        ("DAEB (fees)", "tab Woningtypes blok 5: zet Soort op DAEB; het type telt dan geen koopsom, grondtermijn, bouwtermijnen of extra's, maar fees"),
        ("  invoer fees", "per component (AK fee, bijkomende kosten) een totaal en termijnen als bedrag of % met een kwartaal: '2026 Q4', bouwkwartaal of 'actuals'"),
        ("  in de grafieken", "fees staan in de bouwtermijnengrafiek en in de grafiek per type; de knoppen en de koopsomknop raken ze niet"),
        ("Vorige prognose", "tab Woningtypes blok 4 (start bouw per type) en blok 6 (kwartaal per fee-termijn): lichtrood in de bouwtermijnengrafiek; "
                            "'Ophalen uit FO' bewaart de planning van vóór het ophalen als vorige prognose, 'Vorige prognose uit FO' leest een ouder FO"),
        ("Grafiek per woningtype", "kies een type in de blauwe cel boven de grafiek: grondtermijn, bouwtermijnen en extra's (koop) of fees (DAEB), met de cumulatieve lijn"),
        ("  presentatie", "de macro maakt per woningtype een eigen dia (kopie van dia 8 in het sjabloon)"),
        ("Type toevoegen/verwijderen", "kolommen invoegen of verwijderen op tab Woningtypes (drie) en tab Invoer (twee), of de knoppen op tab Woningtypes (.xlsm); "
                                       "het model leest de blokken op positie, dus alles schuift mee"),
        ("Naar PowerPoint", "knop rechts (alleen in de .xlsm): opent het sjabloon, vult teksten, tabellen en grafieken en bewaart een nieuwe presentatie naast dit bestand"),
        ("Rijen (kwartalen)", "op tab Invoer mag je rijen verwijderen of invoegen; lege rijen tellen niet mee, het model pakt de gevulde rijen op volgorde (maximaal 60)"),
        ("Ander project", "tab Invoer en tab Woningtypes leegmaken en opnieuw vullen; grafieken en KPI's volgen het aantal periodes vanzelf"),
        ("Kringverwijzingen", "geen; elke kolom rekent met de invoer, de typeblokken, kolommen links ervan of de vorige rij"),
    ]
    for i, (kop, tekst) in enumerate(uitleg):
        put(ws, f"B{r + 1 + i}", kop or None, f=F_SECTION)
        put(ws, f"E{r + 1 + i}", tekst, f=F_NOTE)
    return ws


# =============================================================================
#  POWERPOINT
# =============================================================================
def bouw_powerpoint(wb, data):
    ws = wb.create_sheet("PowerPoint")
    ws.sheet_properties.codeName = "shPowerPoint"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B7"
    c_t0 = ws[f"{''.join(ch for ch in LY.PP_TABEL_SC if ch.isalpha())}1"].column
    widths = {"A": 2, "B": 5, "C": 30, "D": 70, "E": 11, "F": 20, "O": 3, "Z": 3, "BK": 3, "BQ": 3, L(c_t0): 30}
    widths[LY.PP_BLOK["bt"]] = 30                       # label 'type · termijn'
    for c in range(7, c_t0):
        widths.setdefault(L(c), 13)
    for c in range(c_t0 + 1, c_t0 + 8):
        widths[L(c)] = 14
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[7].height = 32
    put(ws, "B1", "Gegevens voor PowerPoint", f=F_TITLE)
    put(ws, "B2", "Grafiek bijwerken: blauw kader selecteren → kopiëren → in PowerPoint rechtsklik op de grafiek → Gegevens bewerken → "
                  "klik cel A1 → Plakken speciaal → Waarden. Geen macro nodig.", f=F_NOTE)
    put(ws, "B3", "Sneller: de knop 'Naar PowerPoint' (alleen in de .xlsm) doet dit in één keer: sjabloon openen, teksten, tabellen en grafieken "
                  "vullen, nieuwe presentatie bewaren. Kolom F = de vorm op de dia.", f=F_NOTE)
    put(ws, "B5", "KPI-TEKSTEN VOOR DE DIA'S", f=F_SECTION)
    for col, t in (("B", "Dia"), ("C", "Onderdeel"), ("D", "Tekst voor de dia"), ("E", "Richting"), ("F", "Vorm op de dia")):
        put(ws, f"{col}7", t, f=F_HDR, fl=FL_HDR, al=AL_LEFT_WRAP)

    dal, dal_v, eind, eind_v = hm("dal_basis"), hm("dal_vorig"), hm("eind_basis"), hm("eind_vorig")
    bep, bevp = hm("be_basis_pos"), hm("be_vorig_pos")
    rb = hm("rente_in_basis")
    kpi = [
        (2, "Kopje", f'="KERNCIJFERS · STAND PER "&UPPER({hm("kw_actuals")})', None, "KPI_KICKER"),
        (2, "Titel", "=Dashboard!$B$2", None, "KPI_TITEL"),
        (2, "Ondertitel", f'="Projectresultaat "&{eur_mln(eind)}&" · marge "&IFERROR(FIXED({eind}/{hm("opbr_totaal")}*100,1),"0")&"% · opbrengsten €"'
                          f'&FIXED({hm("opbr_totaal")}/1000000,1)&" mln · kosten €"&FIXED({hm("kosten_totaal")}/1000000,1)&" mln · "&{hm("woningen")}&" woningen"'
                          f'&IF({rb}=1," · incl. rente","")', None, "KPI_SUBTITEL"),
        (2, "Max. financieringsbehoefte", f"={eur_m(dal)}", None, "KPI1_WAARDE"),
        (2, "  toelichting", f'="dieptepunt in "&{hm("dal_basis_kw")}', None, "KPI1_SUB"),
        (2, "  verschil met vorige prognose",
         f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose",IF(ABS(({dal}-{dal_v}))<50000,"gelijk aan vorige prognose",'
         f'"€"&FIXED(ABS(({dal}-{dal_v}))/1000000,1)&"M "&IF(({dal}-{dal_v})<0,"dieper","minder diep")&" dan vorige prognose"))',
         f'=IF(OR({hm("vorig_aanwezig")}=0,ABS(({dal}-{dal_v}))<50000),"gelijk",IF(({dal}-{dal_v})<0,"slechter","beter"))', "KPI1_VERSCHIL"),
        (2, "  vorige prognose", f'=IF({hm("vorig_aanwezig")}=0,"","was "&{eur_m(dal_v)})', None, "KPI1_WAS"),
        (2, "Break-even", f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', None, "KPI2_WAARDE"),
        (2, "  toelichting", "eerste kwartaal met positief saldo", None, "KPI2_SUB"),
        (2, "  verschil met vorige prognose",
         f'=IF(OR({bep}=0,{bevp}=0),"geen vergelijking met vorige prognose",IF(({bep}-{bevp})=0,"gelijk aan vorige prognose",'
         f'ABS(({bep}-{bevp}))&" kw "&IF(({bep}-{bevp})>0,"later","eerder")&" dan vorige prognose"))',
         f'=IF(OR({bep}=0,{bevp}=0,({bep}-{bevp})=0),"gelijk",IF(({bep}-{bevp})>0,"slechter","beter"))', "KPI2_VERSCHIL"),
        (2, "  vorige prognose", f'=IF({bevp}=0,"","was "&{hm("be_vorig")})', None, "KPI2_WAS"),
        (2, "Projectresultaat", f"={eur_m(eind)}", None, "KPI3_WAARDE"),
        (2, "  toelichting", f'="marge "&IFERROR(FIXED({eind}/{hm("opbr_totaal")}*100,1),"0")&"% · eindsaldo in "&{hm("laatste_kw")}', None, "KPI3_SUB"),
        (2, "  verschil met vorige prognose",
         f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose",IF(ABS(({eind}-{eind_v}))<50000,"gelijk aan vorige prognose",'
         f'"€"&FIXED(ABS(({eind}-{eind_v}))/1000000,1)&"M "&IF(({eind}-{eind_v})>0,"hoger","lager")&" dan vorige prognose"))',
         f'=IF(OR({hm("vorig_aanwezig")}=0,ABS(({eind}-{eind_v}))<50000),"gelijk",IF(({eind}-{eind_v})>0,"beter","slechter"))', "KPI3_VERSCHIL"),
        (2, "  vorige prognose", f'=IF({hm("vorig_aanwezig")}=0,"","was "&{eur_m(eind_v)})', None, "KPI3_WAS"),
        (2, "Cumulatieve cashflow · label", f'="CUM. CASHFLOW PER "&UPPER({hm("kw_actuals")})', None, "KPI4_LABEL"),
        (2, "Cumulatieve cashflow", f"={eur_m(hm('stand_nu'))}", None, "KPI4_WAARDE"),
        (2, "  toelichting", f'=IF({hm("vorig_aanwezig")}=0,"geen vorige prognose","vorige prognose "&{eur_m(hm("vorig_nu"))})', None, "KPI4_SUB"),
        (2, "Bandbreedte eindsaldo", f'={eur_m(hm("eind_down"))}&" tot "&{eur_m(hm("eind_up"))}', None, "KPI5_WAARDE"),
        (2, "  toelichting", "downside tot upside", None, "KPI5_SUB"),
        (2, "Verkocht vóór start bouw", f'=FIXED({hm("verk_voor_start_pct")}*100,0)&"%"',
         f'=IF({hm("verk_voor_start_pct")}>={dabs("norm")},"op norm","onder norm")', "KPI6_WAARDE"),
        (2, "  toelichting", f'="norm "&FIXED({dabs("norm")}*100,0)&"% · "&{hm("verk_voor_start")}&" van "&{hm("won_met_start")}&" woningen"', None, "KPI6_SUB"),
        (2, "Woningen getransporteerd", f'={hm("transport_actuals")}&" van "&{hm("woningen")}', None, "KPI7_WAARDE"),
        (2, "  toelichting", f'="t/m "&{hm("kw_actuals")}&" · verkocht "&{hm("verkocht_actuals")}', None, "KPI7_SUB"),
        (3, "Titel", f'="Dieptepunt "&{eur_mln(dal)}&" in "&{hm("dal_basis_kw")}&", eindsaldo "&{eur_mln(eind)}&" in "&{hm("laatste_kw")}', None, "CF_TITEL"),
        (3, "Ondertitel", f'="Stand per "&{hm("kw_actuals")}&": "&{eur_mln(hm("stand_nu"))}&" · balken: opbrengsten en kosten per kwartaal · lijn: cumulatieve cashflow"'
                          f'&IF({rb}=1," · incl. "&{hm("rente_tekst")},"")', None, "CF_SUBTITEL"),
        (3, "Blokje stand · label", f'="STAND "&UPPER({hm("kw_actuals")})', None, "CF_NU_LABEL"),
        (3, "Blokje stand", f"={eur_m(hm('stand_nu'))}", None, "CF_NU_WAARDE"),
        (3, "Blokje dieptepunt", f'={eur_m(dal)}&" · "&{hm("dal_basis_kw")}', None, "CF_DAL_WAARDE"),
        (3, "Blokje eindsaldo", f'={eur_m(eind)}&" · "&{hm("laatste_kw")}', None, "CF_EIND_WAARDE"),
        (3, "Blokje break-even", f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', None, "CF_BE_WAARDE"),
        (4, "Titel", f'="Eindsaldo tussen "&{eur_mln(hm("eind_down"))}&" en "&{eur_mln(hm("eind_up"))}&", basis "&{eur_mln(eind)}', None, "SC_TITEL"),
        (4, "Ondertitel", f'="Basis = eigen cashflow"&IF({rb}=1," + rente","")&" · band = ruimte tussen downside en upside · "'
                          f'&IF({hm("model_aan")}=1,"verschuiving via grondtermijn en bouwtermijnen per woningtype","verschuiving van alle prognose-opbrengsten")'
                          f'&IF({hm("rente")}>0," · "&{hm("rente_tekst")},"")', None, "SC_SUBTITEL"),
        (5, "Titel", f'="Bouwtermijnen per woningtype · start bouw "&{hm("start_tekst")}&" (vroegste type)"', None, "BT_TITEL"),
        (5, "Ondertitel", f'="Typekleur = huidige planning · lichtrood = vorige prognose · grijs = gerealiseerd t/m "&{hm("kw_actuals")}'
                          f'&" · "&{hm("bt_uniek")}&" termijnen, "&{hm("bt_types")}&" woningtypes"', None, "BT_SUBTITEL"),
        (6, "Titel", f'={hm("verkocht_actuals")}&" van "&{hm("woningen")}&" woningen verkocht, "&{hm("transport_actuals")}&" getransporteerd"', None, "VP_TITEL"),
        (6, "Ondertitel", f'="Stand per "&{hm("kw_actuals")}&" · per kwartaal en woningtype · uitverkocht in "&{hm("uitverkocht")}'
                          f'&" · laatste transport in "&{hm("alles_transport")}', None, "VP_SUBTITEL"),
        (7, "Titel", f'="Uitverkocht in "&{hm("uitverkocht")}&", alles getransporteerd in "&{hm("alles_transport")}', None, "VO_TITEL"),
        (7, "Uitverkocht in", f"={hm('uitverkocht')}", None, "VO_STAT1_WAARDE"),
        (7, "Alles getransporteerd in", f"={hm('alles_transport')}", None, "VO_STAT2_WAARDE"),
        (7, "Laatste kwartaal", f"={hm('laatste_kw')}", None, "VO_STAT3_WAARDE"),
        (8, "Titel", f"={hm('pt_titel')}", None, "PT_TITEL"),
        (8, "Ondertitel", f"={hm('pt_subtitel')}", None, "PT_SUBTITEL"),
        # rente: vijfde kleine kaart op dia 2 (sjabloon v10)
        (2, "Rentelasten · label", f'=IF({hm("rente")}=0,"RENTE",IF({rb}=1,"RENTELASTEN (IN BASIS)","RENTELASTEN (BEREKEND)"))', None, "KPI8_LABEL"),
        (2, "Rentelasten", f'=IF({hm("rente")}=0,"geen",{eur_m(hm("rente_basis_tot"))})', None, "KPI8_WAARDE"),
        (2, "  toelichting", f'=IF({hm("rente")}=0,"zet een jaarrente op het Dashboard",FIXED({hm("rente")}*100,1)&"% "&IF({hm("rente_tm_start")}=1,"t/m start bouw","hele looptijd")'
                             f'&" · down "&{eur_m(hm("rente_down_tot"))}&" · up "&{eur_m(hm("rente_up_tot"))})', None, "KPI8_SUB"),
    ]
    for i, (dia, onderdeel, tekst, richting, vorm) in enumerate(kpi):
        r = LY.PP_KPI_ROW1 + i
        put(ws, f"B{r}", dia, f=F_NOTE, bd=BD_ROW)
        put(ws, f"C{r}", onderdeel, f=F_NOTE, bd=BD_ROW)
        put(ws, f"D{r}", tekst, f=F_BOLD9, al=AL_LEFT_TOP, bd=BD_ROW)
        put(ws, f"E{r}", richting, f=F_NOTE, bd=BD_ROW)
        put(ws, f"F{r}", vorm, f=F_NOTE8, bd=BD_ROW)
    r_s = int(LY.PP_CEL_SJABLOON[1:])
    assert LY.PP_KPI_ROW1 + len(kpi) - 1 <= r_s - 2, "KPI-regels lopen tot in de instellingen (de rij erboven moet leeg blijven)"
    put(ws, f"C{r_s}", "Sjabloon (pptx)", f=F_NOTE)
    put(ws, LY.PP_CEL_SJABLOON, data.params.get("sjabloon") or None, f=F_INPUT, fl=FL_INPUT, al=AL_LEFT_TOP)
    put(ws, f"E{r_s}", "leeg = Kwartaal_Template_cashflow_v14.pptx in dezelfde map als dit bestand; anders het volledige pad", f=F_NOTE8)
    put(ws, f"C{r_s + 1}", "Naam nieuwe presentatie", f=F_NOTE)
    put(ws, LY.PP_CEL_NAAM, f'="Cashflow update Q"&{dabs("actuals_kw")}&" "&{dabs("actuals_jaar")}', f=F_CALC9, al=AL_LEFT_TOP)
    put(ws, f"E{r_s + 1}", "de knop zet er datum en tijd achter en bewaart naast dit bestand", f=F_NOTE8)

    # ---- grafiekblokken (kop in rij 7, één rij per periode) ------------------------------
    blokken = [
        (LY.PP_BLOK["cf"], "DIA 3 · CASHFLOW PER KWARTAAL · € mln", "grafiek cashflow · plakken: G7 t/m N, laatste periode (kolom N = scenariolijn; sjabloon v11)",
         ["Kwartaal", "Voorfinanciering", "Positief saldo", "Opbrengsten", "Kosten", "Cumulatieve cashflow", "Vorige prognose", f"={hm('lbl_scenario')}"],
         [FIX["kwartaal"], M["g_voorfinanciering"], M["g_positief_saldo"], M["g_opbrengsten"], M["g_kosten"], M["g_stand"], M["g_vorige"], M["g_scenario"]], "0.0"),
        (LY.PP_BLOK["sc"], "DIA 4 · SCENARIO'S · € mln", "grafiek scenario's · plakken: P7 t/m Y, laatste periode",
         ["Kwartaal", "Band onder", "Bandbreedte", "Downside", "Upside", "Basis", f"={hm('lbl_nu')}", f"={hm('lbl_dal')}", f"={hm('lbl_eind_up')}", f"={hm('lbl_eind_down')}"],
         [FIX["kwartaal"], M["g_band_onder"], M["g_bandbreedte"], M["g_downside"], M["g_upside"], M["g_stand"], M["g_punt_nu"], M["g_punt_dal"], M["g_eind_upside"], M["g_eind_downside"]], "0.0"),
        (LY.PP_BLOK["bt"], "DIA 5 · BOUWTERMIJNEN · tijd in kwartaalindex (jaar × 4 + kwartaal)",
         f"grafiek bouwtermijnen · koprijen jaar/kwartaal, daarna per termijn een baan per type · plakken: AA7 t/m {L(26 + LY.PP_BT_KOL)}, "
         f"laatste rij ({hm('bt_n')} rijen) · '(uit)' = reeks niet in de presentatie",
         ["Termijn"] + [LY.BT_SEG_NAMEN.get(x, f"=Model!${LY.BT[x]}$7") for x in LY.BT_REEKSEN],
         [LY.BT["label"]] + [LY.BT[x] for x in LY.BT_REEKSEN], "0.00"),
        (LY.PP_BLOK["vp"], "DIA 6 · VERKOOP EN TRANSPORT PER WONINGTYPE · aantal woningen", f"twee grafieken (verkocht / getransporteerd) lezen dit blok · plakken: {LY.PP_BLOK['vp']}7 t/m {L(ws[LY.PP_BLOK['vp'] + '1'].column + 25)}, laatste periode",
         ["Kwartaal"] + [f"=Model!${m_col('gv', k)}$7" for k in range(1, N_TYPES + 1)] + [f"=Model!${m_col('gt', k)}$7" for k in range(1, N_TYPES + 1)]
         + ["Verkocht", "Getransporteerd", f"={hm('lbl_uitverkocht')}", f"={hm('lbl_alles_transport')}", "Realisatie"],
         [FIX["kwartaal"]] + [m_col("gv", k) for k in range(1, N_TYPES + 1)] + [m_col("gt", k) for k in range(1, N_TYPES + 1)]
         + [M["g_verkocht"], M["g_getransporteerd"], M["g_punt_uitverkocht"], M["g_punt_alles_transport"], M["g_realisatie"]], FMT_INT),
        (LY.PP_BLOK["vo"], "DIA 7 · VAN VERKOOP NAAR OMZET · cumulatief %", f"grafiek verkoop naar omzet · plakken: {LY.PP_BLOK['vo']}7 t/m {L(ws[LY.PP_BLOK['vo'] + '1'].column + 4)}, laatste periode",
         ["Kwartaal", "Verkocht", "Getransporteerd", "Opbrengsten ontvangen", "Kosten gemaakt"],
         [FIX["kwartaal"], M["g_verkocht_pct"], M["g_getransporteerd_pct"], M["g_opbrengsten_pct"], M["g_kosten_pct"]], FMT_PCT0),
        (LY.PP_BLOK["pt"], "DIA 8 · OPBRENGSTEN PER WONINGTYPE · € mln", f"grafiek per woningtype (het type uit de keuzecel op het Dashboard; de macro maakt per type een "
         f"kopie van dia 8) · plakken: {LY.PP_BLOK['pt']}7 t/m {L(ws[LY.PP_BLOK['pt'] + '1'].column + LY.PP_PT_KOL - 1)}, laatste periode · '(uit)' = reeks niet in de presentatie",
         ["Kwartaal"] + [f'=IF({hm("pt_lbl_" + x)}="","(uit)",{hm("pt_lbl_" + x)})' for x in LY.PT_REEKSEN] + ["Realisatie"],
         [FIX["kwartaal"]] + [LY.PT[x] for x in LY.PT_REEKSEN] + [M["g_realisatie"]], "0.00"),
    ]
    # het blok van de grafiek per type begint bij de eerste getoonde periode (pt_start) en telt pt_n rijen
    offset_van = {LY.PP_BLOK["pt"]: f"{hm('pt_start')}-1"}
    n_ref_van = {LY.PP_BLOK["pt"]: hm("pt_n")}
    blauw = side("medium", PP_BORDER)
    for col0, titel, sub, koppen, bronnen, nf in blokken:
        c0 = ws[f"{col0}1"].column
        put(ws, f"{col0}4", titel, f=F_SECTION)
        put(ws, f"{col0}5", sub, f=F_NOTE8)
        n = len(koppen)
        bt = col0 == LY.PP_BLOK["bt"]                                   # bouwtermijnenblok: BT_N rijen uit de bouwtermijnentabel
        rijen = LY.BT_N if bt else LY.N_PERIODS
        n_ref = hm("bt_n") if bt else n_ref_van.get(col0, hm("n"))
        offset = offset_van.get(col0, "")
        r1, rN = (LY.BT_ROW1, LY.BT_ROWN) if bt else (ROW1, ROWN)
        for j, kop in enumerate(koppen):
            c = L(c0 + j)
            bd = Border(top=blauw, left=blauw if j == 0 else None, right=blauw if j == n - 1 else None)
            put(ws, f"{c}7", kop or None, f=F_HDR, fl=FL_HDR, al=AL_LEFT_WRAP if j == 0 else AL_RIGHT_WRAP, bd=bd)
        for i in range(rijen):
            r = LY.PP_KPI_ROW1 + i
            nr = i + 1
            for j, bron in enumerate(bronnen):
                c = L(c0 + j)
                pos = f"{nr}+{offset}" if offset else f"{nr}"
                if j == 0:
                    formule = f'=IF({nr}>{n_ref},"",INDEX(Model!${bron}${r1}:${bron}${rN},{pos}))'
                else:
                    formule = f'=IF({nr}>{n_ref},NA(),INDEX(Model!${bron}${r1}:${bron}${rN},{pos}))'
                bd = Border(left=blauw if j == 0 else None, right=blauw if j == n - 1 else None, bottom=blauw if i == rijen - 1 else None)
                put(ws, f"{c}{r}", formule, f=F_CALC9, nf=None if j == 0 else nf, al=AL_LEFT_TOP if j == 0 else None, bd=bd)
    ws.conditional_formatting.add(f"G{LY.PP_KPI_ROW1}:{L(c_t0 - 1)}{LY.PP_KPI_ROW1 + LY.BT_N - 1}",
                                  FormulaRule(formula=[f"ISERROR(G{LY.PP_KPI_ROW1})"], font=Font(name=ARIAL, color="D0D3D8")))

    # ---- tabellen (kolom T0 = eerste tabelkolom) --------------------------------------------
    T0 = "".join(ch for ch in LY.PP_TABEL_SC if ch.isalpha())
    T1, T2, T3, T4 = (L(c_t0 + j) for j in (1, 2, 3, 4))
    put(ws, f"{T0}4", "DIA 4 · TABEL SCENARIO'S", f=F_SECTION)
    for col, t in ((T0, None), (T1, "Downside"), (T2, "Basis"), (T3, "Upside")):
        put(ws, f"{col}7", t, f=F_HDR, fl=FL_HDR, al=AL_LEFT_WRAP if col == T0 else AL_RIGHT_WRAP)

    def knop(ref, nf_kw=False):
        if nf_kw:
            return f'=IF({ref}=0,"0 kw",IF({ref}>0,"+","{MINUS}")&ABS({ref})&" kw")'
        return f'=IF({ref}=0,"0%",IF({ref}>0,"+","{MINUS}")&FIXED(ABS({ref})*100,0)&"%")'

    sc = [
        ("Eindsaldo", f"={eur_m(hm('eind_down'))}", f"={eur_m(eind)}", f"={eur_m(hm('eind_up'))}"),
        ("Diepste dal", f"={eur_m(hm('dal_down'))}", f"={eur_m(dal)}", f"={eur_m(hm('dal_up'))}"),
        ("Diepste dal in", f"={hm('dal_down_kw')}", f"={hm('dal_basis_kw')}", f"={hm('dal_up_kw')}"),
        ("Break-even", f'=IF({hm("be_down")}="","niet bereikt",{hm("be_down")})', f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})',
         f'=IF({hm("be_up")}="","niet bereikt",{hm("be_up")})'),
        ("KNOPPEN", None, None, None),
        ("Verschuiving", knop(dabs("shift_down"), True), "–", knop(dabs("shift_up"), True)),
        ("Opbrengsten", knop(dabs("opbr_down")), "–", knop(dabs("opbr_up"))),
        ("Kosten", knop(dabs("kosten_down")), "–", knop(dabs("kosten_up"))),
    ]
    for i, (t, a, b, c) in enumerate(sc):
        r = 8 + i
        put(ws, f"{T0}{r}", t, f=F_NOTE)
        for col, v in ((T1, a), (T2, b), (T3, c)):
            put(ws, f"{col}{r}", v, f=F_CALC9, al=AL_RIGHT)
    # extra knoppen en de scenariolijn: niet in de tabel op dia 4 (die heeft 9 rijen), wel om over te nemen
    r_extra = LY.PP_TABEL_VT_ROW + N_TYPES + 3
    put(ws, f"{T0}{r_extra}", "DIA 4 · EXTRA KNOPPEN EN SCENARIOLIJN (niet in de tabel op de dia)", f=F_SECTION)
    for col, t in ((T0, None), (T1, "Downside"), (T2, "Basis"), (T3, "Upside"), (T4, f"={hm('scn_naam')}")):
        put(ws, f"{col}{r_extra + 1}", t, f=F_HDR, fl=FL_HDR, al=AL_LEFT_WRAP if col == T0 else AL_RIGHT_WRAP)
    dal_scn = f'{eur_m(hm("dal_scn"))}&" · "&{hm("dal_scn_kw")}'
    extra = [
        ("Verschuiving", knop(dabs("shift_down"), True), "–", knop(dabs("shift_up"), True), knop(dabs("shift_scn"), True)),
        ("Uitstel start bouw", knop(dabs("uitstel_down"), True), "–", knop(dabs("uitstel_up"), True), knop(dabs("uitstel_scn"), True)),
        ("Opbrengsten", knop(dabs("opbr_down")), "–", knop(dabs("opbr_up")), knop(dabs("opbr_scn"))),
        ("Kosten", knop(dabs("kosten_down")), "–", knop(dabs("kosten_up")), knop(dabs("kosten_scn"))),
        ("Koopsom (VON)", knop(dabs("koopsom_down")), "–", knop(dabs("koopsom_up")), knop(dabs("koopsom_scn"))),
        ("Jaarrente", f'=FIXED({hm("rente_pct_down")}*100,1)&"%"', f'=FIXED({hm("rente")}*100,1)&"%"', f'=FIXED({hm("rente_pct_up")}*100,1)&"%"',
         f'=FIXED({hm("rente_pct_scn")}*100,1)&"%"'),
        ("Rente", f'={eur_m(hm("rente_down_tot"), 2)}', f'=IF({rb}=1,{eur_m(hm("rente_basis_tot"), 2)},"–")', f'={eur_m(hm("rente_up_tot"), 2)}',
         f'={eur_m(hm("rente_scn_tot"), 2)}'),
        ("Eindsaldo", f"={eur_m(hm('eind_down'))}", f"={eur_m(eind)}", f"={eur_m(hm('eind_up'))}", f"={eur_m(hm('eind_scn'))}"),
        ("Diepste dal", f"={eur_m(hm('dal_down'))}", f"={eur_m(dal)}", f"={eur_m(hm('dal_up'))}", f"={dal_scn}"),
        ("Break-even", f'=IF({hm("be_down")}="","niet bereikt",{hm("be_down")})', f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})',
         f'=IF({hm("be_up")}="","niet bereikt",{hm("be_up")})', f'=IF({hm("be_scn")}="","niet bereikt",{hm("be_scn")})'),
        ("Rente-instelling", f'={hm("rente_tekst")}', None, None, f'=IF({hm("scn_aan")}=1,"lijn aan","lijn uit")'),
    ]
    for i, (t, a, b, c, d_) in enumerate(extra):
        r = r_extra + 2 + i
        put(ws, f"{T0}{r}", t, f=F_NOTE)
        for col, v in ((T1, a), (T2, b), (T3, c), (T4, d_)):
            put(ws, f"{col}{r}", v, f=F_CALC9, al=AL_RIGHT)

    put(ws, f"{T0}{LY.PP_TABEL_VT_ROW - 1}", "DIA 5 · TABEL WONINGTYPES", f=F_SECTION)
    r0 = LY.PP_TABEL_VT_ROW
    for j, t in enumerate(["Type", "Aantal", "Verkocht t/m nu", "Getransporteerd t/m nu", "Start bouw", "Vóór start bouw", "In %", "Norm"]):
        put(ws, f"{L(c_t0 + j)}{r0}", t, f=F_HDR, fl=FL_HDR, al=AL_LEFT_WRAP if j == 0 else AL_RIGHT_WRAP)
    for k in range(1, N_TYPES + 1):
        r = r0 + k
        vc, tc, bv = m_col("vcum", k), m_col("tcum", k), m_col("verv", k)
        aantal = f"N(Model!{vc}$6)"
        leeg = f'IF({aantal}=0,"",'
        start = f"Model!{bv}$6"
        pos = hm("pos_actuals")
        cel = lambda j: f"{L(c_t0 + j)}{r}"  # noqa: E731
        put(ws, cel(0), f'={leeg}{LY.wt_ref(LY.WT_R_NAAM, k)})', f=F_CALC9, al=AL_LEFT_TOP)
        put(ws, cel(1), f'={leeg}{aantal})', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, cel(2), f'={leeg}IF({pos}=0,0,N(INDEX(Model!${vc}${ROW1}:${vc}${ROWN},{pos}))))', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, cel(3), f'={leeg}IF({pos}=0,0,N(INDEX(Model!${tc}${ROW1}:${tc}${ROWN},{pos}))))', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, cel(4), f'={leeg}IF({start}>=99999,"–","Q"&(MOD({start}-1,4)+1)&" {APOS}"&RIGHT(INT(({start}-1)/4),2)))', f=F_CALC9, al=AL_RIGHT)
        put(ws, cel(5), f'={leeg}IF({start}>=99999,"–",Model!{vc}$5))', f=F_CALC9, nf=FMT_INT, al=AL_RIGHT)
        put(ws, cel(6), f'={leeg}IF({start}>=99999,"–",FIXED(Model!{vc}$5/{aantal}*100,0)&"%"))', f=F_CALC9, al=AL_RIGHT)
        put(ws, cel(7), f'={leeg}IF({start}>=99999,"–",FIXED({dabs("norm")}*100,0)&"%"))', f=F_CALC9, al=AL_RIGHT)
    return ws
