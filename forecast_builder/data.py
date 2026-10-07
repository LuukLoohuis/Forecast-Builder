"""Projectgegevens: wat er in de invoercellen komt te staan.

De builder vult het werkboek met een ProjectData. Die komt uit:
  - een bestaand Cashflow_scenario-werkboek (oude indeling met vijf types naast elkaar,
    of de nieuwe indeling met typeblokken), via `lees_werkboek`;
  - of het ingebouwde voorbeeld (`voorbeeld`), zodat een lege build altijd werkt.
"""
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from . import layout as LY

PERIODE_KOLOMMEN = ["jaar", "kw", "kosten", "pct_kosten", "omzet", "pct_omzet", "cf", "cf1000", "cf_vorig"]  # Invoer B..J

STANDAARD_TERMIJNEN = [
    "Start bouw",
    "Na het leggen van de fundering",
    "Casco gereed",
    "Na gereedkomen buitenmetselwerk",
    "Oplevering woning",
]

STANDAARD_EXTRA = ["Kopersmeerwerk 25%", "Kopersmeerwerk 75%", "Kadastrale kosten / rentes", "Overige opbrengsten", "Verschillen", ""]

PARAM_STANDAARD = {
    "actuals_jaar": 2026, "actuals_kw": 3,
    "rente": 0.03, "rente_tm": "start bouw", "rente_basis": "nee",
    "model_aan": "ja",
    "shift_down": 2, "shift_up": -2,
    "uitstel_down": 4, "uitstel_up": 0,
    "opbr_down": -0.03, "opbr_up": 0.02,
    "kosten_down": 0.03, "kosten_up": 0.0,
    # scenariolijn (gele lijn in de cashflowgrafiek): standaard verkoop en transport twee kwartalen eerder
    "shift_scn": -2, "uitstel_scn": 0, "opbr_scn": 0.0, "kosten_scn": 0.0, "scn_aan": "ja", "scn_naam": "Scenario",
    # eigen jaarrente per scenario (None = algemene jaarrente) en koopsom (VON-prijs) in %
    "rente_pct_down": None, "rente_pct_up": None, "rente_pct_scn": None, "koopsom_down": 0.0, "koopsom_up": 0.0, "koopsom_scn": 0.0,
    "norm": 0.7,
    "sjabloon": None,
}


@dataclass
class TypeData:
    naam: str = ""
    aantal: object = None
    koopsom: object = None
    start_jaar: object = None
    start_kw: object = None
    grond_pct: object = None
    termijnen: list = field(default_factory=list)   # [(naam, pct, kw), ...]
    extras: list = field(default_factory=list)      # [(naam, euro per woning, kw of None = bij transport), ...]


@dataclass
class ProjectData:
    periodes: list = field(default_factory=list)    # dicts met PERIODE_KOLOMMEN
    verkocht: list = field(default_factory=list)    # per periode: lijst per type
    transport: list = field(default_factory=list)
    types: list = field(default_factory=list)       # TypeData
    params: dict = field(default_factory=lambda: dict(PARAM_STANDAARD))
    termijn_namen: list = field(default_factory=lambda: list(STANDAARD_TERMIJNEN))
    extra_namen: list = field(default_factory=lambda: list(STANDAARD_EXTRA))

    def type_(self, k):
        """TypeData van blok k (1-based); leeg type als het blok niet gevuld is."""
        return self.types[k - 1] if k - 1 < len(self.types) else TypeData()


def _v(ws, ref):
    v = ws[ref].value
    return None if v == "" else v


def _num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None


# ---------------------------------------------------------------------------
def voorbeeld():
    """Klein fictief project: vier woningtypes, twaalf kwartalen gerealiseerd."""
    p = ProjectData()
    kosten = [250000, 180000, 160000, 220000, 300000, 450000, 900000, 1200000, 2600000, 3100000, 3400000, 2900000,
              4100000, 4600000, 4300000, 3900000, 3200000, 2400000, 1500000, 600000]
    omzet = [0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 0, 3800000, 4900000, 6200000, 6900000, 7400000, 7100000, 5200000, 2600000]
    jaar, kw = 2024, 1
    for i in range(len(kosten)):
        p.periodes.append({"jaar": jaar, "kw": f"Q{kw}", "kosten": kosten[i], "omzet": omzet[i],
                           "pct_kosten": None, "pct_omzet": None, "cf": None, "cf1000": None, "cf_vorig": None})
        kw += 1
        if kw == 5:
            jaar, kw = jaar + 1, 1
    p.types = [
        TypeData("Rijwoning", 40, 385000, 2027, 1, 0.30, [("Start bouw", 0.035, 1), ("Na het leggen van de fundering", 0.07, 2),
                                                           ("Casco gereed", 0.28, 4), ("Na gereedkomen buitenmetselwerk", 0.245, 5), ("Oplevering woning", 0.07, 7)]),
        TypeData("Tweekapper", 20, 495000, 2027, 2, 0.32, [("Start bouw", 0.034, 1), ("Na het leggen van de fundering", 0.068, 2),
                                                            ("Casco gereed", 0.272, 4), ("Na gereedkomen buitenmetselwerk", 0.238, 5), ("Oplevering woning", 0.068, 7)]),
        TypeData("Vrijstaand", 8, 685000, 2027, 3, 0.38, [("Start bouw", 0.031, 1), ("Na het leggen van de fundering", 0.062, 2),
                                                           ("Casco gereed", 0.248, 3), ("Na gereedkomen buitenmetselwerk", 0.217, 4), ("Oplevering woning", 0.062, 6)]),
        TypeData("Appartement", 24, 370000, 2027, 2, 0.22, [("Start bouw", 0.039, 1), ("Na het leggen van de fundering", 0.078, 2),
                                                             ("Casco gereed", 0.312, 5), ("Na gereedkomen buitenmetselwerk", 0.273, 8), ("Oplevering woning", 0.078, 10)]),
    ]
    n = len(p.periodes)
    verkoop = {0: [(8, 10), (4, 5), (2, 4), (6, 4)]}   # type -> (per kwartaal, vanaf periode-index)
    tempo = [(8, 10), (4, 5), (2, 4), (6, 4)]
    cum = [0, 0, 0, 0]
    for i in range(n):
        rij_v, rij_t = [], []
        for t, (per, vanaf) in enumerate(tempo):
            aantal = p.types[t].aantal
            v = min(per, aantal - cum[t]) if i >= vanaf else 0
            cum[t] += v
            rij_v.append(v or None)
        p.verkocht.append(rij_v)
    cum_t = [0, 0, 0, 0]
    for i in range(n):
        rij_t = []
        for t in range(4):
            v = p.verkocht[i - 2][t] if i >= 2 else None
            rij_t.append(v)
        p.transport.append(rij_t)
    p.params["actuals_jaar"], p.params["actuals_kw"] = 2026, 4
    return p


# ---------------------------------------------------------------------------
def lees_werkboek(pad):
    """Leest de invoer uit een bestaand werkboek (oude of nieuwe indeling)."""
    wb = load_workbook(pad, data_only=True, keep_vba=pad.lower().endswith(".xlsm"))
    if "Woningtypes" in wb.sheetnames:
        return _lees_v2(wb)
    return _lees_v1(wb)


def _lees_periodes(ws, p):
    for r in range(LY.ROW1, LY.ROWN + 1):
        rij = {naam: _v(ws, f"{L(2 + i)}{r}") for i, naam in enumerate(PERIODE_KOLOMMEN)}
        if rij["jaar"] is None:
            rij = {naam: None for naam in PERIODE_KOLOMMEN}
        p.periodes.append(rij)


def _lees_v1(wb):
    """Oude indeling: Invoer L:P verkocht, R:V transport, X8:AB12 types, X18:AH27 termijnen; Dashboard W5.. parameters."""
    p = ProjectData()
    ws = wb["Invoer"]
    _lees_periodes(ws, p)
    for r in range(LY.ROW1, LY.ROWN + 1):
        p.verkocht.append([_num(_v(ws, f"{L(12 + t)}{r}")) for t in range(5)])    # L..P
        p.transport.append([_num(_v(ws, f"{L(18 + t)}{r}")) for t in range(5)])   # R..V
    namen = [_v(ws, f"X{r}") for r in range(19, 28)]
    p.termijn_namen = [n for n in namen if n] or list(STANDAARD_TERMIJNEN)
    for t in range(5):
        naam = _v(ws, f"X{8 + t}")
        aantal = _v(ws, f"Y{8 + t}")
        if naam is None and aantal is None:
            continue
        pc, kc = L(25 + 2 * t), L(26 + 2 * t)      # Y/Z, AA/AB, ...
        termijnen = []
        for r in range(19, 28):
            tn, pct, kw = _v(ws, f"X{r}"), _v(ws, f"{pc}{r}"), _v(ws, f"{kc}{r}")
            if tn is None and pct is None and kw is None:
                continue
            termijnen.append((tn, pct, kw))
        p.types.append(TypeData(naam or "", aantal, _v(ws, f"Z{8 + t}"), _v(ws, f"AA{8 + t}"), _v(ws, f"AB{8 + t}"),
                                _v(ws, f"{pc}18"), termijnen))
    d = wb["Dashboard"]
    p.params.update({
        "actuals_jaar": _v(d, "W5"), "actuals_kw": _v(d, "X5"), "model_aan": _v(d, "W8") or "ja",
        "shift_down": _v(d, "W13"), "shift_up": _v(d, "X13"), "opbr_down": _v(d, "W14"), "opbr_up": _v(d, "X14"),
        "kosten_down": _v(d, "W15"), "kosten_up": _v(d, "X15"), "norm": _v(d, "W18"),
    })
    if "PowerPoint" in wb.sheetnames:
        p.params["sjabloon"] = _v(wb["PowerPoint"], LY.PP_CEL_SJABLOON)
    return p


def _lees_v2(wb):
    """Nieuwe indeling: typeblokken op Woningtypes, paren op Invoer, parameters op Dashboard."""
    p = ProjectData()
    ws = wb["Invoer"]
    _lees_periodes(ws, p)
    for r in range(LY.ROW1, LY.ROWN + 1):
        p.verkocht.append([_num(_v(ws, f"{L(LY.in_col(k, 0))}{r}")) for k in range(1, LY.N_TYPES + 1)])
        p.transport.append([_num(_v(ws, f"{L(LY.in_col(k, 1))}{r}")) for k in range(1, LY.N_TYPES + 1)])
    wt = wb["Woningtypes"]
    heeft_extras = isinstance(_v(wt, f"B{LY.WT_R_X_TITEL}"), str) and _v(wt, f"B{LY.WT_R_X_TITEL}").startswith("3 ·")   # blok 3 bestaat sinds v6
    for k in range(1, LY.N_TYPES + 1):
        c0, c1, c2 = (L(LY.wt_col(k, o)) for o in range(3))
        naam = _v(wt, f"{c0}{LY.WT_R_NAAM}")
        aantal = _v(wt, f"{c0}{LY.WT_R_AANTAL}")
        termijnen = []
        for r in range(LY.WT_R_T1, LY.WT_R_TN + 1):
            tn, pct, kw = _v(wt, f"{c0}{r}"), _v(wt, f"{c1}{r}"), _v(wt, f"{c2}{r}")
            if tn is None and pct is None and kw is None:
                continue
            termijnen.append((tn, pct, kw))
        td = TypeData(naam or "", aantal, _v(wt, f"{c0}{LY.WT_R_KOOPSOM}"), _v(wt, f"{c0}{LY.WT_R_STARTJAAR}"),
                      _v(wt, f"{c0}{LY.WT_R_STARTKW}"), _v(wt, f"{c1}{LY.WT_R_GROND}"), termijnen)
        if heeft_extras:
            for r in range(LY.WT_R_X1, LY.WT_R_XN + 1):
                xn, eur, kw = _v(wt, f"{c0}{r}"), _v(wt, f"{c1}{r}"), _v(wt, f"{c2}{r}")
                if xn is None and eur is None and kw is None:
                    continue
                td.extras.append((xn, eur, kw))
        p.types.append(td)
    # lege blokken aan het eind weglaten
    while p.types and not p.types[-1].naam and p.types[-1].aantal is None:
        p.types.pop()
    namen = [t.termijnen[i][0] for t in p.types for i in range(len(t.termijnen)) if t.termijnen[i][0]]
    if namen:
        p.termijn_namen = list(dict.fromkeys(namen))[:LY.N_TERMIJNEN]
    xnamen = [x[0] for t in p.types for x in t.extras if x[0]]
    if xnamen:
        p.extra_namen = list(dict.fromkeys(xnamen))[:LY.N_EXTRA]
    d = wb["Dashboard"]
    # knoppen worden op het label in kolom V gezocht (de rijnummers verschilden per versie); de kolom (W/X/Y) komt uit LY.D.
    # Een knop die het bestand nog niet heeft (oudere versie), houdt de standaardwaarde.
    rijen = {}
    for r in range(1, LY.D_ROW_CONTROLES + 1):
        t = _v(d, f"V{r}")
        if isinstance(t, str) and t.strip() and t.strip() not in rijen:
            rijen[t.strip()] = r
    scn_kolom = _v(d, f"Y{LY.D_ROW_SCENARIO + 1}") == "Scenario"       # oudere v2-bestanden hebben de kolom Scenario nog niet
    paren = {"shift_up": "shift_down", "uitstel_up": "uitstel_down", "opbr_up": "opbr_down", "kosten_up": "kosten_down",
             "rente_pct_up": "rente_pct_down", "koopsom_up": "koopsom_down", "shift_scn": "shift_down", "uitstel_scn": "uitstel_down",
             "opbr_scn": "opbr_down", "kosten_scn": "kosten_down", "rente_pct_scn": "rente_pct_down", "koopsom_scn": "koopsom_down",
             "actuals_kw": "actuals_jaar"}
    for naam, cel in LY.D.items():
        if naam in ("start_project", "dekking", "rente_down", "rente_up", "rente_scn"):
            continue
        if naam.endswith("_scn") and not scn_kolom or naam in ("scn_aan", "scn_naam") and not scn_kolom:
            continue
        label = LY.D_LABELS.get(paren.get(naam, naam))
        rij = rijen.get(label)
        if rij is None:
            continue
        kol = "".join(ch for ch in cel if ch.isalpha())
        p.params[naam] = _v(d, f"{kol}{rij}")
    if "PowerPoint" in wb.sheetnames:
        p.params["sjabloon"] = _v(wb["PowerPoint"], LY.PP_CEL_SJABLOON)
    return p
