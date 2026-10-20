"""Scenario's (sinds v18): start bouw, verkoop en transport per scenario.

Tab 'Scenario's': per scenario (drie naast de Basis) vul je alleen in wat anders is dan de planning: start bouw, verkoop vanaf en
(eventueel) transport vanaf, als kwartaal, en eventueel de grondtermijn als % van de koopsom. Leeg = zoals de planning; een lege
kolom staat uit.

Rekenregels (Model, gebied vanaf layout.VS_COL1; per scenario, per woningtype):
  * verkoop: het verkooptempo uit de planning (tab Invoer) schuift als geheel naar 'verkoop vanaf' (d_v kwartalen); wat al verkocht
    is (t/m actuals) blijft staan;
  * transport: de geplande transportcurve schuift d_t kwartalen ('transport vanaf' min het eerste geplande transport; leeg = mee
    met de start bouw) en is nooit groter dan wat in het scenario verkocht is: een woning gaat niet vóór de verkoop in transport;
  * bouwtermijnen en extra's per bouwkwartaal vervallen vanaf de start bouw van het scenario (d_sb kwartalen verschoven); bij
    transport betaalt de koper de grondtermijn plus alle vervallen termijnen (zoals in het FO, tab 'CF - opbrengsten');
  * grondtermijn: een ander % (bv. 50%) geldt voor alle niet-DAEB-typen, voor de woningen die na de actuals in transport gaan
    (wat al getransporteerd is, is betaald zoals gepland); de bouwtermijnen gaan naar rato omlaag tot samen weer 100%: de koopsom
    blijft gelijk, alleen komt er meer bij transport binnen en minder in de bouwkwartalen. Per woning en kwartaal is het verschil
    koopsom × (grond nieuw − grond gepland) × (1 − vervallen bouwtermijnen / (1 − grond gepland)): bij transport vóór de start
    bouw het hele verschil, aan het eind 0;
  * fees (DAEB) schuiven mee met de start bouw (alleen de prognose);
  * kosten: met posten op tab CF kosten (kosten.py) schuiven alleen de prognosekosten van de posten op 'start bouw' d_sb kwartalen
    mee en die op 'verkoop' d_v kwartalen; zonder posten blijven de kosten gelijk zolang de start bouw gelijk blijft, en schuiven
    anders alle kosten vanaf de geplande start bouw d_sb kwartalen mee (de kosten daarvóór blijven staan);
  * opbrengsten = je eigen cashflow + het verschil volgens de termijnen (zoals de scenariolijn); rente zoals in de basis;
  * de laatste periode krijgt alles wat er nog niet is (zoals bij de knoppen op het Dashboard).
Een scenario zonder wijziging is precies de Basis.
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from . import layout as LY
from .layout import FIX, H, M, N_TYPES, N_VS, ROW1, ROWN, dabs, h, hm, m_col, m_range_abs, sel_rng, tm_rng, vs_col, vs_rng
from .sheets import APOS, MINUS, _comment, eur_m, eur_mln
from .styles import *  # noqa: F401,F403

NR, BRON, KW, KWT, KWV, IDX, FASE, KOS, OPB, STAND, CUM = (FIX[k] for k in
    ("nr", "bron", "kw", "kwartaal", "kwartaal_vol", "idx", "fase", "kosten", "opbr", "stand", "cum_opbr"))


def _rng(col):
    return f"${col}${ROW1}:${col}${ROWN}"


def kw_lbl(x):
    """Formulefragment: kwartaalindex -> 'Q2 ’28' ('–' bij 99999)."""
    return f'IF({x}>=99999,"–","Q"&(MOD({x}-1,4)+1)&" {APOS}"&RIGHT(INT(({x}-1)/4),2))'


def kw_tekst(x):
    """Formulefragment: aantal kwartalen -> '1 kwartaal', '3 kwartalen', '1 jaar', '2 jaar'."""
    return f'IF(MOD({x},4)=0,({x})/4&" jaar",({x})&IF(({x})=1," kwartaal"," kwartalen"))'


def kw_index(ref):
    """Formulefragment: kwartaal in een invoercel -> kwartaalindex (jaar × 4 + kwartaal); 0 = leeg of onbekend.
    Eerst de labels uit de keuzelijst ('Q2 ’28', kolom Kwartaal volledig van het Model), anders tekst als 'Q2 2028', '2028 Q2',
    'Q2-28' of 'Q2’28'."""
    t = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(UPPER(TRIM({ref}&"")),"{APOS}",""),"\'","")," ",""),"-","")'
    q = f'VALUE(IF(LEFT({t},1)="Q",MID({t},2,1),RIGHT({t},1)))'
    y0 = f'VALUE(IF(LEFT({t},1)="Q",MID({t},3,4),LEFT({t},LEN({t})-2)))'
    y = f'IF({y0}<100,2000+{y0},{y0})'
    tekst = f'IFERROR(IF(AND({q}>=1,{q}<=4,{y}>=1990,{y}<=2100),{y}*4+{q},0),0)'
    return (f'IF(TRIM({ref}&"")="",0,IFERROR(INDEX({_rng(IDX)},MATCH(TRIM({ref}&""),{_rng(KWV)},0)),{tekst}))')


# =============================================================================
#  MODEL
# =============================================================================
def bouw_model_scenarios(ws):
    """Kolommen voor de scenario's in het Model (vanaf VS_COL1), de selectie voor de tabel per termijn, de termijnmatrix en de
    hulpcellen vs_* (kolom BY, rijen vanaf 145)."""
    G, SEL, TM = LY.VS_G, LY.SEL, LY.TM
    idx_rng = _rng(IDX)
    n, pa, ia = h("n"), h("pos_actuals"), h("idx_actuals")
    nod = f"(1-{m_range_abs('fee', 5)})"
    koopsom, grond, x_transport = m_range_abs("tcum", 6), m_range_abs("tup", 6), m_range_abs("vx", 6)
    fee_blk = f"${m_col('fee', 1)}${ROW1}:${m_col('fee', N_TYPES)}${ROWN}"
    rb = h("rente_in_basis")
    sb_basis = M["stand_basis"]
    sel, tsel = h("vs_sel"), h("vs_typesel")
    typesel = tm_rng(LY.TM_R_SEL)

    # ---- koppen en breedtes ---------------------------------------------------------------------------------------------
    put(ws, f"{G['kpre']}4", "scenario's (tab Scenario's): gedeeld: kosten vóór / vanaf de geplande start bouw (cumulatief), verkocht en "
                             "getransporteerd niet-DAEB (cumulatief); daarna per scenario de rekenkolommen en vier typeblokken", f=F_NOTE8)
    koppen = {G["kpre"]: "Kosten vóór start bouw cum", G["kpost"]: "Kosten vanaf start bouw cum", G["vnd"]: "Verkocht cum niet-DAEB",
              G["tnd"]: "Transport cum niet-DAEB"}
    namen = {"pv": "Positie verkoop", "pt": "Positie transport", "psb": "Positie start bouw", "qf": "Positie fees", "model": "Model",
             "cumo": "Cum opbrengsten", "opbr": "Opbrengsten", "kcum": "Kosten cum", "kosten": "Kosten", "rente": "Rente",
             "sraw": "Stand (berekend)", "stand": "Stand (getoond)", "pos": "Positief", "tnd": "Transport cum niet-DAEB",
             "g": "Grafiek (€ mln)", "gdal": "Dieptepunt (€ mln)"}
    bloknamen = {"vc": "Verkocht cum", "tc": "Transport cum", "vv": "Termijnen vervallen", "vx": "Extra's vervallen (€/won)"}
    for s in range(1, N_VS + 1):
        put(ws, f"{vs_col(s, 'pv')}4", f"scenario {s}: rij 5/6 per type: verkocht bij start bouw / verkocht t/m actuals (vc), verkocht "
                                       f"totaal / getransporteerd t/m actuals (tc), start bouw (vv)", f=font(8, True, TXT2))
        for naam, t in namen.items():
            koppen[vs_col(s, naam)] = f"S{s} {t}"
        for blok, t in bloknamen.items():
            for k in range(1, N_TYPES + 1):
                koppen[vs_col(s, blok, k)] = f"S{s} {t} {k}"
    put(ws, f"{SEL[('ts', 1)]}4", "tabel per termijn (tab Scenario's, keuze scenario en woningtype): per type transport (rij 6 = start bouw), "
                                  "extra's en fees van het gekozen scenario; cumulatief per onderdeel; inhaal per termijn", f=F_NOTE8)
    for blok, t in (("ts", "Transport cum"), ("xs", "Extra's vervallen"), ("fs", "Fees cum")):
        for k in range(1, N_TYPES + 1):
            koppen[SEL[(blok, k)]] = f"Keuze {t} {k}"
    for naam, t in (("vj", "Verkocht cum"), ("trj", "Getransporteerd cum"), ("gj", "Grondtermijn cum"), ("xj", "Extra's cum"),
                    ("fj", "Fees cum"), ("osel", "Opbrengsten"), ("ksel", "Kosten"), ("ssel", "Stand")):
        koppen[SEL[naam]] = f"Keuze {t}"
    for j in range(1, LY.N_TERMIJNEN + 1):
        koppen[SEL[f"tj{j}"]] = f"Keuze termijn {j} cum"
        koppen[SEL[f"ih{j}"]] = f"Keuze inhaal termijn {j}"
    for col, t in koppen.items():
        put(ws, f"{col}7", t, f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
        ws.column_dimensions[col].width = 11
    ws.column_dimensions[L(LY.VS_COL1 - 1)].width = 3

    # ---- rij 5 en 6 per type ----------------------------------------------------------------------------------------------
    for s in range(1, N_VS + 1):
        d_sb = h(f"vs{s}_d_sb")
        for k in range(1, N_TYPES + 1):
            vc_k, tc_k, bv_k, fe_k = m_col("vcum", k), m_col("tcum", k), m_col("verv", k), m_col("fee", k)
            vcS, tcS, vvS = vs_col(s, "vc", k), vs_col(s, "tc", k), vs_col(s, "vv", k)
            pos_sb = f'COUNTIF({idx_rng},"<="&{vvS}$6)'
            put(ws, f"{vcS}5", f'=IF(OR({vvS}$6>=99999,{fe_k}$5=1,{pos_sb}=0),0,N(INDEX({_rng(vcS)},{pos_sb})))', f=F_CALC, nf=FMT_INT)
            put(ws, f"{vcS}6", f'=IF({pa}=0,0,N(INDEX({_rng(vc_k)},{pa})))', f=F_CALC, nf=FMT_INT)
            put(ws, f"{tcS}5", f'=IF({n}=0,0,N(INDEX({_rng(vc_k)},{n})))', f=F_CALC, nf=FMT_INT)
            put(ws, f"{tcS}6", f'=IF({pa}=0,0,N(INDEX({_rng(tc_k)},{pa})))', f=F_CALC, nf=FMT_INT)
            put(ws, f"{vvS}6", f'=IF({bv_k}$6>=99999,99999,{bv_k}$6+{d_sb})', f=F_CALC, nf="0")
            # rij 5: grondtermijn scenario min gepland (fractie; 0 = zoals gepland, DAEB of geen bouwtermijnen)
            gr, tu_k = h(f"vs{s}_gr"), m_col("tup", k)
            put(ws, f"{vvS}5", f'=IF(OR({gr}<0,N({fe_k}$5)=1,N({tu_k}$6)>=1,N({tu_k}$5)=0),0,{gr}-N({tu_k}$6))', f=F_CALC, nf=FMT_PCT)
    for k in range(1, N_TYPES + 1):
        args = ",".join([f"{m_col('verv', k)}$6"] + [f"{vs_col(s, 'vv', k)}$6" for s in range(1, N_VS + 1)])
        put(ws, f"{SEL[('ts', k)]}6", f"=CHOOSE({sel}+1,{args})", f=F_CALC, nf="0")
        # grondtermijn van het gekozen scenario: rij 5 van xs = verschil grond (fractie), van fs = factor op de bouwtermijnen − 1
        dgs = ",".join(["0"] + [f"{vs_col(s, 'vv', k)}$5" for s in range(1, N_VS + 1)])
        put(ws, f"{SEL[('xs', k)]}5", f"=CHOOSE({sel}+1,{dgs})", f=F_CALC, nf=FMT_PCT)
        tu_k = m_col("tup", k)
        put(ws, f"{SEL[('fs', k)]}5", f"=IF(N({tu_k}$6)>=1,0,-{SEL[('xs', k)]}5/(1-N({tu_k}$6)))", f=F_CALC, nf=FMT_PCT)

    # ---- periode-rijen ----------------------------------------------------------------------------------------------------
    for r in range(ROW1, ROWN + 1):
        p = r - 1
        leeg = f'${NR}{r}=""'
        f = {}
        prog = f'OR({FASE}{r}="Realisatie",${NR}{r}={n})'          # waar: realisatie of laatste periode (dan de planning zelf)
        kp = G["kpre"]
        f[kp] = f'=IF({leeg},"",N({kp}{p})+IF({IDX}{r}<{h("start_idx")},{KOS}{r},0))'
        f[G["kpost"]] = f'=IF({leeg},"",({CUM}{r}-{STAND}{r})-{kp}{r})'
        f[G["vnd"]] = f'=IF({leeg},"",SUMPRODUCT({LY.m_range("vcum", r)},{nod}))'
        f[G["tnd"]] = f'=IF({leeg},"",SUMPRODUCT({LY.m_range("tcum", r)},{nod}))'
        for s in range(1, N_VS + 1):
            c = lambda naam: vs_col(s, naam)  # noqa: E731
            d_v, d_t, d_sb = h(f"vs{s}_d_v"), h(f"vs{s}_d_t"), h(f"vs{s}_d_sb")
            f[c("pv")] = f'=IF({leeg},"",COUNTIF({idx_rng},"<="&({IDX}{r}-{d_v})))'
            f[c("pt")] = f'=IF({leeg},"",COUNTIF({idx_rng},"<="&({IDX}{r}-{d_t})))'
            f[c("psb")] = f'=IF({leeg},"",COUNTIF({idx_rng},"<="&({IDX}{r}-{d_sb})))'
            f[c("qf")] = f'=IF({leeg},"",IF({prog},${NR}{r},MAX({c("psb")}{r},{pa})))'
            for k in range(1, N_TYPES + 1):
                vc_k, tc_k, vx_k, tu_k = m_col("vcum", k), m_col("tcum", k), m_col("vx", k), m_col("tup", k)
                vcS, tcS, vvS, vxS = (vs_col(s, b, k) for b in LY.VS_TB)
                f[vcS] = (f'=IF({leeg},"",IF({prog},N({vc_k}{r}),MAX({vcS}$6,IF({c("pv")}{r}=0,0,'
                          f'N(INDEX({_rng(vc_k)},{c("pv")}{r}))))))')
                schuif = f'IF({c("pt")}{r}=0,0,N(INDEX({_rng(tc_k)},{c("pt")}{r})))'
                # transport: de verschoven planning, nooit meer dan in het scenario verkocht (types zonder geplande verkoop: geen grens)
                f[tcS] = (f'=IF({leeg},"",IF({prog},N({tc_k}{r}),MAX({tcS}$6,IF({tcS}$5=0,{schuif},MIN({schuif},{vcS}{r})))))')
                kw_rng = LY.wt_ref(0, k, 2, rows=(LY.WT_R_T1, LY.WT_R_TN))
                pct_rng = LY.wt_ref(0, k, 1, rows=(LY.WT_R_T1, LY.WT_R_TN))
                # laatste periode: alles wat volgens de planning vervalt (een latere start bouw schuift niets buiten de looptijd)
                f[vvS] = (f'=IF({leeg},"",IF(${NR}{r}={n},N({m_col("verv", k)}{r}),IF({IDX}{r}<{vvS}$6,0,'
                          f'SUMIF({kw_rng},"<="&({IDX}{r}-{vvS}$6+1),{pct_rng})*{tu_k}$5)))')
                x_kw = LY.wt_ref(0, k, 2, rows=(LY.WT_R_X1, LY.WT_R_XN))
                x_eur = LY.wt_ref(0, k, 1, rows=(LY.WT_R_X1, LY.WT_R_XN))
                f[vxS] = (f'=IF({leeg},"",IF(${NR}{r}={n},N({vx_k}{r}),IF({IDX}{r}<{vvS}$6,0,'
                          f'SUMIFS({x_eur},{x_kw},">=1",{x_kw},"<="&({IDX}{r}-{vvS}$6+1)))))')
            tc_r, vv_r, vx_r = vs_rng(s, "tc", r), vs_rng(s, "vv", r), vs_rng(s, "vx", r)
            # grondtermijn van het scenario: alleen voor woningen die na de actuals in transport gaan (tc boven tc t/m actuals)
            na_act = f'(({tc_r}>{vs_rng(s, "tc", 6, abs_=True)})*({tc_r}-{vs_rng(s, "tc", 6, abs_=True)}))'
            gr_extra = (f'SUMPRODUCT({na_act}*{koopsom}*{vs_rng(s, "vv", 5, abs_=True)}'
                        f'*(1-{vv_r}/(1-{grond}+({grond}>=1)))*{nod})')
            f[c("model")] = (f'=IF({leeg},"",SUMPRODUCT({tc_r}*({koopsom}*({grond}+{vv_r})+{x_transport}+{vx_r})*{nod})'
                             f'+IF({h(f"vs{s}_gr")}<0,0,{gr_extra})'
                             f'+IF({c("qf")}{r}=0,0,SUM(INDEX({fee_blk},{c("qf")}{r},0))))')
            f[c("cumo")] = (f'=IF({leeg},"",IF({h("model_aan")}=1,{CUM}{r}+{c("model")}{r}-{M["model_basis"]}{r},'
                            f'IF({prog},{CUM}{r},MAX({h("cum_opbr_actuals")},IF({c("pt")}{r}=0,0,N(INDEX({_rng(CUM)},{c("pt")}{r})))))))')
            f[c("opbr")] = f'=IF({leeg},"",IF({FASE}{r}="Realisatie",{OPB}{r},{c("cumo")}{r}-N({c("cumo")}{p})))'
            # kosten: met posten op tab CF kosten schuiven alleen die op 'start bouw' (d_sb) en 'verkoop' (d_v) mee: basis - eigen
            # cumulatieve curve + verschoven curve (psb / pv = positie van het verschoven kwartaal; vóór de eerste periode 0, na de
            # laatste de laatste); zonder posten schuiven alle kosten vanaf de geplande start bouw mee
            per_post = (f'{CUM}{r}-{STAND}{r}-N({G["cksb"]}{r})+IF({c("psb")}{r}=0,0,N(INDEX({_rng(G["cksb"])},{c("psb")}{r})))'
                        f'-N({G["ckv"]}{r})+IF({c("pv")}{r}=0,0,N(INDEX({_rng(G["ckv"])},{c("pv")}{r})))')
            alles = f'{kp}{r}+IF({c("psb")}{r}=0,0,N(INDEX({_rng(G["kpost"])},{c("psb")}{r})))'
            f[c("kcum")] = (f'=IF({leeg},"",IF({prog},{CUM}{r}-{STAND}{r},IF({h("kf_aan")}=1,'
                            f'IF(AND({d_sb}=0,{d_v}=0),{CUM}{r}-{STAND}{r},{per_post}),IF({d_sb}=0,{CUM}{r}-{STAND}{r},{alles}))))')
            f[c("kosten")] = f'=IF({leeg},"",{c("kcum")}{r}-N({c("kcum")}{p}))'
            eind = f'IF({h("rente_tm_start")}=1,{h(f"vs{s}_sb")},99999)'
            f[c("rente")] = (f'=IF({leeg},"",IF(AND({FASE}{r}="Prognose",{h("rente")}>0,{IDX}{r}<{eind}),'
                             f'MAX(0,-N({c("sraw")}{p}))*{h("rente")}/4,0))')
            f[c("sraw")] = f'=IF({leeg},"",N({c("sraw")}{p})+{c("opbr")}{r}-{c("kosten")}{r}-{c("rente")}{r})'
            f[c("stand")] = f'=IF({leeg},"",IF({rb}=1,{c("sraw")}{r},{c("sraw")}{r}+({STAND}{r}-{M["stand_basis_rente"]}{r})))'
            f[c("pos")] = f'=IF(ISNUMBER({c("stand")}{r}),IF({c("stand")}{r}>0,1,0),0)'
            f[c("tnd")] = f'=IF({leeg},"",SUMPRODUCT({tc_r},{nod}))'
            st = c("stand")
            # grafiek: alleen waar het scenario van de Basis afwijkt (plus de punten ernaast), zodat de lijn daar begint en eindigt
            anders = (f'OR(ABS({st}{r}-{sb_basis}{r})>1,ABS(N({st}{p})-N({sb_basis}{p}))>1,'
                      f'ABS(N({st}{r + 1})-N({sb_basis}{r + 1}))>1)')
            f[c("g")] = f'=IF({leeg},NA(),IF({h(f"vs{s}_aan")}=0,NA(),IF({anders},{st}{r}/1000000,NA())))'
            f[c("gcf")] = f'=IF({h(f"vs{s}_cf")}=1,{c("g")}{r},NA())'              # cashflowgrafiek: alleen met de schakelaar op ja
            f[c("gdal")] = (f'=IF({leeg},NA(),IF(AND({h(f"vs{s}_aan")}=1,${NR}{r}={h(f"vs{s}_dal_pos")},'
                            f'OR(ABS({h(f"vs{s}_dal")}-{h("dal_basis")})>1,{h(f"vs{s}_dal_pos")}<>{h("vs0_dal_pos")})),{st}{r}/1000000,NA()))')
        # ---- selectie (tabel per termijn) ----
        for k in range(1, N_TYPES + 1):
            tcs = ",".join([f"N({m_col('tcum', k)}{r})"] + [f"N({vs_col(s, 'tc', k)}{r})" for s in range(1, N_VS + 1)])
            vxs = ",".join([f"N({m_col('vx', k)}{r})"] + [f"N({vs_col(s, 'vx', k)}{r})" for s in range(1, N_VS + 1)])
            fe_k = m_col("fee", k)
            fes = ",".join([f"N({fe_k}{r})"] + [f'IF({vs_col(s, "qf")}{r}=0,0,N(INDEX({_rng(fe_k)},{vs_col(s, "qf")}{r})))'
                                                for s in range(1, N_VS + 1)])
            f[SEL[("ts", k)]] = f'=IF({leeg},"",CHOOSE({sel}+1,{tcs}))'
            f[SEL[("xs", k)]] = f'=IF({leeg},"",CHOOSE({sel}+1,{vxs}))'
            f[SEL[("fs", k)]] = f'=IF({leeg},"",CHOOSE({sel}+1,{fes}))'
        ts_r, ts_p = sel_rng("ts", r), sel_rng("ts", p)
        start = sel_rng("ts", 6, abs_=True)
        keuze_vc = "+".join([f"({sel}=0)*{LY.m_range('vcum', r)}"] + [f"({sel}={s})*{vs_rng(s, 'vc', r)}" for s in range(1, N_VS + 1)])
        f[SEL["vj"]] = f'=IF({leeg},"",SUMPRODUCT({typesel}*({keuze_vc})))'
        f[SEL["trj"]] = f'=IF({leeg},"",SUMPRODUCT({typesel},{ts_r}))'
        tc_act = vs_rng(1, "tc", 6, abs_=True)                     # getransporteerd t/m actuals (gelijk in elk scenario)
        over_r = f"(({ts_r}>{tc_act})*({ts_r}-{tc_act}))"
        over_p = f"(({ts_p}>{tc_act})*({ts_p}-{tc_act}))" if r > ROW1 else "0"
        dg_sel, f_sel = sel_rng("xs", 5, abs_=True), sel_rng("fs", 5, abs_=True)
        f[SEL["gj"]] = (f'=IF({leeg},"",SUMPRODUCT({ts_r},{tm_rng(LY.TM_R_AG)},{typesel})'
                        f'+SUMPRODUCT({over_r}*{koopsom}*{dg_sel}*{typesel}))')
        f[SEL["xj"]] = f'=IF({leeg},"",SUMPRODUCT({ts_r}*({x_transport}+{sel_rng("xs", r)})*{nod}*{typesel}))'
        f[SEL["fj"]] = f'=IF({leeg},"",SUMPRODUCT({sel_rng("fs", r)},{typesel}))'
        kies = lambda basis, naam: ",".join([f"{basis}{r}"] + [f"{vs_col(s, naam)}{r}" for s in range(1, N_VS + 1)])  # noqa: E731
        f[SEL["osel"]] = f'=IF({leeg},"",CHOOSE({sel}+1,{kies(OPB, "opbr")}))'
        f[SEL["ksel"]] = f'=IF({leeg},"",CHOOSE({sel}+1,{kies(KOS, "kosten")}))'
        f[SEL["ssel"]] = f'=IF({leeg},"",CHOOSE({sel}+1,{kies(sb_basis, "stand")}))'
        for j in range(1, LY.N_TERMIJNEN + 1):
            a, kw = tm_rng(LY.TM_R_A + j - 1), tm_rng(LY.TM_R_KW + j - 1)
            vervallen = f"(({IDX}{r}-{start}-{kw}+1)>=0)*({kw}>=1)"
            f[SEL[f"tj{j}"]] = f'=IF({leeg},"",SUMPRODUCT(({ts_r}+{over_r}*{f_sel})*{a}*{typesel}*{vervallen}))'
            erbij = (f"({ts_r}-{ts_p}+({over_r}-{over_p})*{f_sel})" if r > ROW1 else f"({ts_r}+{over_r}*{f_sel})")
            al_vervallen = f"(({IDX}{r}-{start}-{kw}+1)>0)*({kw}>=1)"
            f[SEL[f"ih{j}"]] = f'=IF({leeg},"",SUMPRODUCT({erbij}*{a}*{typesel}*{al_vervallen}))'
        for col, formule in f.items():
            put(ws, f"{col}{r}", formule, f=F_CALC, nf=FMT_INT)
        for s in range(1, N_VS + 1):
            for k in range(1, N_TYPES + 1):
                ws[f"{vs_col(s, 'vv', k)}{r}"].number_format = FMT_PCT
            for naam in ("g", "gdal"):
                ws[f"{vs_col(s, naam)}{r}"].number_format = "0.00"

    # ---- termijnmatrix: per termijn j en type k naam, bedrag per woning (€, niet-DAEB) en bouwkwartaal ------------------------
    put(ws, f"{TM[1]}{LY.TM_R_NAAM - 1}", "termijnmatrix (tabel per termijn): naam, € per woning en bouwkwartaal per termijn (rij) en type (kolom); "
                                          f"rij {LY.TM_R_SEL} = gekozen typen, rij {LY.TM_R_AG} = grondtermijn per woning", f=F_NOTE8)
    for k in range(1, N_TYPES + 1):
        put(ws, f"{TM[k]}7", f"Termijnen type {k}", f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
        ws.column_dimensions[TM[k]].width = 11
    put(ws, f"{TM['label']}7", "Rijnaam", f=F_HDR8, fl=FL_HDR, al=AL_RIGHT_WRAP)
    ws.column_dimensions[TM["label"]].width = 30
    for j in range(1, LY.N_TERMIJNEN + 1):
        rn, ra, rk = LY.TM_R_NAAM + j - 1, LY.TM_R_A + j - 1, LY.TM_R_KW + j - 1
        for k in range(1, N_TYPES + 1):
            aantal = f"N({LY.wt_ref(LY.WT_R_AANTAL, k)})"
            naam_jk, pct_jk, kw_jk = (LY.wt_ref(LY.WT_R_T1 + j - 1, k, o) for o in range(3))
            put(ws, f"{TM[k]}{rn}", f'=IF({aantal}=0,"",{naam_jk}&"")', f=F_CALC)
            put(ws, f"{TM[k]}{ra}", f'=IF(N({m_col("fee", k)}$5)=1,0,N({m_col("tcum", k)}$6)*N({pct_jk})*{m_col("tup", k)}$5)', f=F_CALC, nf=FMT_INT)
            put(ws, f"{TM[k]}{rk}", f'=IF(ISNUMBER({kw_jk}),{kw_jk},0)', f=F_CALC, nf="0")
        rij = tm_rng(rn)
        eerste = ArrayFormula(f"{TM['label']}{rn}", f'=IF({tsel}=0,IFERROR(INDEX({rij},MATCH(TRUE,{rij}<>"",0)),""),INDEX({rij},{tsel}))')
        put(ws, f"{TM['label']}{rn}", eerste, f=F_CALC)
        put(ws, f"{TM['label']}{ra}", f'=IF({TM["label"]}{rn}<>"",{TM["label"]}{rn},IF(SUM({tm_rng(ra)})>0,"Termijn {j}",""))', f=F_CALC)
    # ---- bouwkosten per termijn (Woningtypes blok 7): per termijn j en type k naam, bedrag (€; 0 zonder bouwkwartaal), bouwkwartaal --
    put(ws, f"{TM[1]}{LY.TM_R_BKN - 1}", "bouwkosten per termijn (tab Woningtypes blok 7, wat de aannemer krijgt): naam, bedrag (€) en "
                                         "bouwkwartaal per termijn (rij) en type (kolom)", f=F_NOTE8)
    for j in range(1, LY.N_TERMIJNEN + 1):
        rn, ra, rk = LY.TM_R_BKN + j - 1, LY.TM_R_BKA + j - 1, LY.TM_R_BKK + j - 1
        for k in range(1, N_TYPES + 1):
            naam_jk, eur_jk, kw_jk = (LY.wt_ref(LY.WT_R_B1 + j - 1, k, o) for o in range(3))
            put(ws, f"{TM[k]}{rn}", f'=IF(TRIM({naam_jk}&"")="","",TRIM({naam_jk}&""))', f=F_CALC)
            put(ws, f"{TM[k]}{ra}", f'=IF(ISNUMBER({kw_jk}),N({eur_jk}),0)', f=F_CALC, nf=FMT_INT)
            put(ws, f"{TM[k]}{rk}", f'=IF(ISNUMBER({kw_jk}),{kw_jk},0)', f=F_CALC, nf="0")
        rij = tm_rng(rn)
        eerste = ArrayFormula(f"{TM['label']}{rn}", f'=IF({tsel}=0,IFERROR(INDEX({rij},MATCH(TRUE,{rij}<>"",0)),""),INDEX({rij},{tsel}))')
        put(ws, f"{TM['label']}{rn}", eerste, f=F_CALC)
        put(ws, f"{TM['label']}{ra}", f'=IF({TM["label"]}{rn}<>"",{TM["label"]}{rn},IF(SUM({tm_rng(ra)})>0,"Bouwkosten-termijn {j}",""))', f=F_CALC)
    for k in range(1, N_TYPES + 1):
        put(ws, f"{TM[k]}{LY.TM_R_SEL}", f"=IF(OR({tsel}=0,{tsel}={k}),1,0)", f=F_CALC, nf="0")
        put(ws, f"{TM[k]}{LY.TM_R_SELND}", f'=IF(AND(OR({tsel}=0,{tsel}={k}),N({m_col("fee", k)}$5)<>1),1,0)', f=F_CALC, nf="0")
        put(ws, f"{TM[k]}{LY.TM_R_AG}", f'=IF(N({m_col("fee", k)}$5)=1,0,N({m_col("tcum", k)}$6)*N({m_col("tup", k)}$6))', f=F_CALC, nf=FMT_INT)
    put(ws, f"{TM['label']}{LY.TM_R_SEL}", "← gekozen typen (1/0)", f=F_NOTE8)
    put(ws, f"{TM['label']}{LY.TM_R_AG}", "← grondtermijn per woning (€)", f=F_NOTE8)
    put(ws, f"{TM['label']}{LY.TM_R_SELND}", "← gekozen en niet-DAEB (1/0), voor koper − aannemer", f=F_NOTE8)

    # ---- hulpcellen -------------------------------------------------------------------------------------------------------
    kwv = _rng(KWV)
    G_ = lambda naam: _rng(G[naam])  # noqa: E731
    S = lambda s, veld: h(f"vs{s}_{veld}")  # noqa: E731
    alle = range(1, N_VS + 1)
    sb_in, ws_naam = lambda s: LY.sc_in(s, "sb"), lambda s: LY.sc_in(s, "naam")  # noqa: E731
    min_aan = lambda veld, leeg: ",".join(f"IF({S(s, 'aan')}=1,{S(s, veld)},{leeg})" for s in alle)  # noqa: E731
    lo = f'MIN(0,{h("dal_basis")},{min_aan("dal", 0)})'
    hi = f'MAX(0,{h("top_basis")},' + ",".join(f'IF({S(s, "aan")}=1,MAX({_rng(vs_col(s, "stand"))}),0)' for s in alle) + ")"
    spreid = f"MAX(1,({hi}-{lo})/1000000)"                 # € mln; minstens 1, ook in een leeg werkboek een normale as
    ruw = f"({spreid}/5)"
    mag = f"10^INT(LOG10({ruw}))"
    stap = f"IF({ruw}/{mag}<=1,1,IF({ruw}/{mag}<=2,2,IF({ruw}/{mag}<=5,5,10)))*{mag}"
    titel_s = lambda s: (f'UPPER(LEFT({S(s, "naam")},1))&MID({S(s, "naam")},2,200)&": financiering tot "&{eur_mln(S(s, "dal"))}'  # noqa: E731
                         f'&IF({S(s, "d_be")}="",IF({S(s, "be_pos")}=0,", geen break-even",""),IF({S(s, "d_be")}=0,", break-even gelijk",'
                         f'", break-even "&{kw_tekst("ABS(" + S(s, "d_be") + ")")}&IF({S(s, "d_be")}>0," later"," eerder")))')
    eind_min = "MIN(" + ",".join(f"IF({S(s, 'aan')}=1,{S(s, 'eind')},1E+300)" for s in alle) + ")"
    eind_max = "MAX(" + ",".join(f"IF({S(s, 'aan')}=1,{S(s, 'eind')},-1E+300)" for s in alle) + ")"
    dal_min = "MIN(" + ",".join(f"IF({S(s, 'aan')}=1,{S(s, 'dal')},1E+300)" for s in alle) + ")"
    hulp = {
        "vs_vnd_act": ("Scenario's: verkocht niet-DAEB t/m actuals", f'=IF({pa}=0,0,N(INDEX({G_("vnd")},{pa})))'),
        "vs_tnd_act": ("Scenario's: getransporteerd niet-DAEB t/m actuals", f'=IF({pa}=0,0,N(INDEX({G_("tnd")},{pa})))'),
        "vs_vv_pos": ("Scenario's: positie eerste geplande verkoop na actuals", f'=COUNTIF({G_("vnd")},"<="&{h("vs_vnd_act")})+1'),
        "vs_vv_plan": ("Scenario's: eerste geplande verkoop (kwartaalindex; 99999 = geen)",
                       f'=IF({h("vs_vv_pos")}>{n},99999,INDEX({idx_rng},{h("vs_vv_pos")}))'),
        "vs_tr_pos": ("Scenario's: positie eerste gepland transport na actuals", f'=COUNTIF({G_("tnd")},"<="&{h("vs_tnd_act")})+1'),
        "vs_tr_plan": ("Scenario's: eerste gepland transport (kwartaalindex; 99999 = geen)",
                       f'=IF({h("vs_tr_pos")}>{n},99999,INDEX({idx_rng},{h("vs_tr_pos")}))'),
        "vs_gap": ("Scenario's: eerste transport na start bouw in de planning (kw)",
                   f'=IF(OR({h("vs_tr_plan")}>=99999,{h("start_idx")}>=99999),0,{h("vs_tr_plan")}-{h("start_idx")})'),
        "vs_aan": ("Scenario's: aantal ingevuld", "=" + "+".join(S(s, "aan") for s in alle)),
        "vs_pp": ("Scenario's: Naar PowerPoint (1/0)", f'=IF(LOWER(TRIM({LY.SC_BLAD}!${LY.SC_CEL_PP[0]}${LY.SC_CEL_PP[1:]}&""))="nee",0,1)'),
        "vs_pp_aan": ("Scenario's: dia 4 in PowerPoint (1/0)", f'=IF(AND({h("vs_pp")}=1,{h("vs_aan")}>0),1,0)'),
        "vs_cf": ("Scenario's: lijnen in de cashflowgrafiek (1/0; alleen 'ja' zet ze aan)",
                  f'=IF(LOWER(TRIM({LY.SC_BLAD}!${LY.SC_CEL_CF[0]}${LY.SC_CEL_CF[1:]}&""))="ja",1,0)'),
        "vs_stap": ("Scenario-grafiek: stap van de waarde-as (€ mln)", f"={stap}"),
        "vs_min": ("Scenario-grafiek: as-minimum (€ mln)", f'={h("vs_stap")}*INT(({lo}/1000000-0.08*{spreid})/{h("vs_stap")})'),
        "vs_max": ("Scenario-grafiek: as-maximum (€ mln)", f'=-{h("vs_stap")}*INT(-({hi}/1000000+0.04*{spreid})/{h("vs_stap")})'),
        "vs_slechtst": ("Scenario met de diepste financiering (0 = geen)",
                        f'=IF({h("vs_aan")}=0,0,' + "".join(f'IF(AND({S(s, "aan")}=1,{S(s, "dal")}={dal_min}),{s},' for s in alle)
                        + "0" + ")" * N_VS + ")"),
        "vs_titel": ("Scenario's: titel (dia en tab)",
                     f'=IF({h("vs_aan")}=0,"Nog geen scenario: vul op tab Scenario\'s een kwartaal in",'
                     f'CHOOSE({h("vs_slechtst")},{",".join(titel_s(s) for s in alle)}))'),
        "vs_sb_mee": ("Scenario's: een scenario verschuift de start bouw (1/0)",
                      "=IF(OR(" + ",".join(f"AND({S(s, 'aan')}=1,{S(s, 'd_sb')}<>0)" for s in alle) + "),1,0)"),
        "vs_subtitel": ("Scenario's: ondertitel",
                        f'="Cumulatieve cashflow per scenario · transport nooit vóór de verkoop · "&{h("kf_tekst_kort")}'),
        "vs_eind_tekst": ("Scenario's: eindsaldo",
                          f'="eindsaldo Basis "&{eur_m(h("eind_basis"))}&IF({h("vs_aan")}=0,"",IF(AND(ABS({eind_min}-{h("eind_basis")})<50000,'
                          f'ABS({eind_max}-{h("eind_basis")})<50000)," · gelijk in alle scenario\'s"," · scenario\'s "&{eur_m(eind_min)}&" tot "&{eur_m(eind_max)}))'),
        "vs_sel": ("Tabel per termijn: gekozen scenario (0 = Basis)",
                   f'=IFERROR(IF(LEFT(TRIM({LY.SC_BLAD}!${LY.SC_CEL_SEL[0]}${LY.SC_CEL_SEL[1:]}&""),8)="Scenario",'
                   f'MIN({N_VS},MAX(0,VALUE(MID(TRIM({LY.SC_BLAD}!${LY.SC_CEL_SEL[0]}${LY.SC_CEL_SEL[1:]}&""),10,1)))),0),0)'),
        "vs_typesel": ("Tabel per termijn: gekozen woningtype (0 = alle)",
                       f'=IF(OR(TRIM({LY.SC_BLAD}!${LY.SC_CEL_TYPE[0]}${LY.SC_CEL_TYPE[1:]}&"")="",'
                       f'LOWER(TRIM({LY.SC_BLAD}!${LY.SC_CEL_TYPE[0]}${LY.SC_CEL_TYPE[1:]}&""))="alle types"),0,'
                       f'IFERROR(MATCH(TRIM({LY.SC_BLAD}!${LY.SC_CEL_TYPE[0]}${LY.SC_CEL_TYPE[1:]}&""),${m_col("gv", 1)}$7:${m_col("gv", N_TYPES)}$7,0),0))'),
        "vs0_leg": ("Legenda Basis", f'="Basis"&IF({h("vs_vv_plan")}>=99999,""," · verkoop "&{kw_lbl(h("vs_vv_plan"))})'),
        "vs0_sub": ("Kaart Basis: verkoop, transport, rente",
                    f'="verkoop "&{kw_lbl(h("vs_vv_plan"))}&" · transport "&{kw_lbl(h("vs_tr_plan"))}&" · rente "&{eur_m(h("rente_basis_tot"))}'),
        "vs0_delta": ("Kaart Basis: regel onderaan", f'="planning uit het werkboek · eindsaldo "&{eur_m(h("eind_basis"))}'),
        "vs0_lbl_dal": ("Label dieptepunt Basis", f"={eur_m(h('dal_basis'))}"),
        "vs_kicker": ("Scenario's: kopje", f'="VERKOOPSCENARIO\'S · START BOUW "&UPPER({kw_lbl(h("start_idx"))})'),
        "vs0_dal_pos": ("Positie dieptepunt Basis", f'=IFERROR(MATCH({h("dal_basis")},{_rng(M["stand_basis"])},0),0)'),
        "vs_nu_pos": ("Scenario-grafiek: positie 'nu' in het tijdvak (vanaf het actuals-jaar)", f'=IF({pa}=0,0,{pa}-{h("vp_start")}+1)'),
        "vs_tr_tekst": ("Scenario's: regel transport automatisch",
                        f'="transport automatisch: niet eerder na start bouw dan in de planning ("&IF({h("vs_tr_plan")}>=99999,"geen transport gepland",'
                        f'IF({h("vs_gap")}=0,"zelfde kwartaal als start bouw",{h("vs_gap")}&" kw na start bouw"))&")"'),
    }
    for s in alle:
        in_sb, in_vv, in_tv = S(s, "in_sb"), S(s, "in_vv"), S(s, "in_tv")
        eff = lambda x: f"IF({x}>{ia},{x},0)"  # noqa: E731
        ref = {v: LY.sc_in(s, v) for v in ("naam", "sb", "vv", "tv", "gr")}
        melding = "&".join(
            f'IF(AND(TRIM({ref[v]}&"")<>"",{S(s, "in_" + v)}=0)," · {wat}: onbekend kwartaal","")'
            f'&IF(AND({S(s, "in_" + v)}>0,{S(s, "in_" + v)}<={ia})," · {wat} ligt in de realisatie","")'
            for v, wat in (("sb", "start bouw"), ("vv", "verkoop"), ("tv", "transport")))
        melding += (f'&IF(AND({eff(in_sb)}>0,{h("start_idx")}<={ia})," · start bouw is al geweest","")'
                    f'&IF(AND({eff(in_sb)}>0,{h("start_idx")}>=99999)," · geen start bouw in de planning","")'
                    f'&IF(OR({S(s, "in_gr")}=-2,{S(s, "in_gr")}>1)," · grondtermijn: geen % tussen 0 en 100","")')
        d_sb, d_v, d_t = S(s, "d_sb"), S(s, "d_v"), S(s, "d_t")
        naam_in = f'TRIM({ref["naam"]}&"")'
        auto = (f'IF({d_sb}<>0,"Start bouw "&{kw_lbl(S(s, "sb"))},"")&IF({d_v}<>0,IF({d_sb}<>0," · verkoop ","Verkoop ")&{kw_lbl(S(s, "vv"))},"")'
                f'&IF({eff(in_tv)}>0,IF(OR({d_sb}<>0,{d_v}<>0)," · transport ","Transport ")&{kw_lbl(S(s, "tv"))},"")'
                f'&IF({S(s, "gr")}>=0,IF(OR({d_sb}<>0,{d_v}<>0,{eff(in_tv)}>0)," · grond ","Grondtermijn ")&FIXED({S(s, "gr")}*100,0)&"%","")')
        be_tekst = (f'IF({S(s, "d_be")}="",IF({S(s, "be_pos")}=0,"geen break-even","break-even "&{S(s, "be")}),IF({S(s, "d_be")}=0,'
                    f'"break-even gelijk","break-even "&{kw_tekst("ABS(" + S(s, "d_be") + ")")}&IF({S(s, "d_be")}>0," later"," eerder")))')
        st_rng = _rng(vs_col(s, "stand"))
        hulp.update({
            f"vs{s}_in_sb": (f"Scenario {s}: start bouw (invoer, kwartaalindex)", f"={kw_index(ref['sb'])}"),
            f"vs{s}_in_vv": (f"Scenario {s}: verkoop vanaf (invoer)", f"={kw_index(ref['vv'])}"),
            f"vs{s}_in_tv": (f"Scenario {s}: transport vanaf (invoer)", f"={kw_index(ref['tv'])}"),
            f"vs{s}_in_gr": (f"Scenario {s}: grondtermijn (invoer als fractie; -1 = leeg, -2 = onleesbaar)",
                             f'=IF(TRIM({ref["gr"]}&"")="",-1,IFERROR(IF(ISNUMBER({ref["gr"]}),IF({ref["gr"]}>1,{ref["gr"]}/100,{ref["gr"]}),'
                             f'VALUE(SUBSTITUTE(TRIM({ref["gr"]}&""),"%",""))/100),-2))'),
            f"vs{s}_gr": (f"Scenario {s}: grondtermijn (fractie; -1 = zoals gepland)",
                          f'=IF(AND({S(s, "in_gr")}>=0,{S(s, "in_gr")}<=1),{S(s, "in_gr")},-1)'),
            f"vs{s}_fout": (f"Scenario {s}: controle invoer (leeg = goed)", f'=MID({melding},4,300)'),
            f"vs{s}_aan": (f"Scenario {s}: ingevuld (1/0)", f'=IF(OR({eff(in_sb)}>0,{eff(in_vv)}>0,{eff(in_tv)}>0,{S(s, "gr")}>=0),1,0)'),
            f"vs{s}_d_sb": (f"Scenario {s}: start bouw verschoven (kw)",
                            f'=IF(OR({eff(in_sb)}=0,{h("start_idx")}>=99999,{h("start_idx")}<={ia}),0,{eff(in_sb)}-{h("start_idx")})'),
            f"vs{s}_sb": (f"Scenario {s}: start bouw (kwartaalindex)", f'=IF({h("start_idx")}>=99999,99999,{h("start_idx")}+{d_sb})'),
            f"vs{s}_vv": (f"Scenario {s}: verkoop vanaf (kwartaalindex)", f'=IF({eff(in_vv)}>0,{eff(in_vv)},{h("vs_vv_plan")})'),
            f"vs{s}_d_v": (f"Scenario {s}: verkoop verschoven (kw)", f'=IF(OR({eff(in_vv)}=0,{h("vs_vv_plan")}>=99999),0,{eff(in_vv)}-{h("vs_vv_plan")})'),
            f"vs{s}_tv": (f"Scenario {s}: transport vanaf (kwartaalindex)",
                          f'=IF({eff(in_tv)}>0,{eff(in_tv)},IF({h("vs_tr_plan")}>=99999,99999,{h("vs_tr_plan")}+{d_sb}))'),
            f"vs{s}_d_t": (f"Scenario {s}: transport verschoven (kw)", f'=IF(OR({S(s, "tv")}>=99999,{h("vs_tr_plan")}>=99999),0,{S(s, "tv")}-{h("vs_tr_plan")})'),
            f"vs{s}_naam": (f"Scenario {s}: naam (invoer of automatisch)", f'=IF({naam_in}<>"",{naam_in},IF({auto}="","Zoals Basis",{auto}))'),
            f"vs{s}_kop": (f"Scenario {s}: kop", f'="Scenario {s}"'),
            f"vs{s}_leg": (f"Scenario {s}: legenda (automatische naam met kleine letter, zoals bij de Basis)",
                           f'="Scenario {s} · "&IF({naam_in}<>"",{naam_in},LOWER(LEFT({S(s, "naam")},1))&MID({S(s, "naam")},2,200))'),
            f"vs{s}_dal": (f"Scenario {s}: dieptepunt (€)", f"=MIN({st_rng})"),
            f"vs{s}_dal_pos": (f"Scenario {s}: positie dieptepunt", f"=IFERROR(MATCH({S(s, 'dal')},{st_rng},0),0)"),
            f"vs{s}_dal_kw": (f"Scenario {s}: dieptepunt in", f'=IF({S(s, "dal_pos")}=0,"–",INDEX({kwv},{S(s, "dal_pos")}))'),
            f"vs{s}_be_pos": (f"Scenario {s}: positie break-even", f"=IFERROR(MATCH(1,{_rng(vs_col(s, 'pos'))},0),0)"),
            f"vs{s}_be": (f"Scenario {s}: break-even", f'=IF({S(s, "be_pos")}=0,"",INDEX({kwv},{S(s, "be_pos")}))'),
            f"vs{s}_eind": (f"Scenario {s}: eindsaldo (€)", f"=IF({n}=0,0,INDEX({st_rng},{n}))"),
            f"vs{s}_rente": (f"Scenario {s}: rente (berekend, €)", f"=SUM({_rng(vs_col(s, 'rente'))})"),
            f"vs{s}_vbsb": (f"Scenario {s}: verkocht bij start bouw (aantal)", f"=SUM({vs_rng(s, 'vc', 5, abs_=True)})"),
            f"vs{s}_vbsb_pct": (f"Scenario {s}: verkocht bij start bouw (%)", f'=IFERROR({S(s, "vbsb")}/{h("won_met_start")},0)'),
            f"vs{s}_et_pos": (f"Scenario {s}: positie eerste transport na actuals (0 = geen)",
                              f'=IF(COUNTIF({_rng(vs_col(s, "tnd"))},"<="&{h("vs_tnd_act")})+1>{n},0,COUNTIF({_rng(vs_col(s, "tnd"))},"<="&{h("vs_tnd_act")})+1)'),
            f"vs{s}_et": (f"Scenario {s}: eerste transport", f'=IF({S(s, "et_pos")}=0,"–",INDEX({kwv},{S(s, "et_pos")}))'),
            f"vs{s}_d_fin": (f"Scenario {s}: extra financiering t.o.v. Basis (€)", f'={h("dal_basis")}-{S(s, "dal")}'),
            f"vs{s}_d_be": (f"Scenario {s}: break-even t.o.v. Basis (kw; leeg = geen vergelijking)",
                            f'=IF(OR({S(s, "be_pos")}=0,{h("be_basis_pos")}=0),"",INDEX({idx_rng},{S(s, "be_pos")})-INDEX({idx_rng},{h("be_basis_pos")}))'),
            f"vs{s}_sub": (f"Scenario {s}: kaart, regel verkoop / transport / rente",
                           f'="verkoop "&{kw_lbl(S(s, "vv"))}&" · transport "&{S(s, "et")}&" · rente "&{eur_m(S(s, "rente"))}'
                           f'&IF({S(s, "gr")}>=0," · grond "&FIXED({S(s, "gr")}*100,0)&"%","")'),
            f"vs{s}_delta": (f"Scenario {s}: kaart, verschil met de Basis",
                             f'=IF(ABS({S(s, "d_fin")})<50000,"financiering gelijk aan Basis",{eur_m("ABS(" + S(s, "d_fin") + ")")}'
                             f'&IF({S(s, "d_fin")}>0," extra"," minder")&" financiering")&" · "&{be_tekst}'),
            f"vs{s}_lbl_dal": (f"Scenario {s}: label dieptepunt", f"={eur_m(S(s, 'dal'))}"),
            f"vs{s}_kaart": (f"Scenario {s}: naam op de kaart (dia)", f'=IF({naam_in}<>"","Scenario {s} · "&{naam_in},"Scenario {s}")'),
            f"vs{s}_pp": (f"Scenario {s}: op dia 4 (1/0)", f'=IF(AND({h("vs_pp")}=1,{S(s, "aan")}=1),1,0)'),
            f"vs{s}_cf": (f"Scenario {s}: lijn in de cashflowgrafiek (1/0)", f'=IF(AND({h("vs_cf")}=1,{S(s, "aan")}=1),1,0)'),
        })
    for naam, (label, formule) in hulp.items():
        r = H[naam]
        put(ws, f"{LY.M_HULP_LABEL}{r}", label, f=F_NOTE8)
        put(ws, f"{LY.M_HULP}{r}", formule, f=F_CALC, al=AL_LEFT_TOP)
    for naam in ("vs_min", "vs_max", "vs_stap"):
        ws[f"{LY.M_HULP}{H[naam]}"].number_format = "0.00"
    for s in alle:
        for veld in ("dal", "eind", "rente", "d_fin"):
            ws[f"{LY.M_HULP}{H[f'vs{s}_{veld}']}"].number_format = FMT_INT
        ws[f"{LY.M_HULP}{H[f'vs{s}_vbsb_pct']}"].number_format = FMT_PCT


# =============================================================================
#  TABBLAD SCENARIO'S
# =============================================================================
KAART_COL = [3, 8, 13, 18]          # kaarten: C:F, H:K, M:P, R:U
R_KAART = 5                         # kleurstreep; daaronder zeven regels
R_S1, R_S2 = 14, 27                 # sectie 1 (invoer) en sectie 2 (uitkomst)
R_S4 = 43                           # sectie 4 (tabel per termijn)
GRAFIEK_COL = 12                    # kolom L: grafiek rechts van sectie 1 en 2
CHIP_RIJ = R_S1 + 1


def _pct(v):
    """invoer grondtermijn (0.5, '50%', '50', 50) -> fractie, of None"""
    if v in (None, ""):
        return None
    try:
        x = float(str(v).replace("%", "").replace(",", ".").strip()) if isinstance(v, str) else float(v)
    except ValueError:
        return None
    return x / 100 if x > 1 or (isinstance(v, str) and "%" in v) else x


def _paar(ws, r, k, waarde, f=None, fl=None, al=AL_RIGHT, nf=None, bd=None):
    """Waarde in het samengevoegde paar van Basis (k=0) of scenario k (twee kolommen)."""
    c0 = LY.SC_COL[k]
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 1)
    put(ws, f"{L(c0)}{r}", waarde, f=f or F_CALC9, fl=fl, al=al, nf=nf)
    if fl is not None:
        ws.cell(r, c0 + 1).fill = fl
    if bd is not None:
        ws.cell(r, c0).border = bd
        ws.cell(r, c0 + 1).border = bd
    return f"{L(c0)}{r}"


def _kop_scenarios(ws, r):
    put(ws, f"B{r}", None, fl=FL_HDR)
    for k, t in enumerate(["Basis"] + [f"Scenario {s}" for s in range(1, N_VS + 1)]):
        _paar(ws, r, k, t, f=F_HDR, fl=FL_HDR)
    for k in range(N_VS + 1):                                # kleurstreep = kleur van de lijn
        for c in (LY.SC_COL[k], LY.SC_COL[k] + 1):
            ws.cell(r + 1, c).border = Border(top=side("thick", LY.VS_KLEUREN[k]))
    ws.row_dimensions[r + 1].height = 4


def bouw_tab(wb, data):
    ws = wb.create_sheet(LY.SC_TAB, 1)
    ws.sheet_properties.codeName = "shScenarios"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    ws.column_dimensions["A"].width = 2.5
    ws.column_dimensions["B"].width = 40
    for j in range(LY.SC_N_KW):
        ws.column_dimensions[L(3 + j)].width = 8.6
    S = lambda s, veld: hm(f"vs{s}_{veld}")  # noqa: E731
    aan = lambda s: hm(f"vs{s}_aan")  # noqa: E731
    GRIJS_FL, LICHT = fill("F2F3F5"), "B4BAC1"
    BD_RIJ = Border(bottom=side("thin", LINE2))
    params = data.params or {}
    invoer = list(params.get("vs") or [])

    # ---- kop ------------------------------------------------------------------------------------------------------------------
    put(ws, "B1", f'="SCENARIO\'S · START BOUW, VERKOOP EN TRANSPORT · STAND PER "&UPPER({hm("kw_actuals")})', f=F_NOTE)
    ws.merge_cells("B2:V2")
    ws.row_dimensions[2].height = 24
    put(ws, "B2", f"={hm('vs_titel')}", f=font(14, True, TXT), al=AL_VCENTER)
    ws.merge_cells("B3:V3")
    put(ws, "B3", '="Basis = je planning (tab Invoer en Woningtypes) · bij transport betaalt de koper de grondtermijn plus alle vervallen '
                  'bouwtermijnen, daarna elke termijn in zijn bouwkwartaal · "&' + hm("kf_tekst"), f=font(10, False, TXT2),
        al=Alignment(horizontal="left", vertical="top", wrap_text=True))
    ws.row_dimensions[3].height = 28

    # ---- kaarten ----------------------------------------------------------------------------------------------------------------
    for r, hgt in ((R_KAART, 4), (R_KAART + 1, 18), (R_KAART + 2, 12), (R_KAART + 3, 28), (R_KAART + 4, 14), (R_KAART + 5, 14),
                   (R_KAART + 6, 14), (R_KAART + 7, 14)):
        ws.row_dimensions[r].height = hgt
    put(ws, f"B{R_KAART + 1}", "IN ÉÉN OOGOPSLAG", f=F_SECTION)
    put(ws, f"B{R_KAART + 2}", "één kaart per scenario", f=F_NOTE8)
    put(ws, f"B{R_KAART + 3}", "kleur = lijn in de grafiek en op de dia", f=F_NOTE8, al=AL_VCENTER)
    put(ws, f"B{R_KAART + 4}", "rekent direct mee met blok 1", f=F_NOTE8)
    for k in range(N_VS + 1):
        c0, c3 = KAART_COL[k], KAART_COL[k] + 3
        A = lambda r: f"{L(c0)}{r}"  # noqa: E731
        for c in range(c0, c3 + 1):
            ws.cell(R_KAART, c).fill = fill(LY.VS_KLEUREN[k])
        for r in range(R_KAART + 1, R_KAART + 8):
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c3)
            for c in range(c0, c3 + 1):
                cel = ws.cell(r, c)
                cel.fill = FL_CARD
                cel.border = Border(left=side("thin", LINE) if c == c0 else None, right=side("thin", LINE) if c == c3 else None,
                                    bottom=side("thin", LINE) if r == R_KAART + 7 else None)
        if k == 0:
            regels = [
                ("Basis", font(10, True, TXT)),
                ("MAX. FINANCIERINGSBEHOEFTE", font(7, False, TXT3)),
                (f"={eur_m(hm('dal_basis'))}", font(18, True, TXT)),
                (f'="dieptepunt "&{hm("dal_basis_kw")}', font(8.5, False, TXT)),
                (f'="break-even "&IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', font(8.5, False, TXT)),
                (f'="verkocht bij start bouw "&FIXED({hm("verk_bij_start_pct")}*100,0)&"% · norm "&FIXED({dabs("norm")}*100,0)&"%"',
                 font(8.5, False, TXT)),
                (f"={hm('vs0_sub')}", font(8, False, TXT3)),
            ]
            pct = hm("verk_bij_start_pct")
        else:
            s = k
            tov = (f'IF(ABS({S(s, "d_fin")})<50000,"gelijk aan Basis",{eur_m("ABS(" + S(s, "d_fin") + ")")}'
                   f'&IF({S(s, "d_fin")}>0," meer"," minder")&" dan Basis")')
            be = (f'IF({S(s, "be")}="","niet bereikt",{S(s, "be")})&IF({S(s, "d_be")}="","",IF({S(s, "d_be")}=0," · gelijk aan Basis",'
                  f'" · "&{kw_tekst("ABS(" + S(s, "d_be") + ")")}&IF({S(s, "d_be")}>0," later"," eerder")))')
            regels = [
                (f'=IF({aan(s)}=1,"Scenario {s} · "&{S(s, "naam")},"Scenario {s}")', font(10, True, TXT)),
                (f'=IF({aan(s)}=1,"MAX. FINANCIERINGSBEHOEFTE","")', font(7, False, TXT3)),
                (f'=IF({aan(s)}=1,{eur_m(S(s, "dal"))},"nog leeg")', font(18, True, TXT)),
                (f'=IF({aan(s)}=1,"dieptepunt "&{S(s, "dal_kw")}&" · "&{tov},"kies in blok 1 een kwartaal bij \'Verkoop vanaf\'")', font(8.5, False, TXT)),
                (f'=IF({aan(s)}=1,"break-even "&{be},"of bij \'Start bouw\'")', font(8.5, False, TXT)),
                (f'=IF({aan(s)}=1,"verkocht bij start bouw "&FIXED({S(s, "vbsb_pct")}*100,0)&"% · norm "&FIXED({dabs("norm")}*100,0)&"%","")',
                 font(8.5, False, TXT)),
                (f'=IF({aan(s)}=1,{S(s, "sub")},"")', font(8, False, TXT3)),
            ]
            pct = S(s, "vbsb_pct")
            # niet ingevuld: grijze streep en grijze tekst
            for c in range(c0, c3 + 1):
                ws.conditional_formatting.add(f"{L(c)}{R_KAART}", FormulaRule(formula=[f"{aan(s)}=0"],
                                              fill=PatternFill(fill_type="solid", bgColor="E3E6EA", fgColor="E3E6EA")))
            ws.conditional_formatting.add(f"{A(R_KAART + 1)}:{A(R_KAART + 7)}", FormulaRule(formula=[f"{aan(s)}=0"],
                                          font=Font(name=ARIAL, color=LICHT), fill=PatternFill(fill_type="solid", bgColor="FFFFFF", fgColor="FFFFFF")))
        for i, (tekst, fnt) in enumerate(regels):
            put(ws, A(R_KAART + 1 + i), tekst, f=fnt, al=AL_LEFT if i != 2 else AL_VCENTER)
        rv = A(R_KAART + 6)                                  # verkocht bij start bouw: groen op norm, rood eronder
        ws.conditional_formatting.add(rv, FormulaRule(formula=[f'AND(LEN({rv})>0,{pct}>={dabs("norm")})'], font=Font(name=ARIAL, color=GREEN)))
        ws.conditional_formatting.add(rv, FormulaRule(formula=[f'AND(LEN({rv})>0,{pct}<{dabs("norm")})'], font=Font(name=ARIAL, color=RED)))

    # ---- 1 · scenario maken ------------------------------------------------------------------------------------------------------
    put(ws, f"B{R_S1}", "1 · SCENARIO MAKEN · vul alleen in wat anders is dan de planning · leeg = automatisch", f=F_SECTION)
    _kop_scenarios(ws, R_S1 + 1)
    # geplande grondtermijn: gewogen naar koopsom over de niet-DAEB-typen (per type kan hij verschillen)
    ks, nd = f"Model!{LY.m_range_abs('tcum', 6)}", f"(1-Model!{LY.m_range_abs('fee', 5)})"
    grond_plan = f"IFERROR(SUMPRODUCT({ks},Model!{LY.m_range_abs('tup', 6)},{nd})/SUMPRODUCT({ks},{nd}),0)"
    rijen = [
        (LY.SC_R_NAAM, LY.SC_LABELS["naam"], "Basis", "naam"),
        (LY.SC_R_SB, LY.SC_LABELS["sb"], f"={kw_lbl(hm('start_idx'))}", "sb"),
        (LY.SC_R_VV, LY.SC_LABELS["vv"], f"={kw_lbl(hm('vs_vv_plan'))}", "vv"),
        (LY.SC_R_TV, LY.SC_LABELS["tv"], f"={kw_lbl(hm('vs_tr_plan'))}", "tv"),
        (LY.SC_R_GR, LY.SC_LABELS["gr"], f"={grond_plan}", "gr"),
    ]
    dv_kw = DataValidation(type="list", formula1="vs_kwartalen", allow_blank=True)
    dv_kw.showErrorMessage = False                           # ook zelf typen mag ('Q2 2028', '2028 Q2')
    dv_kw.promptTitle, dv_kw.prompt = "Kwartaal", "Kies een kwartaal na de actuals, of typ bijvoorbeeld Q2 2028. Leeg = zoals de planning."
    dv_kw.showInputMessage = True
    ws.add_data_validation(dv_kw)
    dv_gr = DataValidation(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True)
    dv_gr.showErrorMessage = False
    dv_gr.promptTitle, dv_gr.prompt = "Grondtermijn", "Bijvoorbeeld 50%. Koopsom blijft gelijk, de bouwtermijnen gaan naar rato omlaag. Leeg = zoals gepland."
    dv_gr.showInputMessage = True
    ws.add_data_validation(dv_gr)
    for r, label, basis, veld in rijen:
        put(ws, f"B{r}", label, f=F_CALC9, al=AL_LEFT, bd=BD_RIJ)
        _paar(ws, r, 0, basis, f=font(9, False, TXT2), fl=GRIJS_FL, bd=BD_RIJ)
        for s in range(1, N_VS + 1):
            waarde = invoer[s - 1].get(veld) if s - 1 < len(invoer) and isinstance(invoer[s - 1], dict) else None
            if veld == "gr":
                waarde = _pct(waarde)
            cel = _paar(ws, r, s, waarde, f=F_INPUT, fl=FL_INPUT, al=AL_RIGHT, bd=BD_RIJ)
            ws[cel].number_format = "0%" if veld == "gr" else "@"
            if veld == "gr":
                dv_gr.add(cel)
            elif veld != "naam":
                dv_kw.add(cel)
        if veld == "gr":
            ws[f"{L(LY.SC_COL[0])}{r}"].number_format = "0%"
    r = LY.SC_R_GR + 1                                      # toelichting over de breedte van blok 1, op twee regels
    ws.merge_cells(f"B{r}:{L(LY.SC_COL[-1] + 1)}{r}")
    put(ws, f"B{r}", f'="Grijs = je planning (niet aanpassen) · blauw = invoer: kies een kwartaal of typ bijvoorbeeld Q2 2028 · een lege kolom '
                     f'staat uit · "&{hm("vs_tr_tekst")}', f=F_NOTE8, al=Alignment(horizontal="left", vertical="top", wrap_text=True))
    ws.row_dimensions[r].height = 22
    put(ws, f"B{LY.SC_R_PP}", LY.SC_LABELS["pp"], f=F_CALC9, al=AL_LEFT, bd=BD_RIJ)
    cel_pp = _paar(ws, LY.SC_R_PP, 0, params.get("vs_pp") or "ja", f=F_INPUT, fl=FL_INPUT, al=AL_CENTER, bd=BD_RIJ)
    assert cel_pp == LY.SC_CEL_PP
    dv_jn = DataValidation(type="list", formula1='"ja,nee"', allow_blank=True)
    ws.add_data_validation(dv_jn)
    dv_jn.add(cel_pp)
    c_noot = L(LY.SC_COL[1])
    ws.merge_cells(f"{c_noot}{LY.SC_R_PP}:{L(LY.SC_COL[-1] + 1)}{LY.SC_R_PP}")
    put(ws, f"{c_noot}{LY.SC_R_PP}", "ja = de ingevulde scenario's op de dia Verkoopscenario's · nee = alleen hier in Excel · "
                                     "zonder ingevuld scenario komt er niets in PowerPoint", f=F_NOTE8,
        al=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    ws.row_dimensions[LY.SC_R_PP].height = 24
    # lijnen in de cashflowgrafiek (Dashboard en dia 3): los aan/uit, standaard uit (rustiger grafiek)
    put(ws, f"B{LY.SC_R_CF}", LY.SC_LABELS["cf"], f=F_CALC9, al=AL_LEFT, bd=BD_RIJ)
    cel_cf = _paar(ws, LY.SC_R_CF, 0, params.get("vs_cf") or "nee", f=F_INPUT, fl=FL_INPUT, al=AL_CENTER, bd=BD_RIJ)
    assert cel_cf == LY.SC_CEL_CF
    dv_jn.add(cel_cf)
    ws.merge_cells(f"{c_noot}{LY.SC_R_CF}:{L(LY.SC_COL[-1] + 1)}{LY.SC_R_CF}")
    put(ws, f"{c_noot}{LY.SC_R_CF}", "ja = de ingevulde scenario's ook als lijn in de cashflowgrafiek (Dashboard, en dia 3 als Naar "
                                     "PowerPoint op ja staat) · nee = alleen de Basis in die grafiek", f=F_NOTE8,
        al=Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1))
    ws.row_dimensions[LY.SC_R_CF].height = 24
    put(ws, f"B{LY.SC_R_FOUT}", "Controle", f=F_NOTE8, al=AL_LEFT)
    ws.row_dimensions[LY.SC_R_FOUT].height = 24
    for s in range(1, N_VS + 1):
        cel = _paar(ws, LY.SC_R_FOUT, s, f'=IF({S(s, "fout")}="",IF({aan(s)}=1,"✓ ingevuld",""),{S(s, "fout")})', f=font(8, False, GREEN),
                    al=Alignment(horizontal="right", vertical="top", wrap_text=True))
        ws.conditional_formatting.add(cel, FormulaRule(formula=[f'{S(s, "fout")}<>""'], font=Font(name=ARIAL, color=RED, bold=True)))

    # ---- 2 · uitkomst ---------------------------------------------------------------------------------------------------------------
    put(ws, f"B{R_S2}", "2 · UITKOMST PER SCENARIO · rekent direct mee · € mln", f=F_SECTION)
    _kop_scenarios(ws, R_S2 + 1)
    norm = dabs("norm")
    be_txt = lambda s: (f'=IF({S(s, "d_be")}="","–",IF({S(s, "d_be")}=0,"gelijk",{kw_tekst("ABS(" + S(s, "d_be") + ")")}'  # noqa: E731
                        f'&IF({S(s, "d_be")}>0," later"," eerder")))')
    uitkomst = [
        ("Start bouw", f"={kw_lbl(hm('start_idx'))}", lambda s: f"={kw_lbl(S(s, 'sb'))}", None, False),
        ("Verkoop vanaf", f"={kw_lbl(hm('vs_vv_plan'))}", lambda s: f"={kw_lbl(S(s, 'vv'))}", None, False),
        ("Eerste transport", f"={kw_lbl(hm('vs_tr_plan'))}", lambda s: f"={S(s, 'et')}", None, False),
        ("Grondtermijn", f"={grond_plan}", lambda s: f'=IF({S(s, "gr")}<0,{grond_plan},{S(s, "gr")})', "0%", False),
        (f'="Verkocht bij start bouw (norm "&FIXED({norm}*100,0)&"%)"', f"={hm('verk_bij_start_pct')}", lambda s: f"={S(s, 'vbsb_pct')}", "0%", "norm"),
        ("Max. financieringsbehoefte", f"={hm('dal_basis')}/1000000", lambda s: f"={S(s, 'dal')}/1000000", FMT_MLN, True),
        ("Dieptepunt in", f"={hm('dal_basis_kw')}", lambda s: f"={S(s, 'dal_kw')}", None, False),
        ("Break-even", f'=IF({hm("be_basis")}="","niet bereikt",{hm("be_basis")})', lambda s: f'=IF({S(s, "be")}="","niet bereikt",{S(s, "be")})', None, True),
        ("Rente (berekend)", f"={hm('rente_basis_tot')}/1000000", lambda s: f"={S(s, 'rente')}/1000000", FMT_MLN, False),
        ("Eindsaldo", f"={hm('eind_basis')}/1000000", lambda s: f"={S(s, 'eind')}/1000000", FMT_MLN, False),
        ("Extra financiering t.o.v. Basis", "–", lambda s: f"={S(s, 'd_fin')}/1000000", FMT_MLN, False),
        ("Break-even t.o.v. Basis", "–", be_txt, None, False),
    ]
    r = R_S2 + 3
    for label, basis, scen, nf, soort in uitkomst:
        vet = soort is True
        put(ws, f"B{r}", label, f=F_BOLD9 if vet else F_CALC9, al=AL_LEFT, bd=BD_RIJ)
        cellen = [_paar(ws, r, 0, basis, f=F_BOLD9 if vet else F_CALC9, nf=nf, bd=BD_RIJ)]
        for s in range(1, N_VS + 1):
            f = scen(s)
            cel = _paar(ws, r, s, f'=IF({aan(s)}=1,{f[1:]},"–")', f=F_BOLD9 if vet else F_CALC9, nf=nf, bd=BD_RIJ)
            ws.conditional_formatting.add(cel, FormulaRule(formula=[f"{aan(s)}=0"], font=Font(name=ARIAL, color=LICHT)))
            cellen.append(cel)
        if soort == "norm":
            for k, cel in enumerate(cellen):
                pct = hm("verk_bij_start_pct") if k == 0 else S(k, "vbsb_pct")
                voorw = "" if k == 0 else f"{aan(k)}=1,"
                ws.conditional_formatting.add(cel, FormulaRule(formula=[f"AND({voorw}{pct}>={norm})"], font=Font(name=ARIAL, bold=True, color=GREEN)))
                ws.conditional_formatting.add(cel, FormulaRule(formula=[f"AND({voorw}{pct}<{norm})"], font=Font(name=ARIAL, bold=True, color=RED)))
        r += 1

    # ---- 3 · grafiek (rechts): legenda als gekleurde blokjes, grafiek eronder (charts.py zet hem op GRAFIEK_ANKER) ----------------
    put(ws, f"{L(GRAFIEK_COL)}{R_S1}", "3 · CUMULATIEVE CASHFLOW PER SCENARIO · stip = dieptepunt · grijs = gerealiseerd · € mln", f=F_SECTION)
    ws.row_dimensions[CHIP_RIJ].height = 16
    for k in range(N_VS + 1):
        c0 = GRAFIEK_COL + 3 * k
        ws.merge_cells(start_row=CHIP_RIJ, start_column=c0, end_row=CHIP_RIJ, end_column=c0 + 2)
        tekst = f"={hm('vs0_leg')}" if k == 0 else f'=IF({aan(k)}=1,{S(k, "leg")},"")'
        cel = f"{L(c0)}{CHIP_RIJ}"
        put(ws, cel, tekst, f=font(8, True, "FFFFFF"), fl=fill(LY.VS_KLEUREN[k]), al=Alignment(horizontal="center", vertical="center", shrink_to_fit=True))
        for c in range(c0, c0 + 3):
            ws.cell(CHIP_RIJ, c).fill = fill(LY.VS_KLEUREN[k])
        if k:
            ws.conditional_formatting.add(f"{cel}:{L(c0 + 2)}{CHIP_RIJ}", FormulaRule(formula=[f"{aan(k)}=0"],
                                          fill=PatternFill(fill_type="solid", bgColor="FFFFFF", fgColor="FFFFFF")))

    # ---- 4 · opbrengsten per termijn en per kwartaal ------------------------------------------------------------------------------
    r0 = R_S4
    put(ws, f"B{r0}", "4 · OPBRENGSTEN PER TERMIJN EN PER KWARTAAL · zoals FO 'CF - opbrengsten' · x € 1.000", f=F_SECTION)
    put(ws, f"B{r0 + 1}", "Scenario", f=F_CALC9, al=AL_RIGHT)
    ws.merge_cells(f"{LY.SC_CEL_SEL}:F{r0 + 1}")
    put(ws, LY.SC_CEL_SEL, params.get("vs_sel") or "Scenario 1", f=F_INPUT, fl=FL_INPUT, al=AL_LEFT)
    for c in range(4, 7):
        ws.cell(r0 + 1, c).fill = FL_INPUT
    dv_sel = DataValidation(type="list", formula1='"Basis,' + ",".join(f"Scenario {s}" for s in range(1, N_VS + 1)) + '"', allow_blank=False)
    ws.add_data_validation(dv_sel)
    dv_sel.add(LY.SC_CEL_SEL)
    put(ws, f"H{r0 + 1}", "Woningtype", f=F_CALC9, al=AL_RIGHT)
    ws.merge_cells(f"{LY.SC_CEL_TYPE}:L{r0 + 1}")
    put(ws, LY.SC_CEL_TYPE, params.get("vs_type") or "alle types", f=F_INPUT, fl=FL_INPUT, al=AL_LEFT)
    for c in range(10, 13):
        ws.cell(r0 + 1, c).fill = FL_INPUT
    # keuzelijst woningtypes: hulpkolom (verborgen) met 'alle types' en de typenamen
    hk = L(3 + LY.SC_N_KW + 3)
    ws.column_dimensions[hk].hidden = True
    put(ws, f"{hk}{r0}", "keuzelijst woningtype", f=F_NOTE8)
    put(ws, f"{hk}{r0 + 1}", "alle types", f=F_NOTE8)
    for k in range(1, N_TYPES + 1):
        put(ws, f"{hk}{r0 + 1 + k}", f'=IF(N(Model!${m_col("vcum", k)}$6)>0,Model!${m_col("gv", k)}$7,"")', f=F_NOTE8)
    dv_type = DataValidation(type="list", formula1=f"${hk}${r0 + 1}:${hk}${r0 + 1 + N_TYPES}", allow_blank=True)
    ws.add_data_validation(dv_type)
    dv_type.add(LY.SC_CEL_TYPE)
    put(ws, f"N{r0 + 1}", f'="← kies scenario en woningtype · '
                         f'rijnamen van "&IF({hm("vs_typesel")}=0,"het eerste type met die termijn","het gekozen type")', f=F_NOTE8, al=AL_VCENTER)

    r_pos, r_mark, r_kop = r0 + 2, r0 + 3, r0 + 4
    ws.row_dimensions[r_pos].hidden = True
    put(ws, f"B{r_pos}", "positie in het Model (verborgen hulprij)", f=F_NOTE8)
    sel = hm("vs_sel")
    sb_sel = f"CHOOSE({sel}+1,{hm('start_idx')}," + ",".join(S(s, "sb") for s in range(1, N_VS + 1)) + ")"
    vv_sel = f"CHOOSE({sel}+1,{hm('vs_vv_plan')}," + ",".join(S(s, "vv") for s in range(1, N_VS + 1)) + ")"
    et_sel = (f"CHOOSE({sel}+1,{hm('vs_tr_plan')}," + ",".join(f'IF({S(s, "et_pos")}=0,99999,INDEX(Model!${IDX}${ROW1}:${IDX}${ROWN},{S(s, "et_pos")}))'
                                                             for s in range(1, N_VS + 1)) + ")")
    model = lambda col, pos: f"INDEX(Model!${col}${ROW1}:${col}${ROWN},{pos})"  # noqa: E731
    put(ws, f"B{r_kop}", "x € 1.000", f=F_HDR, fl=FL_HDR, al=AL_LEFT)
    for j in range(LY.SC_N_KW):
        c = L(3 + j)
        put(ws, f"{c}{r_pos}", f'=IF({hm("vp_start")}+{j}>{hm("n")},"",{hm("vp_start")}+{j})', f=F_NOTE8)
        pos = f"{c}${r_pos}"
        idx = model(IDX, pos)
        put(ws, f"{c}{r_mark}", f'=IF({pos}="","",MID(IF({idx}={sb_sel}," · start bouw","")&IF({idx}={vv_sel}," · verkoop","")'
                                f'&IF({idx}={et_sel}," · 1e transport",""),4,60))', f=font(7, True, "B5541F"),
            al=Alignment(horizontal="center", vertical="bottom", wrap_text=True))
        put(ws, f"{c}{r_kop}", f'=IF({pos}="","",{model(KWV, pos)})', f=F_HDR, fl=FL_HDR, al=AL_RIGHT)
    ws.row_dimensions[r_mark].height = 22
    ws.conditional_formatting.add(f"C{r_kop}:{L(2 + LY.SC_N_KW)}{r_kop}",
                                  FormulaRule(formula=[f'AND(C${r_pos}<>"",C${r_pos}<={hm("pos_actuals")})'],
                                              fill=PatternFill(fill_type="solid", bgColor=GREY_HDR, fgColor=GREY_HDR)))

    def verschil(col, pos, deel=1000):
        """per kwartaal uit een cumulatieve Modelkolom"""
        return f'=IF({pos}="","",({model(col, pos)}-IF({pos}=1,0,{model(col, pos + "-1")}))/{deel})'

    tabel = [("Verkocht (aantal woningen)", LY.SEL["vj"], 1, '0;−0;"–"', None),
             ("Getransporteerd (aantal woningen)", LY.SEL["trj"], 1, '0;−0;"–"', None),
             None,
             ("Grondtermijn (koopsom, bij transport)", LY.SEL["gj"], 1000, None, None)]
    for j in range(1, LY.N_TERMIJNEN + 1):
        tabel.append((f"=Model!${LY.TM['label']}${LY.TM_R_A + j - 1}", LY.SEL[f"tj{j}"], 1000, None, LY.SEL[f"ih{j}"]))
    tabel += [("Extra's (kopersmeerwerk, kadaster, overige)", LY.SEL["xj"], 1000, None, None),
              ("Fees (DAEB)", LY.SEL["fj"], 1000, None, None)]
    FMT_K = f'#,##0;{MINUS}#,##0;"–"'
    r = r_kop + 1
    r_eerste = None
    rijen_koper = []                                       # rijen van de bouwtermijnen van de koper (voor koper - aannemer)
    for regel in tabel:
        if regel is None:
            r += 1
            continue
        label, col, deel, nf, ih = regel
        if ih:
            rijen_koper.append(r)
        put(ws, f"B{r}", label, f=F_CALC, al=AL_LEFT)
        if r_eerste is None and deel == 1000:
            r_eerste = r
        for j in range(LY.SC_N_KW):
            c = L(3 + j)
            pos = f"{c}${r_pos}"
            put(ws, f"{c}{r}", verschil(col, pos, deel), f=F_CALC, nf=nf or FMT_K, al=AL_RIGHT)
        if ih:
            rng = f"C{r}:{L(2 + LY.SC_N_KW)}{r}"
            ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(C${r_pos}<>"",N(INDEX(Model!${ih}${ROW1}:${ih}${ROWN},C${r_pos}))>0.5)'],
                                                           fill=PatternFill(fill_type="solid", bgColor="FCE4D6", fgColor="FCE4D6")))
        r += 1
    r_laatste = r - 1
    put(ws, f"B{r}", "Totaal volgens termijnen", f=font(8, True, TXT), al=AL_LEFT)
    for j in range(LY.SC_N_KW):
        c = L(3 + j)
        put(ws, f"{c}{r}", f'=IF({c}${r_pos}="","",SUM({c}{r_eerste}:{c}{r_laatste}))', f=font(8, True, TXT), nf=FMT_K, al=AL_RIGHT,
            bd=Border(top=side("thin", LINE)))
    ws[f"B{r}"].border = Border(top=side("thin", LINE))
    r += 1
    for label, col, teken, vet in (("Opbrengsten in de cashflow", LY.SEL["osel"], 1, False), ("Kosten", LY.SEL["ksel"], -1, False),
                                   ("Cumulatieve cashflow", LY.SEL["ssel"], 1, True)):
        put(ws, f"B{r}", label, f=font(8, vet, TXT), al=AL_LEFT)
        for j in range(LY.SC_N_KW):
            c = L(3 + j)
            put(ws, f"{c}{r}", f'=IF({c}${r_pos}="","",{teken}*{model(col, c + "$" + str(r_pos))}/1000)', f=font(8, vet, TXT), nf=FMT_K, al=AL_RIGHT)
        r += 1
    put(ws, f"B{r}", "oranje = termijn was al vervallen en komt bij transport in één keer binnen (inhaal) · grijs = gerealiseerd · opbrengsten "
                     "in de cashflow = je eigen cashflow plus het verschil volgens de termijnen", f=F_NOTE8)
    # grijze kolommen voor de gerealiseerde kwartalen
    ws.conditional_formatting.add(f"C{r_kop + 1}:{L(2 + LY.SC_N_KW)}{r - 1}",
                                  FormulaRule(formula=[f'AND(C${r_pos}<>"",C${r_pos}<={hm("pos_actuals")})'],
                                              fill=PatternFill(fill_type="solid", bgColor="F2F3F5", fgColor="F2F3F5")))
    # ---- bouwkosten per termijn (aannemer, Woningtypes blok 7) voor hetzelfde scenario en woningtype, naast wat de koper betaalt ----
    r += 2
    put(ws, f"B{r}", "BOUWKOSTEN PER TERMIJN (AANNEMER) · zelfde scenario en woningtype · tab Woningtypes blok 7 · schuift mee met de "
                     "start bouw · x € 1.000", f=F_SECTION)
    r += 1
    r_bk1 = r
    tsel_ = hm("vs_typesel")
    ia = hm("idx_actuals")
    sel_ts = f"Model!{sel_rng('ts', 6, abs_=True)}"                 # start bouw per type in het gekozen scenario
    basis_ts = f"Model!{LY.m_range_abs('verv', 6)}"                 # start bouw per type in de Basis
    gekozen = f"Model!{tm_rng(LY.TM_R_SEL)}"
    gekozen_nd = f"Model!{tm_rng(LY.TM_R_SELND)}"
    blok_a = f"Model!${LY.TM[1]}${LY.TM_R_BKA}:${LY.TM[N_TYPES]}${LY.TM_R_BKA + LY.N_TERMIJNEN - 1}"
    blok_k = f"Model!${LY.TM[1]}${LY.TM_R_BKK}:${LY.TM[N_TYPES]}${LY.TM_R_BKK + LY.N_TERMIJNEN - 1}"

    def kwartaal(ts, tb, kw):
        """kwartaalindex van een deel in het gekozen scenario: in de Basis start bouw + kw - 1; wat daar t/m actuals valt
        (gerealiseerd) blijft staan, de rest schuift mee met de start bouw van het scenario, niet vóór het kwartaal na actuals"""
        pb = f"({tb}+{kw}-1)"
        q = f"({pb}+({pb}>{ia})*({ts}-{tb}))"
        return f"({q}+({pb}>{ia})*({q}<={ia})*({ia}+1-{q}))"

    # alle types: één rij per type (de termijnen van de typen lopen niet gelijk op); één type: één rij per termijn
    for j in range(1, LY.N_TERMIJNEN + 1):
        bedrag, kw = f"Model!{tm_rng(LY.TM_R_BKA + j - 1)}", f"Model!{tm_rng(LY.TM_R_BKK + j - 1)}"
        bed_k = f"Model!${LY.TM[j]}${LY.TM_R_BKA}:${LY.TM[j]}${LY.TM_R_BKA + LY.N_TERMIJNEN - 1}"
        kw_k = f"Model!${LY.TM[j]}${LY.TM_R_BKK}:${LY.TM[j]}${LY.TM_R_BKK + LY.N_TERMIJNEN - 1}"
        ts_k, tb_k = f"Model!${LY.SEL[('ts', j)]}$6", f"Model!${m_col('verv', j)}$6"
        put(ws, f"B{r}", f'=IF({tsel_}=0,IF(SUM({bed_k})=0,"",Model!${m_col("gv", j)}$7&""),Model!${LY.TM["label"]}${LY.TM_R_BKA + j - 1})',
            f=F_CALC, al=AL_LEFT)
        for jj in range(LY.SC_N_KW):
            c = L(3 + jj)
            idx = model(IDX, f"{c}${r_pos}")
            per_type = f"SUMPRODUCT({bed_k}*({kwartaal(ts_k, tb_k, kw_k)}={idx}))"
            per_termijn = f"SUMPRODUCT({gekozen}*{bedrag}*({kwartaal(sel_ts, basis_ts, kw)}={idx}))"
            put(ws, f"{c}{r}", f'=IF(OR({c}${r_pos}="",$B{r}=""),"",IF({tsel_}=0,{per_type},{per_termijn})/1000)', f=F_CALC, nf=FMT_K,
                al=AL_RIGHT)
        r += 1
    r_bkn = r - 1

    def aannemer_nd(c, op):
        """bouwkosten van de gekozen niet-DAEB-typen in (op '=') of t/m (op '<=') het kwartaal van kolom c, x 1.000"""
        return f"SUMPRODUCT({gekozen_nd}*{blok_a}*({kwartaal(sel_ts, basis_ts, blok_k)}{op}{model(IDX, c + '$' + str(r_pos))}))/1000"

    regels = [("Totaal bouwkosten volgens termijnen", lambda c: f"SUM({c}{r_bk1}:{c}{r_bkn})", True),
              ("Bouwtermijnen koper − aannemer", lambda c: f"SUM({c}{rijen_koper[0]}:{c}{rijen_koper[-1]})-{aannemer_nd(c, '=')}", False)]
    for label, formule, vet in regels:
        put(ws, f"B{r}", label, f=font(8, vet, TXT), al=AL_LEFT, bd=Border(top=side("thin", LINE)) if vet else None)
        for jj in range(LY.SC_N_KW):
            c = L(3 + jj)
            put(ws, f"{c}{r}", f'=IF({c}${r_pos}="","",{formule(c)})', f=font(8, vet, TXT), nf=FMT_K, al=AL_RIGHT,
                bd=Border(top=side("thin", LINE)) if vet else None)
        r += 1
    # cumulatief vanaf het begin van het project: ook wat vóór deze tabel valt (de SEL-kolommen van de koper zijn al cumulatief)
    put(ws, f"B{r}", "  idem, cumulatief vanaf het begin", f=font(8, True, TXT), al=AL_LEFT)
    for jj in range(LY.SC_N_KW):
        c = L(3 + jj)
        pos = f"{c}${r_pos}"
        koper_cum = "+".join(f"N({model(LY.SEL[f'tj{j}'], pos)})" for j in range(1, LY.N_TERMIJNEN + 1))
        put(ws, f"{c}{r}", f'=IF({pos}="","",({koper_cum})/1000-{aannemer_nd(c, "<=")})', f=font(8, True, TXT), nf=FMT_K, al=AL_RIGHT)
    ws.conditional_formatting.add(f"C{r}:{L(2 + LY.SC_N_KW)}{r}", FormulaRule(formula=[f'AND(ISNUMBER(C{r}),C{r}<-0.5)'], font=Font(name=ARIAL, bold=True, color=RED)))
    r += 1
    ws.conditional_formatting.add(f"C{r_bk1}:{L(2 + LY.SC_N_KW)}{r - 1}",
                                  FormulaRule(formula=[f'AND(C${r_pos}<>"",C${r_pos}<={hm("pos_actuals")})'],
                                              fill=PatternFill(fill_type="solid", bgColor="F2F3F5", fgColor="F2F3F5")))
    put(ws, f"B{r}", "alle types: een rij per type · één type: een rij per termijn · gerealiseerde delen blijven in hun kwartaal · aannemer = "
                     "bouwkosten in hun bouwkwartaal", f=F_NOTE8)
    r += 1
    put(ws, f"B{r}", "koper = de bouwtermijnen in de tabel hierboven (zonder grond, extra's en fees) · koper − aannemer zonder DAEB-typen (die "
                     "hebben geen bouwtermijnen van kopers) · rood = de bouw is tot dan voorgefinancierd", f=F_NOTE8)
    r += 2
    for titel, basis_col, naam in (("VERGELIJKING · opbrengsten in de cashflow per kwartaal · x € 1.000", OPB, "opbr"),
                                   ("VERGELIJKING · cumulatieve cashflow · x € 1.000", M["stand_basis"], "stand")):
        put(ws, f"B{r}", titel, f=F_SECTION)
        r += 1
        for k in range(N_VS + 1):
            ws.cell(r, 1).fill = fill(LY.VS_KLEUREN[k])
            label = f"={hm('vs0_leg')}" if k == 0 else f'=IF({aan(k)}=1,{S(k, "leg")},"Scenario {k} (leeg)")'
            put(ws, f"B{r}", label, f=F_CALC, al=AL_LEFT, bd=BD_RIJ)
            col = basis_col if k == 0 else vs_col(k, naam)
            for j in range(LY.SC_N_KW):
                c = L(3 + j)
                v = f"{model(col, c + '$' + str(r_pos))}/1000"
                formule = f'=IF({c}${r_pos}="","",{v})' if k == 0 else f'=IF(OR({c}${r_pos}="",{aan(k)}=0),"",{v})'
                put(ws, f"{c}{r}", formule, f=F_CALC, nf=FMT_K, al=AL_RIGHT, bd=BD_RIJ)
            if k:
                ws.conditional_formatting.add(f"B{r}", FormulaRule(formula=[f"{aan(k)}=0"], font=Font(name=ARIAL, color=LICHT)))
            r += 1
        r += 1
    ws.freeze_panes = None
    ws.print_area = f"A1:{L(2 + LY.SC_N_KW)}{r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


GRAFIEK_ANKER = f"{L(GRAFIEK_COL)}{CHIP_RIJ + 1}"
