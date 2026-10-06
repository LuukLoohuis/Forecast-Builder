"""Opmaak: dezelfde kleuren en lettertypen als het bestaande werkboek."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ARIAL = "Arial"
# kleuren
NAVY = "17365D"        # koppen
BLUE_DARK = "002060"   # basis
RED = "C0392F"
GREEN = "12864E"
TXT = "14171A"
TXT2 = "5B6169"
TXT3 = "6E747C"
GREY_HDR = "9AA3AD"
GREY_TXT = "B4BAC1"
CARD = "F7F8FA"
LINE = "C9CDD3"
LINE2 = "E3E6EA"
INPUT_BLUE_FONT = "0000FF"
INPUT_BLUE_FILL = "F3F8FF"
INPUT_YELLOW_FONT = "8A6100"
INPUT_YELLOW_FILL = "FFF2B2"
PP_BORDER = "2F6FDE"

FMT_INT = '#,##0;\\-#,##0;"-"'
FMT_MLN = '"€"0.0"M";"−€"0.0"M";"€"0.0"M"'
FMT_PCT = "0.0%"
FMT_PCT0 = "0%"
FMT_KW = "\\+0;\\−0;0"
FMT_PCT_SIGN = "\\+0%;\\−0%;0%"


def font(size=9, bold=False, color=TXT, italic=False):
    return Font(name=ARIAL, size=size, bold=bold, italic=italic, color=color)


def fill(color):
    return PatternFill(fill_type="solid", fgColor=color)


def side(style="thin", color=None):
    return Side(style=style, color=color)


def put(ws, ref, value=None, *, f=None, fl=None, nf=None, al=None, bd=None):
    """Schrijft een waarde/formule en opmaak in één keer."""
    c = ws[ref]
    if value is not None:
        c.value = value
    if f is not None:
        c.font = f
    if fl is not None:
        c.fill = fl
    if nf is not None:
        c.number_format = nf
    if al is not None:
        c.alignment = al
    if bd is not None:
        c.border = bd
    return c


def style(ws, ref, *, f=None, fl=None, nf=None, al=None, bd=None):
    return put(ws, ref, None, f=f, fl=fl, nf=nf, al=al, bd=bd)


# veelgebruikte combinaties
F_TITLE = font(12, True, TXT)
F_NOTE = font(9, False, TXT3)
F_NOTE8 = font(8, False, TXT3)
F_SECTION = font(9, True, TXT2)
F_HDR = font(9, True, "FFFFFF")
F_HDR8 = font(8, True, "FFFFFF")
F_INPUT = font(9, False, INPUT_BLUE_FONT)
F_GREY = font(9, False, GREY_TXT)
F_CALC = font(8, False, TXT)
F_CALC9 = font(9, False, TXT)
F_BOLD9 = font(9, True, TXT)
F_YELLOW = font(11, True, INPUT_YELLOW_FONT)

FL_HDR = fill(NAVY)
FL_HDR_GREY = fill(GREY_HDR)
FL_INPUT = fill(INPUT_BLUE_FILL)
FL_YELLOW = fill(INPUT_YELLOW_FILL)
FL_CARD = fill(CARD)

AL_LEFT = Alignment(horizontal="left", vertical="center")
AL_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
AL_RIGHT = Alignment(horizontal="right", vertical="center")
AL_RIGHT_WRAP = Alignment(horizontal="right", vertical="center", wrap_text=True)
AL_CENTER = Alignment(horizontal="center", vertical="center")
AL_CENTER_CONT = Alignment(horizontal="centerContinuous", vertical="center")
AL_VCENTER = Alignment(vertical="center")
AL_LEFT_TOP = Alignment(horizontal="left")

BD_SECTION = Border(bottom=side("thin", LINE))
BD_ROW = Border(bottom=side("thin", LINE2))
