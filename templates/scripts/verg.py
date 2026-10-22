import sys; sys.path.insert(0, sys.path[0])
from hstyle import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.chart import BarChart, Reference

SWKC, WBC = PETROL, "C98A00"
wb = Workbook()
v = wb.active; v.title = "Vergelijking"
for col, w in zip("BCDEFGH", [38, 16, 16, 20, 4, 40, 16]): v.column_dimensions[col].width = w
brand(v, 8, "Vergelijking SWK  vs.  Woningborg  |  Garantie- en Waarborgregeling",
      "Blauwe cijfers op geel = invoer  ·  Tarieven staan op de tabbladen 'Tarieven SWK' en 'Tarieven Woningborg'")

def name(n, ref): wb.defined_names[n] = DefinedName(n, attr_text=ref)

# KPI-tegels
tile(v, 5, 2, 2, "TOTAAL SWK", "=C25", EUR, SWKC, SWKC)
tile(v, 5, 3, 4, "TOTAAL WONINGBORG", "=D25", EUR, WBC, WBC)
tile(v, 5, 5, 6, "VERSCHIL", "=ABS(E25)", EUR)
tile(v, 5, 7, 8, "VOORDELIGST", '=IF(C25=D25,"Gelijk",IF(C25<D25,"SWK","Woningborg"))', "General")
v["E8"] = '=IF(MAX(C25:D25)=0,"",TEXT(ABS(E25)/MAX(C25:D25),"0.0%")&" goedkoper")'
v["E8"].font = font(size=9, italic=True, color=MUTED); v["E8"].alignment = Alignment(horizontal="center")
v.conditional_formatting.add("G6", FormulaRule(formula=['G6="SWK"'], font=Font(color=SWKC, bold=True)))
v.conditional_formatting.add("G6", FormulaRule(formula=['G6="Woningborg"'], font=Font(color=WBC, bold=True)))

# Invoer
bar(v, 10, 2, 4, "PROJECT"); bar(v, 10, 7, 8, "KEUZES")
v["B11"] = "Projectnaam"; v.merge_cells("C11:D11"); inp(v["C11"]); v["C11"] = "Voorbeeldproject"
v["D11"].fill = fill(INPUT); v["D11"].border = BOTTOM
v["B12"] = "Aantal woningen"; inp(v["D12"], "#,##0"); v["D12"] = 1
v["B13"] = "Aanneemsom (excl. btw)"; inp(v["D13"], EUR); v["D13"] = 420000
v["B14"] = '="Btw ("&TEXT(C14,"0%")&")"'; inp(v["C14"], "0%", center=True); v["C14"] = 0.21
v["D14"] = "=D13*C14"; calc(v["D14"], EUR)
v["B15"] = "Aanneemsom (incl. btw)"; v["B15"].font = font(bold=True, color=NAVY)
v["D15"] = "=D13+D14"; calc(v["D15"], EUR, bold=True)
for r in (11, 12, 13, 14): v.cell(r, 2).font = font()
keuzes = [(11, "Hoogbouw (> 13 bouwlagen)", "KeuzeHoogbouw"),
          (12, "Zakelijke opdrachtgever (verlengde garantie)", "KeuzeZakelijk"),
          (13, "Prognose bouwtijd > 3 jaar", "KeuzeBouwtijd")]
dvj = DataValidation(type="list", formula1='"ja,nee"', allow_blank=False); v.add_data_validation(dvj)
for r, lab, n in keuzes:
    v.cell(r, 7, lab).font = font(); c = v.cell(r, 8, "nee"); inp(c, center=True); dvj.add(c.coordinate)
    name(n, f"Vergelijking!$H${r}")
name("AantalWoningen", "Vergelijking!$D$12"); name("AanneemsomIncl", "Vergelijking!$D$15")

# Vergelijkingstabel
bar(v, 18, 2, 8, "KOSTENVERGELIJKING")
hdr = ["Kostenpost", "SWK", "Woningborg", "Verschil (WB − SWK)", "", "Voordeligst", ""]
for i, h in enumerate(hdr):
    c = v.cell(19, 2 + i, h); c.font = font(size=9, bold=True, color="FFFFFF" if i in (1, 2) else NAVY)
    c.alignment = Alignment(horizontal="left" if i == 0 else "right" if i < 4 else "center", wrap_text=True)
    c.border = Border(bottom=Side(style="medium", color=NAVY))
v["C19"].fill = fill(SWKC); v["D19"].fill = fill(WBC)
posts = [(20, "Provisie Sectie I (Afbouw)", "F12"), (21, "Provisie Sectie II (Kwaliteit)", "F13"),
         (22, "Opslagen", "H23"), (23, "Overige kosten", "H33")]
# Provisies verwijzen naar de bedragkolom G van de tariefbladen
posts = [(r, l, {"F12": "G12", "F13": "G13"}.get(x, x)) for r, l, x in posts]
for r, lab, ref in posts:
    v.cell(r, 2, lab).font = font()
    calc(v.cell(r, 3, f"='Tarieven SWK'!{ref}"), EUR, color="008000")
    calc(v.cell(r, 4, f"='Tarieven Woningborg'!{ref}"), EUR, color="008000")
    calc(v.cell(r, 5, f"=D{r}-C{r}"), EUR, bg=None)
    c = v.cell(r, 7, f'=IF(C{r}=D{r},"–",IF(C{r}<D{r},"SWK","Woningborg"))'); c.font = font(size=9, bold=True); c.alignment = Alignment(horizontal="center")
rows = [(26, "Kosten per woning", "=IFERROR(C25/AantalWoningen,0)", "=IFERROR(D25/AantalWoningen,0)", False),
        (27, "In % van aanneemsom incl. btw", "=IFERROR(C25/AanneemsomIncl,0)", "=IFERROR(D25/AanneemsomIncl,0)", False)]
for r, lab, fs, fw, b in rows:
    v.cell(r, 2, lab).font = font(bold=b, color=NAVY if b else "000000")
    fmt = "0.00%" if r == 27 else EUR
    calc(v.cell(r, 3, fs), fmt, bold=b); calc(v.cell(r, 4, fw), fmt, bold=b)
    calc(v.cell(r, 5, f"=D{r}-C{r}"), fmt, bold=b, bg=None)
# Totaal
for c in range(2, 9):
    v.cell(25, c).fill = fill(GOLD)
    v.cell(25, c).border = Border(top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
v["B25"] = "Totaal kosten"; v["B25"].font = font(size=11, bold=True, color=NAVY)
for col, f in (("C", "=SUM(C20:C23)"), ("D", "=SUM(D20:D23)"), ("E", "=D25-C25")):
    v[f"{col}25"] = f; v[f"{col}25"].font = font(size=11, bold=True, color=NAVY); v[f"{col}25"].number_format = EUR
v["G25"] = '=IF(C25=D25,"–",IF(C25<D25,"SWK","Woningborg"))'; v["G25"].font = font(size=11, bold=True); v["G25"].alignment = Alignment(horizontal="center")
v.row_dimensions[25].height = 24
for rng, f0 in (("G20:G25", "G20"),):
    v.conditional_formatting.add(rng, FormulaRule(formula=[f'{f0}="SWK"'], font=Font(color="FFFFFF", bold=True), fill=fill(SWKC)))
    v.conditional_formatting.add(rng, FormulaRule(formula=[f'{f0}="Woningborg"'], font=Font(color="FFFFFF", bold=True), fill=fill(WBC)))
v.conditional_formatting.add("E20:E27", FormulaRule(formula=["E20<0"], font=Font(color="38761D")))
note(v["B28"], "Positief verschil = Woningborg duurder; negatief = Woningborg goedkoper. Bedragen zoals de tarieflijsten, excl. eventuele btw.")

# Controles
bar(v, 32, 2, 4, "CONTROLES"); bar(v, 32, 7, 8, "AANDACHTSPUNTEN")
checks = [(33, "Invoer compleet", '=IF(AND(D13>0,D12>0),"OK","AANVULLEN")'),
          (34, "Keuzes geldig", '=IF(AND(OR(H11="ja",H11="nee"),OR(H12="ja",H12="nee"),OR(H13="ja",H13="nee")),"OK","CONTROLEREN")'),
          (35, "Tarieven SWK bevestigd", "='Tarieven SWK'!H5"),
          (36, "Tarieven Woningborg bevestigd", "='Tarieven Woningborg'!H5")]
for r, lab, f in checks:
    v.cell(r, 2, lab).font = font(size=9, color=MUTED)
    c = v.cell(r, 4, f); c.font = font(size=9, bold=True, color=NAVY); c.alignment = Alignment(horizontal="center")
    for cc in range(2, 5): v.cell(r, cc).border = BOTTOM
v.conditional_formatting.add("D33:D36", FormulaRule(formula=['D33="OK"'], fill=fill(OKF)))
v.conditional_formatting.add("D33:D36", FormulaRule(formula=['D33<>"OK"'], fill=fill(BADF)))
pts = ["• Woningborg-tarieven zijn voorbeeldwaarden tot ze zijn bevestigd.",
       "• Technische beoordeling (transformatie, > 25 lagen) niet inbegrepen.",
       "• Alleen volledige regeling; geen separaat verkochte parkeerplaatsen."]
for i, t in enumerate(pts):
    v.merge_cells(start_row=33 + i, start_column=7, end_row=33 + i, end_column=8); note(v.cell(33 + i, 7), t)

ch = BarChart(); ch.type = "bar"; ch.grouping = "clustered"; ch.title = "Kosten per post (excl. btw)"
ch.add_data(Reference(v, min_col=3, max_col=4, min_row=19, max_row=23), titles_from_data=True)
ch.set_categories(Reference(v, min_col=2, min_row=20, max_row=23))
for s, col in zip(ch.series, (SWKC, WBC)):
    s.graphicalProperties.solidFill = col; s.graphicalProperties.line.solidFill = col
ch.x_axis.scaling.orientation = "maxMin"; ch.y_axis.numFmt = '€ #,##0'; ch.y_axis.majorGridlines = None
ch.legend.position = "b"; ch.height = 7; ch.width = 24
v.add_chart(ch, "B39")
for row in v.iter_rows(min_row=4, max_row=37, min_col=2, max_col=8):
    for c in row:
        a = c.alignment; c.alignment = Alignment(horizontal=a.horizontal, vertical="center", indent=a.indent)
    v.row_dimensions[row[0].row].height = v.row_dimensions[row[0].row].height or 19
v.freeze_panes = "A10"
print_setup(v, "A1:H54", "Heijmans Vastgoed | Vergelijking SWK – Woningborg")
v.page_setup.fitToHeight = 1

# ---------- Tariefbladen ----------
COND = '"Altijd,Hoogbouw,Zakelijk,Bouwtijd > 3 jaar,Nooit"'
def active(ref):
    return (f'IF({ref}="Altijd",1,IF({ref}="Hoogbouw",--(KeuzeHoogbouw="ja"),IF({ref}="Zakelijk",--(KeuzeZakelijk="ja"),'
            f'IF({ref}="Bouwtijd > 3 jaar",--(KeuzeBouwtijd="ja"),0))))')

def sheet(title, color, src, ok, basis, opslag, overig, warn=None):
    s = wb.create_sheet(f"Tarieven {title}")
    s.sheet_properties.tabColor = color
    for col, w in zip("BCDEFGHI", [36, 12, 14, 22, 20, 15, 15, 48]): s.column_dimensions[col].width = w
    brand(s, 9, f"Tarieven {title}  |  Garantie- en Waarborgregeling",
          "Blauwe cijfers op geel = tarief/keuze aanpassen  ·  Lichtblauw = berekening (rekent met de invoer op 'Vergelijking')")
    s["B5"] = "Bron / tarieflijst"; inp(s["C5"]); s["C5"] = src; s.merge_cells("C5:F5")
    for cc in range(4, 7): s.cell(5, cc).fill = fill(INPUT); s.cell(5, cc).border = BOTTOM
    s["I5"] = ok; inp(s["I5"], center=True)
    s["H5"] = '=IF(I5="ja","OK","CONTROLEREN")'; calc(s["H5"]); s["H5"].alignment = Alignment(horizontal="center")
    s["G5"] = "Bevestigd?"; s["G5"].font = font(bold=True); s["G5"].alignment = Alignment(horizontal="right")
    dv = DataValidation(type="list", formula1='"ja,nee"'); s.add_data_validation(dv); dv.add("I5")
    s.conditional_formatting.add("H5", FormulaRule(formula=['H5="OK"'], fill=fill(OKF)))
    s.conditional_formatting.add("H5", FormulaRule(formula=['H5<>"OK"'], fill=fill(BADF)))
    if warn:
        s.merge_cells("B7:I7"); s["B7"] = warn
        s["B7"].font = font(bold=True, color="FFFFFF"); s["B7"].fill = fill("C0392B"); s["B7"].alignment = Alignment(vertical="center", indent=1)
        s.row_dimensions[7].height = 22

    def heads(r, labels):
        for i, h in enumerate(labels):
            c = s.cell(r, 2 + i, h); c.font = font(size=9, bold=True, color=NAVY)
            c.alignment = Alignment(horizontal="left" if i in (0, len(labels) - 1) else "right")
            c.border = Border(bottom=Side(style="medium", color=NAVY))
    # Basis
    bar(s, 9, 2, 9, "BASISPROVISIES", color)
    heads(10, ["Kostenpost", "Tarief", "Grondslag %", "", "Grondslag (€)", "Bedrag (€)", "", "Toelichting"])
    for i, (lab, pct, factor, tip) in enumerate(basis):
        r = 12 + i
        s.cell(r, 2, lab).font = font()
        inp(s.cell(r, 3, pct), PCT4); inp(s.cell(r, 4, factor), "0%")
        calc(s.cell(r, 6, f"=AanneemsomIncl*D{r}"), EUR, bg=None, color=MUTED)
        calc(s.cell(r, 7, f"=F{r}*C{r}"), EUR)
        note(s.cell(r, 9), tip)
    s["B11"] = "(grondslag = % van de aanneemsom incl. btw)"; s["B11"].font = font(size=8, italic=True, color=MUTED)
    # Opslagen
    bar(s, 15, 2, 9, "OPSLAGEN", color)
    heads(16, ["Omschrijving", "Tarief", "Berekend over", "Van toepassing bij", "Grondslag (€)", "Actief", "Bedrag (€)", "Toelichting"])
    dvo = DataValidation(type="list", formula1='"Sectie I,Sectie II,Aanneemsom incl. btw"'); s.add_data_validation(dvo)
    dvc = DataValidation(type="list", formula1=COND); s.add_data_validation(dvc)
    for i in range(6):
        r = 17 + i
        lab, pct, over, cond, tip = opslag[i] if i < len(opslag) else (None, None, None, None, None)
        inp(s.cell(r, 2, lab)); s.cell(r, 2).font = font(color="0000FF")
        inp(s.cell(r, 3, pct), "0.0%"); inp(s.cell(r, 4, over)); inp(s.cell(r, 5, cond))
        dvo.add(f"D{r}"); dvc.add(f"E{r}")
        calc(s.cell(r, 6, f'=IF(D{r}="Sectie I",$G$12,IF(D{r}="Sectie II",$G$13,IF(D{r}="Aanneemsom incl. btw",AanneemsomIncl,0)))'), EUR, bg=None, color=MUTED)
        c = s.cell(r, 7, f'=IF(B{r}="","",IF({active(f"E{r}")}=1,"ja","nee"))'); calc(c, bg=None, color=MUTED); c.alignment = Alignment(horizontal="center")
        calc(s.cell(r, 8, f'=IF(G{r}="ja",F{r}*C{r},0)'), EUR)
        note(s.cell(r, 9), tip or "")
    s["G23"] = "Totaal"; s["G23"].font = font(bold=True); s["G23"].alignment = Alignment(horizontal="right")
    s["H23"] = "=SUM(H17:H22)"; calc(s["H23"], EUR, bold=True)
    # Overige kosten
    bar(s, 25, 2, 9, "OVERIGE KOSTEN", color)
    heads(26, ["Omschrijving", "Bedrag", "Per", "Van toepassing bij", "Aantal", "Actief", "Bedrag (€)", "Toelichting"])
    dvp = DataValidation(type="list", formula1='"woning,project"'); s.add_data_validation(dvp)
    for i in range(6):
        r = 27 + i
        lab, bed, per, cond, tip = overig[i] if i < len(overig) else (None, None, None, None, None)
        inp(s.cell(r, 2, lab)); s.cell(r, 2).font = font(color="0000FF")
        inp(s.cell(r, 3, bed), EUR); inp(s.cell(r, 4, per)); inp(s.cell(r, 5, cond))
        dvp.add(f"D{r}"); dvc.add(f"E{r}")
        calc(s.cell(r, 6, f'=IF(D{r}="woning",AantalWoningen,IF(D{r}="project",1,0))'), '#,##0', bg=None, color=MUTED)
        c = s.cell(r, 7, f'=IF(B{r}="","",IF({active(f"E{r}")}=1,"ja","nee"))'); calc(c, bg=None, color=MUTED); c.alignment = Alignment(horizontal="center")
        calc(s.cell(r, 8, f'=IF(G{r}="ja",F{r}*C{r},0)'), EUR)
        note(s.cell(r, 9), tip or "")
    s["G33"] = "Totaal"; s["G33"].font = font(bold=True); s["G33"].alignment = Alignment(horizontal="right")
    s["H33"] = "=SUM(H27:H32)"; calc(s["H33"], EUR, bold=True)
    # Totaal
    for cc in range(2, 10):
        s.cell(35, cc).fill = fill(GOLD); s.cell(35, cc).border = Border(top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
    s["B35"] = f"Totaal kosten {title}"; s["B35"].font = font(size=11, bold=True, color=NAVY)
    s["H35"] = "=G12+G13+H23+H33"; s["H35"].font = font(size=11, bold=True, color=NAVY); s["H35"].number_format = EUR
    s.row_dimensions[35].height = 24
    for row in s.iter_rows(min_row=4, max_row=36, min_col=2, max_col=9):
        for c in row:
            a = c.alignment; c.alignment = Alignment(horizontal=a.horizontal, vertical="center", indent=a.indent)
        s.row_dimensions[row[0].row].height = s.row_dimensions[row[0].row].height or 19
    s.freeze_panes = "A9"
    print_setup(s, "A1:I36", f"Heijmans Vastgoed | Tarieven {title}")
    s.page_setup.fitToHeight = 1
    return s

sheet("SWK", SWKC, "SWK Rekenmodel 2026 – Berekening Garantie- en Waarborgregeling", "ja",
      [("Provisie Sectie I (Afbouw)", 0.005198, 0.2, "Over 20% van de aanneemsom incl. btw"),
       ("Provisie Sectie II (Kwaliteit)", 0.00038, 1, "Over de aanneemsom incl. btw")],
      [("Opslag hoogbouw", 0.303, "Sectie II", "Hoogbouw", "> 13 bouwlagen"),
       ("Opslag verlengde garantie zakelijk", 0.1, "Sectie II", "Zakelijk", "Alleen bij zakelijke opdrachtgever"),
       ("Opslag prognose bouwtijd > 3 jaar", 0.1, "Sectie I", "Bouwtijd > 3 jaar", "Start bouw tot 1e oplevering")],
      [("Bureaukosten, voorschot", 750, "woning", "Altijd", "Voorschot per woning")])
sheet("Woningborg", WBC, "VOORBEELD – vul hier de actuele Woningborg-tarieflijst in", "nee",
      [("Provisie insolventiedeel / afbouw", 0.005, 0.2, "VOORBEELDWAARDE – controleren"),
       ("Provisie herstel-/kwaliteitsdeel", 0.0004, 1, "VOORBEELDWAARDE – controleren")],
      [("Opslag hoogbouw", 0.3, "Sectie II", "Hoogbouw", "VOORBEELDWAARDE – controleren"),
       ("Opslag zakelijke opdrachtgever", 0.1, "Sectie II", "Zakelijk", "VOORBEELDWAARDE – controleren"),
       ("Opslag lange bouwtijd", 0.1, "Sectie I", "Bouwtijd > 3 jaar", "VOORBEELDWAARDE – controleren")],
      [("Registratie-/bureaukosten", 700, "woning", "Altijd", "VOORBEELDWAARDE – controleren")],
      warn="⚠  VOORBEELDTARIEVEN – dit zijn géén officiële Woningborg-tarieven. Vervang ze en zet 'Bevestigd' op ja.")


# Toelichting
t = wb.create_sheet("Toelichting")
for col, w in zip("BC", [4, 110]): t.column_dimensions[col].width = w
brand(t, 3, "Toelichting", "Hoe werkt deze vergelijking?")
lines = [("1", "Vul op 'Vergelijking' het project in: aantal woningen, aanneemsom excl. btw, btw-% en de drie keuzes (ja/nee)."),
         ("2", "Controleer op 'Tarieven SWK' en 'Tarieven Woningborg' de percentages en bedragen; zet 'Bevestigd' op ja als ze kloppen."),
         ("3", "Opslagen: kies waarover ze worden berekend (Sectie I, Sectie II of aanneemsom) en wanneer ze gelden (Altijd, of gekoppeld aan een keuze)."),
         ("4", "Overige kosten: bedrag per woning of per project, eventueel gekoppeld aan een keuze."),
         ("", ""),
         ("i", "SWK-tarieven zijn overgenomen uit het rekenmodel 2026 (Sectie I 0,5198% over 20%, Sectie II 0,0380%, opslagen 30,3% / 10% / 10%, bureaukosten € 750)."),
         ("!", "Woningborg-tarieven zijn VOORBEELDWAARDEN in dezelfde opbouw, geen officiële tarieven. De controle op 'Vergelijking' blijft rood tot ze bevestigd zijn."),
         ("i", "Sectie I/II zijn SWK-termen; gebruik bij Woningborg de twee basisregels voor hun equivalent (insolventie-/afbouwdeel en herstel-/kwaliteitsdeel).")]
for i, (a, b) in enumerate(lines):
    t.cell(5 + i, 2, a).font = font(bold=True, color=NAVY); t.cell(5 + i, 3, b).font = font()
    t.row_dimensions[5 + i].height = 20

v.sheet_properties.tabColor = GOLD
wb.save(sys.argv[1])
