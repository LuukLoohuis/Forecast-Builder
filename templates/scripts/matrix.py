import sys, math; sys.path.insert(0, sys.path[0])
from hstyle import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

SWKC, WBC = PETROL, "C98A00"
STAT = {"vergelijkbaar": "D9EAD3", "indicatief": "FDE9B8", "niet vergelijkbaar": "F4CCCC"}
wb = Workbook()

def table(ws, r0, headers, widths, rows, color=PETROL, wrapcols=None, status_col=None):
    for i, (h, w) in enumerate(zip(headers, widths)):
        ws.column_dimensions[L(2 + i)].width = w
        c = ws.cell(r0, 2 + i, h); c.font = font(size=9, bold=True, color="FFFFFF"); c.fill = fill(color)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[r0].height = 30
    for j, row in enumerate(rows):
        r = r0 + 1 + j; lines = 1
        for i, val in enumerate(row):
            c = ws.cell(r, 2 + i, val); c.font = font(size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BOTTOM
            if j % 2: c.fill = fill("F7F8FA")
            if isinstance(val, str):
                lines = max(lines, sum(math.ceil(max(len(p), 1) / max(widths[i] * 1.15, 1)) for p in val.split("\n")))
        ws.row_dimensions[r].height = min(12.5 * lines + 4, 400)
        if status_col is not None:
            s = str(row[status_col]).lower()
            for k, col in STAT.items():
                if s.startswith(k):
                    ws.cell(r, 2 + status_col).fill = fill(col); ws.cell(r, 2 + status_col).font = font(size=9, bold=True)
        for i, val in enumerate(row):
            if isinstance(val, str) and val.lower().startswith(("niet aangetroffen", "projectofferte nodig", "schriftelijke bevestiging nodig")):
                ws.cell(r, 2 + i).font = font(size=9, bold=True, color="B42318")
    ws.auto_filter.ref = f"B{r0}:{L(1 + len(headers))}{r0 + len(rows)}"
    return r0 + len(rows) + 1

def sheet(title, h1, sub, ncols, tab=None):
    ws = wb.create_sheet(title) if wb.active.title != "Sheet" else wb.active
    if ws.title == "Sheet": ws.title = title
    brand(ws, ncols + 1, h1, sub)
    if tab: ws.sheet_properties.tabColor = tab
    return ws

def para(ws, r, text, c2, bold=False, size=10, color="000000", h=None, fillc=None):
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=c2)
    c = ws.cell(r, 2, text); c.font = font(bold=bold, size=size, color=color)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if fillc:
        for cc in range(2, c2 + 1): ws.cell(r, cc).fill = fill(fillc)
    ws.row_dimensions[r].height = h or max(16, 13 * math.ceil(len(text) / 150))

# ===================== 1. Samenvatting =====================
s = sheet("1 Samenvatting", "SWK vs. Woningborg  |  Tarieven- en dekkingsanalyse Heijmans Vastgoed",
          "Peildatum 22-10-2026  ·  Statuslabels: vergelijkbaar / indicatief / niet vergelijkbaar  ·  Bron-ID's verwijzen naar tabblad '2 Bronnen'", 8, GOLD)
for col, w in zip("BCDEFGHI", [26, 18, 18, 18, 18, 18, 18, 18]): s.column_dimensions[col].width = w
r = 5
bar(s, r, 2, 9, "MANAGEMENTSAMENVATTING"); r += 1
summ = [
 ("Conclusie", "Er is op dit moment géén definitieve prijsconclusie mogelijk. Voor geen enkele productcategorie zijn er officiële SWK- en Woningborg-tarieven uit hetzelfde jaar. Woningborg: officiële tariefbrief 2027 voor Heijmans Vastgoed B.V. (B1). SWK: alleen een intern Heijmans-rekenmodel '2026' zonder bronvermelding (B2). Een officiële SWK-tariefbrief is niet aangetroffen."),
 ("Wel vergelijkbaar qua functie", "Particuliere nieuwbouw (koper krijgt insolventie- én herstelwaarborg): SWK Garantie- en waarborgregeling 2024 ↔ Woningborg GWR Nieuwbouw 2024. Status: INDICATIEF. Redenen: ander tariefjaar (2026 vs 2027), SWK-tarief uit intern model, andere grondslag (SWK: aanneemsom; WB 'met grond': koop- én aanneemsom), en andere dekking (insolventie SWK max 100% vs WB max 60% van de aanneemsom; beide via zoekresultaat, te verifiëren)."),
 ("Indicatieve uitkomst", "Rekenvoorbeeld 1 (zelfgekozen invoer, zie tab 6): bij verkoop met grond is Woningborg circa € 184 per woning lager excl. btw. Bij verkoop zonder grond zijn ze nagenoeg gelijk (SWK circa € 8 lager). De SWK-bureaukosten van € 750 per woning wegen daarbij zwaar. Het verschil hangt sterk af van de grondprijs, omdat Woningborg 'met grond' over de grond meerekent en SWK niet. Btw over de SWK-provisie is niet vastgesteld, dus een totaal incl. btw voor SWK ontbreekt."),
 ("Niet vergelijkbaar", "Kavelbouw, zakelijk (alleen insolventie én insolventie + herstel), transformatie (particulier en zakelijk) en houtbouw: voor SWK is geen tarief aangetroffen. Een SWK-tarief 'alleen afbouw' mag niet worden afgeleid door Sectie II weg te laten uit het volledige pakket."),
 ("Hoogbouw", "Woningborg: hoger tarief 'vanaf 13 lagen boven straatpeil' (B1 p.3/5). SWK-rekenmodel: opslag 30,3% over Sectie II bij '> 13 bouwlagen' (B2). Dat is een afwijkende grens: 13 lagen valt bij WB wél, bij SWK niet onder hoogbouw. Deze grens is in officiële SWK-bronnen niet aangetroffen. SWK-documenten noemen hoogbouw alleen als technisch risico bij circa 70 meter (O3). Laat de grens en de telwijze schriftelijk bevestigen."),
 ("Btw", "Woningborg: 'Over alle tarieven wordt btw in rekening gebracht' (B1 p.6), zonder percentage. SWK: niet aangetroffen. Btw is daarom apart gepresenteerd en niet als 0% behandeld. Of de btw aftrekbaar is, is niet onderbouwd en dus niet aangenomen."),
 ("Correctie eerdere modellen", "Het eerder gemaakte vergelijkingstemplate zette de btw over de SWK-provisie op 0% en nam de SWK-hoogbouwgrens (> 13 lagen) over uit het interne model. Beide zijn niet bevestigd. Het template mag pas voor besluitvorming worden gebruikt nadat de vragen op tab 7 zijn beantwoord."),
 ("Eerste acties", "1) Officiële SWK-tariefbrief 2027 voor Heijmans Vastgoed opvragen (incl. zakelijk, afbouw, transformatie, kavelbouw, houtbouw, bijdragen). 2) Btw over de SWK-provisie laten bevestigen. 3) Hoogbouwgrens en telwijze bij beide partijen schriftelijk vastleggen. 4) Obligo per project bij beide financieel screeners opvragen."),
 ("Beperking onderzoek", "De netwerkinstellingen van deze cloudomgeving blokkeerden swk.nl, cms.swk.nl en woningborg.nl. Officiële online regelingen zijn daarom alleen via zoekmachine-uittreksels gecontroleerd (bronstatus 'via zoekresultaat'). Artikelnummers en pagina's van die bronnen zijn daardoor niet vastgesteld en moeten in de PDF worden nagelezen. De twee bijgevoegde documenten zijn wel volledig gelezen."),
]
for k, t in summ:
    s.cell(r, 2, k).font = font(bold=True, color=NAVY); s.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
    s.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
    c = s.cell(r, 3, t); c.font = font(); c.alignment = Alignment(wrap_text=True, vertical="top")
    s.row_dimensions[r].height = 13 * math.ceil(len(t) / 135) + 6
    for cc in range(2, 10): s.cell(r, cc).border = BOTTOM
    r += 1
r += 1
bar(s, r, 2, 9, "STATUS PER VERGELIJKING"); r += 1
r = table(s, r, ["Categorie", "SWK", "Woningborg", "Status", "Reden", "", "", ""], [26, 18, 18, 18, 18, 18, 18, 18], [
 ("Particuliere nieuwbouw met grond", "SWK-1", "WB-1", "indicatief", "Jaar 2026 vs 2027; SWK uit intern model; grondslag en insolventiedekking verschillen"),
 ("Particuliere nieuwbouw zonder grond", "SWK-2", "WB-3", "indicatief", "Zelfde grondslag (aanneemsom incl. btw), maar jaar, bron en dekking verschillen"),
 ("Hoogbouw particulier", "SWK-3", "WB-2 / WB-4", "indicatief", "Grens verschilt (WB ≥ 13 lagen, SWK > 13 bouwlagen intern); telwijze onbekend"),
 ("Kavelbouw", "SWK-4", "WB-5", "niet vergelijkbaar", "SWK-tarief niet aangetroffen"),
 ("Zakelijk – alleen insolventie/afbouw", "SWK-5", "WB-9 / WB-10", "niet vergelijkbaar", "SWK-tarief niet aangetroffen; dekking verschilt (SWK 20%, WB 30%)"),
 ("Zakelijk – insolventie + herstel", "SWK-6", "WB-7 / WB-8", "niet vergelijkbaar", "SWK-tarief en -grondslag zakelijk niet aangetroffen"),
 ("Transformatie particulier", "SWK-7", "WB-6", "niet vergelijkbaar", "SWK-tarief niet aangetroffen"),
 ("Transformatie zakelijk", "SWK-8", "WB-11 / WB-12", "niet vergelijkbaar", "SWK-tarief niet aangetroffen"),
 ("Houtbouw", "SWK-9", "WB-13", "niet vergelijkbaar", "Beide: projectofferte nodig"),
], status_col=3)
for rr in range(r - 10, r): s.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=9)

# ===================== 2. Bronnen =====================
b = sheet("2 Bronnen", "Bronnenregister en afwijkingen", "Officiële en toepasselijke tariefinformatie gaat vóór interne modellen  ·  'Via zoekresultaat' = niet volledig nagelezen", 13)
src_rows = [
 ("B1", "Tarieven 2027 (tariefbrief + bijlage 'Uw tarieven 2027 Woningborg Verzekeren')", "Woningborg N.V. / Woningborg Verzekeren", "Officiële tariefbrief (bedrijfsspecifiek)", "Certificaten Nieuwbouw Particulier, Kavelbouw, Transformatie Particulier, Zakelijk (nieuwbouw en transformatie)", "Heijmans Vastgoed B.V., inschrijfnr. 4889, concernnr. 2000335", "2027", "Plannen geaccepteerd op of na 1-1-2027 (p.1, p.6)", "Planacceptatie door Woningborg; transformatie: haalbaarheidsonderzoek, bouwplaatsbezoeken en geluidsmetingen (p.4)", "18-09-2026", "p.1–2 brief; p.3 particulier; p.4 transformatie + bijkomend; p.5 zakelijk; p.6 grondslag/btw/houtbouw/bijdragen/diensten", "Bijlage gebruiker: Tarief_2027_-_04889_-_2000335.pdf", "Volledig gelezen"),
 ("B2", "Rekenmodel 2026 – SWK Berekening Garantie- en Waarborgregeling (tabblad 'SWK berekening')", "Heijmans Vastgoed (intern); tarieven toegeschreven aan SWK", "Intern rekenmodel", "SWK Garantie- en waarborgregeling (volledig pakket)", "Niet vermeld", "2026 (alleen uit bestandsnaam)", "Niet vermeld", "'Alleen plannen waarop de volledige GWR van toepassing is en geen separaat verkochte parkeerplaatsen' (A26–A27)", "Niet vermeld", "Cellen B10–B18 (tarieven), E10–E16 (toelichting), A23–A27 (aandachtspunten)", "Bijlage gebruiker: Rekenmodel_2026_-_SWK_Berekening...xlsx", "Volledig gelezen; geen bron voor tarieven"),
 ("O1", "SWK Garantie- en waarborgregeling 2024", "SWK", "Waarborgvoorwaarden", "GWR 2024 (insolventie- + herstelwaarborg)", "Algemeen", "n.v.t. (voorwaarden)", "Plannen onder GWR 2024", "Zie regeling", "02-2024", "Artikel/pagina niet vastgesteld", "https://cms.swk.nl/app/uploads/2024/02/SWK-Garantie-en-waarborgregeling-2024.pdf", "Via zoekresultaat"),
 ("O2", "SWK Waarborgregeling Afbouw 2024", "SWK", "Waarborgvoorwaarden", "Alleen afbouwwaarborg, zakelijke opdrachtgever", "Algemeen", "n.v.t.", "Niet vastgesteld", "Zie regeling", "02-2024", "Artikel/pagina niet vastgesteld", "https://cms.swk.nl/app/uploads/2024/02/SWK-Waarborgregeling-Afbouw-2024.pdf", "Via zoekresultaat"),
 ("O3", "Toelichting op planaanmelding (versies 01-01-2024 en 01-03-2025)", "SWK", "Toelichting / procedure", "Planaanmelding GWR", "Algemeen", "2024/2025", "n.v.t.", "Technisch bureau verplicht bij hoogbouw (~70 m), transformatie, renovatie, innovatieve bouw", "01-03-2025", "Niet vastgesteld", "https://cms.swk.nl/app/uploads/2025/03/25-03-01-Toelichting-Planaanmelding-2025.pdf", "Via zoekresultaat"),
 ("O4", "Toelichting op opbouw aanneemsom 2025", "SWK", "Toelichting / definitie", "Definitie aanneemsom; eerste termijn", "Algemeen", "2025", "n.v.t.", "Aanneemsom = bouwkosten + bijkomende kosten; bedragen incl. btw", "01-03-2025", "Niet vastgesteld", "https://cms.swk.nl/app/uploads/2025/03/25-03-01-Toelichting-op-opbouw-aanneemsom-2025.pdf", "Via zoekresultaat"),
 ("O5", "Standaard termijnregeling 2024 (EGH en APP)", "SWK", "Termijnregeling", "Termijnschema's", "Algemeen", "2024", "n.v.t.", "—", "01-01-2024", "Niet gelezen", "https://cms.swk.nl/app/uploads/2023/11/24-01-01-Standaard-termijnregeling-2024-EGH.pdf", "Alleen vindplaats"),
 ("O6", "Productpagina's Huurwoningen / Zakelijk vastgoed", "SWK", "Productinformatie", "Zakelijk: keuze afbouw óf afbouw + herstel", "Algemeen", "n.v.t.", "n.v.t.", "—", "Niet vermeld", "—", "https://swk.nl/producten/huurwoningen/", "Via zoekresultaat"),
 ("O7", "Woningborg Garantie- en waarborgregeling Nieuwbouw 2024 (met bijlage 01-01-2026)", "Woningborg", "Waarborgvoorwaarden", "GWRN 2024", "Algemeen", "n.v.t.", "Plannen aangemeld vanaf 1-1-2024", "Zie regeling", "01-01-2026 (bijlage)", "Artikel/pagina niet vastgesteld", "https://www.woningborg.nl/media/dn5jf1zf/gwrn-2024-met-bijlage-01-01-2026.pdf", "Via zoekresultaat"),
 ("O8", "Woningborg Garantie- en waarborgregeling Transformatie 2024", "Woningborg", "Waarborgvoorwaarden", "GWR Transformatie 2024", "Algemeen", "n.v.t.", "Plannen vanaf 1-1-2024", "Zie regeling", "2024", "Niet vastgesteld", "https://www.woningborg.nl/media/hc0igtte/wb-garantie-en-waarborg-transformatie-2024-incl-bijlage-a.pdf", "Via zoekresultaat"),
 ("O9", "Woningborg zakelijk certificaat – brochure opdrachtgevers (2025) en productpagina zakelijke regeling", "Woningborg", "Productinformatie", "Zakelijk certificaat", "Algemeen", "2025", "n.v.t.", "—", "2025", "Niet vastgesteld", "https://www.woningborg.nl/media/cgrlt3ae/wb_brochure-zakelijk-certificaat_opdrachtgevers-edit-2025.pdf", "Via zoekresultaat"),
 ("O10", "Woningborg model-(koop-)aannemingsovereenkomsten 2026 en FAQ 'De overeenkomst'", "Woningborg", "Modelovereenkomst / FAQ", "Termijnregeling", "Algemeen", "2026", "n.v.t.", "Alleen volledig gepresteerde termijnen; standaardtermijnregelingen via portaal", "01-01-2026", "Niet vastgesteld", "https://www.woningborg.nl/kennisbank/veelgestelde-vragen-de-overeenkomst", "Via zoekresultaat"),
 ("O11", "Woningborg FAQ 'Het Woningborg-portaal' en kennisbank 'Opleveren onder de zakelijke waarborgregeling'", "Woningborg", "FAQ", "Obligo", "Algemeen", "n.v.t.", "n.v.t.", "—", "Niet vermeld", "—", "https://www.woningborg.nl/kennisbank/veelgestelde-vragen-het-woningborg-portaal", "Via zoekresultaat"),
]
r = table(b, 5, ["ID", "Document", "Aanbieder", "Soort bron", "Product", "Toepasselijke onderneming", "Tariefjaar", "Ingangsdatum", "Acceptatievoorwaarden", "Datum", "Vindplaats (pagina/cel/artikel)", "Link / bestand", "Bronstatus"],
          [6, 34, 20, 18, 26, 22, 12, 22, 34, 12, 30, 44, 18], src_rows)
r += 1
bar(b, r, 2, 14, "AFWIJKINGEN TUSSEN BRONNEN"); r += 1
table(b, r, ["Nr.", "Onderwerp", "Bron A", "Bron B", "Afwijking", "Voorrang / gevolg", "", "", "", "", "", "", ""],
      [6, 34, 20, 18, 26, 22, 12, 22, 34, 12, 30, 44, 18], [
 ("A1", "Tariefjaar", "B1: 2027", "B2: 2026", "Geen SWK-tarief 2027 beschikbaar", "Vergelijking hooguit indicatief; jaarverschil tonen"),
 ("A2", "Status SWK-tarieven", "B2 (intern)", "Geen officiële SWK-tariefbrief", "Tarieven in B2 zijn niet herleidbaar tot een SWK-document", "Schriftelijke bevestiging nodig"),
 ("A3", "Hoogbouwgrens", "B1: 'vanaf 13 lagen boven straatpeil'", "B2: '> 13 bouwlagen'; O3: ~70 m (technisch)", "13 lagen = hoogbouw bij WB, niet bij SWK (B2); 'bouwlagen' en 'lagen boven straatpeil' zijn mogelijk niet gelijk geteld", "Per aanbieder apart toepassen; SWK-grens laten bevestigen"),
 ("A4", "Grondslag particulier", "B1: (koop-/)aanneemsom incl. meer-/minderwerk en btw; 'met grond' = incl. grond", "B2/O4: aanneemsom incl. btw = bouwkosten + bijkomende kosten (grond apart)", "WB rekent bij 'met grond' over de grond; SWK niet. Meer-/minderwerk bij SWK niet vermeld", "Grondslag per aanbieder apart berekenen"),
 ("A5", "Btw over kosten", "B1 p.6: btw over alle tarieven", "B2/SWK-bronnen: niets vermeld", "Onbekend of SWK btw rekent", "Btw apart tonen; SWK-btw = schriftelijke bevestiging nodig (niet 0%)"),
 ("A6", "Insolventiedekking particulier", "O1: SWK max 100% aanneemsom", "O7: WB max 60% aanneemsom (2020: 100%)", "Dekkingsverschil voor de koper", "Benoemen als resterend dekkingsverschil; nooit als premie gebruiken"),
 ("A7", "Herstelwaarborg-maxima", "O1: € 140.000 / € 21.000 / € 119.000 × n (prijspeil 2024)", "O7: € 159.368 / € 24.629 / € 134.739 × n", "Verschillende bedragen en mogelijk verschillend prijspeil", "Prijspeil per bedrag laten bevestigen"),
 ("A8", "Zakelijke afbouwdekking", "O2: SWK Afbouw 20% aanneemsom", "O9: WB zakelijk insolventie tot 30%", "Dekking verschilt; SWK-tarief ontbreekt", "Niet vergelijkbaar"),
 ("A9", "Sectie I 'over 20% van de aanneemsom'", "B2", "O2 (afbouwdekking 20%)", "Mogelijk verband, niet aangetoond", "20% niet als dekkings- of obligopercentage gebruiken"),
 ("A10", "Eerder vergelijkingstemplate (deze sessie)", "Template: SWK-btw 0%, hoogbouw via lagen > 13", "Deze analyse", "Aannames niet onderbouwd", "Template corrigeren na beantwoording vragen (tab 7)"),
])
b.freeze_panes = "C6"

# ===================== 3. Producten & tarieven =====================
p = sheet("3 Producten en tarieven", "Product- en tarievenmatrix", "Ontbrekend tarief ≠ 0  ·  Effectief SWK-tarief = 20% × 0,5198% + 0,0380% = 0,14196% van de aanneemsom incl. btw (exclusief opslagen en vaste kosten)", 20)
NA, PO, SB = "Niet aangetroffen", "Projectofferte nodig", "Schriftelijke bevestiging nodig"
WBH = "Insolventie max 60% aanneemsom (O7); herstel max € 159.368 incl. btw per huis; app.: privé € 24.629, gemeensch. € 134.739 × aantal privédelen; eigen risico koper € 321 (O7)"
WBP = "Koop- en aanneemsom incl. grond, incl. meer-/minderwerk en btw (B1 p.3, p.6)"
WBA = "Aanneemsom incl. meer-/minderwerk en btw (B1 p.3, p.6)"
WBZ = "Aanneemsom excl. btw (B1 p.6)"
WBZD = "Insolventie tot 30% van de aanneemsom; herstel tot max. 10 jaar (O9). Maximumbedrag herstel: niet aangetroffen"
BTW_WB = "Ja – 'Over alle tarieven wordt btw in rekening gebracht' (B1 p.6); percentage niet vermeld"
rows = [
 ("WB-1", "Woningborg", "Woningborg-certificaat Nieuwbouw Particulier (GWRN 2024)", "Particuliere verkrijger", "Nieuwbouw EGW/appartementen tot 13 lagen boven straatpeil; verkoop met grond", "Insolventie- + herstelwaarborg", WBH, "2027", "0,227%", "%", WBP, "Ja", "Ja", BTW_WB, "Variabel", "Geen; hoogbouw = apart tarief (WB-2)", "Plannen geaccepteerd op/na 1-1-2027; erfpacht met afgekochte canon = 'met grond'", "B1 p.3, p.6; O7", "SWK-1 – indicatief"),
 ("WB-2", "Woningborg", "Idem, hoogbouw", "Particuliere verkrijger", "Vanaf 13 lagen boven straatpeil; verkoop met grond", "Insolventie- + herstelwaarborg", WBH, "2027", "0,250%", "%", WBP, "Ja", "Ja", BTW_WB, "Variabel", "Hogere tarief vervangt basistarief (+5% t.o.v. 2026, B1 p.1)", "Grens 'vanaf 13 lagen boven straatpeil'; telwijze verder niet gedefinieerd", "B1 p.1, p.3", "SWK-3 – indicatief"),
 ("WB-3", "Woningborg", "Woningborg-certificaat Nieuwbouw Particulier", "Particuliere verkrijger", "Tot 13 lagen; verkoop zonder grond", "Insolventie- + herstelwaarborg", WBH, "2027", "0,321%", "%", WBA, "Nee", "Ja", BTW_WB, "Variabel", "Geen", "'Zonder grond' = grond door derde verkocht of erfpacht zonder afkoop canon", "B1 p.3", "SWK-2 – indicatief"),
 ("WB-4", "Woningborg", "Idem, hoogbouw", "Particuliere verkrijger", "Vanaf 13 lagen; verkoop zonder grond", "Insolventie- + herstelwaarborg", WBH, "2027", "0,354%", "%", WBA, "Nee", "Ja", BTW_WB, "Variabel", "Hogere tarief vervangt basistarief", "Zie WB-2", "B1 p.3", "SWK-3 – indicatief"),
 ("WB-5", "Woningborg", "Woningborg-certificaat Kavelbouw", "Particuliere opdrachtgever", "Kavelbouw", "Insolventie- + herstelwaarborg (aanname: GWRN 2024; te bevestigen)", "Schriftelijke bevestiging nodig (welke regeling/limieten voor kavelbouw)", "2027", "0,380%", "%", WBA, "Nee", "Ja", BTW_WB, "Variabel", "Geen hoogbouwonderscheid vermeld", "—", "B1 p.3", "SWK-4 – niet vergelijkbaar"),
 ("WB-6", "Woningborg", "Woningborg-certificaat Transformatie Particulier (GWR Transformatie 2024)", "Particuliere verkrijger", "Transformatie, met én zonder grond", "Insolventie- + herstelwaarborg; bestaand gebouw valt onder dekking", "Insolventie max 60% aanneemsom; herstel max € 107.766 per huis (O8)", "2027", "0,321%", "%", "Koop- en aanneemsom incl. grond, meer-/minderwerk en btw – ook bij verkoop zonder grond (B1 p.4)", "Ja", "Ja", BTW_WB, "Variabel + bijkomende kosten (bezoeken, metingen)", "Tarief –10% t.o.v. 2026 (B1 p.1)", "Haalbaarheidsonderzoek (kosten WB); acceptatie alleen met bouwplaatsbezoeken en geluidsmeting(en) (B1 p.4)", "B1 p.1, p.4; O8", "SWK-7 – niet vergelijkbaar"),
 ("WB-7", "Woningborg", "Woningborg-certificaat Zakelijk – module insolventie- en herstelwaarborg", "Zakelijke opdrachtgever", "Nieuwbouw tot 13 lagen (huur, onverkochte koop, evt. met kantoor/bedrijf/parkeren; ook renovatie)", "Insolventie + herstel", WBZD, "2027", "0,568%", "%", WBZ, "Nee", "Nee", BTW_WB, "Variabel", "Hoogbouw = WB-8", "Waarborgregeling Zakelijk", "B1 p.5, p.6; O9", "SWK-6 – niet vergelijkbaar"),
 ("WB-8", "Woningborg", "Idem", "Zakelijke opdrachtgever", "Nieuwbouw vanaf 13 lagen", "Insolventie + herstel", WBZD, "2027", "0,757%", "%", WBZ, "Nee", "Nee", BTW_WB, "Variabel", "Hogere tarief", "Idem", "B1 p.5", "SWK-6 – niet vergelijkbaar"),
 ("WB-9", "Woningborg", "Woningborg-certificaat Zakelijk – module insolventiewaarborg", "Zakelijke opdrachtgever", "Nieuwbouw tot 13 lagen", "Alleen insolventie", "Insolventie tot 30% aanneemsom (O9)", "2027", "0,343%", "%", WBZ, "Nee", "Nee", BTW_WB, "Variabel", "Hoogbouw = WB-10", "Idem", "B1 p.5; O9", "SWK-5 – niet vergelijkbaar"),
 ("WB-10", "Woningborg", "Idem", "Zakelijke opdrachtgever", "Nieuwbouw vanaf 13 lagen", "Alleen insolventie", "Insolventie tot 30% aanneemsom (O9)", "2027", "0,458%", "%", WBZ, "Nee", "Nee", BTW_WB, "Variabel", "Hogere tarief", "Idem", "B1 p.5", "SWK-5 – niet vergelijkbaar"),
 ("WB-11", "Woningborg", "Zakelijk Transformatie – insolventie + herstel", "Zakelijke opdrachtgever", "Transformatie; bestaand gebouw valt WEL onder dekking", "Insolventie + herstel incl. bestaand gebouw", WBZD, "2027", "0,903%", "%", "Koop- en aanneemsom excl. btw (B1 p.5, p.6)", "Ja", "Nee", BTW_WB, "Variabel + bijkomende kosten transformatie", "Geen lagenonderscheid vermeld", "Bijkomende kosten als Transformatie Particulier (B1 p.5)", "B1 p.5", "SWK-8 – niet vergelijkbaar"),
 ("WB-12", "Woningborg", "Zakelijk Transformatie – alleen insolventie", "Zakelijke opdrachtgever", "Transformatie; bestaand gebouw valt NIET onder dekking", "Alleen insolventie", "Insolventie tot 30% aanneemsom (O9; te bevestigen voor transformatie)", "2027", "0,343%", "%", WBZ, "Nee", "Nee", BTW_WB, "Variabel", "Geen lagenonderscheid vermeld", "—", "B1 p.5", "SWK-8 – niet vergelijkbaar"),
 ("WB-13", "Woningborg", "Alle certificaten – houtbouw", "Alle", "Hoofddraagconstructie in hout (CLT, LVL, glulam)", "Als onderliggend certificaat", "Als onderliggend certificaat", "2027", PO + " (verhoogd tarief mogelijk na technische beoordeling)", "—", "Als onderliggend certificaat", "—", "—", BTW_WB, "Variabel + aanvullende bouwplaatsbezoeken (offerte)", "Mogelijk verhoogd tarief", "Aanvullende bouwplaatsbezoeken altijd nodig (B1 p.2, p.6)", "B1 p.2, p.6", "SWK-9 – niet vergelijkbaar"),
 ("SWK-1", "SWK", "SWK Garantie- en waarborgregeling 2024 (volledig pakket)", "Particuliere verkrijger", "Nieuwbouw EGW/appartementen; verkoop met grond", "Insolventie(afbouw)- + herstelwaarborg", "Insolventie max 100% aanneemsom (O1, te verifiëren); herstel max € 140.000 incl. btw per EGW; app.: privé € 21.000, gemeensch. € 119.000 × aantal privédelen (prijspeil 2024, geïndexeerd; O1)", "2026 (intern model B2)", "Sectie I 0,5198% over 20% van de aanneemsom + Sectie II 0,0380% over de aanneemsom (effectief 0,14196%) + € 750 bureaukosten per woning", "% + €", "Aanneemsom incl. btw (B2). Volgens O4: aanneemsom = bouwkosten + bijkomende kosten. Meer-/minderwerk: niet aangetroffen", "Nee (grond niet in aanneemsom volgens O4; bevestigen)", "Ja (B2)", SB, "Variabel + vast per woning (voorschot)", "Hoogbouw, verlengde garantie zakelijk, bouwtijd > 3 jaar (zie tab 4)", "Alleen volledige GWR; geen separaat verkochte parkeerplaatsen (B2 A26–27)", "B2; O1; O4", "WB-1 – indicatief"),
 ("SWK-2", "SWK", "Idem", "Particuliere verkrijger", "Nieuwbouw; verkoop zonder grond", "Idem", "Idem", "2026 (intern model B2)", "Geen apart tarief aangetroffen; B2 maakt geen onderscheid met/zonder grond", "—", "Aanneemsom incl. btw (B2)", "Nee", "Ja (B2)", SB, "Idem", "Idem", "—", "B2", "WB-3 – indicatief"),
 ("SWK-3", "SWK", "Idem – hoogbouw", "Particuliere verkrijger", "Hoogbouw", "Idem", "Idem", "2026 (intern model B2)", "Opslag 30,3% over Sectie II bij '> 13 bouwlagen' (B2); in officiële bronnen niet aangetroffen", "% van Sectie II", "Sectie II-bedrag", "—", "—", SB, "Variabel", "30,3% over Sectie II", "Technisch bureau verplicht bij hoogbouw (~70 m), kosten deelnemer (O3); B2: technische beoordeling > 25 bouwlagen", "B2; O3", "WB-2/WB-4 – indicatief"),
 ("SWK-4", "SWK", "Kavelbouw", "Particuliere opdrachtgever", "Kavelbouw", NA, NA, NA, NA, "—", NA, NA, NA, NA, NA, NA, NA, "—", "WB-5 – niet vergelijkbaar"),
 ("SWK-5", "SWK", "SWK Waarborgregeling Afbouw 2024", "Zakelijke opdrachtgever", "Nieuwbouw", "Alleen afbouwwaarborg", "Afbouw 20% van de aanneemsom; beroep binnen 6 maanden na oplevering; WKA max 10%, WAV max 5% (O2)", NA, SB + " (tarief niet aangetroffen; niet afleiden uit Sectie I)", "—", NA, NA, NA, NA, NA, NA, "Zie O2", "O2; O6", "WB-9/WB-10 – niet vergelijkbaar"),
 ("SWK-6", "SWK", "SWK Garantie- en waarborgregeling 2024 voor zakelijke opdrachtgever", "Zakelijke opdrachtgever", "Nieuwbouw (huur/zakelijk)", "Afbouw + herstel", "Zie SWK-1 (te bevestigen voor zakelijk)", NA, SB + " (B2 noemt alleen 'opslag verlengde garantie zakelijk' 10% over Sectie II; betekenis onduidelijk)", "—", NA, NA, NA, NA, NA, "Verlengde garantie zakelijk 10% over Sectie II (B2)", "—", "O6; B2", "WB-7/WB-8 – niet vergelijkbaar"),
 ("SWK-7", "SWK", "SWK Garantie- en waarborgregeling 2024 – transformatie", "Particuliere verkrijger", "Transformatie", "Afbouw + herstel (GWR geldt ook voor transformatie, O3)", NA, NA, NA, "—", NA, NA, NA, NA, NA, NA, "Technisch bureau verplicht, kosten deelnemer (O3)", "O3", "WB-6 – niet vergelijkbaar"),
 ("SWK-8", "SWK", "Transformatie zakelijk", "Zakelijke opdrachtgever", "Transformatie", NA, NA, NA, NA, "—", NA, NA, NA, NA, NA, NA, NA, "—", "WB-11/WB-12 – niet vergelijkbaar"),
 ("SWK-9", "SWK", "Houtbouw", "Alle", "Houtbouw", NA, NA, NA, PO, "—", NA, NA, NA, NA, NA, NA, "Mogelijk 'innovatieve bouw' → technisch bureau (O3); niet bevestigd", "O3", "WB-13 – niet vergelijkbaar"),
]
hdr = ["ID", "Aanbieder", "Product", "Doelgroep", "Projecttype", "Dekking", "Dekkingslimiet", "Tariefjaar", "Tarief", "Eenheid (%/‰/€)", "Berekeningsgrondslag", "Grond inbegrepen?", "Btw in grondslag?", "Btw over kosten?", "Vaste/variabele kosten", "Opslagen", "Uitsluitingen/voorwaarden", "Bron", "Vergelijkbaar met"]
wid = [7, 11, 26, 14, 24, 20, 34, 12, 26, 10, 28, 12, 10, 24, 16, 22, 28, 12, 18]
r = table(p, 5, hdr, wid, rows)
for rr in range(6, r):
    col = WBC if str(p.cell(rr, 2).value).startswith("WB") else SWKC
    p.cell(rr, 2).font = font(size=9, bold=True, color=col)
    v = str(p.cell(rr, 20).value)
    for k, c in STAT.items():
        if k in v: p.cell(rr, 20).fill = fill(c)
p.freeze_panes = "D6"

# ===================== 4. Bijkomende kosten =====================
k = sheet("4 Bijkomende kosten", "Matrix bijkomende kosten", "Vaste bijdragen op ondernemingsniveau niet aan één project toerekenen zonder verdeelsleutel  ·  Voorschot ≠ extra kosten bovenop werkelijke kosten", 12)
krows = [
 ("K-WB1", "Woningborg", "Bouwplaatsbezoek", "€ 1.041,00", "€ per bezoek", "Aantal bezoeken volgens offerte", "Excl. btw (B1 p.4)", "Transformatie (particulier en zakelijk); houtbouw (aanvullende bezoeken)", "Variabel (offerte)", "Rechtstreeks door WB gefactureerd; niet in premie inbegrepen", "B1 p.4, p.5, p.6", "Officieel"),
 ("K-WB2", "Woningborg", "Geluidsmeting", "€ 911,00", "€ per meting", "Aantal metingen volgens offerte", "Excl. btw (B1 p.4)", "Transformatie", "Variabel (offerte)", "Idem", "B1 p.4", "Officieel"),
 ("K-WB3", "Woningborg", "Haalbaarheidsonderzoek transformatie (eerste opname)", "€ 0 voor Heijmans", "—", "—", "—", "Transformatie", "—", "Expliciet 'voor rekening van Woningborg Verzekeren' (dus geen ontbrekend bedrag)", "B1 p.4", "Officieel"),
 ("K-WB4", "Woningborg", "Jaarlijkse bijdrage", "€ 644,00", "€ per ingeschreven onderneming per jaar", "Onderneming", "Niet vermeld bij de post; p.6: btw over alle tarieven", "Elke ingeschreven onderneming", "Vast (bedrijfsniveau)", "Niet per project tellen; alleen via verdeelsleutel over projecten", "B1 p.6", "Officieel"),
 ("K-WB5", "Woningborg", "Inschrijfgeld", "€ 4.018,00; binnen concern € 1.133,00 per extra inschrijving", "€ eenmalig per onderneming", "Onderneming", "Idem", "Nieuwe inschrijving (Heijmans Vastgoed is al ingeschreven: nr. 4889)", "Eenmalig", "Voor bestaand project waarschijnlijk n.v.t.", "B1 p.6", "Officieel"),
 ("K-WB6", "Woningborg", "Opstellen depotovereenkomst", "€ 1.500,-", "€ eenmalig", "Per verzoek", "Niet vermeld; bevestigen", "Op verzoek", "Eenmalig", "—", "B1 p.6", "Officieel"),
 ("K-WB7", "Woningborg", "Opstellen afbouwverklaring", "€ 500,-", "€ eenmalig", "Per verzoek", "Niet vermeld; bevestigen", "Op verzoek", "Eenmalig", "—", "B1 p.6", "Officieel"),
 ("K-WB8", "Woningborg", "Laten vervallen van plan na risicobeoordeling", "€ 2.500,-", "€ eenmalig", "Per plan", "Niet vermeld; bevestigen", "Plan vervalt na risicobeoordeling", "Eenmalig", "—", "B1 p.6", "Officieel"),
 ("K-WB9", "Woningborg", "Verzoek planacceptatie bij < 70% verkocht", "€ 1.500,-", "€ eenmalig", "Per verzoek", "Niet vermeld; bevestigen", "Voorverkoop onder 70%", "Eenmalig", "—", "B1 p.6", "Officieel"),
 ("K-WB10", "Woningborg", "Verzoek aanpassing standaard termijnregeling", "€ 1.250,-", "€ eenmalig", "Per verzoek", "Niet vermeld; bevestigen", "Afwijkende termijnregeling; facturatie pas als werk op bouwplaats gereed is", "Eenmalig", "—", "B1 p.6", "Officieel"),
 ("K-WB11", "Woningborg", "Hoogbouw", "In tarief (WB-2, WB-4, WB-8, WB-10)", "%", "Zie tarievenmatrix", "Btw over tarief", "Vanaf 13 lagen boven straatpeil", "Variabel", "Geen aparte opslag; niet dubbel tellen", "B1 p.3, p.5", "Officieel"),
 ("K-WB12", "Woningborg", "Houtbouw – verhoogd tarief", PO, "—", "—", "—", "Hoofddraagconstructie in hout", "—", "Na technische beoordeling", "B1 p.6", "Officieel"),
 ("K-WB13", "Woningborg", "Bijzonder/aanvullend obligo, voorschotten, nacalculatie", NA, "—", "—", "—", "—", "—", "—", "—", "—"),
 ("K-SWK1", "SWK", "Bureaukosten", "€ 750", "€ per woning, 'voorschot'", "Aantal woningen", SB, "Altijd (B2)", "Vast per woning (voorschot)", "Of er nacalculatie volgt, is niet bekend: tel werkelijke kosten niet óók bovenop het voorschot", "B2 B18", "Intern model"),
 ("K-SWK2", "SWK", "Hoogbouwopslag", "30,3%", "% van Sectie II-provisie", "Sectie II-bedrag", SB, "'> 13 bouwlagen' (B2); telwijze niet aangetroffen", "Variabel", "Grens wijkt af van WB (≥ 13 lagen boven straatpeil)", "B2 B14/E14", "Intern model"),
 ("K-SWK3", "SWK", "Opslag verlengde garantie zakelijk", "10%", "% van Sectie II", "Sectie II-bedrag", SB, "'Alleen bij zakelijke opdrachtgever' (B2); inhoud verlengde garantie niet aangetroffen", "Variabel", "—", "B2 B15/E15", "Intern model"),
 ("K-SWK4", "SWK", "Opslag lange bouwtijd", "10%", "% van Sectie I", "Sectie I-bedrag", SB, "Prognose bouwtijd > 3 jaar, start bouw tot 1e oplevering (B2)", "Variabel", "—", "B2 B16/E16", "Intern model"),
 ("K-SWK5", "SWK", "Technische beoordeling / technisch bureau", PO, "—", "—", "—", "Hoogbouw (~70 m), transformatie, renovatie, innovatieve bouw (O3); B2: transformatie of > 25 bouwlagen", "Variabel", "Kosten voor deelnemer (O3)", "O3; B2 A23", "Via zoekresultaat"),
 ("K-SWK6", "SWK", "Beoordeling afwijkende termijnregeling", NA, "—", "—", "—", "Afwijkende termijnregeling (B2 A24)", "—", "B2 noemt het als aandachtspunt zonder bedrag", "B2 A24", "Intern model"),
 ("K-SWK7", "SWK", "Deelnemersbijdrage / jaarlijkse bijdrage / inschrijfgeld", NA, "—", "—", "—", "—", "—", "—", "—", "—"),
 ("K-SWK8", "SWK", "Depotovereenkomst, afbouwverklaring, vervallen plan, planacceptatie < x% verkocht", NA, "—", "—", "—", "—", "—", "—", "—", "—"),
 ("K-SWK9", "SWK", "Bijzonder/aanvullend obligo, nacalculatie", NA, "—", "—", "—", "—", "—", "—", "—", "—"),
 ("K-X1", "Beide", "Inspectiekosten / eigen risico koper", "SWK € 115 incl. btw (O1); WB eigen risico € 321 incl. btw (O7)", "€", "—", "Incl. btw", "Kosten tussen koper en ondernemer/waarborg; géén certificaatkosten voor Heijmans", "—", "Niet meetellen in certificaatkosten", "O1; O7", "Via zoekresultaat"),
]
kt = table(k, 5, ["ID", "Aanbieder", "Post", "Bedrag / tarief", "Eenheid", "Grondslag", "Btw-behandeling", "Toepassingsvoorwaarde", "Vast / variabel / eenmalig", "Dubbeltelling / opmerking", "Bron", "Bronstatus"],
           [9, 11, 30, 20, 18, 18, 22, 34, 16, 34, 12, 14], krows)
k.freeze_panes = "D6"

# ===================== 5. Dekking, termijnen, obligo =====================
d = sheet("5 Dekking-termijn-obligo", "Vier onderdelen apart: kosten · dekking · termijnregeling · obligo",
          "Een obligo- of dekkingspercentage is nooit een premiepercentage", 5)
drows = [
 ("1. Certificaatkosten", "Wat betaalt Heijmans?", "Provisie/premie + opslagen + bureaukosten; zie tabs 3 en 4. Bron: intern model 2026 (B2); btw onbekend", "Premie = tarief × grondslag; zie tabs 3 en 4. Bron: officiële brief 2027 (B1); btw wordt gerekend", "Indicatief (jaar/bron)"),
 ("2. Garantiedekking – insolventie particulier", "Wat krijgt de koper bij insolventie vóór oplevering?", "Max 100% van de aanneemsom; SWK kiest tussen afbouw en terugbetaling betaalde termijnen + rente (O1, via zoekresultaat)", "Max 60% van de aanneemsom (2020: 100%); WB kiest de voor WB gunstigste optie; uitkering kan lager zijn dan 60% (O7, via zoekresultaat)", "Dekkingsverschil – benoemen"),
 ("2. Garantiedekking – herstel particulier", "Gebreken na oplevering", "Max € 140.000 incl. btw per EGW; € 21.000 privédeel; € 119.000 × aantal privédelen gemeensch. (prijspeil 2024, jaarlijks geïndexeerd) (O1)", "Max € 159.368 incl. btw per huis; € 24.629 privédeel; € 134.739 × aantal gemeensch.; CPI-indexatie; bedragen vast bij afgifte certificaat; eigen risico € 321 (O7)", "Prijspeil per bedrag laten bevestigen"),
 ("2. Garantiedekking – termijnen", "Duur", "6 jaar; ernstige gebreken 10 jaar; ingang 3 maanden na oplevering (O1-brochure)", "6 jaar; ernstige gebreken 10 jaar; ingang 3 maanden na oplevering (O7)", "Gelijk op hoofdlijnen"),
 ("2. Garantiedekking – zakelijk", "Opdrachtgever", "Afbouw: 20% aanneemsom (O2). Afbouw + herstel: GWR 2024 (O6), limieten zakelijk niet aangetroffen", "Insolventie tot 30% aanneemsom; herstel tot max. 10 jaar (O9)", "Niet vergelijkbaar"),
 ("2. Garantiedekking – transformatie", "Bestaand gebouw", "GWR 2024 geldt ook voor transformatie (O3); limieten niet aangetroffen", "Herstel max € 107.766 per huis; insolventie max 60% (O8); bestaand gebouw onder dekking (B1 p.4)", "Niet vergelijkbaar"),
 ("3. Termijnregeling", "Wanneer mag Heijmans factureren?", "Standaard termijnregelingen 2024 EGH/APP (O5). Eerste termijn volgens staffel van reeds betaalde bijkomende kosten, exclusief winst, rente en risico; AK telt mee tot 1% bij KAO (O4)", "Standaardtermijnregelingen via portaal; alleen volledig gepresteerde termijnen; aanpassing € 1.250, facturatie pas als werk op bouwplaats gereed is (O10; B1 p.6)", "Apart beoordelen; geen kostenpost"),
 ("3. Termijnregeling – let op", "Afleidingen", "Woonlagen in een termijnschema zijn géén premie- of hoogbouwgrens", "Idem", "—"),
 ("4. Obligo", "Hoeveel waarborgcapaciteit neemt het project in?", NA + " (geen publieke informatie over obligo per project of per onderneming)", "Obligo per onderneming in portaal: huidig/toekomstig benodigd en uitstaand; vrijval na oplevering; vragen via financieel screener (O11). Omvang per project: niet aangetroffen", "Bij beide opvragen"),
]
table(d, 5, ["Onderdeel", "Vraag", "SWK", "Woningborg", "Beoordeling"], [32, 30, 60, 60, 26], drows)

# ===================== 6. Rekenvoorbeelden =====================
e = sheet("6 Rekenvoorbeelden", "Rekenvoorbeelden (alleen functioneel vergelijkbare pakketten)",
          "GELE CELLEN = ZELFGEKOZEN VOORBEELDINVOER, geen projectgegevens  ·  Tarieven: Woningborg 2027 (B1), SWK 2026 (B2, intern)  ·  Status: INDICATIEF", 8, GOLD)
for col, w in zip("BCDEFGHI", [52, 18, 18, 18, 4, 44, 18, 18]): e.column_dimensions[col].width = w
bar(e, 5, 2, 9, "VOORBEELDINVOER (zelfgekozen – aanpassen)")
inputs = [(6, "Aantal woningen", 20, "#,##0"), (7, "Aanneemsom per woning, excl. btw", 350000, EUR),
          (8, "Koopsom grond per woning, excl. btw", 75000, EUR), (9, "Btw-tarief op aanneemsom (voorbeeld)", 0.21, "0%"),
          (10, "Btw-tarief op grond (voorbeeld; afhankelijk van levering)", 0.21, "0%"),
          (11, "Btw over Woningborg-premie (percentage niet in B1; voorbeeld)", 0.21, "0%"),
          (12, "Btw over SWK-provisie – SCENARIO (onbekend)", 0.21, "0%")]
for rr, lab, val, fmt in inputs:
    e.cell(rr, 2, lab).font = font(); c = e.cell(rr, 4, val); inp(c, fmt)
e["F12"] = "Niet bevestigd: dit is een scenario, geen aanname. Zie vraag S-3."; e["F12"].font = font(size=9, italic=True, color="B42318")
e["B13"] = "Aanneemsom incl. btw per woning"; e["D13"] = "=D7*(1+D9)"; calc(e["D13"], EUR)
e["B14"] = "Koop- en aanneemsom incl. btw per woning"; e["D14"] = "=D7*(1+D9)+D8*(1+D10)"; calc(e["D14"], EUR)
for rr in (13, 14): e.cell(rr, 2).font = font(bold=True)

def block(r0, title, status, reason, wb_rate, wb_base_label, wb_base, swk_extra=None):
    bar(e, r0, 2, 9, title)
    e.cell(r0 + 1, 2, "Status"); e.cell(r0 + 1, 3, status).font = font(bold=True)
    e.cell(r0 + 1, 3).fill = fill(STAT[status]); e.merge_cells(start_row=r0 + 1, start_column=4, end_row=r0 + 1, end_column=9)
    e.cell(r0 + 1, 4, reason).font = font(size=9); e.cell(r0 + 1, 4).alignment = Alignment(wrap_text=True)
    e.row_dimensions[r0 + 1].height = 28
    h = r0 + 2
    for i, t in enumerate(["Per woning", "SWK", "Woningborg", "Verschil (WB − SWK)", "", "Berekening / bron"]):
        c = e.cell(h, 2 + i, t); c.font = font(size=9, bold=True, color=NAVY); c.border = Border(bottom=Side(style="medium", color=NAVY))
    a = h + 1
    lines = [
     ("Grondslag", "=$D$13", wb_base, f"SWK: aanneemsom incl. btw (B2) · WB: {wb_base_label}"),
     ("Sectie I (0,5198% × 20% × aanneemsom incl. btw)", f"=C{a}*0.2*0.005198", None, "B2 B10/C10"),
     ("Sectie II (0,0380% × aanneemsom incl. btw)", f"=C{a}*0.00038", None, "B2 B12"),
     ("Opslagen SWK", swk_extra or "=0", None, "B2 B14–B16"),
     ("Premie Woningborg (tarief × grondslag)", None, f"=D{a}*{wb_rate}", f"B1: {wb_rate*100:.3f}%".replace(".", ",")),
     ("Bureaukosten SWK (voorschot)", "=750", None, "B2 B18"),
     ("Totaal excl. btw", f"=SUM(C{a+1}:C{a+5})", f"=SUM(D{a+1}:D{a+5})", ""),
     ("Btw over kosten", "=\"onbekend\"", f"=D{a+6}*$D$11", "SWK: schriftelijke bevestiging nodig · WB: B1 p.6"),
     ("Totaal incl. btw", "=\"onbekend\"", f"=D{a+6}+D{a+7}", ""),
     ("SWK incl. btw – SCENARIO bij btw volgens D12", f"=C{a+6}*(1+$D$12)", None, "Scenario, geen feit"),
     ("Projecttotaal excl. btw (× aantal woningen)", f"=C{a+6}*$D$6", f"=D{a+6}*$D$6", "Bijdragen op ondernemingsniveau (K-WB4) niet meegeteld"),
    ]
    for i, (lab, fs, fw, src) in enumerate(lines):
        rr = a + i
        e.cell(rr, 2, lab).font = font(bold=lab.startswith(("Totaal", "Projecttotaal")))
        if fs: calc(e.cell(rr, 3, fs), EUR2, bold=lab.startswith("Totaal"))
        if fw: calc(e.cell(rr, 4, fw), EUR2, bold=lab.startswith("Totaal"))
        if fs and fw and "onbekend" not in fs: calc(e.cell(rr, 5, f"=D{rr}-C{rr}"), EUR2, bg=None)
        note(e.cell(rr, 7, src) if False else e.cell(rr, 7), src)
        e.cell(rr, 3).alignment = e.cell(rr, 4).alignment = Alignment(horizontal="right")
    return a + len(lines) + 1

r = 17
r = block(r, "VOORBEELD 1 – Particuliere nieuwbouw MET grond, < 13 lagen  (SWK-1 ↔ WB-1)", "indicatief",
          "Functie gelijk (insolventie + herstel voor de koper). Verschillen: tariefjaar 2026/2027, SWK-tarief intern, grondslag (WB incl. grond), insolventiedekking SWK 100% vs WB 60%, btw SWK onbekend.",
          0.00227, "koop- en aanneemsom incl. btw (B1 p.3)", "=$D$14")
r = block(r + 1, "VOORBEELD 2 – Particuliere nieuwbouw ZONDER grond, < 13 lagen  (SWK-2 ↔ WB-3)", "indicatief",
          "Grondslag gelijk (aanneemsom incl. btw). Overige verschillen als voorbeeld 1.",
          0.00321, "aanneemsom incl. btw (B1 p.3)", "=$D$13")
r = block(r + 1, "VOORBEELD 3 – Hoogbouw MET grond, 14 lagen (beide hoogbouw)  (SWK-3 ↔ WB-2)", "indicatief",
          "Bij 14 lagen vallen beide onder hoogbouw. Bij precies 13 lagen: WB wel (≥ 13), SWK niet (> 13; intern). Grens en telwijze SWK niet officieel bevestigd.",
          0.0025, "koop- en aanneemsom incl. btw (B1 p.3)", "=$D$14", swk_extra=None)
# hoogbouwopslag in voorbeeld 3: rij 'Opslagen SWK' = 30,3% × Sectie II
a3 = r - 12
e.cell(a3 + 3, 3).value = f"=C{a3+2}*0.303"
note(e.cell(a3 + 3, 7), "B2 B14: 30,3% × Sectie II")
bar(e, r + 1, 2, 9, "NIET BEREKEND – niet vergelijkbaar")
nb = ["Zakelijk alleen insolventie (WB-9: 0,343% × aanneemsom excl. btw) ↔ SWK Afbouw: SWK-tarief niet aangetroffen. SWK-tarief niet afleiden uit Sectie I.",
      "Zakelijk insolventie + herstel (WB-7: 0,568%) ↔ SWK zakelijk: tarief niet aangetroffen.",
      "Kavelbouw, transformatie, houtbouw: SWK-tarief niet aangetroffen / projectofferte nodig."]
for i, t in enumerate(nb): para(e, r + 2 + i, "• " + t, 9, size=9)
e.freeze_panes = "A5"

# ===================== 7. Ontbrekend & vragen =====================
q = sheet("7 Ontbrekend en vragen", "Ontbrekende gegevens en concrete vragen", "Pas na beantwoording kan de status 'indicatief' naar 'vergelijkbaar'", 5)
qrows = [
 ("S-1", "SWK", "Officiële tarieven 2027 voor Heijmans Vastgoed", "Wilt u de tariefbrief 2027 (en 2026) voor Heijmans Vastgoed sturen, met alle producten: GWR particulier (met/zonder grond), kavelbouw, Waarborgregeling Afbouw, GWR zakelijk, transformatie en renovatie?", "Hoog"),
 ("S-2", "SWK", "Herkomst rekenmodel B2", "Zijn Sectie I 0,5198% over 20% van de aanneemsom, Sectie II 0,0380%, de opslagen 30,3%/10%/10% en € 750 bureaukosten de tarieven voor 2026? Welke grondslag en welk document horen daarbij?", "Hoog"),
 ("S-3", "SWK", "Btw", "Wordt over de provisie, de opslagen en de bureaukosten btw berekend? Zo ja, welk percentage?", "Hoog"),
 ("S-4", "SWK", "Grondslag", "Welke aanneemsom geldt voor de provisie: incl. of excl. btw, incl. meer-/minderwerk, incl. of excl. grond en bijkomende kosten? Wordt achteraf nagecalculeerd op de definitieve aanneemsom?", "Hoog"),
 ("S-5", "SWK", "Hoogbouwgrens", "Vanaf hoeveel bouwlagen geldt de hoogbouwopslag? Hoe worden lagen geteld (boven straatpeil, kelder, technische ruimte/dakopbouw)? Hoe verhoudt dit zich tot de technische grens van circa 70 m?", "Hoog"),
 ("S-6", "SWK", "Bureaukosten", "Is € 750 per woning een voorschot met nacalculatie? Zo ja, op welke basis wordt verrekend?", "Middel"),
 ("S-7", "SWK", "Zakelijk", "Wat is het tarief voor de Waarborgregeling Afbouw 2024 (20% dekking) en voor de GWR 2024 bij een zakelijke opdrachtgever? Wat dekt de 'verlengde garantie zakelijk'?", "Hoog"),
 ("S-8", "SWK", "Vaste bijdragen en diensten", "Zijn er een deelnemersbijdrage, jaarlijkse bijdrage, inschrijfgeld of kosten voor depot, afbouwverklaring, afwijkende termijnregeling, planacceptatie bij lage voorverkoop of vervallen plannen?", "Middel"),
 ("S-9", "SWK", "Technische beoordeling", "Wat kost het technisch bureau bij transformatie, hoogbouw en houtbouw (tarief of bandbreedte)?", "Middel"),
 ("S-10", "SWK", "Dekking", "Bevestig de maximale insolventiedekking (100% van de aanneemsom?) en de herstelmaxima met prijspeil voor plannen uit 2027.", "Middel"),
 ("S-11", "SWK", "Obligo", "Hoeveel obligo neemt een project in (methode en percentage) en wat is de huidige obligoruimte van Heijmans Vastgoed?", "Middel"),
 ("W-1", "Woningborg", "Btw-percentage", "Welk btw-percentage wordt over premies en extra dienstverlening gerekend? Zijn de bedragen voor extra dienstverlening (p.6) incl. of excl. btw?", "Hoog"),
 ("W-2", "Woningborg", "Telwijze 13 lagen", "Hoe worden 'lagen boven straatpeil' geteld (begane grond, dakopbouw, technische verdieping, half verdiept)? Geldt het hoogbouwtarief per gebouw of per plan?", "Hoog"),
 ("W-3", "Woningborg", "Kavelbouw", "Onder welke regeling en met welke dekkingslimieten valt het kavelbouwcertificaat?", "Laag"),
 ("W-4", "Woningborg", "Zakelijk dekking", "Bevestig 30% insolventiedekking en de herstellimieten voor de zakelijke regeling, ook bij transformatie 'alleen insolventie'.", "Middel"),
 ("W-5", "Woningborg", "Houtbouw", "Hoe hoog is de houtbouwtoeslag (bandbreedte) en hoeveel bouwplaatsbezoeken zijn gebruikelijk?", "Middel"),
 ("W-6", "Woningborg", "Obligo", "Hoeveel obligo neemt een project in (methode) en hoe valt het vrij (bij oplevering of na de herstelperiode)?", "Middel"),
 ("W-7", "Woningborg", "Dekking particulier", "Bevestig 60% insolventiedekking en € 159.368 / € 24.629 / € 134.739 herstelmaxima met prijspeil voor plannen uit 2027.", "Middel"),
 ("W-8", "Woningborg", "Premieherberekening", "Wordt de premie herberekend bij meer-/minderwerk of een gewijzigde koopsom na planacceptatie?", "Middel"),
 ("X-1", "Intern", "Btw-aftrek", "Kan Heijmans Vastgoed de btw over waarborgkosten aftrekken? Dit bepaalt of de vergelijking excl. of incl. btw moet gebeuren (fiscale afdeling).", "Hoog"),
]
table(q, 5, ["ID", "Aan", "Ontbrekend gegeven", "Concrete vraag", "Prioriteit"], [7, 12, 30, 110, 10], qrows)

# ===================== 8. Checklist =====================
c = sheet("8 Checklist", "Checklist projectcontroller: zijn twee offertes echt vergelijkbaar?",
          "Vul per vraag ja/nee in (kolom E)  ·  Eén 'nee' bij een blokkerende vraag = niet vergelijkbaar", 5)
items = [
 ("Zelfde functie: dekken beide offertes dezelfde waarborgen (insolventie + herstel, of alleen insolventie)?", "Blokkerend"),
 ("Zelfde doelgroep en regeling (particulier vs zakelijk, nieuwbouw vs transformatie vs kavelbouw)?", "Blokkerend"),
 ("Zijn beide tarieven officieel en voor Heijmans Vastgoed (tariefbrief/offerte, geen intern model)?", "Blokkerend"),
 ("Zelfde tariefjaar en ingangsdatum (planacceptatie)? Zo nee: jaarverschil getoond en geen definitieve conclusie?", "Blokkerend"),
 ("Grondslag per aanbieder correct bepaald (grond, meer-/minderwerk, bijkomende kosten, btw)?", "Blokkerend"),
 ("Hoogbouwgrens per aanbieder apart getoetst, met de juiste telwijze van lagen?", "Blokkerend"),
 ("Opslagen (hoogbouw, bouwtijd, zakelijk, houtbouw) volledig en niet dubbel meegenomen?", "Blokkerend"),
 ("Btw apart getoond (excl. btw | btw | incl. btw) en btw-aftrek onderbouwd?", "Blokkerend"),
 ("Vaste bijdragen op ondernemingsniveau niet volledig aan één project toegerekend?", "Belangrijk"),
 ("Voorschotten en werkelijke kosten (nacalculatie) niet dubbel geteld?", "Belangrijk"),
 ("Bijkomende kosten (technische beoordeling, bouwplaatsbezoeken, metingen, depot, termijnregeling) meegenomen?", "Belangrijk"),
 ("Resterende dekkingsverschillen (insolventie %, herstelmaxima, prijspeil) benoemd?", "Belangrijk"),
 ("Termijnregeling apart beoordeeld (cashflow), niet als kostenpost?", "Belangrijk"),
 ("Obligobeslag per aanbieder opgevraagd en apart beoordeeld (niet als premie)?", "Belangrijk"),
 ("Ontbrekende tarieven gemarkeerd als 'niet aangetroffen' of 'offerte nodig' (nooit als 0)?", "Blokkerend"),
 ("Elke tarief- en voorwaardebron voorzien van document, datum en pagina/artikel?", "Belangrijk"),
]
for col, w in zip("BCDEF", [6, 96, 14, 12, 4]): c.column_dimensions[col].width = w
for i, h in enumerate(["Nr.", "Controlevraag", "Type", "Ja/nee"]):
    x = c.cell(5, 2 + i, h); x.font = font(size=9, bold=True, color="FFFFFF"); x.fill = fill(PETROL)
dv = DataValidation(type="list", formula1='"ja,nee"', allow_blank=True); c.add_data_validation(dv)
for i, (t, typ) in enumerate(items):
    rr = 6 + i
    c.cell(rr, 2, i + 1).font = font(); c.cell(rr, 3, t).font = font(); c.cell(rr, 3).alignment = Alignment(wrap_text=True, vertical="center")
    c.cell(rr, 4, typ).font = font(size=9, bold=typ == "Blokkerend", color="B42318" if typ == "Blokkerend" else MUTED)
    inp(c.cell(rr, 5), center=True); dv.add(c.cell(rr, 5).coordinate)
    for cc in range(2, 6): c.cell(rr, cc).border = BOTTOM
    c.row_dimensions[rr].height = 20
end = 6 + len(items) - 1
rr = end + 2
c.cell(rr, 3, "Uitkomst").font = font(bold=True, size=11, color=NAVY)
c.merge_cells(start_row=rr, start_column=4, end_row=rr, end_column=5)
c.cell(rr, 4, f'=IF(COUNTIF(E6:E{end},"")>0,"NOG INVULLEN",IF(COUNTIFS(D6:D{end},"Blokkerend",E6:E{end},"nee")>0,"NIET VERGELIJKBAAR",IF(COUNTIF(E6:E{end},"nee")>0,"INDICATIEF","VERGELIJKBAAR")))')
c.cell(rr, 4).font = font(bold=True, size=11); c.cell(rr, 4).alignment = Alignment(horizontal="center")
for t, col in (("VERGELIJKBAAR", OKF), ("INDICATIEF", "FDE9B8"), ("NIET VERGELIJKBAAR", BADF)):
    c.conditional_formatting.add(f"D{rr}", FormulaRule(formula=[f'D{rr}="{t}"'], fill=fill(col)))

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 90
    print_setup(ws, None, f"Heijmans Vastgoed | SWK vs. Woningborg | {ws.title}")
    ws.print_area = None
    ws.page_setup.fitToHeight = 0
wb.save(sys.argv[1])
