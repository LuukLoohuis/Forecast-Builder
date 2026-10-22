import sys; sys.path.insert(0, sys.path[0])
from hstyle import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

wb = Workbook(); ws = wb.active; ws.title = "SWK berekening"
for col, w in zip("BCDEFGH", [38, 13, 16, 9, 15, 14, 56]): ws.column_dimensions[col].width = w
brand(ws, 8, "Berekening provisie SWK  |  Garantie- en Waarborgregeling",
      "Tarieven 2026  ·  Blauwe cijfers op geel = invoer  ·  Grijze cijfers = SWK-tarief  ·  Lichtblauw = berekening")

# KPI-tegels
tile(ws, 5, 2, 2, "TOTAAL KOSTEN SWK", "=F30", EUR)
tile(ws, 5, 3, 4, "KOSTEN PER WONING", "=G30", EUR)
tile(ws, 5, 5, 6, "% AANNEEMSOM INCL. BTW", "=IFERROR(F30/D14,0)", "0.00%")
tile(ws, 5, 7, 8, "STATUS", '=IF(AND(D35="OK",D36="OK",D37="OK"),"OK","CONTROLEREN")', "General")
ws.conditional_formatting.add("G6", FormulaRule(formula=['G6="OK"'], font=Font(color="38761D", bold=True)))
ws.conditional_formatting.add("G6", FormulaRule(formula=['G6<>"OK"'], font=Font(color="CC0000", bold=True)))

# Invoer
bar(ws, 9, 2, 8, "INVOER")
rows = [
    (10, "Projectnaam", None, None, "Optioneel; verschijnt niet in de berekening"),
    (11, "Aantal woningen", 1, "#,##0", None),
    (12, "Aanneemsom (excl. btw)", 420000, EUR, "cf. definitie SWK (zie ook bijlage 3)"),
]
for r, lab, val, fmt, tip in rows:
    ws.cell(r, 2, lab).font = font()
    c = ws.cell(r, 4, val); inp(c, fmt)
    if r == 10:
        ws.merge_cells("D10:F10")
        for cc in (5, 6): ws.cell(r, cc).fill = fill(INPUT); ws.cell(r, cc).border = BOTTOM
    if tip: note(ws.cell(r, 8), tip)
ws["B13"] = '="Btw ("&TEXT(C13,"0%")&")"'; ws["B13"].font = font()
inp(ws["C13"], "0%", center=True); ws["C13"] = 0.21
ws["D13"] = "=D12*C13"; calc(ws["D13"], EUR)
ws["B14"] = "Aanneemsom (incl. btw)"; ws["B14"].font = font(bold=True, color=NAVY)
ws["D14"] = "=SUM(D12:D13)"; calc(ws["D14"], EUR, bold=True)
ws["B15"] = "Aanneemsom per woning (incl. btw)"; ws["B15"].font = font(color=MUTED)
ws["D15"] = "=IFERROR(D14/D11,0)"; calc(ws["D15"], EUR, color=MUTED)

# Berekening
bar(ws, 17, 2, 8, "BEREKENING")
for i, h in enumerate(["Kostenpost", "Tarief 2026", "Grondslag", "Keuze", "Bedrag", "Per woning", "Toelichting"]):
    c = ws.cell(18, 2 + i, h); c.font = font(size=9, bold=True, color=NAVY)
    c.alignment = Alignment(horizontal="left" if i in (0, 6) else "right" if i != 3 else "center")
    c.border = Border(bottom=Side(style="medium", color=NAVY))

def group(r, text):
    ws.cell(r, 2, text).font = font(size=9, bold=True, italic=True, color=PETROL)

def line(r, lab, tar, tfmt, grond, gfmt, keuze, bedrag, tip, indent=False):
    c = ws.cell(r, 2, lab); c.font = font(color=MUTED if indent else "000000", size=9 if indent else 10)
    if indent: c.alignment = Alignment(indent=2)
    tarief(ws.cell(r, 3, tar), tfmt)
    if grond: calc(ws.cell(r, 4, grond), gfmt, bg=None, color=MUTED)
    if keuze: inp(ws.cell(r, 5, keuze), center=True)
    if bedrag:
        calc(ws.cell(r, 6, bedrag), EUR)
        calc(ws.cell(r, 7, f"=IFERROR(F{r}/$D$11,0)"), EUR, color=MUTED, bg=None)
    note(ws.cell(r, 8), tip)

group(19, "Basisprovisies")
line(20, "Provisie Sectie I (Afbouw)", 0.005198, PCT4, "=D14*C21", EUR, None, "=D20*C20", "Over 20% van de aanneemsom incl. btw")
line(21, "↳ grondslag Sectie I", 0.2, "0%", None, None, None, None, "Aandeel van de aanneemsom waarover Sectie I wordt berekend", indent=True)
line(22, "Provisie Sectie II (Kwaliteit)", 0.00038, PCT4, "=D14", EUR, None, "=D22*C22", "Over de aanneemsom incl. btw")
group(23, "Opslagen en keuzes")
line(24, "Opslag hoogbouw", 0.303, "0.0%", "=F22", EUR, "nee", '=IF(E24="ja",D24*C24,0)', "Opslag over Sectie II; > 13 bouwlagen")
line(25, "Opslag verlengde garantie zakelijk", 0.1, "0.0%", "=F22", EUR, "nee", '=IF(E25="ja",D25*C25,0)', "Opslag over Sectie II; alleen bij zakelijke opdrachtgever")
line(26, "Opslag prognose bouwtijd > 3 jaar", 0.1, "0.0%", "=F20", EUR, "nee", '=IF(E26="ja",D26*C26,0)', "Opslag over Sectie I; start bouw tot 1e oplevering > 3 jaar")
group(27, "Overige kosten")
line(28, "Bureaukosten, voorschot", 750, EUR, "=D11", '#,##0 "won."', None, "=D28*C28", "Per woning")
for r in (20, 22, 24, 25, 26, 28):
    ws.cell(r, 4).alignment = Alignment(horizontal="right")

# Totaal
for c in range(2, 9):
    ws.cell(30, c).fill = fill(GOLD)
    ws.cell(30, c).border = Border(top=Side(style="medium", color=NAVY), bottom=Side(style="medium", color=NAVY))
ws["B30"] = "Totaal kosten SWK"; ws["B30"].font = font(size=11, bold=True, color=NAVY)
ws["F30"] = "=SUM(F20:F28)"; ws["F30"].font = font(size=11, bold=True, color=NAVY); ws["F30"].number_format = EUR
ws["G30"] = "=IFERROR(F30/D11,0)"; ws["G30"].font = font(size=11, bold=True, color=NAVY); ws["G30"].number_format = EUR
ws["H30"] = '="= "&TEXT(IFERROR(F30/D14,0),"0.00%")&" van de aanneemsom incl. btw"'; ws["H30"].font = font(size=9, bold=True, color=NAVY)
ws.row_dimensions[30].height = 24

dv = DataValidation(type="list", formula1='"ja,nee"', allow_blank=False, showErrorMessage=True,
                    errorTitle="Ongeldige keuze", error="Kies ja of nee.")
ws.add_data_validation(dv); dv.add("E24:E26")

# Controles + aandachtspunten
bar(ws, 33, 2, 4, "CONTROLES"); bar(ws, 33, 6, 8, "AANDACHTSPUNTEN")
checks = [(35, "Invoer compleet", '=IF(AND(D12>0,D11>0),"OK","AANVULLEN")'),
          (36, "Keuzes geldig", '=IF(AND(OR(E24="ja",E24="nee"),OR(E25="ja",E25="nee"),OR(E26="ja",E26="nee")),"OK","CONTROLEREN")'),
          (37, "Rekenstatus", '=IF(AND(F30>=0,D14>0),"OK","CONTROLEREN")')]
for r, lab, f in checks:
    ws.cell(r, 2, lab).font = font(size=9, color=MUTED)
    c = ws.cell(r, 4, f); c.font = font(size=9, bold=True, color=NAVY); c.alignment = Alignment(horizontal="center")
    for cc in range(2, 5): ws.cell(r, cc).border = BOTTOM
ws.conditional_formatting.add("D35:D37", FormulaRule(formula=['D35="OK"'], fill=fill(OKF)))
ws.conditional_formatting.add("D35:D37", FormulaRule(formula=['D35<>"OK"'], fill=fill(BADF)))
pts = ["• Technische beoordeling, bijvoorbeeld bij transformatie of hoogbouw (> 25 bouwlagen).",
       "• Beoordeling van een afwijkende termijnregeling."]
for i, t in enumerate(pts):
    ws.merge_cells(start_row=35 + i, start_column=6, end_row=35 + i, end_column=8)
    note(ws.cell(35 + i, 6), t)

ws.merge_cells("B40:H40")
c = ws["B40"]; c.value = ("ⓘ  Deze berekening geldt alleen voor plannen waarop de volledige Garantie- en Waarborgregeling "
                         "van toepassing is en waarin geen parkeerplaatsen separaat worden verkocht.")
c.font = font(size=9, italic=True, color=MUTED); c.alignment = Alignment(vertical="center", indent=1)
for cc in range(2, 9): ws.cell(40, cc).fill = fill(INPUT)
ws.row_dimensions[40].height = 26
for r in range(9, 41):
    if ws.row_dimensions[r].height is None: ws.row_dimensions[r].height = 19

ws.freeze_panes = "A9"
for row in ws.iter_rows(min_row=4, max_row=40, min_col=2, max_col=8):
    for c in row:
        a = c.alignment
        c.alignment = Alignment(horizontal=a.horizontal, vertical="center", indent=1 if c.column == 8 else a.indent, wrap_text=a.wrap_text)
print_setup(ws, "A1:H40", "Heijmans Vastgoed | SWK-berekening")
ws.page_setup.fitToHeight = 1
wb.save(sys.argv[1])
