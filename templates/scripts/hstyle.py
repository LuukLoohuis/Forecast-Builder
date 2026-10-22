# Gedeelde Heijmans-huisstijl (afgeleid van het bestaande SWK-rekenmodel)
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
F = "Arial"
NAVY, PETROL, GOLD = "111F39", "003B5C", "FEB900"
CALC, LABEL, INPUT, MUTED, LINE = "EAF3F8", "F2F3F4", "FFF3C4", "4A4F55", "D0D5DD"
OKF, BADF = "D9EAD3", "F4CCCC"
EUR = '€ #,##0;[Red](€ #,##0);-'
EUR2 = '€ #,##0.00;[Red](€ #,##0.00);-'
PCT4 = '0.0000%'
thin = Side(style="thin", color=LINE)
BOTTOM = Border(bottom=thin)
def font(**k):
    k.setdefault("name", F); k.setdefault("size", 10); return Font(**k)
def fill(c): return PatternFill("solid", fgColor=c)
def bar(ws, row, c1, c2, text, color=PETROL, fg="FFFFFF", size=10, height=20):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.cell(row, c1, text).font = font(bold=True, color=fg, size=size)
    ws.cell(row, c1).alignment = Alignment(vertical="center", indent=1)
    for c in range(c1, c2 + 1): ws.cell(row, c).fill = fill(color)
    ws.row_dimensions[row].height = height
def brand(ws, c2, title, subtitle):
    bar(ws, 1, 2, c2, "HEIJMANS  |  VASTGOED", NAVY, GOLD, 12, 26)
    bar(ws, 2, 2, c2, title, GOLD, NAVY, 16, 32)
    ws.merge_cells(start_row=3, start_column=2, end_row=3, end_column=c2)
    ws.cell(3, 2, subtitle).font = font(size=9, italic=True, color=MUTED)
    ws.cell(3, 2).alignment = Alignment(indent=1)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2
def inp(c, fmt=None, center=False):
    c.font = font(bold=True, color="0000FF"); c.fill = fill(INPUT); c.border = BOTTOM
    if fmt: c.number_format = fmt
    if center: c.alignment = Alignment(horizontal="center")
def tarief(c, fmt):
    c.font = font(color="666666"); c.fill = fill(LABEL); c.border = BOTTOM; c.number_format = fmt
def calc(c, fmt=None, bold=False, color="000000", bg=CALC):
    c.font = font(bold=bold, color=color)
    if bg: c.fill = fill(bg)
    c.border = BOTTOM
    if fmt: c.number_format = fmt
def note(c, text, size=9):
    c.value = text; c.font = font(size=size, color=MUTED); c.alignment = Alignment(vertical="center", wrap_text=False)
def tile(ws, row, c1, c2, label, formula, fmt, accent=PETROL, valcolor=NAVY):
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    ws.merge_cells(start_row=row + 1, start_column=c1, end_row=row + 2, end_column=c2)
    ws.cell(row, c1, label).font = font(size=9, bold=True, color="FFFFFF")
    ws.cell(row, c1).alignment = Alignment(horizontal="center", vertical="center")
    v = ws.cell(row + 1, c1, formula); v.font = font(size=18, bold=True, color=valcolor); v.number_format = fmt
    v.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(c1, c2 + 1):
        ws.cell(row, c).fill = fill(accent)
        for r in (row + 1, row + 2): ws.cell(r, c).fill = fill(CALC)
        ws.cell(row + 2, c).border = Border(bottom=Side(style="medium", color=accent))
def print_setup(ws, area, footer):
    ws.print_area = area
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.oddFooter.center.text = footer; ws.oddFooter.right.text = "Pagina &P van &N"
    ws.oddFooter.center.font = ws.oddFooter.right.font = "Arial"
