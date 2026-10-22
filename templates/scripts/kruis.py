import sys; sys.path.insert(0, sys.path[0])
from hstyle import *
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Side

wb = Workbook(); ws = wb.active; ws.title = "Vergelijkbaarheid"
brand(ws, 7, "Vergelijkbaarheid Woningborg-categorie × SWK-scenario", "Tabel zoals aangeleverd; symbolen en voetnootcijfers ongewijzigd overgenomen")
for col, w in zip("BCDEFG", [34, 14, 16, 20, 16, 18]): ws.column_dimensions[col].width = w
hdr = ["Woningborg-categorie", "S1 Volledig", "S2 Alleen afbouw", "S3 Zakelijk gecombineerd", "S4 Transformatie", "S5 Kavel/bijzonder"]
rows = [
 ("W1 Particulier", "🟠 I", "🔴 X", "🔴 X", "—", "?"),
 ("W2 Particulier nieuwbouw", "🟠 I", "🔴 X", "🔴 X", "—", "?"),
 ("W3 Particulier nieuwbouw,", "🟠 I¹", "🔴 X", "🔴 X", "—", "?"),
 ("W4 Particulier nieuwbouw, zonder", "🟠 I¹", "🔴 X", "🔴 X", "—", "?"),
 ("W5 Particulier", "?", "🔴 X", "🔴 X", "—", "?²"),
 ("W6 Part", "?", "🔴 X", "🔴 X", "?²", "?"),
 ("W7 Zakelijk", "🟠 I³", "🔴 X", "🟠 I³", "—", "?"),
 ("W8 Zakelijk nieuwbouw", "🟠 I¹³", "🔴 X", "🟠 I¹³", "—", "?"),
 ("W9", "🔴 X", "🟠 O", "🔴 X", "—", "?"),
 ("W10", "🔴 X", "🟠 O¹", "🔴 X", "—", "?"),
 ("W11 Zakelijk", "?", "🔴 X", "?", "?²", "?"),
 ("W12", "🔴 X", "🟠 O⁴", "🔴 X", "?²", "?"),
]
for i, h in enumerate(hdr):
    c = ws.cell(5, 2 + i, h); c.font = font(bold=True, color="FFFFFF", size=10); c.fill = fill(PETROL)
    c.alignment = Alignment(horizontal="left" if i == 0 else "center", vertical="center", wrap_text=True)
ws.row_dimensions[5].height = 32
ORANGE, RED, GREY = "FDE2C4", "F4CCCC", "F2F3F4"
for j, row in enumerate(rows):
    r = 6 + j
    for i, v in enumerate(row):
        c = ws.cell(r, 2 + i, v); c.border = Border(bottom=Side(style="thin", color=LINE), right=Side(style="thin", color=LINE))
        if i == 0:
            c.font = font(bold=True, color=NAVY)
        else:
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = font(bold=True)
            if v.startswith("🟠"): c.fill = fill(ORANGE)
            elif v.startswith("🔴"): c.fill = fill(RED)
            elif v.startswith("?") or v == "—": c.fill = fill(GREY)
    ws.row_dimensions[r].height = 20
note(ws["B19"], "Rijlabels, legenda (I / X / O / ? / —) en voetnoten ¹ ² ³ ⁴ zijn ongewijzigd zoals aangeleverd; de toelichting bij de voetnoten ontbreekt nog.")
ws.freeze_panes = "C6"
print_setup(ws, "A1:G19", "Heijmans Vastgoed | Vergelijkbaarheid SWK – Woningborg"); ws.page_setup.fitToHeight = 1
wb.save(sys.argv[1])
