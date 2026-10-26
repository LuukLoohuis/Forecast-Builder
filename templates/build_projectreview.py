"""Genereert templates/Projectreview.xlsx: een invulbaar projectreview-template
met dropdowns, datumvalidatie, automatische berekeningen en statuskleuren."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).with_name("Projectreview.xlsx")

NAVY = "17375E"
LIGHT = "DCE6F1"
INPUT = "FFFDE7"
FONT = "Arial"

thin = Side(style="thin", color="A6A6A6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

LISTS = {
    "lst_Status": ["🟢 Op schema", "🟠 Aandacht", "🔴 Kritiek"],
    "lst_Fase": ["Initiatief", "Haalbaarheid", "Ontwerp", "Voorbereiding",
                 "Verkoop", "Realisatie", "Oplevering", "Nazorg"],
    "lst_Soort": ["Grondgebonden", "Appartementen", "Gemengd", "Transformatie"],
    "lst_Personen": ["Ontwikkelaar 1", "Ontwikkelaar 2", "Ontwikkelaar 3",
                     "Projectleider", "Verkoopmanager"],
    "lst_Type": ["Actualiteit", "Kans", "Risico", "Besluit"],
    "lst_Actie": ["⬜ Open", "🔄 Bezig", "✅ Gereed"],
    "lst_Kwartaal": [f"Q{q} {y}" for y in range(2025, 2033) for q in range(1, 5)],
}

wb = Workbook()
ws = wb.active
ws.title = "Review"
lst = wb.create_sheet("Lijsten")
info = wb.create_sheet("Toelichting")

# ---------- Lijsten (bewerkbaar door gebruiker) ----------
for col, (name, values) in enumerate(LISTS.items(), start=1):
    letter = lst.cell(row=1, column=col).column_letter
    head = lst.cell(row=1, column=col, value=name.replace("lst_", ""))
    head.font = Font(name=FONT, bold=True, color="FFFFFF")
    head.fill = PatternFill("solid", fgColor=NAVY)
    for r, v in enumerate(values, start=2):
        c = lst.cell(row=r, column=col, value=v)
        c.font = Font(name=FONT)
        c.fill = PatternFill("solid", fgColor=INPUT)
    ref = f"Lijsten!${letter}$2:${letter}${len(values) + 1}"
    wb.defined_names[name] = DefinedName(name, attr_text=ref)
    lst.column_dimensions[letter].width = 20
lst["I1"] = "Pas de gele cellen aan. Extra waarde nodig? Voeg een rij in midden in de lijst (rechtsklik > Invoegen), dan groeit de dropdown mee."
lst["I1"].font = Font(name=FONT, italic=True, color="595959")

# ---------- Review: opmaakhelpers ----------
widths = {"A": 20, "B": 17, "C": 24, "D": 17, "E": 20, "F": 24}
for k, v in widths.items():
    ws.column_dimensions[k].width = v


def style(rng, *, bold=False, fill=None, color="000000", size=10, wrap=True, align="left"):
    for row in ws[rng]:
        for c in row:
            c.font = Font(name=FONT, bold=bold, color=color, size=size)
            if fill:
                c.fill = PatternFill("solid", fgColor=fill)
            c.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal=align)
            c.border = BORDER


def label(cell, text, merge_to=None):
    if merge_to:
        ws.merge_cells(f"{cell}:{merge_to}")
    ws[cell] = text
    style(f"{cell}:{merge_to or cell}", bold=True, fill=LIGHT)


def field(cell, merge_to=None, fmt=None, align="left"):
    if merge_to:
        ws.merge_cells(f"{cell}:{merge_to}")
    style(f"{cell}:{merge_to or cell}", fill=INPUT, color="1F5FA8", bold=True, align=align)
    if fmt:
        ws[cell].number_format = fmt


def calc(cell, formula, fmt=None, merge_to=None):
    if merge_to:
        ws.merge_cells(f"{cell}:{merge_to}")
    ws[cell] = formula
    style(f"{cell}:{merge_to or cell}", bold=True, align="right")
    if fmt:
        ws[cell].number_format = fmt


def section(row, text):
    ws.merge_cells(f"A{row}:F{row}")
    ws[f"A{row}"] = text
    style(f"A{row}:F{row}", bold=True, fill=NAVY, color="FFFFFF", size=11)
    ws.row_dimensions[row].height = 20


def dv_list(name, *ranges):
    dv = DataValidation(type="list", formula1=f"={name}", allow_blank=True,
                        showErrorMessage=True, errorTitle="Ongeldige keuze",
                        error="Kies een waarde uit de lijst (beheer op tabblad Lijsten).")
    for r in ranges:
        dv.add(r)
    ws.add_data_validation(dv)


# ---------- Titel + projectkop ----------
ws.merge_cells("A1:F1")
ws["A1"] = '=IF(B2="","[Leeg]",B2)'
style("A1:F1", bold=True, fill=NAVY, color="FFFFFF", size=16)
ws.row_dimensions[1].height = 30

label("A2", "Projectnaam"); field("B2", "D2")
label("E2", "Status"); field("F2")
label("A3", "Reviewdatum"); field("B3", fmt="d-m-yyyy")
label("C3", "SAP WBS"); field("D3")
label("E3", "Soort"); field("F3")
label("A4", "Projectfase"); field("B4")
label("C4", "Verantwoordelijke 1"); field("D4")
label("E4", "Verantwoordelijke 2"); field("F4")

dv_list("lst_Status", "F2")
dv_list("lst_Soort", "F3")
dv_list("lst_Fase", "B4")
dv_list("lst_Personen", "D4", "F4")
date_dv = DataValidation(type="date", operator="greaterThan", formula1="DATE(2000,1,1)",
                         allow_blank=True, showErrorMessage=True, errorTitle="Ongeldige datum",
                         error="Voer een datum in, bijv. 9-9-2026.")
ws.add_data_validation(date_dv)
date_dv.add("B3")

# ---------- 1. Kerngegevens ----------
section(6, "1. KERNGEGEVENS")
for col, t in zip("ABCDEF", ["Woningen", "Aantal", "Verkocht", "% verkocht", "Planning", "Kwartaal"]):
    label(f"{col}7", t)
label("A8", "DAEB"); field("B8", fmt="0", align="right"); field("C8", fmt="0", align="right")
label("A9", "Niet-DAEB"); field("B9", fmt="0", align="right"); field("C9", fmt="0", align="right")
label("A10", "Totaal")
calc("B10", "=SUM(B8:B9)", "0")
calc("C10", "=SUM(C8:C9)", "0")
for r in (8, 9, 10):
    calc(f"D{r}", f'=IF(N(B{r})=0,"",C{r}/B{r})', "0%")
for r, t in zip((8, 9, 10), ["Verkoopstart", "Start bouw", "Oplevering"]):
    label(f"E{r}", t); field(f"F{r}")
dv_list("lst_Kwartaal", "F8:F10")
label("A11", "Verkooptempo (p/mnd)"); field("B11", fmt="0.0", align="right")
label("C11", "Begin volgend jaar in verkoop"); field("D11", fmt="0", align="right")
label("E11", "Belangrijkste mijlpaal"); field("F11")

num_dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0",
                        allow_blank=True, showErrorMessage=True, error="Voer een getal ≥ 0 in.")
ws.add_data_validation(num_dv)
num_dv.add("B8:C9"); num_dv.add("B11"); num_dv.add("D11")

# ---------- 2. Actualiteiten / review ----------
section(13, "2. ACTUALITEITEN / REVIEW")
label("A14", "Type"); label("B14", "Onderwerp"); label("C14", "Stand van zaken", "D14")
label("E14", "Status"); label("F14", "Besluit / aandachtspunt")
ACT_ROWS = range(15, 23)
for r in ACT_ROWS:
    field(f"A{r}"); field(f"B{r}"); field(f"C{r}", f"D{r}"); field(f"E{r}"); field(f"F{r}")
dv_list("lst_Type", f"A{ACT_ROWS[0]}:A{ACT_ROWS[-1]}")
dv_list("lst_Status", f"E{ACT_ROWS[0]}:E{ACT_ROWS[-1]}")

# ---------- 3. Actielijst ----------
section(24, "3. ACTIELIJST")
label("A25", "Status"); label("B25", "Actie", "C25"); label("D25", "Eigenaar")
label("E25", "Deadline"); label("F25", "Signaal (automatisch)")
ACTION_ROWS = range(26, 36)
for r in ACTION_ROWS:
    field(f"A{r}"); field(f"B{r}", f"C{r}"); field(f"D{r}"); field(f"E{r}", fmt="d-m-yyyy")
    # Peildatum = reviewdatum, of vandaag als die leeg is
    calc(f"F{r}",
         f'=IF(B{r}="","",IF(ISNUMBER(SEARCH("Gereed",A{r})),"✅ Afgerond",'
         f'IF(E{r}="","⚪ Geen deadline",IF(E{r}<IF($B$3="",TODAY(),$B$3),"🔴 Verlopen",'
         f'IF(E{r}-IF($B$3="",TODAY(),$B$3)<=14,"🟠 Binnen 2 weken","🟢 Op schema")))))')
    ws[f"F{r}"].alignment = Alignment(vertical="center", horizontal="left")
first, last = ACTION_ROWS[0], ACTION_ROWS[-1]
dv_list("lst_Actie", f"A{first}:A{last}")
dv_list("lst_Personen", f"D{first}:D{last}")
date_dv.add(f"E{first}:E{last}")

label("A36", "Samenvatting", "B36")
calc("C36", f'=COUNTIFS(B{first}:B{last},"<>",A{first}:A{last},"<>*Gereed*")&" open actie(s)"', merge_to="D36")
calc("E36", f'=COUNTIF(F{first}:F{last},"*Verlopen*")&" verlopen"', merge_to="F36")

# ---------- 4. Vragen ----------
section(38, "4. VRAGEN VOOR VOLGENDE REVIEW")
for i, r in enumerate(range(39, 44), start=1):
    ws[f"A{r}"] = f"❓ Vraag {i}"
    style(f"A{r}:A{r}", bold=True, fill=LIGHT)
    field(f"B{r}", f"F{r}")

# ---------- Voorwaardelijke opmaak ----------
GREEN, ORANGE, RED, GREY = "C6EFCE", "FFE0B2", "FFC7CE", "E7E6E6"


def cf_text(rng, text, fill, strike=False):
    tl = rng.split(":")[0]
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'ISNUMBER(SEARCH("{text}",{tl}))'],
        fill=PatternFill("solid", fgColor=fill, bgColor=fill),
        font=Font(strike=strike, color="7F7F7F") if strike else None, stopIfTrue=True))


for rng in ["F2", f"E{ACT_ROWS[0]}:E{ACT_ROWS[-1]}"]:
    cf_text(rng, "Op schema", GREEN); cf_text(rng, "Aandacht", ORANGE); cf_text(rng, "Kritiek", RED)
sig = f"F{first}:F{last}"
cf_text(sig, "Verlopen", RED); cf_text(sig, "Binnen 2", ORANGE)
cf_text(sig, "Op schema", GREEN); cf_text(sig, "Afgerond", GREY)
# Afgeronde acties grijs doorgehaald
ws.conditional_formatting.add(f"A{first}:E{last}", FormulaRule(
    formula=[f'ISNUMBER(SEARCH("Gereed",$A{first}))'],
    fill=PatternFill("solid", fgColor=GREY, bgColor=GREY),
    font=Font(strike=True, color="7F7F7F")))
# Verkocht > aantal = invoerfout
ws.conditional_formatting.add("C8:C9", FormulaRule(
    formula=["AND(ISNUMBER(C8),ISNUMBER(B8),C8>B8)"],
    fill=PatternFill("solid", fgColor=RED, bgColor=RED)))
# Planning in onlogische volgorde (start bouw vóór verkoopstart, oplevering vóór start bouw)
for r in (9, 10):
    ws.conditional_formatting.add(f"F{r}", FormulaRule(
        formula=[f"AND(F{r}<>\"\",F{r-1}<>\"\","
                 f"IFERROR(MATCH(F{r},lst_Kwartaal,0)<MATCH(F{r-1},lst_Kwartaal,0),FALSE))"],
        fill=PatternFill("solid", fgColor=ORANGE, bgColor=ORANGE)))

# ---------- Afdrukken / weergave ----------
for r in range(2, 44):
    if ws.row_dimensions[r].height is None:
        ws.row_dimensions[r].height = 18
ws.freeze_panes = "A2"
ws.sheet_view.showGridLines = False
ws.print_area = "A1:F43"
ws.page_setup.orientation = "portrait"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True

# ---------- Toelichting ----------
info.column_dimensions["A"].width = 110
lines = [
    ("Projectreview – toelichting", True),
    ("", False),
    ("Lichtgele cellen zijn invulvelden. Witte cellen met vette tekst rekenen automatisch.", False),
    ("Klik op een geel veld met een pijltje voor een dropdown (status, fase, type, eigenaar, kwartaal).", False),
    ("Dropdownwaarden (o.a. namen van ontwikkelaars) beheer je op het tabblad 'Lijsten'.", False),
    ("", False),
    ("Automatisch:", True),
    ("• Titel = projectnaam; totaal woningen, totaal verkocht en % verkocht.", False),
    ("• Statuskleuren groen/oranje/rood bij de projectstatus en bij actualiteiten.", False),
    ("• Actielijst: signaal 'Verlopen' / 'Binnen 2 weken' t.o.v. de reviewdatum (of vandaag als die leeg is).", False),
    ("• Acties op 'Gereed' worden grijs doorgehaald; samenvatting telt open en verlopen acties.", False),
    ("• Waarschuwing (rood) als verkocht > aantal; oranje als planning in onlogische volgorde staat.", False),
    ("", False),
    ("Voorbeeld van een ingevulde actie:", True),
    ("Status: ⬜ Open | Actie: Offerte installateur opvragen | Eigenaar: Ontwikkelaar 1 | Deadline: 1-10-2026", False),
    ("", False),
    ("Tip: in OneNote via Invoegen > Spreadsheet > Bestaande Excel-spreadsheet dit bestand insluiten.", False),
]
for i, (t, b) in enumerate(lines, start=1):
    c = info.cell(row=i, column=1, value=t)
    c.font = Font(name=FONT, bold=b, size=14 if i == 1 else 10)

wb.save(OUT)
print(OUT)
