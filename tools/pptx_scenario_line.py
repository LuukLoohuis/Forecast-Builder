#!/usr/bin/env python3
"""Voeg de scenariolijn (gele lijn) toe aan de cashflowgrafiek op dia 3 van het kwartaalsjabloon (v10 -> v11).

Gebruik:
    python3 tools/pptx_scenario_line.py INVOER.pptx UITVOER.pptx [UITVOER2.pptx ...]

Bewerking (python-pptx + XML, zelfde werkwijze als tools/pptx_rente_card.py):
  * Op dia 3 wordt het grafiekvak CF_GRAFIEK gezocht en het chart-part geopend.
  * In de lijngroep (<c:lineChart>) komt een nieuwe reeks:
      - naam  = celverwijzing Sheet1!$H$1 (strCache 'Scenario'), zodat de legenda de kop
                uit het gegevensblad toont ('Scenario', de gekozen naam of 'Scenario (uit)');
      - categorieen = dezelfde A-kolom als de andere reeksen (deep copy van <c:cat>);
      - waarden = Sheet1!$H$2:$H$<zelfde einde als de andere reeksen>, numCache met alleen
                formatCode en ptCount (geen punten: de macro vult de kolom);
      - lijn amber EDA100, 2,25 pt, ronde uiteinden, geen markering, smooth 0;
      - idx/order = hoogste bestaande + 1.
  * Het ingebedde werkboek (ppt/embeddings/Werkblad_cf.xlsx) krijgt kolom H met kop 'Scenario'
    en lege cellen, zodat 'Gegevens bewerken' in PowerPoint een kloppende tabel toont.
  * Controles: de legenda heeft geen legendEntry/delete voor de nieuwe idx, en de reeksnaam
    staat niet vast via <c:v> maar via <c:strRef>.

De VBA-macro (vba/CashflowNaarPowerPoint_v10.bas, VulGrafiek) schrijft het blok G7:N.. van tab
PowerPoint in A1:H(n+1) van het gegevensblad en zet per reeks het bereik op basis van de
kolomletter in de SERIES-formule; kolom H wordt dus automatisch gevuld.
"""
import argparse
import copy
import io
import shutil
import sys

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

SLIDE_IDX = 2                 # dia 3
VORM = "CF_GRAFIEK"
KOP = "Scenario"
KOLOM = "H"
KLEUR = "EDA100"              # amber, zelfde als chartxml.AMBER
LIJN_EMU = 28575              # 2,25 pt
NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def q(tag):
    pre, local = tag.split(":")
    return "{%s}%s" % (NS[pre], local)


def el(tag, **attrs):
    e = etree.Element(q(tag))
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def sub(parent, tag, **attrs):
    e = el(tag, **attrs)
    parent.append(e)
    return e


def zoek_grafiek(slide):
    for sh in slide.shapes:
        if sh.name == VORM:
            if not sh.has_chart:
                sys.exit(f"{VORM} op dia {SLIDE_IDX + 1} is geen grafiek.")
            return sh
    sys.exit(f"{VORM} niet gevonden op dia {SLIDE_IDX + 1}.")


def splits_ref(f):
    """'Sheet1!$G$2:$G$33' -> ('Sheet1', 'G', 2, 33)"""
    blad, bereik = f.rsplit("!", 1)
    van, tot = bereik.split(":")
    kol = van.replace("$", "").rstrip("0123456789")
    r1 = int(van.replace("$", "")[len(kol):])
    r2 = int(tot.replace("$", "").lstrip("ABCDEFGHIJKLMNOPQRSTUVWXYZ"))
    return blad, kol, r1, r2


def maak_reeks(voorbeeld, idx, order, blad, r1, r2):
    ser = el("c:ser")
    sub(ser, "c:idx", val=idx)
    sub(ser, "c:order", val=order)
    tx = sub(ser, "c:tx")
    sr = sub(tx, "c:strRef")
    sub(sr, "c:f").text = f"{blad}!${KOLOM}$1"
    sc = sub(sr, "c:strCache")
    sub(sc, "c:ptCount", val=1)
    pt = sub(sc, "c:pt", idx=0)
    sub(pt, "c:v").text = KOP
    sp = sub(ser, "c:spPr")
    ln = sub(sp, "a:ln", w=LIJN_EMU, cap="rnd")
    sub(sub(ln, "a:solidFill"), "a:srgbClr", val=KLEUR)
    sub(ln, "a:prstDash", val="solid")
    sub(ln, "a:round")
    sub(sub(ser, "c:marker"), "c:symbol", val="none")
    ser.append(copy.deepcopy(voorbeeld.find("c:cat", NS)))
    val = sub(ser, "c:val")
    nr = sub(val, "c:numRef")
    sub(nr, "c:f").text = f"{blad}!${KOLOM}${r1}:${KOLOM}${r2}"
    nc = sub(nr, "c:numCache")
    vb_fmt = voorbeeld.find("c:val/c:numRef/c:numCache/c:formatCode", NS)
    sub(nc, "c:formatCode").text = vb_fmt.text if vb_fmt is not None else "General"
    sub(nc, "c:ptCount", val=r2 - r1 + 1)
    sub(ser, "c:smooth", val=0)
    return ser


def werk_werkboek_bij(chart_part, r2):
    wb_part = chart_part.chart_workbook.xlsx_part
    if wb_part is None:
        sys.exit("De grafiek heeft geen ingebed werkboek.")
    wb = load_workbook(io.BytesIO(wb_part.blob))
    ws = wb.worksheets[0]
    if ws[f"{KOLOM}1"].value not in (None, KOP):
        sys.exit(f"{KOLOM}1 in het ingebedde werkboek is al in gebruik: {ws[KOLOM + '1'].value!r}")
    ws[f"{KOLOM}1"] = KOP
    for r in range(2, r2 + 1):
        ws[f"{KOLOM}{r}"] = None
    ws.column_dimensions[KOLOM].width = ws.column_dimensions["G"].width or 20
    buf = io.BytesIO()
    wb.save(buf)
    chart_part.chart_workbook.update_from_xlsx_blob(buf.getvalue())
    return ws.title


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieen van de uitvoer")
    args = ap.parse_args()

    prs = Presentation(args.invoer)
    slide = prs.slides[SLIDE_IDX]
    shape = zoek_grafiek(slide)
    chart = shape.chart
    cs = chart._chartSpace
    line = cs.find(".//c:plotArea/c:lineChart", NS)
    if line is None:
        sys.exit("Geen lijngroep (<c:lineChart>) in CF_GRAFIEK.")
    alle = cs.findall(".//c:plotArea/*/c:ser", NS)
    for s in alle:
        f = s.find("c:tx/c:strRef/c:f", NS)
        if f is not None and f.text.endswith(f"!${KOLOM}$1"):
            sys.exit(f"Er is al een reeks op kolom {KOLOM}; niets gedaan.")
    idx = max(int(s.find("c:idx", NS).get("val")) for s in alle) + 1
    order = max(int(s.find("c:order", NS).get("val")) for s in alle) + 1
    voorbeeld = line.findall("c:ser", NS)[-1]
    blad, _, r1, r2 = splits_ref(voorbeeld.find("c:val/c:numRef/c:f", NS).text)
    cat_f = voorbeeld.find("c:cat/*/c:f", NS).text
    assert splits_ref(cat_f)[1:] == ("A", r1, r2), cat_f

    ser = maak_reeks(voorbeeld, idx, order, blad, r1, r2)
    voorbeeld.addnext(ser)          # na de laatste reeks, voor dLbls/marker/axId

    # legenda: geen verborgen entry voor de nieuwe reeks
    legend = cs.find(".//c:chart/c:legend", NS)
    verborgen = [e.find("c:idx", NS).get("val") for e in legend.findall("c:legendEntry", NS)
                 if e.find("c:delete", NS) is not None and e.find("c:delete", NS).get("val") in ("1", "true")]
    assert str(idx) not in verborgen, verborgen
    assert ser.find("c:tx/c:v", NS) is None and ser.find("c:tx/c:strRef/c:f", NS) is not None

    bladnaam = werk_werkboek_bij(chart.part, r2)
    assert bladnaam == blad, (bladnaam, blad)

    prs.save(args.uitvoer)
    for extra in args.extra:
        shutil.copyfile(args.uitvoer, extra)
    print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
    print(f"  nieuwe reeks idx={idx} order={order} naam={blad}!${KOLOM}$1 waarden={blad}!${KOLOM}${r1}:${KOLOM}${r2}")
    print(f"  verborgen legenda-entries (idx): {', '.join(verborgen) or '-'}")
    for s in cs.findall(".//c:plotArea/*/c:ser", NS):
        grp = etree.QName(s.getparent()).localname
        print(f"  {grp:10s} idx={s.find('c:idx', NS).get('val')} order={s.find('c:order', NS).get('val')} "
              f"{s.find('c:tx/c:strRef/c:f', NS).text:14s} {s.find('c:val/c:numRef/c:f', NS).text}")


if __name__ == "__main__":
    main()
