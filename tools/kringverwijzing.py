"""Statische controle op kringverwijzingen, zoals Excel die meldt: elke formule hangt af van álle cellen in de bereiken die ze
noemt (ook bij INDEX over een bereik), dus een formule in zo'n bereik die (indirect) van die formule afhangt, is een kring.

Gebruik: python3 tools/kringverwijzing.py werkboek.xlsx  (geeft de kringen; exit 1 als er een is)."""
import re
import sys
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

REF = re.compile(r"^(?:(?P<blad>'[^']+'|[A-Za-z0-9_.]+)!)?(?P<ref>\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?|\$?[A-Z]{1,3}:\$?[A-Z]{1,3}|\$?\d+:\$?\d+)$")


def _rect(ref, max_row, max_col):
    """(r1, c1, r2, c2) van een bereik; hele kolommen/rijen begrensd op het gebruikte gebied."""
    ref = ref.replace("$", "")
    if re.match(r"^[A-Z]{1,3}:[A-Z]{1,3}$", ref):
        a, b = ref.split(":")
        return 1, column_index_from_string(a), max_row, column_index_from_string(b)
    if re.match(r"^\d+:\d+$", ref):
        a, b = ref.split(":")
        return int(a), 1, int(b), max_col
    c1, r1, c2, r2 = range_boundaries(ref)
    return r1, c1, r2, c2


def kringen(pad):
    wb = load_workbook(pad)
    grenzen = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
    namen = {}
    for naam, dn in wb.defined_names.items():
        namen[naam.lower()] = dn.attr_text
    formules = {}                                  # (blad, r, c) -> formule
    per_blad = defaultdict(dict)                   # blad -> {(r, c): node}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cel in row:
                v = cel.value
                if hasattr(v, "text"):             # ArrayFormula
                    v = v.text
                if isinstance(v, str) and v.startswith("="):
                    formules[(ws.title, cel.row, cel.column)] = v
                    per_blad[ws.title][(cel.row, cel.column)] = (ws.title, cel.row, cel.column)

    def bereiken(blad, formule, diepte=0):
        uit = []
        try:
            tokens = Tokenizer(formule).items
        except Exception:
            return uit
        for t in tokens:
            if t.type == Token.OPERAND and t.subtype == Token.RANGE:
                m = REF.match(t.value)
                if m:
                    b = (m.group("blad") or blad).strip("'")
                    if b in grenzen:
                        uit.append((b, _rect(m.group("ref"), *grenzen[b])))
                elif t.value.lower() in namen and diepte < 3:
                    uit += bereiken(blad, "=" + namen[t.value.lower()], diepte + 1)
        return uit

    graaf = {}
    for node, formule in formules.items():
        blad = node[0]
        doelen = set()
        for b, (r1, c1, r2, c2) in bereiken(blad, formule):
            cellen = per_blad[b]
            if (r2 - r1 + 1) * (c2 - c1 + 1) <= len(cellen):
                for r in range(r1, r2 + 1):
                    for c in range(c1, c2 + 1):
                        if (r, c) in cellen:
                            doelen.add(cellen[(r, c)])
            else:
                for (r, c), n in cellen.items():
                    if r1 <= r <= r2 and c1 <= c <= c2:
                        doelen.add(n)
        graaf[node] = doelen
    # Tarjan: sterk samenhangende componenten
    index, laag, stapel, op_stapel, sccs = {}, {}, [], set(), []
    teller = [0]

    def sterk(v):
        werk = [(v, iter(graaf[v]))]
        index[v] = laag[v] = teller[0]
        teller[0] += 1
        stapel.append(v)
        op_stapel.add(v)
        while werk:
            node, it = werk[-1]
            for w in it:
                if w not in index:
                    index[w] = laag[w] = teller[0]
                    teller[0] += 1
                    stapel.append(w)
                    op_stapel.add(w)
                    werk.append((w, iter(graaf[w])))
                    break
                elif w in op_stapel:
                    laag[node] = min(laag[node], index[w])
            else:
                werk.pop()
                if werk:
                    laag[werk[-1][0]] = min(laag[werk[-1][0]], laag[node])
                if laag[node] == index[node]:
                    scc = []
                    while True:
                        w = stapel.pop()
                        op_stapel.discard(w)
                        scc.append(w)
                        if w == node:
                            break
                    if len(scc) > 1 or node in graaf[node]:
                        sccs.append(scc)

    for v in graaf:
        if v not in index:
            sterk(v)
    return [(scc, {n: formules[n] for n in scc[:6]}) for scc in sccs]


def adres(n):
    return f"{n[0]}!{get_column_letter(n[2])}{n[1]}"


if __name__ == "__main__":
    uit = kringen(sys.argv[1])
    if not uit:
        print("geen kringverwijzingen")
        sys.exit(0)
    for scc, voorbeeld in uit:
        print(f"kring van {len(scc)} cellen, o.a. {', '.join(adres(n) for n in scc[:8])}")
        for n, f in voorbeeld.items():
            print(f"   {adres(n)}: {f[:160]}")
    sys.exit(1)
