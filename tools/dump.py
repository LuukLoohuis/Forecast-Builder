"""Hulpmiddel: toon berekende waarden van een (herberekende) werkmap."""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1], data_only=True)
for spec in sys.argv[2:]:
    sheet, rng = spec.split("!")
    ws = wb[sheet]
    print(f"== {spec}")
    for row in ws[rng]:
        vals = []
        for c in row:
            v = c.value
            if isinstance(v, float):
                v = round(v, 2)
            vals.append("" if v is None else str(v)[:28])
        if any(vals):
            print(" | ".join(vals))
