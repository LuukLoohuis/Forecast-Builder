"""Scenario-tests voor het RDFO-prototype.

Bouwt de werkmap, past kopieën aan zoals een gebruiker zou doen (bron vervangen, vastleggen in Historie,
parameters wijzigen), laat LibreOffice herberekenen en controleert de uitkomsten tegen een onafhankelijke
pandas-berekening op de bronbestanden.

Gebruik:  python tools/test_rdfo.py <werkmap-dir> [pad-naar-recalc.py]
Geeft een lijst (test, bewijs, GESLAAGD/MISLUKT) terug en schrijft die naar <werkmap-dir>/testresultaten.json.
"""
import json
import re
import shutil
import subprocess
import sys
import warnings
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_rdfo as B  # noqa: E402

warnings.filterwarnings("ignore")
WORK = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/rdfo_tests")
RECALC = sys.argv[2] if len(sys.argv) > 2 else "recalc.py"
WORK.mkdir(parents=True, exist_ok=True)
TOL = 0.01
results = []


def record(name, ok, evidence):
    results.append((name, evidence, "GESLAAGD" if ok else "MISLUKT"))
    print(("OK   " if ok else "FAIL ") + name + " — " + evidence)


def recalc(path):
    out = subprocess.run([sys.executable, RECALC, str(path), "600"], capture_output=True, text=True)
    info = json.loads(out.stdout)
    # bewuste #N/A: gaten in de BOR%-lijn van de dashboardgrafiek (versie zonder vastgelegde waarde)
    info["total_errors"] = sum(1 for k, x in info.get("error_summary", {}).items() for loc in x["locations"]
                               if not (k == "#N/A" and re.match(r"Outlook![P-W]\d+$", loc)))
    return info, load_workbook(path, data_only=True)


def name_ref(wb, name):
    sheet, cell = wb.defined_names[name].attr_text.split("!")
    return sheet, cell.replace("$", "")


def set_name(wb, name, value):
    s, c = name_ref(wb, name)
    wb[s][c] = value


def v(wbv, ref):
    s, c = ref.split("!")
    return wbv[s][c].value


def scenario(tag, mutate=None, base=None):
    """Kopie van de basis (formules), mutatie toepassen, opslaan, herberekenen."""
    src = base or WORK / "base.xlsx"
    p = WORK / f"{tag}.xlsx"
    if mutate is None:
        shutil.copy(src, p)
    else:
        wb = load_workbook(src)
        mutate(wb)
        wb.save(p)
    info, wbv = recalc(p)
    return p, info, wbv


def rd(key, colm):
    return f"RDFO!{colm}{B.RDFO_ROWS[key]}"


def col_of(ws, header):
    for c in ws[1]:
        if c.value == header:
            return c.column
    raise KeyError(header)


def row_of(ws, colidx, value):
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=colidx).value == value:
            return r
    raise KeyError(value)


# ------------------------------------------------------------------ onafhankelijke verwachting
D = B.DATA
z = pd.read_excel(D / B.SRC_Z_FILE)
e = pd.read_excel(D / B.SRC_E_FILE)
hist = pd.read_excel(D / B.HIST_FILE, sheet_name=None)
mp = hist["Projectmapping"]
ph = hist["Prognosehistorie"]
wh = hist["Woninghistorie"]
z["fy"] = -z["Bedrijfsopbrengsten LJ"] - z["OP.BOB 2026"]
z["bor"] = -z["Bruto (omzet) result. LJ"] - z["OP.BOR 2026"]
EXP_FY = z["fy"].sum()
EXP_BOR = z["bor"].sum()
EXP_BOR_ABS = z["bor"].abs().sum()
EXP_ACT = (-z["Bedrijfsopbrengsten LJ"]).sum()
EXP_FY27 = (-z["OP.BOB 2027"]).sum()
EXP_VERS = ph[ph.Doeljaar == 2026].groupby("Versie")["Omzet FY EUR"].sum().to_dict()
EXP_VERS27 = ph[ph.Doeljaar == 2027].groupby("Versie")["Omzet FY EUR"].sum().to_dict()
EXP_WON = int(e["Aantal woningen"].sum())
EXP_WON_WEIGHTED = float((e["Aantal woningen"] * e["Aandeel Heijm."].str.replace(",", ".").astype(float) / 100).sum())

# ------------------------------------------------------------------ basis
base = WORK / "base.xlsx"
B.build(str(base))
p, info, w0 = scenario("t01_basis")
ok = info.get("total_errors") == 0 and v(w0, "Controle!D4") == 0 and v(w0, "Controle!D5") == 0
record("Basis: werkmap rekent zonder formulefouten en controles", ok,
       f"{info.get('total_formulas')} formules, {info.get('total_errors')} fouten; Controle: {v(w0, 'Controle!D4')} FOUT, {v(w0, 'Controle!D5')} WAARSCHUWING")

got = {c: v(w0, rd("omz_proj", c)) for c in "CDEFH"}
exp = {"C": EXP_VERS["Budget"], "D": EXP_VERS["Q1"], "E": EXP_VERS["Q2"], "F": EXP_VERS["Q3"], "H": EXP_FY}
ok = all(abs(got[c] * 1000 - exp[c]) < TOL for c in exp)
record("RDFO-versiekolommen = fictieve historie; upload = bron (onafhankelijk nagerekend)", ok,
       "Budget/Q1/Q2/Q3/upload (EUR x 1.000): " + " / ".join(f"{got[c]:,.1f}" for c in "CDEFH"))

verschil = sum(abs(w0["Concept"][f"U{r}"].value or 0) for r in range(B.C0, B.C1 + 1))
ok = abs(v(w0, rd("omz_proj", "F")) - v(w0, rd("omz_proj", "H"))) < TOL and abs(v(w0, rd("bor_tot", "F")) - v(w0, rd("bor_tot", "H"))) < TOL \
    and verschil < TOL
record("Q3-snapshot sluit aan op de bronbestanden", ok,
       f"Q3 vastgelegd {v(w0, rd('omz_proj', 'F')):,.1f} = upload {v(w0, rd('omz_proj', 'H')):,.1f}; som |verschil| alle 280 conceptregels = 0")

ok = v(w0, rd("omz_ops", "G")) == "n.b." and v(w0, rd("bor_tot", "G")) == "n.b." and v(w0, rd("omz_proj", "K")) * 1000 - EXP_ACT < TOL
record("Ontbrekende Q4 en ontbrekende historische actuals niet als nul", ok,
       f"Q4-kolom toont '{v(w0, rd('omz_ops', 'G'))}'; realisatie YTD Q3 = {v(w0, rd('omz_proj', 'K')):,.1f}; vastgelegde versies zonder realisatie (blijven n.b.): {sum(1 for r in range(5, 10) if (w0['CALC_Versies'][f'C{r}'].value or 0) > 0 and w0['CALC_Versies'][f'D{r}'].value == 0)}")

# ------------------------------------------------------------------ nieuwe bronupload
DELTA = 0.1 * 2352000 + 100000


def new_upload(wb):
    ws = wb["Bron_ZPMW03"]
    r = row_of(ws, col_of(ws, "WBS-element"), "G.900005.8.9090")
    for h in ["OP.BOB 2026-Q4", "OP.BOB 2026", "Kosten 2026-Q4", "Kosten 2026", "OP.BOR 2026-Q4", "OP.BOR 2026"]:
        c = col_of(ws, h)
        ws.cell(row=r, column=c).value = ws.cell(row=r, column=c).value * 1.1
    ws.cell(row=r, column=col_of(ws, "Bedrijfsopbrengsten LJ")).value -= 100000
    ws.cell(row=r, column=col_of(ws, "Kosten LJ")).value += 100000


p, info, w2 = scenario("t02_nieuwe_upload", new_upload)
same = all(abs(v(w2, rd(k, c)) - v(w0, rd(k, c))) < TOL for k in ("omz_proj", "bor_tot", "won_fy") for c in "CDEF")
moved = abs((v(w2, rd("omz_proj", "H")) - v(w0, rd("omz_proj", "H"))) * 1000 - DELTA) < TOL
warn = v(w2, "Controle!D4") == 0 and any(w2["Controle"][f"G{r}"].value == "WAARSCHUWING" and "snapshot" in str(w2["Controle"][f"D{r}"].value) for r in range(9, 90))
record("Nieuwe bronupload verandert Budget/Q1/Q2/Q3 niet; alleen upload-concept beweegt", same and moved and warn and info.get("total_errors") == 0,
       f"Budget–Q3 ongewijzigd: {same}; upload +{(v(w2, rd('omz_proj', 'H')) - v(w0, rd('omz_proj', 'H'))):,.1f} (verwacht +{DELTA/1000:,.1f}); "
       f"Controle meldt afwijking upload vs vastgelegde Q3 als WAARSCHUWING: {warn}")


# ------------------------------------------------------------------ Q4 vastleggen (handmatige plakstap gesimuleerd)
def q4_params(wb):
    set_name(wb, "P_UploadVersie", "Q4")
    set_name(wb, "P_ImportID", "IMP-2026-Q4-TEST")


p3, info, w3 = scenario("t03a_q4_concept", q4_params)
n = v(w3, "Concept!E3")
block = [[w3["Concept"].cell(row=r, column=c).value for c in range(1, 19)] for r in range(8, 8 + n)]
revs = {row[6] for row in block}


def paste(wb, rows, times=1):
    ws = wb["Historie"]
    r0 = max(r for r in range(B.L0, B.L1 + 1) if ws[f"J{r}"].value not in (None, "")) + 1
    for t in range(times):
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                ws.cell(row=r0 + t * len(rows) + i, column=1 + j, value=val)
    return r0


def q4_frozen(wb, times=1, select=True):
    q4_params(wb)
    paste(wb, block, times)
    if select:
        set_name(wb, "Sel_Versie", "Q4")


p, info, w4 = scenario("t03b_q4_vastgelegd", q4_frozen)
ok = n == 280 and revs == {0} and abs(v(w4, rd("omz_proj", "G")) - v(w0, rd("omz_proj", "H"))) < TOL \
    and all(abs(v(w4, rd("omz_proj", c)) - v(w0, rd("omz_proj", c))) < TOL for c in "CDEF") \
    and v(w4, "CALC_Versies!B11") == "Vastgesteld" and abs(v(w4, rd("omz_proj", "N"))) < TOL
record("Kwartaalvastlegging Q4: Concept → Historie (dezelfde kopieerstap als de macro)", ok,
       f"{n} regels, revisie {sorted(revs)}; Q4-kolom = {v(w4, rd('omz_proj', 'G')):,.1f}; Budget–Q3 ongewijzigd; geselecteerd = {v(w4, 'CALC_Versies!B11')}")


def q4_then_new_upload(wb):
    q4_frozen(wb)
    new_upload(wb)


p, info, w5 = scenario("t03c_q4_daarna_nieuwe_upload", q4_then_new_upload)
ok = abs(v(w5, rd("omz_proj", "G")) - v(w4, rd("omz_proj", "G"))) < TOL and abs((v(w5, rd("omz_proj", "H")) - v(w4, rd("omz_proj", "H"))) * 1000 - DELTA) < TOL
record("Vastgelegde Q4 is onveranderlijk: latere upload herberekent haar niet", ok,
       f"Q4 vóór/na nieuwe upload: {v(w4, rd('omz_proj', 'G')):,.1f} / {v(w5, rd('omz_proj', 'G')):,.1f}; upload-concept {v(w5, rd('omz_proj', 'H')):,.1f}")

p, info, w6 = scenario("t04_dubbele_import", lambda wb: q4_frozen(wb, times=2))
ok = abs(v(w6, rd("omz_proj", "G")) - v(w4, rd("omz_proj", "G"))) < TOL and v(w6, "Historie!M3") == 280 and v(w6, "Controle!D4") == 0
record("Dubbele import veroorzaakt geen dubbeltelling", ok,
       f"Q4 na 2x plakken: {v(w6, rd('omz_proj', 'G')):,.1f} (1x: {v(w4, rd('omz_proj', 'G')):,.1f}); {v(w6, 'Historie!M3')} regels als dubbel genegeerd en gemeld")


# ------------------------------------------------------------------ correctie met revisie
def correction(wb, toel="Test: correctie omzet Q2 na review"):
    q2 = float(ph[(ph["WBS-element"] == "G.900001.8.9090") & (ph.Versie == "Q2") & (ph.Doeljaar == 2026)]["Omzet FY EUR"].iloc[0])
    m = {r[0]: r for r in hist["Projectmapping"].itertuples(index=False)}["G.900001.8.9090"]
    row = ["COR-2026-Q2-01", "handmatige correctie", B.IMPORTDATUM, "Vastgesteld", 2026, "Q2", 1, pd.Timestamp(2026, 6, 30).to_pydatetime(),
           2026, "G.900001.8.9090", m[1], m[2], m[5], "J", B.K_OMZ, q2 + 100000, "EUR", toel]
    paste(wb, [row])


p, info, w7 = scenario("t05_correctie_revisie", correction)
d = (v(w7, rd("omz_ops", "E")) - v(w0, rd("omz_ops", "E"))) * 1000
hdr = v(w7, "RDFO!E6")
ok = abs(d - 100000) < TOL and "●" in str(hdr) and v(w7, "Controle!D4") == 0 and all(abs(v(w7, rd("omz_ops", c)) - v(w0, rd("omz_ops", c))) < TOL for c in "CDF")
record("Correctie op vastgestelde versie = nieuwe revisie, oude regel blijft, geen dubbeltelling", ok,
       f"Q2 opstal +{d:,.0f} EUR (correctie +100.000, niet +5,5 mln); kolomkop gemarkeerd met ●; overige versies ongewijzigd")
p, info, w8 = scenario("t05b_correctie_zonder_toelichting", lambda wb: correction(wb, toel=None))
ok = v(w8, "Controle!D4") >= 1
record("Correctie zonder toelichting wordt als FOUT gemeld", ok, f"Controle: {v(w8, 'Controle!D4')} FOUT")

# ------------------------------------------------------------------ woningen en verlies
ws = w0["Woningen"]
row = next(r for r in range(1, ws.max_row + 1) if ws[f"B{r}"].value == "Meerhovenstad, Het Kompas")
ok = v(w0, "Woningen!B5") == EXP_WON and ws[f"F{row}"].value == 24 and ws[f"E{row}"].value == 50 and abs(ws[f"G{row}"].value - 48) < TOL
record("Woningen worden niet dubbel gewogen (aandeelbasis blijft aandeelbasis)", ok,
       f"Totaal {v(w0, 'Woningen!B5')} = brontotaal {EXP_WON} (met nogmaals aandeel zou {EXP_WON_WEIGHTED:,.0f} zijn); deelneming Het Kompas 50%: 24 st. (bruto afgeleid 48)")

ws = w0["Projecten"]
r8 = next(r for r in range(10, 120) if ws[f"A{r}"].value == "G.900008.8.9090")
ok = ws[f"N{r8}"].value < 0 and ws[f"R{r8}"].value == "Verlies" and abs(v(w0, rd("bor_tot", "H")) * 1000 - EXP_BOR) < TOL
record("Verlies blijft zichtbaar (geen ABS)", ok,
       f"Lindehaven, De Brink BOR {ws[f'N{r8}'].value:,.2f} (x1.000) gemarkeerd 'Verlies'; RDFO BOR {v(w0, rd('bor_tot', 'H')) * 1000:,.0f} = signed som {EXP_BOR:,.0f} (ABS zou {EXP_BOR_ABS:,.0f} geven)")
pk = next(r for r in range(1, w0["Woningen"].max_row + 1) if w0["Woningen"][f"B{r}"].value == "Boswaard, Parkeergebouw")
record("Parkeergebouw zonder woningen: geen deling door nul", info.get("total_errors") == 0 and w0["Woningen"][f"F{pk}"].value == "n.v.t.",
       f"Projectmatrix toont '{w0['Woningen'][f'F{pk}'].value}', % verkocht leeg; 0 formulefouten")


# ------------------------------------------------------------------ alias en kolomvolgorde
def alias_and_order(wb):
    ws = wb["Bron_ZPMW03"]
    ws.cell(row=1, column=col_of(ws, "Naam maatschappij")).value = "Bedrijfsnaam"
    a, b = col_of(ws, "OP.BOB 2026"), col_of(ws, "Risicoklasse")
    for r in range(1, ws.max_row + 1):
        ws.cell(row=r, column=a).value, ws.cell(row=r, column=b).value = ws.cell(row=r, column=b).value, ws.cell(row=r, column=a).value
    we = wb["Bron_Woningetalage"]
    a, b = col_of(we, "Huidig jaar"), col_of(we, "Volgend jaar")
    for r in range(1, we.max_row + 1):
        we.cell(row=r, column=a).value, we.cell(row=r, column=b).value = we.cell(row=r, column=b).value, we.cell(row=r, column=a).value


p, info, w9 = scenario("t06_alias_kolomvolgorde", alias_and_order)
ok = abs(v(w9, rd("omz_proj", "H")) - v(w0, rd("omz_proj", "H"))) < TOL and v(w9, rd("won_fy", "H")) == v(w0, rd("won_fy", "H")) \
    and v(w9, "Controle!D4") == 0 and v(w9, "VERW_Project!" + B.VCOLS["I_Bedrijfsnaam"] + "6") == "Demo Ontwikkeling BV"
record("Alias Bedrijfsnaam/Naam maatschappij en gewijzigde kolomvolgorde", ok,
       "Header hernoemd + 'OP.BOB 2026' en 'Huidig jaar' van positie gewisseld: uitkomsten identiek, 0 FOUT")


# ------------------------------------------------------------------ nieuw en vervallen project
def new_and_dropped(wb):
    ws = wb["Bron_ZPMW03"]
    wc = col_of(ws, "WBS-element")
    rdrop = row_of(ws, wc, "G.900020.8.9090")
    for c in range(1, ws.max_column + 1):
        ws.cell(row=rdrop, column=c).value = ws.cell(row=2, column=c).value
    ws.cell(row=rdrop, column=wc).value = "G.900099.8.9090"
    ws.cell(row=rdrop, column=col_of(ws, "Projectdefinitie")).value = "G.900099"


p, info, w10 = scenario("t07_nieuw_vervallen", new_and_dropped)
wc = w10["Controle"]
txt = {wc[f"D{r}"].value: (wc[f"E{r}"].value, wc[f"G{r}"].value) for r in range(9, 90)}
nieuw = txt["Nieuwe projecten: ZPMW03-WBS niet in Mapping"]
verv = txt["Vervallen / ontbrekend: Mapping-WBS niet in ZPMW03"]
lists = [wc.cell(row=r, column=c).value for r in range(60, 130) for c in range(2, 10)]
ok = nieuw == (1, "FOUT") and verv == (1, "WAARSCHUWING") and "G.900099.8.9090" in lists and "Boswaard, Parkeergebouw" in lists
record("Nieuwe en vervallen projecten verschijnen als uitzondering (niet stil weg)", ok,
       f"nieuw: {nieuw}; vervallen: {verv}; uitzonderingslijst bevat G.900099 en Boswaard, Parkeergebouw")


# ------------------------------------------------------------------ doeljaar 2027 en demo-aannames
p, info, w11 = scenario("t08_doeljaar_2027", lambda wb: set_name(wb, "Sel_Doeljaar", 2027))
ok = abs(v(w11, rd("omz_proj", "F")) * 1000 - EXP_VERS27["Q3"]) < TOL and abs(v(w11, rd("omz_proj", "H")) * 1000 - EXP_FY27) < TOL \
    and v(w11, rd("omz_proj", "K")) == "n.v.t." and abs(v(w11, rd("omz_proj", "L")) - v(w11, rd("omz_proj", "J"))) < TOL
record("Doeljaar 2027: jaarprognoses per versie, geen realisatie (n.v.t.), to go = FY", ok,
       f"Q3 2027 = {v(w11, rd('omz_proj', 'F')):,.1f}; upload 2027 = {v(w11, rd('omz_proj', 'H')):,.1f}; realisatie '{v(w11, rd('omz_proj', 'K'))}'")

p, info, w12 = scenario("t09_demo_aannames", lambda wb: set_name(wb, "P_Demo", "Ja"))
ebit = v(w12, rd("ebit", "H")) - v(w12, rd("bor_tot", "H"))
ok = abs(v(w12, rd("ak", "H")) + 1400) < TOL and abs(ebit - (-1400 + 300 + 150)) < TOL and v(w12, "Controle!D5") >= 1 and v(w0, rd("ak", "H")) == "n.b."
record("Fictieve demo-aannames alleen zichtbaar na expliciete keuze (en dan gemeld)", ok,
       f"P_Demo=Nee: AK-dekking '{v(w0, rd('ak', 'H'))}'; P_Demo=Ja: AK-dekking {v(w12, rd('ak', 'H')):,.0f} en Controle-waarschuwing")


# ------------------------------------------------------------------ stoplicht
ws = w0["Stoplicht"]
rows_sl = {ws[f"A{r}"].value: r for r in range(B.SL0, B.SL1 + 1) if ws[f"A{r}"].value}
r_wg = rows_sl["Duinwaard, De Watergaarde"]      # handmatig 0% verkocht, netcongestie Laag
r_pk = rows_sl["Boswaard, Parkeergebouw"]        # geen woningen, netcongestie Hoog
codes = lambda w, r: [w["Stoplicht"][f"{c}{r}"].value for c in "IJK"]  # noqa: E731
hdrs = [w0["Stoplicht"].cell(row=7, column=c).value for c in range(1, 12)]
ok = codes(w0, r_wg) == [1, 3, 2] and codes(w0, r_pk) == [0, 3, 1] and not any("RvS" in str(h) or "Totaal" in str(h) for h in hdrs)
record("Stoplicht (verkoop, omgevingsvergunning, netcongestie): kleur volgt uit invoer en uitgangspunten", ok,
       f"Watergaarde (0% verkocht, net Laag) → {codes(w0, r_wg)}; Parkeergebouw (geen woningen, net Hoog) → {codes(w0, r_pk)} (0 = leeg/n.v.t.)")


def stoplicht_rules(wb):
    st = wb["Stoplicht"]
    st["U16"] = 0.0           # oranje vanaf 0% ...
    st["U15"] = 0.0           # ... en groen vanaf 0% → verkoop altijd groen
    for r in range(1, 60):    # netcongestie 'Laag' → groen
        if st[f"S{r}"].value == "Netcongestie":
            for rr in range(r + 2, r + 7):
                if st[f"T{rr}"].value == "Laag":
                    st[f"U{rr}"] = "Groen"


p, info, w13 = scenario("t10_stoplicht_uitgangspunten", stoplicht_rules)
ok = codes(w13, r_wg) == [3, 3, 3]
record("Stoplicht: uitgangspunten aanpassen kleurt alle lichten direct mee", ok,
       f"Groengrens 0% en netcongestie Laag = groen → Watergaarde {codes(w13, r_wg)}")

# ------------------------------------------------------------------ orderintake
zz = pd.read_excel(D / B.SRC_Z_FILE)
zz["oi"] = pd.to_datetime(zz["Orderintake"])
exp_oi = float(-zz[(zz.oi.dt.year == 2026) & (zz.oi <= pd.Timestamp(2026, 9, 30))]["PEW opbrengst"].sum())
exp_n = int(((zz.oi.dt.year == 2026) & (zz.oi <= pd.Timestamp(2026, 9, 30))).sum())
got_oi = v(w0, "Dashboard!Q8") * 1e6
ow = w0["Outlook"]
orow = {ow[f"B{r}"].value: r for r in range(5, 30) if ow[f"B{r}"].value}
oi_years = {y: float(-zz[zz.oi.dt.year == y]["PEW opbrengst"].sum()) for y in (2025, 2026)}
ok = (abs(got_oi - exp_oi) < 1 and str(v(w0, "Dashboard!Q9")).startswith(str(exp_n))
      and all(abs(ow[f"{c}{orow['Orderintake PEW-omzet']}"].value * 1000 - oi_years[y]) < 1 for c, y in (("C", 2025), ("D", 2026))))
record("Orderintake per jaar/kwartaal = bron (kolom Orderintake, PEW opbrengst)", ok,
       f"2026 t/m peildatum: {got_oi:,.0f} EUR in {exp_n} projecten (onafhankelijk: {exp_oi:,.0f})")

# ------------------------------------------------------------------ outlook: alle jaren
ee = pd.read_excel(D / B.SRC_E_FILE)
won28 = int(ee[2028].sum())
won26 = int(ee["Huidig jaar"].sum())
omz27 = float(-zz["OP.BOB 2027"].sum())
omz28 = float(-zz["OP.BOB 2028"].sum())
ok = (ow[f"F{orow['Woningverkopen (aantal)']}"].value == won28 and ow[f"D{orow['Woningverkopen (aantal)']}"].value == won26
      and abs(ow[f"E{orow['Omzet totaal']}"].value * 1000 - omz27) < 1 and abs(ow[f"F{orow['Omzet totaal']}"].value * 1000 - omz28) < 1
      and ow["F5"].value == "2028" and ow["G5"].value == "2029")
record("Outlook: omzet, woningverkopen en orderintake voor elk jaar (2025 t/m 2032)", ok,
       f"woningverkopen 2026/2028 = {won26}/{won28}; omzet 2027/2028 = {omz27/1e6:,.1f}/{omz28/1e6:,.1f} mln (= bron); 2029+ in bron 0")

# ------------------------------------------------------------------ PowerPoint-inhoud en macro-verpakking
tp_, to_ = B.DIA_BLOCKS["T_PROJ1"], B.DIA_BLOCKS["T_OUT"]
ok = v(w0, "PPT!B14") == 20 and str(v(w0, "PPT!B6")).startswith("Omzet FY 2026") and v(w0, "Dias!C4") == v(w0, "Dashboard!B8") \
    and v(w0, f"Dias!B{tp_+1}") == "Lindehaven, Parkzoom fase 1"
record("PowerPoint-inhoud staat klaar (tabbladen PPT en Dias)", ok,
       f"{v(w0, 'PPT!B14')} projecten in de projecttabel; kernboodschap: '{v(w0, 'PPT!B6')}'")
gew_toel = v(w0, f"Dias!M{tp_-1}")
gew_jaar = [v(w0, f"Dias!{c}{to_-1}") for c in "CDEFGHIJ"]
ok = gew_toel == 0 and gew_jaar[:4] == [1, 1, 1, 1] and gew_jaar[4:] == [0, 0, 0, 0]
record("PowerPoint: lege toelichtingskolom en jaren zonder cijfers vallen weg", ok,
       f"gewicht toelichting = {gew_toel}; gewichten jaarkolommen 2025-2032 = {gew_jaar}")

import zipfile  # noqa: E402
from oletools.olevba import VBA_Parser  # noqa: E402
import inject_cache  # noqa: E402
import package_xlsm  # noqa: E402

inj = WORK / "pkg_cache.xlsx"
inject_cache.main(str(base), str(WORK / "t01_basis.xlsx"), str(inj))
xlsm = WORK / "pkg.xlsm"
n_btn = package_xlsm.main(str(inj), str(xlsm))
vp = VBA_Parser(str(xlsm))
mods = {n: src for (_f, _s, n, src) in vp.extract_macros()}
subs = set()
for src in mods.values():
    for ln in src.splitlines():
        if ln.startswith("Public Sub "):
            subs.add(ln.split()[2].split("(")[0])
zx = zipfile.ZipFile(xlsm)
macros_in_buttons = set()
for nme in zx.namelist():
    if nme.startswith("xl/drawings/drawing"):
        macros_in_buttons |= {m.split("!")[1] for m in __import__("re").findall(r'macro="([^"]+)"', zx.read(nme).decode()) if "!" in m}
ct = zx.read("[Content_Types].xml").decode()
ok = ("modRDFO.bas" in mods and macros_in_buttons <= subs and len(macros_in_buttons) == 6 and n_btn == len(B.BUTTONS)
      and "macroEnabled.main+xml" in ct and "shStoplicht.cls" in mods)
record("Macro-bestand: VBA-project leesbaar, alle knoppen verwijzen naar bestaande macro's", ok,
       f"{len(mods)} VBA-modules (oletools); {n_btn} knoppen → {sorted(macros_in_buttons)}. Let op: uitvoering in Excel zelf is hier niet getest.")

errs = [r for r in results if r[2] != "GESLAAGD"]
(WORK / "testresultaten.json").write_text(json.dumps(results, ensure_ascii=False, indent=1, default=str))
print(f"\n{len(results) - len(errs)}/{len(results)} geslaagd")
sys.exit(1 if errs else 0)
