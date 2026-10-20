"""Sjabloon v18 -> v19: minder tekst en scenariolijnen met meer contrast.

    python3 tools/pptx_v19_kort.py INVOER_v18.pptx UITVOER_v19.pptx [--render MAP] [--simuleer doorgerekend.xlsx UIT.pptx]

* Weg (krijgen sinds macro v21 geen tekst meer; de macro haalt ze ook uit een v18-sjabloon): de ondertitels van dia 2, 3, 4 en 7
  (KPI_SUBTITEL, CF_SUBTITEL, VS_SUBTITEL, VP_SUBTITEL), CF_VOET, KPI4_SUB, de regel onderaan de scenariokaarten (VS_Kk_DELTA)
  en de voetnoot van dia 4 (VS_NOOT).
* Scenariolijnen in layout.VS_KLEUREN met een eigen lijnstijl (layout.VS_STREEP): in CF_GRAFIEK (dia 3, kolom I..K) en
  VS_GRAFIEK (dia 4, kolom C..E; de dieptepunten in kolom G..I in dezelfde kleur), de accentstrepen VS_Kk_ACCENT en de
  legendalijnen VS_LEGk_LIJN.
De overige dia's en vormen blijven zoals ze zijn.

--simuleer: doet op de nieuwe v19 na wat de macro met dia 2, 3 en 4 doet (zie pptx_v17_kerncijfers en pptx_v18_scenarios)."""
import argparse
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import layout as LY  # noqa: E402
from pptx_v17_kerncijfers import Gegevens as GegevensDia2, simuleer as simuleer_dia2  # noqa: E402
from pptx_v18_scenarios import Gegevens, _kolom, render_nl, simuleer  # noqa: E402

WEG = ["KPI_SUBTITEL", "KPI4_SUB", "CF_SUBTITEL", "CF_VOET", "VS_SUBTITEL", "VS_NOOT", "VP_SUBTITEL"] + \
      [f"VS_K{k}_DELTA" for k in range(LY.N_VS + 1)]
CF_KOL = {kol: s for s, kol in zip(range(1, LY.N_VS + 1), "IJK")}
VS_KOL = {**{kol: s for s, kol in zip(range(1, LY.N_VS + 1), "CDE")}, **{kol: s for s, kol in zip(range(1, LY.N_VS + 1), "GHI")}}


def _kleur(el, kleur):
    """alle srgbClr onder el op kleur (vult, lijn en markering)"""
    for c in el.iter(qn("a:srgbClr")):
        c.set("val", kleur)


def _streep(ln, streep):
    for oud in ln.findall(qn("a:prstDash")):
        ln.remove(oud)
    d = ln.makeelement(qn("a:prstDash"), {"val": streep})
    # volgorde in a:ln: fill, prstDash, ... (na de vulling)
    fill = ln.find(qn("a:solidFill"))
    if fill is not None:
        fill.addnext(d)
    else:
        ln.insert(0, d)


def kleur_reeksen(chart, kolommen):
    """reeksen op de gegeven kolommen (kolom -> scenario) in de scenariokleur; lijnen ook in de scenariostijl"""
    n = 0
    for ser in chart._chartSpace.xpath(".//c:plotArea/*/c:ser"):
        f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        s = kolommen.get(_kolom(f.text) if f is not None else "")
        if not s:
            continue
        sppr = ser.find(qn("c:spPr"))
        ln = sppr.find(qn("a:ln")) if sppr is not None else None
        if ln is not None and ln.find(qn("a:solidFill")) is not None:
            _kleur(ln, LY.VS_KLEUREN[s])
            ln.set("w", str(int(2.25 * 12700)))
            _streep(ln, LY.VS_STREEP[s])
        marker = ser.find(qn("c:marker") + "/" + qn("c:spPr"))
        if marker is not None:
            fill = marker.find(qn("a:solidFill"))
            if fill is not None:
                _kleur(fill, LY.VS_KLEUREN[s])
        n += 1
    return n


def bewerk(prs):
    weg, gekleurd = [], 0
    for slide in prs.slides:
        for sh in list(slide.shapes):
            if sh.name in WEG:
                sh._element.getparent().remove(sh._element)
                weg.append(sh.name)
        for sh in slide.shapes:
            if sh.name == "CF_GRAFIEK" and sh.has_chart:
                gekleurd += kleur_reeksen(sh.chart, CF_KOL)
            elif sh.name == "VS_GRAFIEK" and sh.has_chart:
                gekleurd += kleur_reeksen(sh.chart, VS_KOL)
            for k in range(1, LY.N_VS + 1):
                if sh.name == f"VS_K{k}_ACCENT":
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = RGBColor.from_string(LY.VS_KLEUREN[k])
                elif sh.name == f"VS_LEG{k}_LIJN":
                    ln = sh._element.find(".//" + qn("a:ln"))
                    if ln is not None:
                        _kleur(ln, LY.VS_KLEUREN[k])
                        _streep(ln, LY.VS_STREEP[k])
                        ln.set("w", str(int(2.25 * 12700)))
    return weg, gekleurd


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("--render", default="", help="map voor PNG's (LibreOffice + pdftoppm); leeg = niet renderen")
    ap.add_argument("--simuleer", nargs=2, metavar=("WERKBOEK", "UIT_PPTX"), default=None,
                    help="doe op de nieuwe v19 na wat de macro met dia 2, 3 en 4 doet en bewaar het als UIT_PPTX")
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    if not any(sh.name == "VS_GRAFIEK" for s in prs.slides for sh in s.shapes):
        sys.exit("Geen VS_GRAFIEK gevonden: is dit een v18-sjabloon?")
    weg, gekleurd = bewerk(prs)
    prs.save(args.uitvoer)
    print(f"Opgeslagen: {args.uitvoer}\n  weg: {', '.join(weg) or '-'}\n  scenarioreeksen herkleurd: {gekleurd}")
    if args.render:
        pngs = render_nl(args.uitvoer, args.render)
        print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd"))
    if args.simuleer:
        prs2 = Presentation(args.uitvoer)
        simuleer_dia2(prs2, GegevensDia2(args.simuleer[0]))
        simuleer(prs2, Gegevens(args.simuleer[0]))
        prs2.save(args.simuleer[1])
        print(f"  simulatie: {args.simuleer[1]}")
        if args.render:
            pngs = render_nl(args.simuleer[1], os.path.join(args.render, "simulatie"))
            print("  PNG's simulatie: " + (", ".join(pngs) if pngs else "niet gerenderd"))


if __name__ == "__main__":
    main()
