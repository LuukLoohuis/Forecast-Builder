#!/usr/bin/env python3
"""Kwartaalsjabloon v13 -> v14: dia 8 'Opbrengsten per woningtype' (PT_GRAFIEK) en herkleurde dia 5 en 6.

Gebruik:
    python3 tools/pptx_v14_daeb.py INVOER_v13.pptx UITVOER_v14.pptx [UITVOER2.pptx ...] --voorbeeld doorgerekend.xlsx [--render MAP]

Bewerking (python-pptx + lxml; hulpfuncties uit tools/pptx_v12_bouwtermijnen.py en tools/pptx_v13_banen.py):
  * Dia 5 (BT_GRAFIEK) en dia 6 (VP1_/VP2_GRAFIEK): chart-XML opnieuw uit chartxml met de huidige kleuren (vorige prognose
    lichtrood, typekleur 8 olijf); de ingebedde werkboeken blijven, alleen de koppen '(uit)' in het voorbeeld worden 'Type k'.
  * Nieuwe dia 8 (kopie van de tekstvormen van dia 7): PT_KICKER, PT_TITEL, PT_SUBTITEL, PT_KAART (volle breedte) en PT_GRAFIEK
    uit chartxml.type_paneel (gestapelde kolommen grondtermijn / bouwtermijnen / extra's / fee 1 / fee 2, cumulatieve lijn,
    grijze waas, legenda onderaan). Ingebed werkboek: rij 1 de koppen van het blok FM7.. van tab PowerPoint (A Kwartaal, dan
    de zes reeksen, H Realisatie), daaronder de voorbeeldrijen uit het doorgerekende werkboek (--voorbeeld, selectieblok van
    het gekozen type). De macro maakt per woningtype een kopie van deze dia en vult die (VulDiasPerType); dia 8 zelf gaat weg.
  * De oude dia 8 (afsluiting) wordt dia 9; dianummervelden worden bijgewerkt.
"""
import argparse
import copy
import os
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import chartxml as cx, layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import (EMU_CM, VP_LBL, VP_REFS, VP_TYPES, controleer, dianummers_bijwerken, hernoem, nieuwe_grafiek,  # noqa: E402
                                    q, render, set_text_keep_format, shapes_by_name, zet_chart_xml)
from pptx_v13_banen import KOLOMMEN as BT_KOLOMMEN, refs as bt_refs  # noqa: E402

PT_TEKST = {"PT_KICKER": "OPBRENGSTEN PER WONINGTYPE", "PT_TITEL": "Opbrengsten per woningtype",
            "PT_SUBTITEL": "per kwartaal · grondtermijn, bouwtermijnen en extra's (koop) of fees (DAEB) · lijn = cumulatief · grijs = gerealiseerd"}
PT_KOLOMMEN = ["kwartaal"] + LY.PT_REEKSEN + ["realisatie"]
PT_KOL = {naam: get_column_letter(i + 1) for i, naam in enumerate(PT_KOLOMMEN)}
PT_NAMEN_LEEG = {"grond": "Grondtermijn", "bouw": "Bouwtermijnen", "extra": "Extra opbrengsten", "fee1": "Fee 1", "fee2": "Fee 2", "cum": "Cumulatief",
                 "totaal": "Totaal per kwartaal", "eind": "Eindpunt"}
DIA_VO = 6                      # dia 7 (0-gebaseerd): bron van de tekstvormen
DIA_PT_POS = 7                  # nieuwe dia 8 komt op positie 7 in de sldIdLst


def werkboek_van(gf):
    """Het ingebedde werkboek (Sheet1) van een grafiekvak als openpyxl-werkblad."""
    import io
    blob = gf.chart.part.chart_workbook.xlsx_part.blob
    return load_workbook(io.BytesIO(blob)).active


def pt_refs(n):
    r = {naam: f"Sheet1!${PT_KOL[naam]}$2:${PT_KOL[naam]}${n + 1}" for naam in PT_KOLOMMEN}
    r.update({f"naam_{naam}": f"Sheet1!${PT_KOL[naam]}$1" for naam in LY.PT_REEKSEN})
    return r


def werkboek_pt(pad_voorbeeld):
    """Sheet1 met de koppen van het blok op tab PowerPoint en de voorbeeldrijen van het gekozen type (Model, selectieblok)."""
    wb = load_workbook(pad_voorbeeld, data_only=True)
    m, pp = wb["Model"], wb["PowerPoint"]
    n = int(m[f"{LY.M_HULP}{LY.H['n']}"].value or 0)
    c0 = pp[LY.PP_BLOK["pt"] + "1"].column
    koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_PT_KOL)]
    for j, naam in enumerate(LY.PT_REEKSEN, start=1):
        if koppen[j] in (None, "", "(uit)"):
            koppen[j] = PT_NAMEN_LEEG[naam]
    ws = Workbook().active
    ws.title = "Sheet1"
    ws.append([str(k) for k in koppen])
    for i in range(n):
        rij = []
        for j in range(LY.PP_PT_KOL):
            v = pp.cell(LY.PP_KPI_ROW1 + i, c0 + j).value
            rij.append(None if (isinstance(v, str) and v.startswith("#")) else v)
        ws.append(rij)
    ws.column_dimensions["A"].width = 12
    titel = m[f"{LY.M_HULP}{LY.H['pt_titel']}"].value
    sub = m[f"{LY.M_HULP}{LY.H['pt_subtitel']}"].value
    return ws, n, titel, sub


def herkleur_dia5_6(prs):
    """Chart-XML van BT_GRAFIEK, VP1_GRAFIEK en VP2_GRAFIEK opnieuw uit chartxml (kleuren), met de bestaande voorbeeldgegevens."""
    by5 = shapes_by_name(prs.slides[4])
    set_text_keep_format(by5["BT_SUBTITEL"], "typekleur = huidige planning · lichtrood = vorige prognose · grijs = gerealiseerd")
    gf = by5["BT_GRAFIEK"]
    ws = werkboek_van(gf)
    n = ws.max_row - 1
    for j, naam in enumerate(BT_KOLOMMEN, start=1):          # voorbeeldkoppen '(uit)' -> 'Type k' (de macro plakt zijn eigen koppen)
        v = ws.cell(1, j).value
        if v == "(uit)":
            if naam.startswith("cur"):
                ws.cell(1, j).value = f"Type {naam[3:]}"
            elif naam.startswith("prev"):
                ws.cell(1, j).value = f"Type {naam[4:]} · vorige prognose"
            else:
                ws.cell(1, j).value = ""
    t_first = min(v for v in (ws[f"{get_column_letter(BT_KOLOMMEN.index('s0') + 1)}{r}"].value for r in range(2, n + 2)) if isinstance(v, (int, float)))
    s6 = get_column_letter(BT_KOLOMMEN.index("s6") + 1)
    kw1 = get_column_letter(BT_KOLOMMEN.index("kw1") + 1)
    t_end = None
    for r in range(2, n + 2):
        if isinstance(ws[f"{kw1}{r}"].value, (int, float)):      # koprij kwartalen: t_first + aantal kwartaalblokjes
            t_end = ws[f"{s6}{r}"].value + sum(1 for j, naam in enumerate(BT_KOLOMMEN, start=1)
                                               if naam.startswith("kw") and isinstance(ws.cell(r, j).value, (int, float)))
    zet_chart_xml(gf, cx.bouwtermijnen(bt_refs(n), int(t_first), int(t_end) if t_end else None, legenda=True), ws, n)
    by6 = shapes_by_name(prs.slides[5])
    for naam, soort in (("VP1_GRAFIEK", "verkocht"), ("VP2_GRAFIEK", "transport")):
        gf = by6[naam]
        ws = werkboek_van(gf)
        n = ws.max_row - 1
        zet_chart_xml(gf, cx.verkoop_paneel(VP_REFS, VP_LBL, VP_TYPES, soort, hoogte_cm=gf.height / EMU_CM), ws, n)


def dia8_pertype(prs, pad_voorbeeld):
    bron = prs.slides[DIA_VO]
    slide = prs.slides.add_slide(bron.slide_layout)
    for sh in list(slide.shapes):
        sh._element.getparent().remove(sh._element)
    tree = slide.shapes._spTree
    bron_namen = shapes_by_name(bron)
    for el in bron.shapes._spTree:
        if el.tag != q("p:sp"):
            continue
        naam = el.find(".//" + q("p:cNvPr")).get("name")
        if naam in ("VO_KICKER", "VO_TITEL", "VO_SUBTITEL", "VO_KAART") or naam.startswith("Slide Number"):
            tree.append(copy.deepcopy(el))
    by_name = shapes_by_name(slide)
    for oud, nieuw in (("VO_KICKER", "PT_KICKER"), ("VO_TITEL", "PT_TITEL"), ("VO_SUBTITEL", "PT_SUBTITEL"), ("VO_KAART", "PT_KAART")):
        hernoem(by_name[oud], nieuw)
        if nieuw in PT_TEKST:
            set_text_keep_format(by_name[oud], PT_TEKST[nieuw])
    kaart = by_name["VO_KAART"]
    kaart.width = bron_namen["VO_STAT1"].left + bron_namen["VO_STAT1"].width - kaart.left     # volle breedte, zoals VP_KAART
    g = bron_namen["VO_GRAFIEK"]
    box = (g.left, g.top, kaart.left + kaart.width - (g.left - kaart.left) - g.left, g.height)
    ws, n, titel, sub = werkboek_pt(pad_voorbeeld)
    if titel:
        set_text_keep_format(by_name["VO_TITEL"], str(titel))
    if sub:
        set_text_keep_format(by_name["VO_SUBTITEL"], str(sub))
    gf = nieuwe_grafiek(slide, "PT_GRAFIEK", box, XL_CHART_TYPE.COLUMN_STACKED, n, len(LY.PT_REEKSEN) + 1, "chart_pt", "Werkblad_pt")
    zet_chart_xml(gf, cx.type_paneel(pt_refs(n), hoogte_cm=box[3] / EMU_CM, legenda=True), ws, n)
    nummer = by_name.get("Slide Number Placeholder 0")
    if nummer is not None:
        nummer._element.addprevious(gf._element)
    lst = prs.slides._sldIdLst
    el = lst[-1]
    lst.remove(el)
    lst.insert(DIA_PT_POS, el)
    return slide, n, box


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieën van de uitvoer")
    ap.add_argument("--voorbeeld", required=True, help="doorgerekend werkboek (.xlsx) met de voorbeeldgegevens (tab Model en PowerPoint)")
    ap.add_argument("--render", default="", help="map voor PNG's van alle dia's (LibreOffice + pdftoppm); leeg = niet renderen")
    args = ap.parse_args()

    prs = Presentation(args.invoer)
    if len(prs.slides) != 8:
        sys.exit(f"Verwacht 8 dia's in v13, gevonden {len(prs.slides)}.")
    if "PT_GRAFIEK" in shapes_by_name(prs.slides[DIA_PT_POS]):
        sys.exit("PT_GRAFIEK bestaat al in de invoer; niets gedaan.")
    herkleur_dia5_6(prs)
    slide, n, box = dia8_pertype(prs, args.voorbeeld)
    velden = dianummers_bijwerken(prs)
    prs.save(args.uitvoer)
    for extra in args.extra:
        shutil.copyfile(args.uitvoer, extra)
    print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
    print(f"  dia 5 en 6 herkleurd; dia 8: PT_GRAFIEK op x={box[0]} y={box[1]} w={box[2]} h={box[3]} ({n} voorbeeldrijen); dianummervelden: {velden}")
    print("Controle:")
    fouten, charts, emb = controleer(args.uitvoer, {
        5: ["BT_KICKER", "BT_TITEL", "BT_SUBTITEL", "VT_KAART", "BT_GRAFIEK", "VT_TABELKAART", "VT_TABEL"],
        6: ["VP_KICKER", "VP_TITEL", "VP_SUBTITEL", "VP_KAART", "VP1_GRAFIEK", "VP2_GRAFIEK"],
        7: ["VO_TITEL", "VO_GRAFIEK"],
        8: ["PT_KICKER", "PT_TITEL", "PT_SUBTITEL", "PT_KAART", "PT_GRAFIEK"],
    })
    print("  chart-parts: " + ", ".join(charts))
    print("  ingebedde werkboeken: " + ", ".join(emb))
    if fouten:
        sys.exit("FOUTEN:\n  " + "\n  ".join(fouten))
    print("  geen fouten")
    if args.render:
        pngs = render(args.uitvoer, args.render)
        print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd (LibreOffice/pdftoppm ontbreekt)"))


if __name__ == "__main__":
    main()
