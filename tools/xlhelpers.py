"""Stijl- en schrijfhulpjes voor de RDFO-werkmap (openpyxl)."""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

# Huisstijl: donkerblauw + geel, rustige lichte vlakken
NAVY = "14284B"
NAVY2 = "2E4A7D"
YELLOW = "F5B800"
YELLOW_LT = "FFF4CC"
LIGHT = "F3F5F9"
MID = "D9DEE8"
GREY = "6B7385"
RED = "C0392B"
GREEN = "2E7D32"
ORANGE_LT = "FCE4D6"
WHITE = "FFFFFF"

FONT = "Arial"

FMT_EUR = '#,##0;[Red]-#,##0;"–"'
FMT_EUR1 = '#,##0.0;[Red]-#,##0.0;"–"'
FMT_EUR2 = '#,##0.00;[Red]-#,##0.00;"–"'
FMT_RAW = '#,##0.00;[Red]-#,##0.00;0'
FMT_PCT = '0.0%;[Red]-0.0%;"–"'
FMT_INT = '#,##0;[Red]-#,##0;"–"'
FMT_INT0 = '#,##0;[Red]-#,##0;0'
FMT_DATE = 'dd-mm-yyyy'

thin = Side(style="thin", color=MID)
thick_navy = Side(style="medium", color=NAVY)
yellow_side = Side(style="thick", color=YELLOW)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def font(size=9, bold=False, color="000000", italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def col(n):
    return get_column_letter(n)


def add_name(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def base_sheet(ws, tab=None, zoom=90, grid=False):
    ws.sheet_view.showGridLines = grid
    ws.sheet_view.zoomScale = zoom
    if tab:
        ws.sheet_properties.tabColor = tab


def title(ws, text, sub=None, width_cols=14, row=1):
    """Titelbalk: donkerblauw met gele onderrand."""
    for c in range(1, width_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill(NAVY)
        cell.border = Border(bottom=yellow_side)
    t = ws.cell(row=row, column=2, value=text)
    t.font = font(14, True, WHITE)
    t.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 28
    if sub is not None:
        s = ws.cell(row=row + 1, column=2, value=sub)
        s.font = font(9, False, GREY, italic=True)
    ws.column_dimensions["A"].width = 2


def header(ws, row, c1, labels, wrap=True, height=30, color=NAVY):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=c1 + i, value=lab)
        c.fill = fill(color)
        c.font = font(9, True, WHITE)
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=wrap)
    ws.row_dimensions[row].height = height


def section(ws, row, c1, text, ncols=8):
    for c in range(c1, c1 + ncols):
        ws.cell(row=row, column=c).border = Border(bottom=thick_navy)
    cell = ws.cell(row=row, column=c1, value=text)
    cell.font = font(11, True, NAVY)


def put(ws, ref, value, fmt=None, bold=False, color="000000", fillc=None, align=None, italic=False, size=9, wrap=False):
    c = ws[ref] if isinstance(ref, str) else ws.cell(row=ref[0], column=ref[1])
    c.value = value
    c.font = font(size, bold, color, italic)
    if fmt:
        c.number_format = fmt
    if fillc:
        c.fill = fill(fillc)
    if align or wrap:
        c.alignment = Alignment(horizontal=align, vertical="center" if not wrap else "top", wrap_text=wrap)
    return c


def input_style(c, fmt=None):
    c.fill = fill(YELLOW_LT)
    c.font = font(9, False, "1F4E9A")
    c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if fmt:
        c.number_format = fmt


def helper_style(c):
    c.font = font(8, False, GREY)


def row_border(ws, row, c1, c2, top=False, bold=None, fillc=None):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=row, column=c)
        cell.border = Border(bottom=thin, top=thick_navy if top else None)
        if fillc:
            cell.fill = fill(fillc)
        if bold is not None:
            f = cell.font
            cell.font = Font(name=FONT, size=f.size or 9, bold=bold, color=f.color, italic=f.italic)


def dv_list(ws, ref, formula, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    dv.error = "Kies een waarde uit de lijst."
    dv.showErrorMessage = True
    ws.add_data_validation(dv)
    dv.add(ref)
    return dv


def widths(ws, mapping):
    for k, v in mapping.items():
        ws.column_dimensions[k].width = v


def landscape(ws, area=None):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    if area:
        ws.print_area = area


def text_block(ws, start_row, rows, c1=2, widths_=None, wrap_col=None):
    """Schrijft een eenvoudige tabel (lijst van tuples). Eerste tuple = kop."""
    r = start_row
    header(ws, r, c1, list(rows[0]), height=22)
    for tup in rows[1:]:
        r += 1
        for i, v in enumerate(tup):
            c = ws.cell(row=r, column=c1 + i, value=v)
            c.font = font(9)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=thin)
    return r
