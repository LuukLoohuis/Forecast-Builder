"""Sjabloon v14 -> v15.

    python3 tools/pptx_v15_pertype.py INVOER_v14.pptx UITVOER_v15.pptx [UITVOER2.pptx ...] --voorbeeld doorgerekend.xlsx [--render MAP]

  * dia 3 (CF_GRAFIEK): de stippellijn 'Vorige prognose' lichtblauw (chartxml.VORIG_LIJN) in plaats van grijs.
  * dia 8 (PT_GRAFIEK): nieuwe grafiek-XML uit chartxml.type_paneel met de reeksen van het blok op tab PowerPoint (10 kolommen:
    kwartaal, grondtermijn, bouwtermijnen, extra's, fee 1, fee 2, cumulatief, totaal per kwartaal (label boven de stapel),
    eindpunt (label met het totaal), realisatie); het tijdvak begint twee jaar vóór de eerste opbrengst (de macro vult
    alleen de getoonde periodes, Model!BY128).
De vormnamen blijven gelijk; de macro (KOL_PT = 10) vult het blok."""
import argparse
import os
import re
import shutil
import sys

from pptx import Presentation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import chartxml as cx, layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import EMU_CM, controleer, render, shapes_by_name, zet_chart_xml  # noqa: E402
from pptx_v14_daeb import pt_refs, werkboek_pt  # noqa: E402

DIA_CF = 2          # 0-gebaseerd: dia 3 (cashflow per kwartaal)
DIA_PT = 7          # 0-gebaseerd: dia 8 (opbrengsten per woningtype)
OUD_VORIG = cx.GREY_LINE


def herkleur_vorige_prognose(prs):
    """CF_GRAFIEK: de reeks 'Vorige prognose' (grijze stippellijn) lichtblauw."""
    gf = shapes_by_name(prs.slides[DIA_CF])["CF_GRAFIEK"]
    part = gf.chart.part
    xml = part.blob.decode("utf-8")
    sers = re.findall(r"<c:ser>.*?</c:ser>", xml, flags=re.S)
    doel = [s for s in sers if "Vorige prognose" in s]
    assert len(doel) == 1, f"reeks 'Vorige prognose' {len(doel)}x gevonden in CF_GRAFIEK"
    nieuw = doel[0].replace(f'<a:srgbClr val="{OUD_VORIG}"/>', f'<a:srgbClr val="{cx.VORIG_LIJN}"/>')
    assert nieuw != doel[0], "grijze lijnkleur niet gevonden in de reeks 'Vorige prognose'"
    xml = xml.replace(doel[0], nieuw)
    from lxml import etree
    part._element = etree.fromstring(xml.encode("utf-8"))
    return gf


def vernieuw_dia8(prs, pad_voorbeeld):
    """PT_GRAFIEK: nieuwe chart-XML (totaal boven de stapel, eindlabel, tijdvak) en ingebed werkboek met de voorbeeldrijen."""
    slide = prs.slides[DIA_PT]
    gf = shapes_by_name(slide)["PT_GRAFIEK"]
    ws, n, titel, sub = werkboek_pt(pad_voorbeeld)
    hoogte_cm = gf.height / EMU_CM
    zet_chart_xml(gf, cx.type_paneel(pt_refs(n), hoogte_cm=hoogte_cm, legenda=True), ws, n)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieën van de uitvoer")
    ap.add_argument("--voorbeeld", required=True, help="doorgerekend werkboek (.xlsx) met de voorbeeldgegevens (tab Model en PowerPoint)")
    ap.add_argument("--render", default="", help="map voor PNG's van alle dia's (LibreOffice + pdftoppm); leeg = niet renderen")
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    if len(prs.slides) != 9:
        sys.exit(f"Verwacht 9 dia's in v14 (titel + 8), gevonden {len(prs.slides)}.")
    herkleur_vorige_prognose(prs)
    n = vernieuw_dia8(prs, args.voorbeeld)
    prs.save(args.uitvoer)
    for extra in args.extra:
        shutil.copyfile(args.uitvoer, extra)
    print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
    print(f"  dia 3: 'Vorige prognose' lichtblauw ({cx.VORIG_LIJN}); dia 8: PT_GRAFIEK vernieuwd ({n} voorbeeldrijen, {LY.PP_PT_KOL} kolommen)")
    fouten, charts, emb = controleer(args.uitvoer, {
        3: ["CF_GRAFIEK"],
        8: ["PT_KICKER", "PT_TITEL", "PT_SUBTITEL", "PT_KAART", "PT_GRAFIEK"],
    })
    print("  chart-parts: " + ", ".join(charts))
    if fouten:
        sys.exit("FOUTEN:\n  " + "\n  ".join(fouten))
    print("  geen fouten")
    if args.render:
        pngs = render(args.uitvoer, args.render)
        print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd (LibreOffice/pdftoppm ontbreekt)"))


if __name__ == "__main__":
    main()
