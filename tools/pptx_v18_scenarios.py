"""Sjabloon v17 -> v18: dia 4 'Verkoopscenario's' en de scenariolijnen in de cashflowgrafiek (dia 3).

    python3 tools/pptx_v18_scenarios.py INVOER_v17.pptx UITVOER_v18.pptx [UITVOER2.pptx ...] [--voorbeeld doorgerekend.xlsx]
                                        [--render MAP] [--simuleer doorgerekend.xlsx UIT.pptx]

* Dia 3, CF_GRAFIEK: drie lijnreeksen erbij (scenario 1..3, violet / magenta / pruim, 1,75 pt), met naam en waarden op kolom I, J en
  K van het gegevensblad (de macro schrijft het blok van tab PowerPoint, layout.PP_BLOK['cf'], 11 kolommen, in A1 en zet elke
  reeks op de kolom uit zijn SERIES-formule; een reeks met '(uit)' in de kop gaat weg). De legenda toont ze dus alleen als ze er zijn.
* Na dia 3 een nieuwe dia 4 'Verkoopscenario's' (zelfde indeling als dia 2): kopje, titel en ondertitel (VS_KICKER, VS_TITEL,
  VS_SUBTITEL); vier kaarten VS_K0..VS_K3 (Basis en scenario 1..3: KAART, ACCENT in de kleur van de lijn, NAAM, SUB, LABEL, WAARDE,
  DAL, BE, VERKOCHT, DELTA); een legenda (VS_LEGk_LIJN + VS_LEGk); de grafiek VS_GRAFIEK (lijn per scenario, dieptepunt als stip
  met het bedrag als label; tekenvlak layout.VS_PLOT in kader layout.VS_KADER) met het grijze vlak VS_REALISATIE + VS_REAL_LBL
  erachter, VS_EIND_LBL (eindsaldo) en de voetnoot VS_NOOT.
De macro (v18) vult de teksten, haalt de dia weg als er geen scenario in PowerPoint komt, haalt lege kaarten en legenda-regels weg en
verdeelt de rest (SchikDiaScenarios), vult de grafiek en zet de as, het grijze vlak en het eindsaldo-label (SchikGrafiekVS).
Voorbeeldteksten en -grafiek komen uit een doorgerekend werkboek (--voorbeeld; anders bouwt de tool het voorbeeldproject zelf).

--simuleer: doet op de nieuwe v18 na wat de macro met dia 2, 3 en 4 doet, met de waarden uit een doorgerekend werkboek, zodat de
dia's te beoordelen zijn zonder PowerPoint."""
import argparse
import copy
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import controleer, set_text_keep_format, shapes_by_name  # noqa: E402
from pptx_v17_kerncijfers import (GRIJS, GRIJS2, I, LIJN, NAVY, RICHTING_KLEUR, Gegevens as GegevensDia2, _nofill_spPr,  # noqa: E402
                                  lijn, maak_voorbeeld, nl, rect, simuleer as simuleer_dia2, tekstvak)

DIA_CF = 2                      # 0-gebaseerd: dia 3 (cashflow)
DIA_VS = 3                      # nieuwe dia 4
KICKER = "F08A00"
ARCERING = "EBEAE6"
KLEUR = LY.VS_KLEUREN
LIJN_PT = {0: 2.25, 1: 2.0, 2: 2.0, 3: 2.0}
VS_NAMEN = ["jaar", "basis"] + [f"s{s}" for s in range(1, LY.N_VS + 1)] + ["dal0"] + [f"dal{s}" for s in range(1, LY.N_VS + 1)]
assert len(VS_NAMEN) == LY.PP_VS_KOL
CHART_NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
FMT_LABEL = '"€"0.0"M";"−€"0.0"M"'

# ---- raster van dia 4 (inch) ------------------------------------------------------------------------------------------------
MARGE_L, MARGE_R = 0.45, 9.55
KAART_T, KAART_H, KAART_GAP = 1.27, 1.22, 0.12
LEG_T = 2.58
FRAME_L, FRAME_T, FRAME_R, FRAME_B = LY.VS_KADER
PLOT_L, PLOT_T, PLOT_R, PLOT_B = LY.VS_PLOT


# ---- gegevens uit een doorgerekend werkboek --------------------------------------------------------------------------------------
class Gegevens:
    """Teksten (en richting) van dia 3 en 4 zoals de macro ze invult, de blokken 'cf' en 'vs' van tab PowerPoint en de hulpcellen
    die de macro leest."""

    def __init__(self, pad):
        wb = load_workbook(pad, data_only=True)
        m, pp = wb["Model"], wb["PowerPoint"]
        self.tekst, self.richting = {}, {}
        r = LY.PP_KPI_ROW1
        while pp[f"C{r}"].value:
            vorm = pp[f"F{r}"].value
            if pp[f"B{r}"].value in (3, 4) and vorm:
                v = pp[f"D{r}"].value
                self.tekst[vorm] = "" if v is None else nl(str(v))
                self.richting[vorm] = (pp[f"E{r}"].value or "").strip().lower()
            r += 1
        h = lambda naam: m[f"{LY.M_HULP}{LY.H[naam]}"].value  # noqa: E731
        self.n = int(h("n") or 0)
        self.vp_n = int(h("vp_n") or 0)
        self.vs_pp = int(h("vs_pp_aan") or 0)
        self.vs_min, self.vs_max, self.vs_stap = float(h("vs_min") or 0), float(h("vs_max") or 0), float(h("vs_stap") or 0)
        self.vs_nu = int(h("vs_nu_pos") or 0)
        self.blok = {}
        for naam, kol, n in (("cf", LY.PP_CF_KOL, self.n), ("vs", LY.PP_VS_KOL, self.vp_n)):
            c0 = pp[LY.PP_BLOK[naam] + "1"].column
            koppen = ["" if pp.cell(7, c0 + j).value is None else str(pp.cell(7, c0 + j).value) for j in range(kol)]
            rijen = []
            for i in range(n):
                rij = []
                for j in range(kol):
                    v = pp.cell(LY.PP_KPI_ROW1 + i, c0 + j).value
                    rij.append(None if v in (None, "") or (isinstance(v, str) and v.startswith("#")) else v)
                rijen.append(rij)
            self.blok[naam] = (koppen, rijen)


# ---- grafiek-XML -------------------------------------------------------------------------------------------------------------
def _sub(parent, tag, **attrs):
    e = etree.SubElement(parent, qn(tag))
    for k, v in attrs.items():
        e.set(k, str(v))
    return e


def _kolom(f):
    """'Sheet1!$I$2:$I$29' -> 'I'"""
    m = re.search(r"!\$?([A-Z]+)\$?\d", f or "")
    return m.group(1) if m else None


def _cache(parent_ref, waarden, tekst=False):
    """vervangt de str-/numCache onder een strRef/numRef door `waarden`"""
    oud = parent_ref.find(qn("c:strCache" if tekst else "c:numCache"))
    fmt = None
    if oud is not None:
        f = oud.find(qn("c:formatCode"))
        fmt = f.text if f is not None else None
        parent_ref.remove(oud)
    cache = _sub(parent_ref, "c:strCache" if tekst else "c:numCache")
    if not tekst:
        _sub(cache, "c:formatCode").text = fmt or "General"
    _sub(cache, "c:ptCount", val=len(waarden))
    for i, w in enumerate(waarden):
        if w is None or w == "":
            continue
        pt = _sub(cache, "c:pt", idx=i)
        _sub(pt, "c:v").text = str(w) if tekst else repr(float(w))


def vul_grafiek(chart, koppen, rijen):
    """Wat VulGrafiekVorm doet, rechtstreeks in de grafiek-XML: per reeks de kolom uit de waarden-formule, kop als naam, waarden als
    cache; categorieën uit kolom A; reeksen met '(uit)' in de kop gaan weg (VerwijderUitgezetteReeksen)."""
    n = len(rijen)
    cs = chart._chartSpace
    for ser in cs.xpath(".//c:plotArea/*/c:ser"):
        f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        kol = _kolom(f.text if f is not None else "")
        if kol is None:
            continue
        j = ord(kol) - ord("A")
        kop = koppen[j] if j < len(koppen) else ""
        if kop.endswith("(uit)"):
            ser.getparent().remove(ser)
            continue
        blad = f.text.split("!")[0]
        f.text = f"{blad}!${kol}$2:${kol}${n + 1}"
        _cache(f.getparent(), [r[j] if isinstance(r[j], (int, float)) else None for r in rijen])
        tx = ser.find(qn("c:tx") + "/" + qn("c:strRef"))
        if tx is not None:
            _cache(tx, [kop], tekst=True)
        cat = ser.find(qn("c:cat"))
        if cat is not None:
            ref = cat.find(qn("c:strRef")) if cat.find(qn("c:strRef")) is not None else cat.find(qn("c:numRef"))
            if ref is not None:
                ref.find(qn("c:f")).text = f"{blad}!$A$2:$A${n + 1}"
                if ref.tag == qn("c:numRef"):
                    cat.remove(ref)
                    ref = _sub(cat, "c:strRef")
                    _sub(ref, "c:f").text = f"{blad}!$A$2:$A${n + 1}"
                _cache(ref, ["" if r[0] is None else (str(int(r[0])) if isinstance(r[0], float) and r[0] == int(r[0]) else str(r[0]))
                             for r in rijen], tekst=True)


# ---- dia 3: drie lijnreeksen erbij ----------------------------------------------------------------------------------------------
def cashflow_reeksen(prs):
    slide = prs.slides[DIA_CF]
    gf = shapes_by_name(slide).get("CF_GRAFIEK")
    if gf is None or not gf.has_chart:
        sys.exit("CF_GRAFIEK niet gevonden op dia 3")
    cs = gf.chart._chartSpace
    lijnen = cs.find(".//" + qn("c:plotArea") + "/" + qn("c:lineChart"))
    alle = cs.xpath(".//c:plotArea/*/c:ser")
    if any((_kolom(s.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f")).text) or "") in ("I", "J", "K") for s in alle):
        sys.exit("CF_GRAFIEK heeft al reeksen op kolom I..K; is dit al een v18-sjabloon?")
    voorbeeld = [s for s in lijnen.findall(qn("c:ser")) if _kolom(s.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f")).text) == "H"]
    if not voorbeeld:
        sys.exit("geen scenariolijn (kolom H) in CF_GRAFIEK gevonden")
    voorbeeld = voorbeeld[0]
    f_val = voorbeeld.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f")).text
    blad = f_val.split("!")[0]
    r2 = int(re.search(r"(\d+)$", f_val).group(1))
    idx = max(int(s.find(qn("c:idx")).get("val")) for s in alle) + 1
    order = max(int(s.find(qn("c:order")).get("val")) for s in alle) + 1
    na = voorbeeld
    for s, kol in zip(range(1, LY.N_VS + 1), "IJK"):
        ser = copy.deepcopy(voorbeeld)
        ser.find(qn("c:idx")).set("val", str(idx))
        ser.find(qn("c:order")).set("val", str(order))
        idx, order = idx + 1, order + 1
        tx = ser.find(qn("c:tx") + "/" + qn("c:strRef"))
        tx.find(qn("c:f")).text = f"{blad}!${kol}$1"
        _cache(tx, [f"Scenario {s}"], tekst=True)
        ln = ser.find(qn("c:spPr") + "/" + qn("a:ln"))
        ln.set("w", str(int(1.75 * 12700)))
        ln.find(qn("a:solidFill") + "/" + qn("a:srgbClr")).set("val", KLEUR[s])
        f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        f.text = f"{blad}!${kol}$2:${kol}${r2}"
        _cache(f.getparent(), [None] * (r2 - 1))
        na.addnext(ser)
        na = ser
    # ingebed werkboek: koppen in I..K, zodat 'Gegevens bewerken' in PowerPoint een kloppende tabel toont
    deel = gf.chart.part.chart_workbook
    wb = load_workbook(io.BytesIO(deel.xlsx_part.blob))
    ws = wb.worksheets[0]
    for s, kol in zip(range(1, LY.N_VS + 1), "IJK"):
        ws[f"{kol}1"] = f"Scenario {s}"
        ws.column_dimensions[kol].width = ws.column_dimensions["H"].width or 20
    buf = io.BytesIO()
    wb.save(buf)
    deel.update_from_xlsx_blob(buf.getvalue())


# ---- dia 4 ------------------------------------------------------------------------------------------------------------------
def kaart_vormen(slide, k, x, w, g):
    """Kaart k (0 = Basis) op x met breedte w: rand, accentstreep, teksten (maten in inch)."""
    y = KAART_T
    t = lambda naam: g.tekst.get(f"VS_K{k}_{naam}", "")  # noqa: E731
    rect(slide, f"VS_K{k}_KAART", x, y, w, KAART_H, "FFFFFF", LIJN, 0.8, rond=0.03)
    rect(slide, f"VS_K{k}_ACCENT", x, y + 0.12, 0.05, KAART_H - 0.24, KLEUR[k])
    xt, wt = x + 0.17, w - 0.27
    tekstvak(slide, f"VS_K{k}_NAAM", xt, y + 0.08, wt, 0.16, t("NAAM") or ("Basis" if k == 0 else f"Scenario {k}"), 8, NAVY, True, krimp=True)
    tekstvak(slide, f"VS_K{k}_SUB", xt, y + 0.25, wt, 0.14, t("SUB"), 7, GRIJS, krimp=True)
    tekstvak(slide, f"VS_K{k}_WAARDE", xt, y + 0.42, 1.0, 0.30, t("WAARDE"), 16, NAVY, True, krimp=True)
    tekstvak(slide, f"VS_K{k}_LABEL", xt + 1.03, y + 0.44, wt - 1.03, 0.12, t("LABEL") or "MAX. FINANCIERING", 6, GRIJS2, True, 40, krimp=True)
    tekstvak(slide, f"VS_K{k}_DAL", xt + 1.03, y + 0.57, wt - 1.03, 0.13, t("DAL"), 7, GRIJS, krimp=True)
    tekstvak(slide, f"VS_K{k}_BE", xt, y + 0.76, wt, 0.14, t("BE"), 7.5, NAVY, True, krimp=True)
    r = g.richting.get(f"VS_K{k}_VERKOCHT", "")
    tekstvak(slide, f"VS_K{k}_VERKOCHT", xt, y + 0.90, wt, 0.14, t("VERKOCHT"), 7.5, RICHTING_KLEUR.get(r, GRIJS2), krimp=True)
    tekstvak(slide, f"VS_K{k}_DELTA", xt, y + 1.05, wt, 0.13, t("DELTA"), 6.5, GRIJS, krimp=True)


def bouw_dia4(prs, g):
    """Nieuwe dia na dia 3 (zelfde indeling als dia 2), met de vormen van de verkoopscenario's en voorbeeldgegevens uit `g`."""
    slide = prs.slides.add_slide(prs.slides[1].slide_layout)
    lst = prs.slides._sldIdLst
    nieuw = list(lst)[-1]
    lst.remove(nieuw)
    lst.insert(DIA_VS, nieuw)
    for ph in list(slide.placeholders):                      # alleen het dianummer van de indeling blijft
        if ph.placeholder_format.type != 13:
            ph._element.getparent().remove(ph._element)
    tekstvak(slide, "VS_KICKER", 0.45, 0.30, 9.1, 0.17, g.tekst.get("VS_KICKER", "VERKOOPSCENARIO'S"), 7, KICKER, True, 118)
    tekstvak(slide, "VS_TITEL", 0.45, 0.50, 9.1, 0.40, g.tekst.get("VS_TITEL", ""), 17, NAVY, True, krimp=True)
    tekstvak(slide, "VS_SUBTITEL", 0.45, 0.93, 9.1, 0.20, g.tekst.get("VS_SUBTITEL", ""), 8.5, GRIJS, krimp=True)
    n = LY.N_VS + 1
    w = (MARGE_R - MARGE_L - (n - 1) * KAART_GAP) / n
    for k in range(n):
        kaart_vormen(slide, k, MARGE_L + k * (w + KAART_GAP), w, g)
    for k in range(n):                                       # legenda: korte lijn + tekst (grijs, de lijn draagt de kleur)
        x = 0.50 + k * 2.25
        lijn(slide, f"VS_LEG{k}_LIJN", x, LEG_T + 0.07, x + 0.26, LEG_T + 0.07, KLEUR[k], LIJN_PT[k])
        tekstvak(slide, f"VS_LEG{k}", x + 0.33, LEG_T, 1.85, 0.14, g.tekst.get(f"VS_LEG{k}", ""), 7.5, GRIJS, krimp=True)
    # grijs vlak (achter de grafiek) en de grafiek
    koppen, rijen = g.blok["vs"]
    nq = max(2, len(rijen))
    x_nu = PLOT_L + (max(1, g.vs_nu) - 1) / (nq - 1) * (PLOT_R - PLOT_L)
    rect(slide, "VS_REALISATIE", PLOT_L, PLOT_T, max(0.01, x_nu - PLOT_L), PLOT_B - PLOT_T, ARCERING)
    tekstvak(slide, "VS_REAL_LBL", PLOT_L + 0.05, PLOT_B - 0.18, 1.2, 0.14, "gerealiseerd", 6.5, GRIJS2)
    cd = CategoryChartData()
    cd.categories = ["" if r[0] is None else str(int(r[0])) for r in rijen]
    for j, naam in enumerate(VS_NAMEN[1:], start=1):
        cd.add_series(koppen[j] or naam, [r[j] if isinstance(r[j], (int, float)) else None for r in rijen])
    fw, fh = FRAME_R - FRAME_L, FRAME_B - FRAME_T
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, I(FRAME_L), I(FRAME_T), I(fw), I(fh), cd)
    gf.name = "VS_GRAFIEK"
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.size = Pt(7)
    ch.font.color.rgb = RGBColor.from_string(GRIJS)
    reeksen = list(ch.plots[0].series)
    for k, sr in enumerate(reeksen[:n]):
        sr.smooth = False
        sr.format.line.color.rgb = RGBColor.from_string(KLEUR[k])
        sr.format.line.width = Pt(LIJN_PT[k])
        sr.marker.style = XL_MARKER_STYLE.NONE
    for k, sr in enumerate(reeksen[n:]):
        sr.smooth = False
        sr.format.line.fill.background()
        sr.marker.style = XL_MARKER_STYLE.CIRCLE
        sr.marker.size = 8
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = RGBColor.from_string(KLEUR[k])
        sr.marker.format.line.color.rgb = RGBColor.from_string("FFFFFF")
        sr.marker.format.line.width = Pt(1.5)
        dls = sr.data_labels
        dls.show_value = True
        dls.number_format = FMT_LABEL
        dls.number_format_is_linked = False
        dls.position = XL_LABEL_POSITION.BELOW
        dls.font.size = Pt(7.5)
        dls.font.bold = True
        dls.font.color.rgb = RGBColor.from_string("333333")
    ca = ch.category_axis
    ca.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = RGBColor.from_string("C9CCD1")
    ca.tick_labels.font.size = Pt(7)
    ca.tick_labels.font.color.rgb = RGBColor.from_string(GRIJS2)
    va = ch.value_axis
    if g.vs_max > g.vs_min:
        va.minimum_scale, va.maximum_scale = g.vs_min, g.vs_max
        if g.vs_stap > 0:
            va.major_unit = g.vs_stap
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor.from_string("E9EAEC")
    va.major_gridlines.format.line.width = Pt(0.75)
    va.format.line.fill.background()
    va.major_tick_mark = XL_TICK_MARK.NONE
    va.tick_labels.font.size = Pt(7)
    va.tick_labels.font.color.rgb = RGBColor.from_string(GRIJS2)
    va.tick_labels.number_format = '"€"0"M";"−€"0"M";"€0"'
    va.tick_labels.number_format_is_linked = False
    cb = etree.Element(qn("c:crossBetween"))
    cb.set("val", "midCat")
    va._element.find(qn("c:crosses")).addnext(cb)
    cs = ch._chartSpace
    pa = cs.chart.plotArea
    oud = pa.find(qn("c:layout"))
    if oud is not None:
        pa.remove(oud)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", (PLOT_L - FRAME_L) / fw),
                     ("y", (PLOT_T - FRAME_T) / fh), ("w", (PLOT_R - PLOT_L) / fw), ("h", (PLOT_B - PLOT_T) / fh)):
        etree.SubElement(ml, qn(f"c:{tag}")).set("val", f"{val:.5f}" if isinstance(val, float) else val)
    pa.insert(0, lay)
    _nofill_spPr(cs, ("c:txPr", "c:externalData", "c:printSettings"))
    _nofill_spPr(pa, ("c:extLst",))
    # eindsaldo: linksboven in het tekenvlak (daar lopen de lijnen zelden: het tijdvak begint in het actuals-jaar, meestal onder nul)
    tekstvak(slide, "VS_EIND_LBL", PLOT_L + 0.08, PLOT_T + 0.04, 4.0, 0.14, g.tekst.get("VS_EIND_LBL", ""), 7, GRIJS, krimp=True)
    tekstvak(slide, "VS_NOOT", 0.45, 5.03, 9.1, 0.14, g.tekst.get("VS_NOOT", ""), 6.5, GRIJS2, krimp=True)
    return slide


# ---- simulatie van de macro (dia 3 en 4) ------------------------------------------------------------------------------------------
def _vul_teksten(slide, g):
    by = shapes_by_name(slide)
    for vorm, tekst in g.tekst.items():
        sh = by.get(vorm)
        if sh is None or not sh.has_text_frame:
            continue
        set_text_keep_format(sh, tekst)
        r = g.richting.get(vorm, "")
        if r:
            kleur = RGBColor.from_string(RICHTING_KLEUR.get(r, GRIJS2))
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = kleur


def simuleer(prs, g):
    """Dia 3: teksten en de cashflowgrafiek met de scenariolijnen. Dia 4: weg zonder scenario's in PowerPoint; anders teksten,
    SchikDiaScenarios (lege kaarten en legenda-regels weg, de rest verdeeld), de grafiek en SchikGrafiekVS."""
    _vul_teksten(prs.slides[DIA_CF], g)
    koppen, rijen = g.blok["cf"]
    vul_grafiek(shapes_by_name(prs.slides[DIA_CF])["CF_GRAFIEK"].chart, koppen, rijen)
    slide = prs.slides[DIA_VS]
    by = shapes_by_name(slide)
    if "VS_GRAFIEK" not in by:
        return
    if g.vs_pp != 1:
        lst = prs.slides._sldIdLst
        lst.remove(list(lst)[DIA_VS])
        return
    _vul_teksten(slide, g)
    # SchikDiaScenarios
    kaarten = {k: by[f"VS_K{k}_KAART"] for k in range(LY.N_VS + 1) if f"VS_K{k}_KAART" in by}
    blijft = {k: k == 0 or bool(g.tekst.get(f"VS_K{k}_WAARDE")) for k in kaarten}
    leg_l = {k: by[f"VS_LEG{k}_LIJN"].left for k in kaarten if f"VS_LEG{k}_LIJN" in by}
    if not all(blijft.values()):
        oud = {k: (v.left, v.width) for k, v in kaarten.items()}
        links = min(v[0] for v in oud.values())
        rechts = max(v[0] + v[1] for v in oud.values())
        gesorteerd = sorted(oud.values())
        tussen = min(b[0] - (a[0] + a[1]) for a, b in zip(gesorteerd, gesorteerd[1:]))
        for k, ok in blijft.items():
            if not ok:
                for sh in list(slide.shapes):
                    if sh.name.startswith(f"VS_K{k}_") or sh.name in (f"VS_LEG{k}", f"VS_LEG{k}_LIJN"):
                        sh._element.getparent().remove(sh._element)
        over = [k for k in sorted(kaarten) if blijft[k]]
        w = (rechts - links - (len(over) - 1) * tussen) / len(over)
        for t, k in enumerate(over):
            dl = links + t * (w + tussen)
            ol, ow = oud[k]
            for sh in slide.shapes:
                if not sh.name.startswith(f"VS_K{k}_"):
                    continue
                mL, mR = sh.left - ol, ol + ow - (sh.left + sh.width)
                if sh.name.endswith("_ACCENT"):
                    sh.left = Emu(int(dl + mL))
                elif w - mL - mR > 1:
                    sh.left, sh.width = Emu(int(dl + mL)), Emu(int(w - mL - mR))
            dx = leg_l[sorted(leg_l)[t]] - leg_l[k]
            for naam in (f"VS_LEG{k}_LIJN", f"VS_LEG{k}"):
                sh = shapes_by_name(slide).get(naam)
                if sh is not None:
                    sh.left = Emu(int(sh.left + dx))
    # grafiek en SchikGrafiekVS
    by = shapes_by_name(slide)
    gf = by["VS_GRAFIEK"]
    koppen, rijen = g.blok["vs"]
    vul_grafiek(gf.chart, koppen, rijen)
    va = gf.chart.value_axis
    if g.vs_max > g.vs_min:
        va.minimum_scale, va.maximum_scale = g.vs_min, g.vs_max
        if g.vs_stap > 0:
            va.major_unit = g.vs_stap
    vlak, lbl = by.get("VS_REALISATIE"), by.get("VS_REAL_LBL")
    n, pos = len(rijen), g.vs_nu
    plot_l = gf.left + gf.width * LY.VS_PLOT_X
    if pos < 1 or n < 2 or pos > n:
        for sh in (vlak, lbl):
            if sh is not None:
                sh._element.getparent().remove(sh._element)
    elif vlak is not None:
        x_nu = plot_l + gf.width * LY.VS_PLOT_W * (pos - 1) / (n - 1)
        vlak.left, vlak.width = Emu(int(plot_l)), Emu(int(max(1, x_nu - plot_l)))


def render_nl(pad, map_uit, resolutie=110):
    """Dia's als PNG via LibreOffice (Nederlandse landinstelling: komma als decimaalteken) en pdftoppm."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftoppm"):
        return []
    map_uit = os.path.abspath(map_uit)
    os.makedirs(map_uit, exist_ok=True)
    profiel = os.path.join(map_uit, "lo_profiel")
    os.makedirs(os.path.join(profiel, "user"), exist_ok=True)
    with open(os.path.join(profiel, "user", "registrymodifications.xcu"), "w", encoding="utf-8") as fh:
        fh.write('<?xml version="1.0" encoding="UTF-8"?>\n<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
                 'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
                 '<item oor:path="/org.openoffice.Setup/L10N"><prop oor:name="ooSetupSystemLocale" oor:op="fuse"><value>nl-NL</value>'
                 '</prop></item>\n</oor:items>\n')
    subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{profiel}", "--convert-to", "pdf", "--outdir", map_uit, pad],
                   capture_output=True, timeout=300)
    pdf = os.path.join(map_uit, os.path.splitext(os.path.basename(pad))[0] + ".pdf")
    if not os.path.exists(pdf):
        return []
    stam = os.path.join(map_uit, "dia")
    subprocess.run(["pdftoppm", "-png", "-r", str(resolutie), pdf, stam], capture_output=True, timeout=300)
    return sorted(os.path.join(map_uit, f) for f in os.listdir(map_uit) if f.startswith("dia") and f.endswith(".png"))


VERWACHT = (["VS_KICKER", "VS_TITEL", "VS_SUBTITEL", "VS_GRAFIEK", "VS_REALISATIE", "VS_REAL_LBL", "VS_EIND_LBL", "VS_NOOT"]
            + [f"VS_K{k}_{x}" for k in range(LY.N_VS + 1) for x in ("KAART", "ACCENT", "NAAM", "SUB", "LABEL", "WAARDE", "DAL", "BE",
                                                                     "VERKOCHT", "DELTA")]
            + [f"VS_LEG{k}{x}" for k in range(LY.N_VS + 1) for x in ("", "_LIJN")])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieën van de uitvoer")
    ap.add_argument("--voorbeeld", default="", help="doorgerekend werkboek (.xlsx) voor de voorbeeldteksten en -grafiek; leeg = zelf bouwen")
    ap.add_argument("--render", default="", help="map voor PNG's van alle dia's (LibreOffice + pdftoppm); leeg = niet renderen")
    ap.add_argument("--simuleer", nargs=2, metavar=("WERKBOEK", "UIT_PPTX"), default=None,
                    help="doe op de nieuwe v18 na wat de macro met dia 2, 3 en 4 doet en bewaar het als UIT_PPTX")
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    if len(prs.slides) != 9:
        sys.exit(f"Verwacht 9 dia's in v17 (titel + 8), gevonden {len(prs.slides)}.")
    if any(sh.name == "VS_GRAFIEK" for s in prs.slides for sh in s.shapes):
        sys.exit("Dit sjabloon heeft al een dia met VS_GRAFIEK (v18?).")
    with tempfile.TemporaryDirectory() as tmp:
        pad = args.voorbeeld or maak_voorbeeld(tmp)
        g = Gegevens(pad)
        cashflow_reeksen(prs)
        bouw_dia4(prs, g)
        prs.save(args.uitvoer)
        for extra in args.extra:
            shutil.copyfile(args.uitvoer, extra)
        print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
        print(f"  dia 3: drie scenariolijnen (kolom I..K); dia 4 'Verkoopscenario's' ({len(g.blok['vs'][1])} voorbeeldrijen)")
        fouten, charts, emb = controleer(args.uitvoer, {DIA_VS + 1: VERWACHT})
        print("  chart-parts: " + ", ".join(charts))
        if fouten:
            sys.exit("FOUTEN:\n  " + "\n  ".join(fouten))
        print("  geen fouten")
        if args.render:
            pngs = render_nl(args.uitvoer, args.render)
            print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd (LibreOffice/pdftoppm ontbreekt)"))
        if args.simuleer:
            prs2 = Presentation(args.uitvoer)
            simuleer_dia2(prs2, GegevensDia2(args.simuleer[0]))
            simuleer(prs2, Gegevens(args.simuleer[0]))
            prs2.save(args.simuleer[1])
            print(f"  simulatie: {args.simuleer[1]}")
            if args.render:
                pngs = render_nl(args.simuleer[1], os.path.join(args.render, "simulatie"))
                print("  PNG's simulatie: " + (", ".join(pngs) if pngs else "niet gerenderd"))


if __name__ == "__main__":
    main()
