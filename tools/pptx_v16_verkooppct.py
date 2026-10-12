"""Sjabloon v15 -> v16.

    python3 tools/pptx_v16_verkooppct.py INVOER_v15.pptx UITVOER_v16.pptx [UITVOER2.pptx ...] --voorbeeld doorgerekend.xlsx [--render MAP]

  * dia 6 (VP1_GRAFIEK verkocht, VP2_GRAFIEK getransporteerd): nieuwe grafiek-XML uit chartxml.verkoop_paneel: het cumulatieve
    percentage van alle woningen als donkerblauwe lijn op een rechteras 0-100 % met het percentage in vet boven de lijn;
    de aantallen blijven in de blokjes. Ingebed werkboek met de 30 kolommen van het blok op tab PowerPoint (layout.PP_VP_KOL:
    kwartaal, verkocht per type, getransporteerd per type, verkocht, getransporteerd, mijlpaal uitverkocht, mijlpaal laatste
    transport, realisatie, verkocht cumulatief % en label, getransporteerd cumulatief % en label); alleen de periodes vanaf
    1 januari van het actuals-jaar (de macro vult Model!BY132 rijen).
  * VP_SUBTITEL krijgt de voorbeeldtekst met de percentages (de macro overschrijft hem uit tab PowerPoint).
De vormnamen blijven gelijk; de macro (KOL_VP = 30, CEL_VP_N) vult het blok."""
import argparse
import os
import shutil
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import chartxml as cx, layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import EMU_CM, controleer, render, set_text_keep_format, shapes_by_name, ws_kopie, zet_chart_xml  # noqa: E402

DIA_VP = 5          # 0-gebaseerd: dia 6 (verkoop en transport per kwartaal)
VP_NAMEN = (["kwartaal"] + [f"v{k}" for k in range(1, LY.N_TYPES + 1)] + [f"t{k}" for k in range(1, LY.N_TYPES + 1)]
            + ["verkocht", "getransporteerd", "punt_uitverkocht", "punt_alles_transport", "realisatie", "pct", "pct_lbl", "tpct", "tpct_lbl"])
VP_KOL = {naam: get_column_letter(j + 1) for j, naam in enumerate(VP_NAMEN)}
assert len(VP_NAMEN) == LY.PP_VP_KOL and VP_KOL["realisatie"] == "Z" and VP_KOL["pct"] == "AA" and VP_KOL["tpct_lbl"] == "AD"


def vp_refs(n):
    """Reeksverwijzingen in het ingebedde werkboek (rijen 2..n+1); de macro zet ze later op het aantal periodes van het blok."""
    def rng(naam):
        return f"Sheet1!${VP_KOL[naam]}$2:${VP_KOL[naam]}${n + 1}"
    r = {naam: rng(naam) for naam in ("kwartaal", "verkocht", "getransporteerd", "punt_uitverkocht", "punt_alles_transport",
                                      "realisatie", "pct", "pct_lbl", "tpct", "tpct_lbl")}
    lbl = {"uitverkocht": f"Sheet1!${VP_KOL['punt_uitverkocht']}$1", "alles_transport": f"Sheet1!${VP_KOL['punt_alles_transport']}$1"}
    types = [(f"Sheet1!${VP_KOL[f'v{k}']}$1", rng(f"v{k}"), rng(f"t{k}")) for k in range(1, LY.N_TYPES + 1)]
    return r, lbl, types


def werkboek_vp(pad_voorbeeld):
    """Sheet1 met de koppen en de getoonde rijen van het blok van dia 6 (tab PowerPoint, vanaf PP_BLOK['vp']): precies wat de
    macro in het gegevensblad van de grafiek zet. Foutwaarden (#N/A buiten het tijdvak, labels zonder stapel) worden lege cellen."""
    wb = load_workbook(pad_voorbeeld, data_only=True)
    m, pp = wb["Model"], wb["PowerPoint"]
    n = int(m[f"{LY.M_HULP}{LY.H['vp_n']}"].value or 0)
    assert n > 0, "geen getoonde periodes (Model vp_n)"
    c0 = pp[LY.PP_BLOK["vp"] + "1"].column
    koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_VP_KOL)]
    assert koppen[0] == "Kwartaal" and koppen[-1] == "Label transport %", koppen
    ws = Workbook().active
    ws.title = "Sheet1"
    ws.append(["" if k is None else str(k) for k in koppen])
    for i in range(n):
        rij = []
        for j in range(LY.PP_VP_KOL):
            v = pp.cell(LY.PP_KPI_ROW1 + i, c0 + j).value
            rij.append(None if (isinstance(v, str) and v.startswith("#")) else v)
        ws.append(rij)
    ws.column_dimensions["A"].width = 12
    subtitel = None
    for rij in pp.iter_rows(min_row=1, max_row=60, min_col=1, max_col=7):
        if rij[5].value == "VP_SUBTITEL":
            subtitel = rij[3].value if isinstance(rij[3].value, str) else subtitel
    return ws, n, subtitel


def vernieuw_dia6(prs, pad_voorbeeld):
    slide = prs.slides[DIA_VP]
    by_name = shapes_by_name(slide)
    ws, n, subtitel = werkboek_vp(pad_voorbeeld)
    r, lbl, types = vp_refs(n)
    for naam, soort in (("VP1_GRAFIEK", "verkocht"), ("VP2_GRAFIEK", "transport")):
        gf = by_name[naam]
        xml = cx.verkoop_paneel(r, lbl, types, soort, hoogte_cm=gf.height / EMU_CM)
        zet_chart_xml(gf, xml, ws_kopie(ws), n)
    if subtitel and "VP_SUBTITEL" in by_name:
        set_text_keep_format(by_name["VP_SUBTITEL"], subtitel)
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieën van de uitvoer")
    ap.add_argument("--voorbeeld", required=True, help="doorgerekend werkboek (.xlsx) met de voorbeeldgegevens (tab Model en PowerPoint)")
    ap.add_argument("--render", default="", help="map voor PNG's van alle dia's (LibreOffice + pdftoppm); leeg = niet renderen")
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    if len(prs.slides) != 9:
        sys.exit(f"Verwacht 9 dia's in v15 (titel + 8), gevonden {len(prs.slides)}.")
    n = vernieuw_dia6(prs, args.voorbeeld)
    prs.save(args.uitvoer)
    for extra in args.extra:
        shutil.copyfile(args.uitvoer, extra)
    print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
    print(f"  dia 6: VP1_GRAFIEK en VP2_GRAFIEK vernieuwd ({n} voorbeeldrijen, {LY.PP_VP_KOL} kolommen, cumulatief % als lijn met labels)")
    fouten, charts, emb = controleer(args.uitvoer, {6: ["VP_KICKER", "VP_TITEL", "VP_SUBTITEL", "VP_KAART", "VP1_GRAFIEK", "VP2_GRAFIEK"]})
    print("  chart-parts: " + ", ".join(charts))
    if fouten:
        sys.exit("FOUTEN:\n  " + "\n  ".join(fouten))
    print("  geen fouten")
    if args.render:
        pngs = render(args.uitvoer, args.render)
        print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd (LibreOffice/pdftoppm ontbreekt)"))


if __name__ == "__main__":
    main()
