"""Sjabloon v19 -> v20: dia 7 (verkoop en transport) met veel DAEB, en een simulatie van dia 6 en 7.

    python3 tools/pptx_v20_daeb.py INVOER_v19.pptx UITVOER_v20.pptx [--render MAP] [--simuleer doorgerekend.xlsx UIT.pptx]

* VP1_GRAFIEK en VP2_GRAFIEK: een reeks erbij op de %-as (kolom AE resp. AF van het gegevensblad, layout.PP_VP_KOL = 32): de
  markering DAEB, een gouden ruit met de reeksnaam als label links ('DAEB · 187 woningen in één keer'). Alleen met veel DAEB (meer
  dan de helft van de woningen) heeft die kolom een waarde en een naam; anders is de kop '(uit)' en haalt de macro de reeks weg.
  Met veel DAEB tonen de balken en het % alleen de koopwoningen (het Model laat de DAEB-balken weg).
* Rechts op dia 7 drie kaartjes VP_STAT1..3 (kader, LABEL, WAARDE, SUB) in de stijl van dia 8; de kaart VP_KAART en de twee
  grafieken worden smaller (zoals VO_KAART / VO_GRAFIEK).
* Dia 7: een legenda VP_LEGENDA boven de grafieken (■ per type, ◆ DAEB; de macro kleurt de tekens met de kleuren uit kolom E),
  een accentstreep per kaartje (VP_STAT1..3_ACCENT, kleur uit kolom D), het % rechts van het punt en een gestippelde gouden lijn
  van de DAEB-ruit naar de as; de grafiektitels zet de macro uit tab PowerPoint ('Verkocht per kwartaal · koop (36)').
* Projectnaam: op de titeldia COVER_PROJECT en COVER_KWARTAAL, in de dia-indeling de voettekst VOET_TEKST (de macro vult ze).
* De tabel woningtypes (VT_TABELKAART, VT_TABEL) op een eigen dia direct na de bouwtermijnen; de bouwtermijnengrafiek vult dia 6.
* BT_GRAFIEK: elk rijlabel tonen (tickLblSkip = 1) en de reeksen dfee (DAEB-fee in goud), fgap (ruimte) en fl1..fl10 (vak met het
  bedrag als label rechts naast de fee) in de tweede stapel vóór s8; de kolommen achteraan het blok (layout.BT_REEKSEN). De
  bouwtermijnen in twee groepen (koop / DAEB-fees, kopregels, dikke fee-balken) komen uit tab PowerPoint.

--simuleer: doet op de nieuwe v20 na wat de macro met dia 6 en 7 doet (teksten, BT_GRAFIEK met de tijd-as, VP1/VP2 met '(uit)'),
met de waarden uit een doorgerekend werkboek."""
import argparse
import copy
import io
import os
import re
import sys

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import set_text_keep_format, shapes_by_name  # noqa: E402
from pptx_v18_scenarios import _cache, render_nl  # noqa: E402
from pptx_v17_kerncijfers import nl  # noqa: E402

GOUD, GOUD_LABEL, GOUD_VLAK = "EDA100", "7A5300", "FFF6DD"
KOL_DAEB = {"VP1_GRAFIEK": get_column_letter(LY.PP_VP_KOL - 1), "VP2_GRAFIEK": get_column_letter(LY.PP_VP_KOL)}   # AE, AF
KAART_W = 6.6                   # breedte VP_KAART (inch), zoals VO_KAART
STAT_L, STAT_W = 7.25, 2.3      # kaartjes rechts, zoals VO_STAT


def _kolom(f):
    m = re.search(r"!\$?([A-Z]+)\$?\d", f or "")
    return m.group(1) if m else None


def _dia(prs, vorm):
    for i, s in enumerate(prs.slides):
        if any(sh.name == vorm for sh in s.shapes):
            return i, s
    sys.exit(f"{vorm} niet gevonden")


def markering(chart, kol):
    """kopie van de mijlpaalreeks (punt met label) als ruit in goud op de %-as, waarden in kolom `kol`"""
    cs = chart._chartSpace
    if any(_kolom(f.text) == kol for f in cs.iter(qn("c:f"))):
        return False
    mijlpaal = next(s for s in cs.xpath(".//c:plotArea/c:lineChart/c:ser") if s.find(qn("c:marker") + "/" + qn("c:symbol")) is not None
                    and s.find(qn("c:marker") + "/" + qn("c:symbol")).get("val") == "circle")
    pct_groep = next(g for g in cs.xpath(".//c:plotArea/c:lineChart") if g.find(qn("c:axId")).get("val") != mijlpaal.getparent().find(qn("c:axId")).get("val"))
    alle = cs.xpath(".//c:plotArea/*/c:ser")
    ser = copy.deepcopy(mijlpaal)
    ser.find(qn("c:idx")).set("val", str(max(int(s.find(qn("c:idx")).get("val")) for s in alle) + 1))
    ser.find(qn("c:order")).set("val", str(max(int(s.find(qn("c:order")).get("val")) for s in alle) + 1))
    f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
    blad = f.text.split("!")[0]
    r2 = int(re.search(r"(\d+)$", f.text).group(1))
    f.text = f"{blad}!${kol}$2:${kol}${r2}"
    _cache(f.getparent(), [None] * (r2 - 1))
    tx = ser.find(qn("c:tx") + "/" + qn("c:strRef"))
    tx.find(qn("c:f")).text = f"{blad}!${kol}$1"
    _cache(tx, ["DAEB"], tekst=True)
    mk = ser.find(qn("c:marker"))
    mk.find(qn("c:symbol")).set("val", "diamond")
    size = mk.find(qn("c:size"))
    if size is not None:
        size.set("val", "9")
    for c in mk.iter(qn("a:srgbClr")):
        c.set("val", GOUD if c.get("val") not in ("FFFFFF",) else "FFFFFF")
    dl = ser.find(qn("c:dLbls"))
    if dl is not None:
        sp = dl.find(qn("c:spPr"))
        if sp is not None:
            vul = sp.find(qn("a:solidFill") + "/" + qn("a:srgbClr"))
            if vul is not None:
                vul.set("val", GOUD_VLAK)
            ln = sp.find(qn("a:ln") + "//" + qn("a:srgbClr"))
            if ln is not None:
                ln.set("val", GOUD)
        for c in dl.iter(qn("a:srgbClr")):
            if c.getparent().getparent().tag in (qn("a:defRPr"), qn("a:rPr")):
                c.set("val", GOUD_LABEL)
        pos = dl.find(qn("c:dLblPos"))
        if pos is not None:
            pos.set("val", "l")
    # na de laatste reeks van de %-groep (vóór dLbls/marker/axId van de groep)
    laatste = pct_groep.findall(qn("c:ser"))[-1]
    laatste.addnext(ser)
    # ingebed werkboek: kop in de nieuwe kolom
    deel = chart.part.chart_workbook
    wb = load_workbook(io.BytesIO(deel.xlsx_part.blob))
    wb.worksheets[0][f"{kol}1"] = "DAEB"
    buf = io.BytesIO()
    wb.save(buf)
    deel.update_from_xlsx_blob(buf.getvalue())
    return True


def _nieuw_id(slide, el):
    """unieke vorm-id op de dia (een kopie houdt anders de id van het origineel)"""
    ids = [int(c.get("id")) for c in slide.shapes._spTree.iter(qn("p:cNvPr"))]
    for c in el.iter(qn("p:cNvPr")):
        c.set("id", str(max(ids) + 1))
        ids.append(max(ids) + 1)


def kaartjes(prs):
    _, d7 = _dia(prs, "VP1_GRAFIEK")
    _, d8 = _dia(prs, "VO_STAT1")
    by7, by8 = shapes_by_name(d7), shapes_by_name(d8)
    if "VP_STAT1" in by7:
        return False
    by7["VP_KAART"].width = Inches(KAART_W)
    for naam in ("VP1_GRAFIEK", "VP2_GRAFIEK"):
        by7[naam].width = Inches(KAART_W - 0.2)
    boom = d7.shapes._spTree
    for n in range(1, 4):
        for deel in ("", "_LABEL", "_WAARDE"):
            el = copy.deepcopy(by8[f"VO_STAT{n}{deel}"]._element)
            _nieuw_id(d7, el)
            boom.append(el)
        by = shapes_by_name(d7)
        for deel in ("", "_LABEL", "_WAARDE"):
            sh = by[f"VO_STAT{n}{deel}"]
            sh.name = f"VP_STAT{n}{deel}"
        by = shapes_by_name(d7)
        sub = copy.deepcopy(by[f"VP_STAT{n}_LABEL"]._element)
        _nieuw_id(d7, sub)
        boom.append(sub)
        sh = [s for s in d7.shapes if s.name == f"VP_STAT{n}_LABEL"][-1]
        sh.name = f"VP_STAT{n}_SUB"
        w = shapes_by_name(d7)[f"VP_STAT{n}_WAARDE"]
        sh.top = Emu(int(w.top + w.height + Inches(0.06)))
        set_text_keep_format(sh, "toelichting")
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = False
    return True


def _kol_bt(naam):
    return get_column_letter(2 + LY.BT_REEKSEN.index(naam))      # kolom A = label


def _hernummer(cs):
    for i, ser in enumerate(cs.xpath(".//c:plotArea/*/c:ser")):
        ser.find(qn("c:idx")).set("val", str(i))
        ser.find(qn("c:order")).set("val", str(i))


def _zet_kolom(ser, kol, naam=None):
    f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
    blad = f.text.split("!")[0]
    r2 = int(re.search(r"(\d+)$", f.text).group(1))
    f.text = f"{blad}!${kol}$2:${kol}${r2}"
    _cache(f.getparent(), [None] * (r2 - 1))
    tx = ser.find(qn("c:tx"))
    if tx is not None:
        ser.remove(tx)
    tx = ser.makeelement(qn("c:tx"), {})
    ref = etree.SubElement(tx, qn("c:strRef"))
    etree.SubElement(ref, qn("c:f")).text = f"{blad}!${kol}$1"
    _cache(ref, [naam or kol], tekst=True)
    ser.find(qn("c:order")).addnext(tx)


def bt_reeksen(chart):
    """dfee, fgap en fl1..fl10 in de tweede stapel vóór s8"""
    cs = chart._chartSpace
    if any(_kolom(f.text) == _kol_bt("dfee") for f in cs.iter(qn("c:f"))):
        return False
    per_kol = {_kolom(s.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f")).text): s for s in cs.xpath(".//c:plotArea/*/c:ser")}
    s8, cur1, kw1 = per_kol[_kol_bt("s8")], per_kol[_kol_bt("cur1")], per_kol[_kol_bt("kw1")]
    nieuw = []
    d = copy.deepcopy(cur1)
    _zet_kolom(d, _kol_bt("dfee"), "DAEB-fee")
    for c in d.find(qn("c:spPr")).iter(qn("a:srgbClr")):
        c.set("val", GOUD)
    nieuw.append(d)
    g = copy.deepcopy(s8)
    _zet_kolom(g, _kol_bt("fgap"), "·")
    nieuw.append(g)
    for j in range(1, LY.BT_NFL + 1):
        fl = copy.deepcopy(kw1)
        _zet_kolom(fl, _kol_bt(f"fl{j}"), "€0,0M")
        dl = fl.find(qn("c:dLbls"))
        if dl is not None:
            pos = dl.find(qn("c:dLblPos"))
            if pos is not None:
                pos.set("val", "inBase")
            for e in dl.iter(qn("a:defRPr")):
                e.set("sz", "700")
                e.set("b", "1")
            for c in dl.iter(qn("a:srgbClr")):
                c.set("val", GOUD_LABEL)
        sp = fl.find(qn("c:spPr"))
        if sp is not None:
            for c in list(sp):
                if c.tag in (qn("a:solidFill"), qn("a:noFill")):
                    sp.remove(c)
            sp.insert(0, sp.makeelement(qn("a:noFill"), {}))
        nieuw.append(fl)
    for el in nieuw:
        s8.addprevious(el)
    _hernummer(cs)
    deel = chart.part.chart_workbook
    wb = load_workbook(io.BytesIO(deel.xlsx_part.blob))
    ws = wb.worksheets[0]
    ws[f"{_kol_bt('dfee')}1"], ws[f"{_kol_bt('fgap')}1"] = "DAEB-fee", "·"
    for j in range(1, LY.BT_NFL + 1):
        ws[f"{_kol_bt(f'fl{j}')}1"] = "(uit)"
    buf = io.BytesIO()
    wb.save(buf)
    deel.update_from_xlsx_blob(buf.getvalue())
    return True


def vp_opmaak(chart, kol_daeb):
    """% rechts van het punt; gestippelde gouden lijn van de DAEB-ruit naar de as (foutbalk omlaag, 100%)"""
    for ser in chart._chartSpace.xpath(".//c:plotArea/c:lineChart/c:ser"):
        f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        kol = _kolom(f.text)
        dl = ser.find(qn("c:dLbls"))
        if kol == kol_daeb:
            if ser.find(qn("c:errBars")) is None:
                err = etree.fromstring(
                    '<c:errBars xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
                    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><c:errDir val="y"/><c:errBarType val="minus"/>'
                    '<c:errValType val="percentage"/><c:noEndCap val="1"/><c:val val="100"/><c:spPr><a:ln w="15875"><a:solidFill>'
                    f'<a:srgbClr val="{GOUD}"/></a:solidFill><a:prstDash val="dash"/></a:ln></c:spPr></c:errBars>')
                ser.find(qn("c:cat")).addprevious(err)
        elif dl is not None and ser.find(qn("c:marker") + "/" + qn("c:symbol")) is not None \
                and ser.find(qn("c:marker") + "/" + qn("c:symbol")).get("val") == "none":
            pos = dl.find(qn("c:dLblPos"))
            if pos is not None:
                pos.set("val", "r")


def dia7_extra(prs):
    """legenda boven de grafieken en een accentstreep per kaartje"""
    _, d7 = _dia(prs, "VP1_GRAFIEK")
    by = shapes_by_name(d7)
    if "VP_LEGENDA" in by:
        return False
    kaart = by["VP_KAART"]
    tb = d7.shapes.add_textbox(kaart.left + Inches(0.15), kaart.top + Inches(0.06), kaart.width - Inches(0.3), Inches(0.2))
    tb.name = "VP_LEGENDA"
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    run = tf.paragraphs[0].add_run()
    run.text = "■ type 1   ■ type 2   ◆ DAEB"
    run.font.size, run.font.name = Pt(8), "Arial"
    run.font.color.rgb = RGBColor.from_string("5B6169")
    boven = kaart.top + Inches(0.3)
    h = (kaart.top + kaart.height - Inches(0.06) - boven) / 2
    for i, naam in enumerate(("VP1_GRAFIEK", "VP2_GRAFIEK")):
        by[naam].top, by[naam].height = Emu(int(boven + i * h)), Emu(int(h))
    for n in range(1, 4):
        box = by[f"VP_STAT{n}"]
        a = d7.shapes.add_shape(1, box.left, box.top + Inches(0.1), Inches(0.05), box.height - Inches(0.2))
        a.name = f"VP_STAT{n}_ACCENT"
        a.fill.solid()
        a.fill.fore_color.rgb = RGBColor.from_string(GOUD if n == 2 else "002060")
        a.line.fill.background()
    return True


BT_PLOT = (0.30, 0.02, 0.69, 0.96)     # tekenvlak van BT_GRAFIEK (fracties): vaste ruimte links voor de rijlabels


def bt_tekenvlak(chart):
    """vast tekenvlak (manualLayout, inner): de labels links krijgen altijd dezelfde breedte, ook met lange fee-namen"""
    lay = chart._chartSpace.find(".//" + qn("c:plotArea") + "/" + qn("c:layout"))
    for c in list(lay):
        lay.remove(c)
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    etree.SubElement(ml, qn("c:layoutTarget")).set("val", "inner")
    for tag in ("c:xMode", "c:yMode"):
        etree.SubElement(ml, qn(tag)).set("val", "edge")
    for tag, v in zip(("c:x", "c:y", "c:w", "c:h"), BT_PLOT):
        etree.SubElement(ml, qn(tag)).set("val", str(v))


def projectnaam(prs):
    """titeldia: projectnaam en kwartaal als vormen COVER_PROJECT / COVER_KWARTAAL; voettekst ('PROJECT  |  Q3 2026') in de
    dia-indeling als VOET_TEKST (de macro vult ze vanaf tab PowerPoint, met de projectnaam van het Dashboard)"""
    n = 0
    tekst = lambda sh: sh.text_frame.text.strip() if sh.has_text_frame else ""  # noqa: E731
    cover = prs.slides[0]
    vormen = [sh for sh in cover.shapes if sh.has_text_frame and tekst(sh)]
    if not any(sh.name == "COVER_PROJECT" for sh in cover.shapes) and vormen:
        # de projectnaam staat boven de titel 'Financiële update'; het kwartaal is de tekst 'Qn jjjj'
        kwartaal = next((sh for sh in vormen if re.fullmatch(r"Q[1-4] \d{4}", tekst(sh))), None)
        titel = next((sh for sh in vormen if "update" in tekst(sh).lower()), None)
        boven = [sh for sh in vormen if titel is not None and sh.top < titel.top and sh is not kwartaal]
        if boven:
            max(boven, key=lambda sh: sh.top).name = "COVER_PROJECT"
            n += 1
        if kwartaal is not None:
            kwartaal.name = "COVER_KWARTAAL"
            n += 1
    for lay in prs.slide_layouts:
        for sh in lay.shapes:
            if sh.has_text_frame and "|" in sh.text_frame.text and sh.name != "VOET_TEKST":
                sh.name = "VOET_TEKST"
                n += 1
    return n


def tabel_eigen_dia(prs):
    """de tabel woningtypes (VT_TABELKAART + VT_TABEL) naar een eigen dia direct na de bouwtermijnen; de grafiek krijgt de ruimte"""
    i6, d6 = _dia(prs, "BT_GRAFIEK")
    by = shapes_by_name(d6)
    if "VT_TABEL" not in by:
        return False
    nieuw = prs.slides.add_slide(d6.slide_layout)
    for sh in list(nieuw.shapes):
        sh._element.getparent().remove(sh._element)
    boom = nieuw.shapes._spTree
    for naam in ("BT_KICKER", "BT_TITEL"):
        if naam in by:
            el = copy.deepcopy(by[naam]._element)
            boom.append(el)
    for sh in d6.shapes:
        if sh.is_placeholder:
            boom.append(copy.deepcopy(sh._element))
    for naam in ("VT_TABELKAART", "VT_TABEL"):
        el = by[naam]._element
        el.getparent().remove(el)
        boom.append(el)
    bn = shapes_by_name(nieuw)
    bn["BT_KICKER"].name, bn["BT_TITEL"].name = "VT_KICKER", "VT_TITEL"
    bn = shapes_by_name(nieuw)
    set_text_keep_format(bn["VT_KICKER"], "WONINGTYPES")
    set_text_keep_format(bn["VT_TITEL"], "Woningtypes")
    dy = bn["VT_TABELKAART"].top - Inches(1.25)
    for naam in ("VT_TABELKAART", "VT_TABEL"):
        bn[naam].top = Emu(int(bn[naam].top - dy))
    # direct na de bouwtermijnen in de volgorde van de dia's
    lst = prs.slides._sldIdLst
    el = list(lst)[-1]
    lst.remove(el)
    lst.insert(i6 + 1, el)
    # de bouwtermijnen vullen nu de dia
    by["VT_KAART"].top, by["VT_KAART"].height = Inches(1.25), Inches(3.95)
    by["BT_GRAFIEK"].top, by["BT_GRAFIEK"].height = Inches(1.31), Inches(3.83)
    return True


def labels_alle(chart):
    """BT_GRAFIEK: elk rijlabel tonen (met kopregels en DAEB-fees zijn er meer rijen; anders slaat de as labels over)"""
    for ax in chart._chartSpace.xpath(".//c:plotArea/c:catAx"):
        if ax.find(qn("c:tickLblSkip")) is not None:
            continue
        skip = ax.makeelement(qn("c:tickLblSkip"), {"val": "1"})
        na = ax.find(qn("c:noMultiLvlLbl"))
        if na is not None:
            na.addprevious(skip)
        else:
            ax.append(skip)


# ---- simulatie (dia 6 en 7) ----------------------------------------------------------------------------------------------------
class Gegevens:
    def __init__(self, pad):
        wb = load_workbook(pad, data_only=True)
        m, pp = wb["Model"], wb["PowerPoint"]
        self.tekst, self.kolom_e = {}, {}
        r = LY.PP_KPI_ROW1
        while pp[f"C{r}"].value:
            vorm = pp[f"F{r}"].value
            if pp[f"B{r}"].value in (1, 6, 7, 8) and vorm:
                v = pp[f"D{r}"].value
                self.tekst[vorm] = "" if v is None else nl(str(v))
                self.kolom_e[vorm] = "" if pp[f"E{r}"].value is None else str(pp[f"E{r}"].value)
            r += 1
        h = lambda naam: m[f"{LY.M_HULP}{LY.H[naam]}"].value  # noqa: E731
        self.t_first, self.t_end = float(h("t_first") or 0), float(h("t_end") or 0)
        self.blok = {}
        for naam, kol, n in (("bt", LY.PP_BT_KOL, int(h("bt_n") or 0)), ("vp", LY.PP_VP_KOL, int(h("vp_n") or 0))):
            c0 = column_index_from_string(LY.PP_BLOK[naam])
            koppen = ["" if pp.cell(7, c0 + j).value is None else str(pp.cell(7, c0 + j).value) for j in range(kol)]
            rijen = [[None if v in (None, "") or (isinstance(v, str) and v.startswith("#")) else v
                      for v in (pp.cell(LY.PP_KPI_ROW1 + i, c0 + j).value for j in range(kol))] for i in range(n)]
            self.blok[naam] = (koppen, rijen)


def vul(chart, koppen, rijen):
    """zoals VulGrafiekVorm: per reeks de kolom uit de formule, kop als naam ('(uit)' = weg), waarden als cache, categorieën uit A"""
    n = len(rijen)
    for ser in chart._chartSpace.xpath(".//c:plotArea/*/c:ser"):
        f = ser.find(qn("c:val") + "/" + qn("c:numRef") + "/" + qn("c:f"))
        kol = _kolom(f.text if f is not None else "")
        if kol is None:
            continue
        j = column_index_from_string(kol) - 1
        kop = koppen[j] if j < len(koppen) else ""
        if kop.endswith("(uit)"):
            ser.getparent().remove(ser)
            continue
        blad = f.text.split("!")[0]
        f.text = f"{blad}!${kol}$2:${kol}${n + 1}"
        _cache(f.getparent(), [r[j] if isinstance(r[j], (int, float)) else None for r in rijen])
        tx = ser.find(qn("c:tx") + "/" + qn("c:strRef"))
        if tx is not None:
            tkol = _kolom(tx.find(qn("c:f")).text)
            tj = column_index_from_string(tkol) - 1 if tkol else j
            _cache(tx, [koppen[tj] if tj < len(koppen) else kop], tekst=True)
        cat = ser.find(qn("c:cat"))
        if cat is not None:
            ref = cat.find(qn("c:strRef")) if cat.find(qn("c:strRef")) is not None else cat.find(qn("c:numRef"))
            if ref is not None:
                if ref.tag == qn("c:numRef"):
                    cat.remove(ref)
                    ref = etree.SubElement(cat, qn("c:strRef"))
                    etree.SubElement(ref, qn("c:f"))
                ref.find(qn("c:f")).text = f"{blad}!$A$2:$A${n + 1}"
                _cache(ref, ["" if r[0] is None else (str(int(r[0])) if isinstance(r[0], float) and r[0] == int(r[0]) else str(r[0]))
                             for r in rijen], tekst=True)


def kleur_legenda(sh, kleuren):
    """zoals KleurLegenda in de macro: elk ■ en ◆ in de volgende kleur (één run per teken)"""
    lijst = [k for k in kleuren.split(",") if k]
    p_ = sh.text_frame.paragraphs[0]
    tekst = "".join(r.text for r in p_.runs)
    if not p_.runs:
        return
    font = p_.runs[0].font
    size, naam, kleur = font.size, font.name, font.color.rgb if font.color and font.color.type else None
    for r in list(p_.runs)[1:]:
        r._r.getparent().remove(r._r)
    p_.runs[0].text = ""
    stukken, k, buf = [], 0, ""
    for c in tekst:
        if c in "■◆":
            if buf:
                stukken.append((buf, None))
                buf = ""
            stukken.append((c, lijst[k] if k < len(lijst) else None))
            k += 1
        else:
            buf += c
    if buf:
        stukken.append((buf, None))
    for i, (t, kl) in enumerate(stukken):
        r = p_.runs[0] if i == 0 else p_.add_run()
        r.text = t
        r.font.size, r.font.name = size, naam
        if kl:
            r.font.color.rgb = RGBColor.from_string(kl)
        elif kleur is not None:
            r.font.color.rgb = kleur


def simuleer(prs, g):
    # voettekst in de dia-indelingen (de macro: ZoekVorm zoekt ook daar)
    if "VOET_TEKST" in g.tekst:
        for lay in prs.slide_layouts:
            for sh in lay.shapes:
                if sh.name == "VOET_TEKST":
                    set_text_keep_format(sh, g.tekst["VOET_TEKST"])
    for i_dia in range(len(prs.slides)):
        by = shapes_by_name(prs.slides[i_dia])
        for vorm, tekst in g.tekst.items():
            sh = by.get(vorm)
            if sh is None:
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:        # zoals de macro: tekst = grafiektitel
                ch = sh.chart
                ch.has_title = True
                tf = ch.chart_title.text_frame
                tf.text = tekst
                for p_ in tf.paragraphs:
                    for r_ in p_.runs:
                        r_.font.size, r_.font.bold = Pt(8), True
                        r_.font.color.rgb = RGBColor.from_string("5B6169")
            elif vorm.endswith("_ACCENT"):
                sh.fill.solid()
                sh.fill.fore_color.rgb = RGBColor.from_string(tekst)
            elif sh.has_text_frame:
                set_text_keep_format(sh, tekst)
                e = g.kolom_e.get(vorm, "")
                if vorm.endswith("_LEGENDA") or vorm == "BT_SUBTITEL":
                    kleur_legenda(sh, e)
                elif e == "daeb":
                    for p_ in sh.text_frame.paragraphs:
                        for r_ in p_.runs:
                            r_.font.color.rgb = RGBColor.from_string(GOUD_LABEL)
        if "BT_GRAFIEK" in by:
            ch = by["BT_GRAFIEK"].chart
            vul(ch, *g.blok["bt"])
            for ax in ch._chartSpace.xpath(".//c:plotArea/c:valAx"):
                sc = ax.find(qn("c:scaling"))
                for tag, v in (("c:max", g.t_end), ("c:min", g.t_first)):
                    e = sc.find(qn(tag))
                    if e is None:
                        e = etree.SubElement(sc, qn(tag))
                    e.set("val", repr(v))
            for lg in ch._chartSpace.findall(".//" + qn("c:legend")):
                lg.getparent().remove(lg)
            # zoals de macro: rijlabels passend bij het aantal rijen (inch -> pt; tekenvlak ongeveer 90% van de grafiek)
            n = max(1, len(g.blok["bt"][1]))
            letter = max(5.0, min(7.5, 2.4 * by["BT_GRAFIEK"].height / 12700 * 0.9 / n))
            # LibreOffice verbergt rijlabels die dicht op elkaar staan (ook naast lege labels); PowerPoint toont ze allemaal
            # (tickLblSkip = 1). Voor een getrouwe preview: de labels van de as als tekstvakken op de rijen
            gf, rijen = by["BT_GRAFIEK"], g.blok["bt"][1]
            for ax in ch._chartSpace.xpath(".//c:plotArea/c:catAx"):
                pos = ax.find(qn("c:tickLblPos"))
                if pos is not None:
                    pos.set("val", "none")
            x0, y0, w0, h0 = BT_PLOT
            rij_h = gf.height * h0 / max(1, len(rijen))
            for i, rij in enumerate(rijen):
                t = (rij[0] or "").strip() if isinstance(rij[0], str) else ""
                if not t:
                    continue
                tb = prs.slides[i_dia].shapes.add_textbox(gf.left, Emu(int(gf.top + gf.height * y0 + rij_h * (i + 0.5) - Pt(letter) * 0.75)),
                                                          Emu(int(gf.width * x0 - Inches(0.06))), Emu(int(Pt(letter) * 1.5)))
                tf = tb.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                p_ = tf.paragraphs[0]
                p_.alignment = PP_ALIGN.RIGHT
                r_ = p_.add_run()
                r_.text = t
                r_.font.size, r_.font.name = Pt(letter), "Arial"
                kop = t.isupper() and "·" in t
                r_.font.bold = kop
                r_.font.color.rgb = RGBColor.from_string(GOUD_LABEL if t.startswith("DAEB ·") else ("002060" if kop else "666666"))
        for naam in ("VP1_GRAFIEK", "VP2_GRAFIEK"):
            if naam in by:
                vul(by[naam].chart, *g.blok["vp"])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("--render", default="")
    ap.add_argument("--simuleer", nargs=2, metavar=("WERKBOEK", "UIT_PPTX"), default=None)
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    _, d7 = _dia(prs, "VP1_GRAFIEK")
    by = shapes_by_name(d7)
    n = sum(markering(by[naam].chart, kol) for naam, kol in KOL_DAEB.items())
    k = kaartjes(prs)
    for naam, kol in KOL_DAEB.items():
        vp_opmaak(shapes_by_name(d7)[naam].chart, kol)
    dia7_extra(prs)
    projectnaam(prs)
    tabel_eigen_dia(prs)
    _, d6 = _dia(prs, "BT_GRAFIEK")
    labels_alle(shapes_by_name(d6)["BT_GRAFIEK"].chart)
    bt_tekenvlak(shapes_by_name(d6)["BT_GRAFIEK"].chart)
    bt_reeksen(shapes_by_name(d6)["BT_GRAFIEK"].chart)
    prs.save(args.uitvoer)
    print(f"Opgeslagen: {args.uitvoer}\n  DAEB-markering in {n} grafiek(en); kaartjes VP_STAT1..3: {'ja' if k else 'stonden er al'}")
    if args.render:
        render_nl(args.uitvoer, args.render)
    if args.simuleer:
        prs2 = Presentation(args.uitvoer)
        simuleer(prs2, Gegevens(args.simuleer[0]))
        prs2.save(args.simuleer[1])
        print(f"  simulatie: {args.simuleer[1]}")
        if args.render:
            render_nl(args.simuleer[1], os.path.join(args.render, "simulatie"))


if __name__ == "__main__":
    main()
