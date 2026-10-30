"""Maakt een voorbeeldpresentatie zoals de knop 'Maak PowerPoint' die maakt (zelfde opbouw en huisstijl).

De dia-blokken worden via LibreOffice uit het verborgen tabblad Dias gerenderd; titels, kernboodschappen en
aandachtspunten komen uit het tabblad PPT. Bedoeld om de vormgeving te beoordelen zonder Excel/PowerPoint.

Gebruik: python tools/preview_pptx.py <werkmap-met-cache.xlsx|xlsm> <uitvoer.pptx> <werkmap-voor-tijdelijke-bestanden>
"""
import sys
import warnings
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageChops
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import preview  # noqa: E402

warnings.filterwarnings("ignore")
NAVY, GEEL, GRIJS, LICHT, LIJN, WIT = (RGBColor(20, 40, 75), RGBColor(245, 184, 0), RGBColor(107, 115, 133),
                                       RGBColor(243, 245, 249), RGBColor(217, 222, 232), RGBColor(255, 255, 255))
BLOKKEN = ["DIA_KPI", "DIA_RDFO", "DIA_PROJ1", "DIA_PROJ2", "DIA_PROJ3", "DIA_BRUG", "DIA_OUTT", "DIA_OUTG", "DIA_OI", "DIA_WON"]


def trim(path):
    im = Image.open(path).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    im = im.crop(box) if box else im
    im.save(path)
    return im.size


def vlak(sld, x, y, w, h, kleur):
    sh = sld.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = kleur
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def tekst(sld, x, y, w, h, t, grootte, kleur, vet=False, rechts=False):
    tb = sld.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.text = t
    if rechts:
        p.alignment = PP_ALIGN.RIGHT
    f = p.runs[0].font
    f.name, f.size, f.bold = "Arial", Pt(grootte), vet
    f.color.rgb = kleur
    return tb


def main(src, out, workdir):
    wb = load_workbook(src, data_only=True)
    T = lambda n: str(_val(wb, n))  # noqa: E731
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(960), Pt(540)
    blank = prs.slide_layouts[6]

    # dia-blokken renderen (één pagina per blok)
    pngs = preview.main(src, Path(workdir) / "dias", ["Dias"], dpi=150)
    pages = [p for p in pngs]
    # de pdf bevat alle bladen; pak de laatste len(BLOKKEN) pagina's (Dias staat achteraan)
    blok_png = dict(zip(BLOKKEN, pages[-len(BLOKKEN):]))
    for p in blok_png.values():
        trim(p)

    # 1. titel
    s = prs.slides.add_slide(blank)
    vlak(s, 0, 0, 300, 540, NAVY)
    vlak(s, 300, 0, 10, 540, GEEL)
    tekst(s, 40, 60, 230, 30, T("P_Regio"), 14, GEEL, True)
    tekst(s, 40, 470, 230, 40, "Voorbeeld — gemaakt met tools/preview_pptx.py", 9, WIT)
    tekst(s, 360, 190, 560, 110, T("PPT_Titel"), 38, NAVY, True)
    vlak(s, 362, 304, 90, 6, GEEL)
    tekst(s, 360, 324, 560, 30, T("PPT_Sub"), 16, NAVY)
    tekst(s, 360, 356, 560, 20, "Bron: RDFO_prototype.xlsm", 10, GRIJS)

    nr = [1]

    def nieuwe_dia(titel, kern):
        nr[0] += 1
        s = prs.slides.add_slide(blank)
        vlak(s, 0, 0, 960, 6, NAVY)
        vlak(s, 36, 30, 6, 30, GEEL)
        tekst(s, 52, 28, 860, 34, titel, 24, NAVY, True)
        if kern:
            tekst(s, 52, 64, 860, 22, kern, 12, GRIJS)
        ln = s.shapes.add_connector(1, Pt(36), Pt(510), Pt(924), Pt(510))
        ln.line.color.rgb = LIJN
        ln.line.width = Pt(0.75)
        tekst(s, 36, 516, 700, 16, T("PPT_Footer"), 9, GRIJS)
        tekst(s, 880, 516, 44, 16, str(nr[0]), 9, NAVY, True, rechts=True)
        return s

    def blok_dia(titel, kern, blok):
        s = nieuwe_dia(titel, kern)
        png = blok_png[blok]
        w, h = Image.open(png).size
        bw, bh = 888, 408
        if w / h > bw / bh:
            pw, ph = bw, bw * h / w
        else:
            pw, ph = bh * w / h, bh
        s.shapes.add_picture(str(png), Pt(36 + (bw - pw) / 2), Pt(96), Pt(pw), Pt(ph))

    n = int(_val(wb, "PPT_ProjN") or 0)
    delen = min(3, -(-n // 18))
    blok_dia("Kerncijfers", T("PPT_Msg1"), "DIA_KPI")
    blok_dia("Omzet en resultaat", T("PPT_Msg2"), "DIA_RDFO")
    for k in range(1, delen + 1):
        blok_dia("Projecten en stoplicht" + (f" ({k}/{delen})" if delen > 1 else ""), T("PPT_Msg3"), f"DIA_PROJ{k}")
    blok_dia("Prognoseverloop en realisatie", T("PPT_Msg4"), "DIA_BRUG")
    blok_dia("Outlook", T("PPT_Msg5"), "DIA_OUTT")
    blok_dia("Outlook in beeld", T("PPT_Msg5"), "DIA_OUTG")
    blok_dia("Orderintake", T("PPT_Msg6"), "DIA_OI")
    blok_dia("Woningverkopen", T("PPT_Msg7"), "DIA_WON")
    s = nieuwe_dia("Aandachtspunten", "")
    vlak(s, 36, 100, 888, 390, LICHT)
    sheet, rng = wb.defined_names["PPT_Bullets"].attr_text.split("!")
    items = [c.value for row in wb[sheet][rng.replace("$", "")] for c in row if c.value not in (None, "")]
    tb = s.shapes.add_textbox(Pt(60), Pt(120), Pt(840), Pt(360))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, t in enumerate(items or ["Geen bijzonderheden."]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "■  " + str(t)
        p.space_after = Pt(14)
        for r in p.runs:
            r.font.name, r.font.size = "Arial", Pt(16)
            r.font.color.rgb = NAVY
    prs.save(out)
    return len(prs.slides)


def _val(wb, name):
    sheet, ref = wb.defined_names[name].attr_text.split("!")
    v = wb[sheet][ref.replace("$", "")].value
    return "" if v is None else v


if __name__ == "__main__":
    print(main(sys.argv[1], sys.argv[2], sys.argv[3]), "dia's")
