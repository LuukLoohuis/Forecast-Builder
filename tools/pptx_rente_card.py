#!/usr/bin/env python3
"""Voeg op dia 2 van het kwartaalsjabloon een vijfde kleine KPI-kaart (rente) toe.

Gebruik:
    python3 tools/pptx_rente_card.py INVOER.pptx UITVOER.pptx [--sub-pt 7.5]

Bewerking:
  * KPI4..KPI7 (kaart + LABEL/WAARDE/SUB) worden verplaatst en versmald zodat
    vijf kaarten in dezelfde totale breedte passen (tussenruimte 137160 EMU).
  * KPI8 wordt toegevoegd als deep copy van de KPI7-vormen (identieke opmaak),
    direct na KPI7_SUB in de spTree, met nieuwe unieke shape-id's.
  * LABEL-vakken worden twee regels hoog (zelfde middelpunt), zodat lange
    labels mogen omlopen zonder de WAARDE te raken.
  * Lettergroottes in de smallere kaarten: LABEL 6pt, SUB 6,5pt en lange
    WAARDE-teksten (>10 tekens) 11pt, zodat niets buiten de kaart afbreekt.
"""
import argparse
import copy
import sys

from pptx import Presentation
from pptx.util import Pt

SLIDE_IDX = 1          # dia 2
LEFT = 411480          # x van de eerste kaart
TOTAL_W = 8321040      # breedte van VOETNOOT / het kaartenblok
GAP = 137160
N_CARDS = 5
CARD_W = (TOTAL_W - (N_CARDS - 1) * GAP) // N_CARDS   # 1554480
INNER = 164592         # binnenmarge links/rechts
TEXT_W = CARD_W - 2 * INNER                           # 1225296
# LABEL-vak: twee regels hoog, gecentreerd rond het oorspronkelijke middelpunt
# (y 3291840 + 128016/2), zodat lange labels in de smallere kaart mogen omlopen
# zonder over de WAARDE heen te vallen.
LABEL_H = 219456
LABEL_Y = 3291840 + 128016 // 2 - LABEL_H // 2        # 3246120
SMALL = ["KPI4", "KPI5", "KPI6", "KPI7", "KPI8"]
PARTS = ["KAART", "LABEL", "WAARDE", "SUB"]

KPI8_TEXT = {
    "LABEL": "RENTELASTEN",
    "WAARDE": "€0,3M",
    "SUB": "3,0% t/m start bouw · down €0,7M · up €0,3M",
}


def shapes_by_name(slide):
    return {sh.name: sh for sh in slide.shapes}


def set_text_keep_format(shape, text):
    """Vervang de tekst van de eerste run; overige runs/alinea's verwijderen."""
    tf = shape.text_frame
    p = tf.paragraphs[0]
    runs = p.runs
    if not runs:
        p.text = text
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for extra in list(tf.paragraphs[1:]):
        extra._p.getparent().remove(extra._p)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("--sub-pt", type=float, default=6.5,
                    help="lettergrootte (pt) voor de SUB-regels van KPI4..KPI8 (0 = ongewijzigd)")
    ap.add_argument("--label-pt", type=float, default=6.0,
                    help="lettergrootte (pt) voor de LABEL-regels van KPI4..KPI8 (0 = ongewijzigd)")
    ap.add_argument("--waarde-lang-pt", type=float, default=11.0,
                    help="lettergrootte (pt) voor WAARDE-teksten langer dan 10 tekens (0 = ongewijzigd)")
    args = ap.parse_args()

    prs = Presentation(args.invoer)
    slide = prs.slides[SLIDE_IDX]
    sp_tree = slide.shapes._spTree
    by_name = shapes_by_name(slide)

    if "KPI8_KAART" in by_name:
        sys.exit("KPI8 bestaat al in de invoer; niets gedaan.")

    # 1. KPI8 als deep copy van KPI7, direct na KPI7_SUB
    max_id = max(int(el.get("id")) for el in sp_tree.iter()
                 if el.tag.endswith("}cNvPr"))
    anchor = by_name["KPI7_SUB"]._element
    new_elems = []
    for part in PARTS:
        src = by_name[f"KPI7_{part}"]._element
        el = copy.deepcopy(src)
        max_id += 1
        cnvpr = el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
        cnvpr.set("id", str(max_id))
        cnvpr.set("name", f"KPI8_{part}")
        new_elems.append(el)
    for el in reversed(new_elems):
        anchor.addnext(el)          # behoudt volgorde KAART, LABEL, WAARDE, SUB
    by_name = shapes_by_name(slide)

    for part, text in KPI8_TEXT.items():
        set_text_keep_format(by_name[f"KPI8_{part}"], text)

    # 2. Vijf kaarten uitlijnen
    for i, kpi in enumerate(SMALL):
        x = LEFT + i * (CARD_W + GAP)
        card = by_name[f"{kpi}_KAART"]
        card.left, card.width = x, CARD_W
        for part in PARTS[1:]:
            sh = by_name[f"{kpi}_{part}"]
            sh.left, sh.width = x + INNER, TEXT_W
            if part == "LABEL":
                sh.top, sh.height = LABEL_Y, LABEL_H
            pt = {"SUB": args.sub_pt, "LABEL": args.label_pt}.get(part, 0)
            if part == "WAARDE" and len(sh.text_frame.text) > 10:
                pt = args.waarde_lang_pt
            if pt:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(pt)

    prs.save(args.uitvoer)
    print(f"Opgeslagen: {args.uitvoer}")
    for sh in slide.shapes:
        if sh.name.startswith("KPI") and sh.name[3] in "45678":
            print(f"  {sh.shape_id:3d} {sh.name:12s} x={sh.left:8d} y={sh.top:8d} "
                  f"w={sh.width:8d} h={sh.height:7d} rechts={sh.left+sh.width}")


if __name__ == "__main__":
    main()
