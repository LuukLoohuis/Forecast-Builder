"""Sjabloon v16 -> v17: dia 2 'kerncijfers' in het nieuwe ontwerp (tijdlijn).

    python3 tools/pptx_v17_kerncijfers.py INVOER_v16.pptx UITVOER_v17.pptx [UITVOER2.pptx ...] [--voorbeeld doorgerekend.xlsx]
                                          [--render MAP] [--simuleer doorgerekend.xlsx UIT.pptx]

Dia 2 wordt opnieuw opgebouwd (kicker, titel en ondertitel van het sjabloon blijven staan; de rest gaat weg):
  * drie tegels KPI1..KPI3 (max. financieringsbehoefte, break-even, projectresultaat) met een accentstreep in de kleur van het
    moment op de curve (rood, navy, groen), groot getal, toelichting, stip + verschil met de vorige prognose, 'was …';
  * een tijdlijnband met de kleine cashflowgrafiek KPI_GRAFIEK (categorie-lijngrafiek: cum. cashflow navy, vorige prognose
    lichtblauw gestreept, één-punts-reeksen voor nu, dieptepunt, break-even, eindsaldo en €0 met de reeksnaam als label;
    jaartallen bovenaan op de categorie-as), het grijze vlak GRAFIEK_REALISATIE t/m het 'nu'-kwartaal, de tijdlijn TIJDLIJN_LIJN,
    het nu-blok (TL_NU_LABEL 'nu', KPI4_WAARDE, KPI4_SUB) en het bijschrift 'vorige prognose' (GRAFIEK_LEGENDA_VORIG + _LIJN);
  * vier kaarten: verkocht bij start bouw (KPI6, balk KPI6_SPOOR/_BALK met normstreep KPI6_NORM), verkocht / getransporteerd
    (KPI7, twee balken KPI7_SPOOR_V/_BALK_V en _T met de woorden KPI7_LABEL_V/_T), marge (KPI8) en de grootste lasten (KPI9:
    KPI9_NAAM1..3, KPI9_BEDRAG1..3, KPI9_SPOOR1..3/KPI9_BALK1..3).
De voorbeeldteksten, balken en grafiekgegevens komen uit een doorgerekend werkboek (tab PowerPoint, blok 'kc' = layout.PP_BLOK['kc']);
zonder --voorbeeld bouwt de tool het voorbeeldproject met een vorige prognose zelf (LibreOffice nodig). Lettertypen komen uit het
thema van het sjabloon (geen vaste letternaam). De macro (v17) vult de teksten, zet de balken (VulBalken), vult de grafiek vanuit
het blok en schikt het grijze vlak en het nu-blok (SchikGrafiekDia2); zonder lasten haalt SchikDia2 de lastenkaart weg.

--simuleer: doet op een v17-sjabloon na wat de macro doet (teksten, balken, grafiek, schikken) met de waarden uit een ander
doorgerekend werkboek, zodat de dia te beoordelen is zonder PowerPoint (bv. zonder vorige prognose of zonder lasten)."""
import argparse
import copy
import io
import os
import shutil
import subprocess
import sys
import tempfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from forecast_builder import layout as LY  # noqa: E402
from pptx_v12_bouwtermijnen import controleer, render, set_text_keep_format, shapes_by_name  # noqa: E402

DIA_KPI = 1                # 0-gebaseerd: dia 2
# kleuren uit het sjabloon (dia 2 v16): navy, rood, groen, toelichting, labels, kaartlijn, kaartvulling klein
NAVY, ROOD, GROEN, GRIJS, GRIJS2 = "002060", "C0392F", "12864E", "666666", "7A7A7A"
LIJN, KAARTVUL, LICHTBLAUW, NAVY2 = "E4E5E6", "F7F8FA", "9DC3E6", "8FA3C8"
ARCERING, ASLIJN, TIJDLIJN = "EBEAE6", "C9CCD1", "B5B9BF"
RICHTING_KLEUR = {"slechter": ROOD, "onder norm": ROOD, "beter": GROEN, "op norm": GROEN}

# ---- raster (inch) ----------------------------------------------------------------------------------------------------------
MARGE_L, MARGE_R, KOL_GAP = 0.5, 9.5, 0.15
KOL_W = (MARGE_R - MARGE_L - 2 * KOL_GAP) / 3                 # drie kolommen boven (tegels) en bij drie kaarten onder
KAART4_W = (MARGE_R - MARGE_L - 3 * KOL_GAP) / 4              # vier kaarten onder
TEGEL_T, TEGEL_H = 1.20, 1.24
FRAME_L, BAND_T, FRAME_R, BAND_B = LY.KC_KADER                # grafiekkader = tijdlijnband
PLOT_L, PLOT_T, PLOT_R, PLOT_B = LY.KC_PLOT                   # tekenvlak (manualLayout)
KAART_T, KAART_H = 4.12, 0.98                                 # onderrand 5,10: de voetstrook blijft vrij
PAD_TEGEL, PAD_KAART = 0.22, 0.18
NU_BLOK_W = 1.30                                              # nu-blok: 'nu' (0,56, rechts uitgelijnd) + 0,04 + getal (0,70); de sub eronder even breed
KEEP = ("KPI_KICKER", "KPI_TITEL", "KPI_SUBTITEL")           # vormen van dia 2 die blijven (plus de dianummer-placeholder)
KC_NAMEN = ["kwartaal", "stand", "vorig", "nu", "dal", "be", "eind", "nul"]
assert len(KC_NAMEN) == LY.PP_KC_KOL


def I(v):
    """inch -> EMU"""
    return Emu(int(round(v * 914400)))


def zonder_stijl(sh):
    """Verwijdert p:style (thema-lijn/-vulling/-schaduw): LibreOffice tekent anders een schaduw die PowerPoint niet toont."""
    st = sh._element.find(qn("p:style"))
    if st is not None:
        sh._element.remove(st)
    return sh


def tekstvak(slide, naam, l, t, w, h, tekst, pt, kleur, vet=False, spc=0, anker=MSO_ANCHOR.MIDDLE, uitl=PP_ALIGN.LEFT, krimp=False):
    """Eén tekstvak, één alinea, één run; themaletter, vaste grootte, word-wrap aan (maten in inch). krimp=True: 'tekst verkleinen
    bij overloop' (normAutofit), zodat een lange waarde ('niet bereikt', '−12,3% / −11,0%') in PowerPoint kleiner wordt in plaats
    van om te lopen; anders geen autofit."""
    tb = slide.shapes.add_textbox(I(l), I(t), I(w), I(h))
    tb.name = naam
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if krimp else MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anker
    p = tf.paragraphs[0]
    p.alignment = uitl
    r = p.add_run()
    r.text = tekst
    r.font.size, r.font.bold = Pt(pt), vet
    r.font.color.rgb = RGBColor.from_string(kleur)
    if spc:
        r._r.get_or_add_rPr().set("spc", str(spc))
    return tb


def label(slide, naam, l, t, w, tekst, pt=7, spc=60):
    """label in kleine kapitalen met letterafstand, zoals KPI1_LABEL in het sjabloon"""
    return tekstvak(slide, naam, l, t, w, 0.14, tekst, pt, GRIJS2, True, spc)


def rect(slide, naam, l, t, w, h, vul, lijn=None, lijn_pt=0.8, rond=0.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rond else MSO_SHAPE.RECTANGLE, I(l), I(t), I(w), I(h))
    sh.name = naam
    zonder_stijl(sh)
    if rond:
        sh.adjustments[0] = rond
    if vul is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = RGBColor.from_string(vul)
    if lijn is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = RGBColor.from_string(lijn)
        sh.line.width = Pt(lijn_pt)
    sh.shadow.inherit = False
    return sh


def kaart(slide, naam, l, t, w, h, vul="FFFFFF"):
    return rect(slide, naam, l, t, w, h, vul, LIJN, 0.8, rond=0.03)


def stip(slide, naam, cx, cy, d, kleur):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, I(cx - d / 2), I(cy - d / 2), I(d), I(d))
    sh.name = naam
    zonder_stijl(sh)
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(kleur)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def lijn(slide, naam, x1, y1, x2, y2, kleur, pt=0.75, streep=None):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    c.name = naam
    zonder_stijl(c)
    c.line.color.rgb = RGBColor.from_string(kleur)
    c.line.width = Pt(pt)
    if streep is not None:
        c.line.dash_style = streep
    return c


# ---- werkboek --------------------------------------------------------------------------------------------------------------
def maak_voorbeeld(map_, vorig=True, lasten=True):
    """Bouwt het voorbeeldproject (met een vorige prognose en de drie lasten) en rekent het door met LibreOffice; geeft het pad."""
    import build
    from forecast_builder import data as D
    p = D.voorbeeld()
    if vorig:
        stand = 0.0
        for per in p.periodes:
            stand += (per["omzet"] or 0) - (per["kosten"] or 0)
            per["cf_vorig"] = round(stand * (0.93 if stand < 0 else 0.72) / 1000, 1)
    if not lasten:
        p.params["lasten"] = []
    pad = build.bouw(p, os.path.join(map_, "voorbeeld_dia2"), None, caches=False)[0]
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    assert soffice, "LibreOffice (soffice) is nodig om het voorbeeld door te rekenen; geef anders --voorbeeld"
    subprocess.run([soffice, "--headless", f"-env:UserInstallation=file://{map_}/lo_profiel", "--calc", "--convert-to",
                    "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(map_, "calc"), pad], check=True, capture_output=True, timeout=300)
    return os.path.join(map_, "calc", os.path.basename(pad))


def nl(t):
    """decimalen met een komma (LibreOffice rekent hier in een Engelse locale; Excel in het Nederlands schrijft de komma zelf)"""
    import re
    return re.sub(r"(\d)\.(\d)", r"\1,\2", t)


class Gegevens:
    """Alles wat dia 2 uit een doorgerekend werkboek nodig heeft: teksten per vorm (zoals de macro ze invult), balkaandelen,
    het grafiekblok en de hulpcellen."""

    def __init__(self, pad):
        wb = load_workbook(pad, data_only=True)
        m, pp = wb["Model"], wb["PowerPoint"]
        self.tekst, self.richting, self.aandeel = {}, {}, {}
        r = LY.PP_KPI_ROW1
        while pp[f"C{r}"].value:
            vorm = pp[f"F{r}"].value
            if pp[f"B{r}"].value == 2 and vorm:
                v = pp[f"D{r}"].value
                if "_BALK" in vorm or vorm.endswith("_NORM"):
                    self.aandeel[vorm] = float(v) if isinstance(v, (int, float)) else 0.0
                else:
                    self.tekst[vorm] = "" if v is None else nl(str(v))
                    self.richting[vorm] = (pp[f"E{r}"].value or "").strip().lower()
            r += 1
        h = lambda naam: m[f"{LY.M_HULP}{LY.H[naam]}"].value  # noqa: E731
        self.n = int(h("n") or 0)
        self.pos_nu = int(h("pos_actuals") or 0)
        self.kc_min, self.kc_max = float(h("kc_min") or 0), float(h("kc_max") or 0)
        c0 = pp[LY.PP_BLOK["kc"] + "1"].column
        self.koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_KC_KOL)]
        assert self.koppen[0] == "Kwartaal" and self.koppen[-1] == "€0", self.koppen
        self.rijen = []
        for i in range(self.n):
            rij = []
            for j in range(LY.PP_KC_KOL):
                v = pp.cell(LY.PP_KPI_ROW1 + i, c0 + j).value
                rij.append(None if v in (None, "") or (isinstance(v, str) and v.startswith("#")) else v)
            self.rijen.append(rij)

    def werkblad(self):
        """Sheet1 zoals de macro het gegevensblad van KPI_GRAFIEK vult: kop + n rijen (lege cel voor #N/A)."""
        ws = Workbook().active
        ws.title = "Sheet1"
        ws.append(["" if k is None else str(k) for k in self.koppen])
        for rij in self.rijen:
            ws.append(rij)
        return ws

    def kolom(self, naam):
        j = KC_NAMEN.index(naam)
        return [rij[j] for rij in self.rijen]


# ---- dia 2 -----------------------------------------------------------------------------------------------------------------
def _nofill_spPr(parent, before_tags=()):
    spPr = etree.Element(qn("c:spPr"))
    etree.SubElement(spPr, qn("a:noFill"))
    ln = etree.SubElement(spPr, qn("a:ln"))
    etree.SubElement(ln, qn("a:noFill"))
    for tag in before_tags:
        el = parent.find(qn(tag))
        if el is not None:
            el.addprevious(spPr)
            return spPr
    parent.append(spPr)
    return spPr


def _dlbl_opmaak(dLbls, pt, kleur):
    """lettergrootte en kleur van een set datalabels (themaletter); niet afbreken"""
    txPr = dLbls.get_or_add_txPr()
    txPr.bodyPr.set("wrap", "none")
    defRPr = txPr.find(qn("a:p")).find(qn("a:pPr")).find(qn("a:defRPr"))
    defRPr.set("sz", str(int(pt * 100)))
    defRPr.set("b", "0")
    for kind in list(defRPr):
        defRPr.remove(kind)
    sf = etree.SubElement(defRPr, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr")).set("val", kleur)


def x_van(shp, i, n):
    """x (EMU) van kwartaal i (1..n) in de grafiek, zoals de macro rekent: links + breedte × (KC_PLOT_X + KC_PLOT_W × (i−1)/(n−1))"""
    return shp.left + shp.width * (LY.KC_PLOT_X + LY.KC_PLOT_W * (i - 1) / max(1, n - 1))


def grafiek(slide, g):
    """KPI_GRAFIEK: categorie-lijngrafiek met de brede werkmap van het blok (A jaarlabel, B cum. cashflow, C vorige prognose,
    D nu, E dieptepunt, F break-even, G eindsaldo, H €0); reeksen met '(uit)' in de kop krijgen geen waarden (de macro haalt ze weg)."""
    cd = CategoryChartData()
    cd.categories = ["" if v is None else str(int(v)) if isinstance(v, (int, float)) else str(v) for v in g.kolom("kwartaal")]
    for naam, kop in zip(KC_NAMEN[1:], g.koppen[1:]):
        waarden = g.kolom(naam)
        if str(kop).endswith("(uit)"):
            waarden = [None] * g.n
        cd.add_series(str(kop), waarden)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.LINE, I(FRAME_L), I(BAND_T), I(FRAME_R - FRAME_L), I(BAND_B - BAND_T), cd)
    gf.name = "KPI_GRAFIEK"
    ch = gf.chart
    ch.has_title = False
    ch.has_legend = False
    ch.font.size = Pt(7)
    ch.font.color.rgb = RGBColor.from_string(GRIJS)
    reeksen = list(ch.plots[0].series)
    s_stand, s_vorig = reeksen[0], reeksen[1]
    s_stand.smooth = False
    s_stand.format.line.color.rgb = RGBColor.from_string(NAVY)
    s_stand.format.line.width = Pt(2)
    s_stand.marker.style = XL_MARKER_STYLE.NONE
    s_vorig.smooth = False
    s_vorig.format.line.color.rgb = RGBColor.from_string(LICHTBLAUW)
    s_vorig.format.line.width = Pt(1.25)
    s_vorig.format.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    s_vorig.marker.style = XL_MARKER_STYLE.NONE
    # (vulling, rand, labelpositie) van nu, dieptepunt, break-even, eindsaldo; label = reeksnaam op reeksniveau (volgt het punt)
    markeringen = [("FFFFFF", NAVY, None), (ROOD, "FFFFFF", XL_LABEL_POSITION.BELOW),
                   (NAVY, "FFFFFF", XL_LABEL_POSITION.ABOVE), (GROEN, "FFFFFF", XL_LABEL_POSITION.RIGHT)]
    for (vul, rand, pos), sr in zip(markeringen, reeksen[2:6]):
        sr.smooth = False
        sr.format.line.fill.background()
        sr.marker.style = XL_MARKER_STYLE.CIRCLE
        sr.marker.size = 7
        sr.marker.format.fill.solid()
        sr.marker.format.fill.fore_color.rgb = RGBColor.from_string(vul)
        sr.marker.format.line.color.rgb = RGBColor.from_string(rand)
        sr.marker.format.line.width = Pt(1.25)
        if pos is not None:
            dls = sr.data_labels
            dls.show_series_name, dls.show_value = True, False
            dls.position = pos
            _dlbl_opmaak(dls._element, 7, GRIJS)
            dls._element.find(qn("c:showLeaderLines")).set("val", "0")
    s_nul = reeksen[6]
    s_nul.smooth = False
    s_nul.format.line.fill.background()
    s_nul.marker.style = XL_MARKER_STYLE.NONE
    dls = s_nul.data_labels
    dls.show_series_name, dls.show_value = True, False
    dls.position = XL_LABEL_POSITION.ABOVE
    _dlbl_opmaak(dls._element, 7, GRIJS2)
    dls._element.find(qn("c:showLeaderLines")).set("val", "0")

    ca = ch.category_axis
    ca.tick_label_position = XL_TICK_LABEL_POSITION.HIGH          # jaartallen bovenaan, onder de tijdlijn
    ca.major_tick_mark = XL_TICK_MARK.NONE
    ca.minor_tick_mark = XL_TICK_MARK.NONE
    ca.has_major_gridlines = False
    ca.format.line.color.rgb = RGBColor.from_string(ASLIJN)       # de categorie-as kruist op 0: dit is de nullijn
    ca.format.line.width = Pt(0.75)
    ca.tick_labels.font.size = Pt(7)
    ca.tick_labels.font.color.rgb = RGBColor.from_string(GRIJS2)
    va = ch.value_axis
    va.visible = False
    va.has_major_gridlines = False
    va.has_minor_gridlines = False
    va.minimum_scale, va.maximum_scale = g.kc_min, g.kc_max        # de macro zet ze op Model kc_min/kc_max
    cb = etree.Element(qn("c:crossBetween"))                       # eerste/laatste kwartaal op de randen van het tekenvlak
    cb.set("val", "midCat")
    va._element.find(qn("c:crosses")).addnext(cb)
    # tekenvlak vast (fracties van het kader): bovenaan ruimte voor de jaartallen, rechts voor het label 'eindsaldo'
    fw, fh = FRAME_R - FRAME_L, BAND_B - BAND_T
    cs = ch._chartSpace
    pa = cs.chart.plotArea
    oud = pa.find(qn("c:layout"))
    if oud is not None:
        pa.remove(oud)
    lay = etree.Element(qn("c:layout"))
    ml = etree.SubElement(lay, qn("c:manualLayout"))
    for tag, val in (("layoutTarget", "inner"), ("xMode", "edge"), ("yMode", "edge"), ("x", LY.KC_PLOT_X),
                     ("y", (PLOT_T - BAND_T) / fh), ("w", LY.KC_PLOT_W), ("h", (PLOT_B - PLOT_T) / fh)):
        etree.SubElement(ml, qn(f"c:{tag}")).set("val", f"{val:.5f}" if isinstance(val, float) else val)
    pa.insert(0, lay)
    _nofill_spPr(cs, ("c:txPr", "c:externalData", "c:printSettings"))
    _nofill_spPr(pa, ("c:extLst",))
    rc = cs.find(qn("c:roundedCorners"))
    if rc is None:
        rc = etree.Element(qn("c:roundedCorners"))
        cs.insert(0, rc)
    rc.set("val", "0")
    return gf


def kleur_van(g, vorm, standaard):
    return RICHTING_KLEUR.get(g.richting.get(vorm, ""), standaard)


def bouw_dia2(slide, g, vier_kaarten=True):
    """Vervangt de inhoud van dia 2 (behalve kicker, titel, ondertitel en de dianummer-placeholder) door het nieuwe ontwerp."""
    by = shapes_by_name(slide)
    for sh in list(slide.shapes):
        if sh.name in KEEP or sh.is_placeholder:
            continue
        sh._element.getparent().remove(sh._element)
    for naam in KEEP:
        if naam in by and naam in g.tekst:
            set_text_keep_format(by[naam], g.tekst[naam])

    # ---- drie tegels, accentstreep in de kleur van het moment op de curve
    for i, (k, kleur) in enumerate(((1, ROOD), (2, NAVY), (3, GROEN))):
        L, T = MARGE_L + i * (KOL_W + KOL_GAP), TEGEL_T
        kaart(slide, f"KPI{k}_KAART", L, T, KOL_W, TEGEL_H)
        rect(slide, f"KPI{k}_ACCENT", L, T + 0.16, 0.045, TEGEL_H - 0.32, kleur)
        ix, iw = L + PAD_TEGEL, KOL_W - PAD_TEGEL - 0.14
        label(slide, f"KPI{k}_LABEL", ix, T + 0.14, iw, {1: "MAX. FINANCIERINGSBEHOEFTE", 2: "BREAK-EVEN", 3: "PROJECTRESULTAAT"}[k])
        tekstvak(slide, f"KPI{k}_WAARDE", ix - 0.02, T + 0.27, iw, 0.46, g.tekst[f"KPI{k}_WAARDE"], 30, kleur_van(g, f"KPI{k}_WAARDE", kleur), True, krimp=True)
        tekstvak(slide, f"KPI{k}_SUB", ix, T + 0.72, iw, 0.16, g.tekst[f"KPI{k}_SUB"], 8.5, GRIJS)
        rk = kleur_van(g, f"KPI{k}_VERSCHIL", GRIJS2)
        stip(slide, f"KPI{k}_STIP", ix + 0.0425, T + 0.98, 0.085, rk)
        tekstvak(slide, f"KPI{k}_VERSCHIL", ix + 0.15, T + 0.90, iw - 0.15, 0.16, g.tekst[f"KPI{k}_VERSCHIL"], 8, rk, True)
        tekstvak(slide, f"KPI{k}_WAS", ix + 0.15, T + 1.06, iw - 0.15, 0.15, g.tekst[f"KPI{k}_WAS"], 8, GRIJS)

    # ---- tijdlijnband: grijs vlak (gerealiseerd t/m 'nu'), grafiek, tijdlijn, nu-blok, bijschrift vorige prognose
    ph = PLOT_B - PLOT_T
    vlak = rect(slide, "GRAFIEK_REALISATIE", PLOT_L, PLOT_T, 0.01, ph, ARCERING)
    gf = grafiek(slide, g)
    x_nu = x_van(gf, g.pos_nu, g.n) if g.pos_nu >= 1 else gf.left + I(PLOT_L - FRAME_L)
    vlak.width = max(Emu(1), Emu(int(x_nu - vlak.left)))
    lijn(slide, "TIJDLIJN_LIJN", MARGE_L, PLOT_T, MARGE_R, PLOT_T, TIJDLIJN, 0.75)
    xr = x_nu / 914400 - 0.1                                     # rechterkant van het nu-blok (inch), links van de stip
    r1, r2 = PLOT_T + ph - 0.26, PLOT_T + ph - 0.12
    tekstvak(slide, "TL_NU_LABEL", xr - NU_BLOK_W, r1, NU_BLOK_W - 0.74, 0.14, "nu", 8, NAVY, uitl=PP_ALIGN.RIGHT)
    tekstvak(slide, "KPI4_WAARDE", xr - 0.70, r1, 0.70, 0.14, g.tekst["KPI4_WAARDE"], 10, NAVY, True)
    tekstvak(slide, "KPI4_SUB", xr - NU_BLOK_W, r2, NU_BLOK_W, 0.11, g.tekst["KPI4_SUB"], 7, GRIJS, uitl=PP_ALIGN.RIGHT, krimp=True)
    ly = PLOT_T + ph - 0.085
    tekstvak(slide, "GRAFIEK_LEGENDA_VORIG", MARGE_R - 0.95, ly - 0.07, 0.95, 0.14, g.tekst.get("GRAFIEK_LEGENDA_VORIG", ""), 7, GRIJS,
             uitl=PP_ALIGN.RIGHT)
    lijn(slide, "GRAFIEK_LEGENDA_VORIG_LIJN", MARGE_R - 1.25, ly, MARGE_R - 1.01, ly, LICHTBLAUW, 1.25, MSO_LINE_DASH_STYLE.DASH)

    # ---- onderste rij: vier (of drie) lichte kaarten; label, getal, toelichting en onderin de balken
    kaarten = [6, 7, 8, 9] if vier_kaarten else [6, 7, 8]
    kw = KAART4_W if vier_kaarten else KOL_W
    for i, k in enumerate(kaarten):
        L, T = MARGE_L + i * (kw + KOL_GAP), KAART_T
        kaart(slide, f"KPI{k}_KAART", L, T, kw, KAART_H, KAARTVUL)
        ix, iw = L + PAD_KAART, kw - 2 * PAD_KAART
        label(slide, f"KPI{k}_LABEL", ix, T + 0.12, iw, g.tekst[f"KPI{k}_LABEL"], 6, 40)
        if k != 9:
            wk = kleur_van(g, f"KPI{k}_WAARDE", NAVY)
            tekstvak(slide, f"KPI{k}_WAARDE", ix - 0.01, T + 0.27, iw, 0.30, g.tekst[f"KPI{k}_WAARDE"], 16, wk, True, krimp=True)
            tekstvak(slide, f"KPI{k}_SUB", ix, T + 0.60, iw, 0.14, g.tekst[f"KPI{k}_SUB"], 6.5, GRIJS)
        if k == 6:       # pil-balk over de hele breedte met normstreep; de macro: BALK.Width = aandeel × SPOOR.Width
            by6, bh = T + 0.80, 0.07
            rect(slide, "KPI6_SPOOR", ix, by6, iw, bh, LIJN, rond=0.5)
            rect(slide, "KPI6_BALK", ix, by6, max(bh, iw * g.aandeel["KPI6_BALK"]), bh, wk, rond=0.5)
            rect(slide, "KPI6_NORM", ix + iw * g.aandeel["KPI6_NORM"] - 0.01, by6 - 0.035, 0.02, bh + 0.07, NAVY)
        elif k == 7:     # twee benoemde balken met de woorden ervoor
            lw, bh = 0.80, 0.06
            for sfx, y0, kleur_b, woord in (("V", T + 0.74, NAVY, "verkocht"), ("T", T + 0.86, NAVY2, "getransporteerd")):
                tekstvak(slide, f"KPI7_LABEL_{sfx}", ix, y0 - 0.035, lw - 0.04, bh + 0.07, woord, 6, GRIJS)
                rect(slide, f"KPI7_SPOOR_{sfx}", ix + lw, y0, iw - lw, bh, LIJN, rond=0.5)
                rect(slide, f"KPI7_BALK_{sfx}", ix + lw, y0, max(bh, (iw - lw) * g.aandeel[f"KPI7_BALK_{sfx}"]), bh, kleur_b, rond=0.5)
        elif k == 9:     # grootste lasten: drie regels naam / bedrag met een dunne balk (aandeel van het hoogste bedrag)
            for j in range(1, LY.N_LASTEN + 1):
                y0 = T + 0.30 + (j - 1) * 0.21
                tekstvak(slide, f"KPI9_NAAM{j}", ix, y0, iw - 0.6, 0.14, g.tekst.get(f"KPI9_NAAM{j}", ""), 7, GRIJS)
                tekstvak(slide, f"KPI9_BEDRAG{j}", ix + iw - 0.7, y0, 0.7, 0.14, g.tekst.get(f"KPI9_BEDRAG{j}", ""), 7.5, NAVY, True,
                         uitl=PP_ALIGN.RIGHT)
                rect(slide, f"KPI9_SPOOR{j}", ix, y0 + 0.155, iw, 0.04, LIJN, rond=0.5)
                rect(slide, f"KPI9_BALK{j}", ix, y0 + 0.155, max(0.04, iw * g.aandeel.get(f"KPI9_BALK{j}", 0)), 0.04, NAVY, rond=0.5)
    return slide


# ---- simulatie van de macro ---------------------------------------------------------------------------------------------------
def simuleer(prs, g):
    """Doet op dia 2 van een v17-sjabloon na wat de macro doet met de waarden van `g`: VulTeksten (tekst en richtingkleur),
    SchikDia2 (leeg bijschrift weg, lege lastenkaart weg en drie kaarten onder de tegels), VulBalken (balkbreedtes), de grafiek
    (gegevens, '(uit)'-reeksen leeg, as-grenzen) en SchikGrafiekDia2 (grijs vlak en nu-blok)."""
    slide = prs.slides[DIA_KPI]
    by = shapes_by_name(slide)
    for vorm, tekst in g.tekst.items():                                 # VulTeksten
        sh = by.get(vorm)
        if sh is None or not sh.has_text_frame:
            continue
        set_text_keep_format(sh, tekst)
        r = g.richting.get(vorm, "")
        if r:
            kleur = RGBColor.from_string(RICHTING_KLEUR.get(r, GRIJS2))
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = kleur
            if vorm.endswith("_VERSCHIL") and vorm.replace("_VERSCHIL", "_STIP") in by:
                by[vorm.replace("_VERSCHIL", "_STIP")].fill.fore_color.rgb = kleur
    # SchikDia2
    if not g.tekst.get("GRAFIEK_LEGENDA_VORIG"):
        for naam in ("GRAFIEK_LEGENDA_VORIG", "GRAFIEK_LEGENDA_VORIG_LIJN"):
            if naam in by:
                by[naam]._element.getparent().remove(by[naam]._element)
    kaarten = {k: by[f"KPI{k}_KAART"] for k in range(4, 21) if f"KPI{k}_KAART" in by}
    rij_top = max(k.top for k in kaarten.values())
    kaarten = {k: v for k, v in kaarten.items() if v.top >= rij_top - Pt(7)}
    blijft = {k: (f"KPI{k}_WAARDE" in g.tekst) or bool(g.tekst.get(f"KPI{k}_NAAM1")) for k in kaarten}
    if not all(blijft.values()):
        oud = {k: (v.left, v.width) for k, v in kaarten.items()}
        for k, ok in blijft.items():
            if not ok:
                for sh in list(slide.shapes):
                    if sh.name.startswith(f"KPI{k}_"):
                        sh._element.getparent().remove(sh._element)
        over = sorted((k for k, ok in blijft.items() if ok), key=lambda k: oud[k][0])
        links = min(v[0] for v in oud.values())
        rechts = max(v[0] + v[1] for v in oud.values())
        gesorteerd = sorted(oud.values())
        tussen = min(b[0] - (a[0] + a[1]) for a, b in zip(gesorteerd, gesorteerd[1:]))
        if len(over) == 3:
            doel = sorted((by[f"KPI{t}_KAART"].left, by[f"KPI{t}_KAART"].width) for t in (1, 2, 3))
            binnen = by["KPI1_LABEL"].left + by["KPI1_LABEL"].text_frame.margin_left - by["KPI1_KAART"].left
        else:
            w = (rechts - links - (len(over) - 1) * tussen) / len(over)
            doel = [(links + t * (w + tussen), w) for t in range(len(over))]
            binnen = -1
        for k, (dl, db) in zip(over, doel):
            ol, ow = oud[k]
            for sh in slide.shapes:
                if not sh.name.startswith(f"KPI{k}_"):
                    continue
                mL, mR = sh.left - ol, ol + ow - (sh.left + sh.width)
                if "_BALK" in sh.name or sh.name.endswith("_NORM"):
                    sh.left = Emu(int(dl + mL))
                    continue
                tekstvorm = (sh.name.endswith(("_LABEL", "_WAARDE", "_SUB")) or "_NAAM" in sh.name or "_BEDRAG" in sh.name)
                if binnen >= 0 and tekstvorm:
                    mL = binnen - sh.text_frame.margin_left
                    mR = binnen - sh.text_frame.margin_right
                if db - mL - mR > 1:
                    sh.left, sh.width = Emu(int(dl + mL)), Emu(int(db - mL - mR))
    by = shapes_by_name(slide)
    for vorm, aandeel in g.aandeel.items():                             # VulBalken
        aandeel = min(1.0, max(0.0, aandeel))
        sh = by.get(vorm)
        spoor = by.get(vorm.replace("_BALK", "_SPOOR") if "_BALK" in vorm else vorm.replace("_NORM", "_SPOOR"))
        if sh is None or spoor is None:
            continue
        if vorm.endswith("_NORM"):
            sh.left = Emu(int(spoor.left + aandeel * spoor.width - sh.width / 2))
        else:
            sh.left, sh.top, sh.height = spoor.left, spoor.top, spoor.height
            sh.width = Emu(int(max(spoor.height, aandeel * spoor.width)))
            if aandeel * spoor.width < Pt(0.5):
                sh._element.getparent().remove(sh._element)
            elif vorm.endswith("_BALK"):
                r = g.richting.get(vorm.replace("_BALK", "_WAARDE"), "")
                if r:
                    sh.fill.fore_color.rgb = RGBColor.from_string(RICHTING_KLEUR.get(r, GRIJS2))
    # grafiek: nieuwe gegevens (replace_data houdt de opmaak per reeks), '(uit)'-reeksen leeg, as-grenzen
    gf = by["KPI_GRAFIEK"]
    cd = CategoryChartData()
    cd.categories = ["" if v is None else str(int(v)) if isinstance(v, (int, float)) else str(v) for v in g.kolom("kwartaal")]
    for naam, kop in zip(KC_NAMEN[1:], g.koppen[1:]):
        cd.add_series(str(kop), [None] * g.n if str(kop).endswith("(uit)") else g.kolom(naam))
    gf.chart.replace_data(cd)
    gf.chart.value_axis.minimum_scale, gf.chart.value_axis.maximum_scale = g.kc_min, g.kc_max
    # SchikGrafiekDia2: grijs vlak t/m 'nu' en het nu-blok ernaast (rechts van de stip als het links niet past)
    vlak = by.get("GRAFIEK_REALISATIE")
    blok = [by[n] for n in ("TL_NU_LABEL", "KPI4_WAARDE", "KPI4_SUB") if n in by]
    if g.pos_nu < 1 or g.n < 2:
        for sh in [vlak] + blok:
            if sh is not None:
                sh._element.getparent().remove(sh._element)
        return slide
    plot_l = gf.left + gf.width * LY.KC_PLOT_X
    x_nu = x_van(gf, g.pos_nu, g.n)
    x_oud = vlak.left + vlak.width
    vlak.left, vlak.width = Emu(int(plot_l)), Emu(int(max(1, x_nu - plot_l)))
    dx = x_nu - x_oud
    for sh in blok:
        sh.left = Emu(int(sh.left + dx))
    b_links = min(sh.left for sh in blok)
    if b_links < plot_l:
        dx = (x_nu + Pt(7)) - b_links
        for sh in blok:
            sh.left = Emu(int(sh.left + dx))
            for p in sh.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT
        if "KPI4_WAARDE" in by:
            by["KPI4_WAARDE"].left = Emu(int(x_nu + Pt(7) + I(0.2)))
    return slide


VERWACHT = ["KPI_KICKER", "KPI_TITEL", "KPI_SUBTITEL", "KPI1_KAART", "KPI1_ACCENT", "KPI1_LABEL", "KPI1_WAARDE", "KPI1_SUB", "KPI1_STIP",
            "KPI1_VERSCHIL", "KPI1_WAS", "KPI2_WAARDE", "KPI3_WAARDE", "GRAFIEK_REALISATIE", "KPI_GRAFIEK", "TIJDLIJN_LIJN", "TL_NU_LABEL",
            "KPI4_WAARDE", "KPI4_SUB", "GRAFIEK_LEGENDA_VORIG", "GRAFIEK_LEGENDA_VORIG_LIJN", "KPI6_KAART", "KPI6_LABEL", "KPI6_WAARDE",
            "KPI6_SUB", "KPI6_SPOOR", "KPI6_BALK", "KPI6_NORM", "KPI7_KAART", "KPI7_WAARDE", "KPI7_LABEL_V", "KPI7_SPOOR_V", "KPI7_BALK_V",
            "KPI7_LABEL_T", "KPI7_SPOOR_T", "KPI7_BALK_T", "KPI8_KAART", "KPI8_LABEL", "KPI8_WAARDE", "KPI8_SUB", "KPI9_KAART", "KPI9_LABEL",
            "KPI9_NAAM1", "KPI9_BEDRAG1", "KPI9_SPOOR1", "KPI9_BALK1", "KPI9_NAAM3", "KPI9_BEDRAG3", "KPI9_SPOOR3", "KPI9_BALK3"]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("invoer")
    ap.add_argument("uitvoer")
    ap.add_argument("extra", nargs="*", help="extra kopieën van de uitvoer")
    ap.add_argument("--voorbeeld", default="", help="doorgerekend werkboek (.xlsx) voor de voorbeeldteksten en -grafiek; leeg = zelf bouwen")
    ap.add_argument("--render", default="", help="map voor PNG's van alle dia's (LibreOffice + pdftoppm); leeg = niet renderen")
    ap.add_argument("--simuleer", nargs=2, metavar=("WERKBOEK", "UIT_PPTX"), default=None,
                    help="doe op de nieuwe v17 na wat de macro doet met dit doorgerekende werkboek en bewaar het als UIT_PPTX")
    args = ap.parse_args()
    prs = Presentation(args.invoer)
    if len(prs.slides) != 9:
        sys.exit(f"Verwacht 9 dia's in v16 (titel + 8), gevonden {len(prs.slides)}.")
    with tempfile.TemporaryDirectory() as tmp:
        pad = args.voorbeeld or maak_voorbeeld(tmp)
        g = Gegevens(pad)
        bouw_dia2(prs.slides[DIA_KPI], g)
        prs.save(args.uitvoer)
        for extra in args.extra:
            shutil.copyfile(args.uitvoer, extra)
        print(f"Opgeslagen: {args.uitvoer}" + "".join(f"\n            {e}" for e in args.extra))
        print(f"  dia 2 opnieuw opgebouwd ({g.n} voorbeeldrijen in KPI_GRAFIEK, {len(g.tekst)} teksten, {len(g.aandeel)} balken)")
        fouten, charts, emb = controleer(args.uitvoer, {2: VERWACHT})
        print("  chart-parts: " + ", ".join(charts))
        if fouten:
            sys.exit("FOUTEN:\n  " + "\n  ".join(fouten))
        print("  geen fouten")
        if args.render:
            pngs = render(args.uitvoer, args.render)
            print("  PNG's: " + (", ".join(pngs) if pngs else "niet gerenderd (LibreOffice/pdftoppm ontbreekt)"))
        if args.simuleer:
            prs2 = Presentation(args.uitvoer)
            simuleer(prs2, Gegevens(args.simuleer[0]))
            prs2.save(args.simuleer[1])
            print(f"  simulatie: {args.simuleer[1]}")
            if args.render:
                pngs = render(args.simuleer[1], os.path.join(args.render, "simulatie"))
                print("  PNG's simulatie: " + (", ".join(pngs) if pngs else "niet gerenderd"))


if __name__ == "__main__":
    main()
