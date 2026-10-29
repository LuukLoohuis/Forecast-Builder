"""Zet berekende waarden (uit een LibreOffice-herberekende kopie) als cache in de door openpyxl geschreven werkmap.

openpyxl schrijft formules zonder resultaat; Excel rekent ze bij openen opnieuw (fullCalcOnLoad), maar previews
(telefoon, mail, verkenner) tonen dan lege cellen. Dit script vult alleen de <v>-waarde van formulecellen; formules,
opmaak, grafieken, tabellen en validaties blijven exact zoals openpyxl ze schreef.

Gebruik: python tools/inject_cache.py <openpyxl.xlsx> <herberekend.xlsx> <uitvoer.xlsx>
"""
import re
import sys
import warnings
import zipfile

from lxml import etree
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def sheet_files(z):
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels}
    out = {}
    for s in wbx.iter(f"{{{NS}}}sheet"):
        t = target[s.get(f"{{{RNS}}}id")]
        out[s.get("name")] = "xl/" + t.lstrip("/").replace("xl/", "") if not t.startswith("xl/") else t
    return out


def main(src, calc, out):
    vals = load_workbook(calc, data_only=True)
    zin = zipfile.ZipFile(src)
    files = {v: k for k, v in sheet_files(zin).items()}
    zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    filled = 0
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename in files:
            ws = vals[files[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter(f"{{{NS}}}c"):
                f = c.find(f"{{{NS}}}f")
                if f is None:
                    continue
                vnode = c.find(f"{{{NS}}}v")
                if vnode is None:
                    vnode = etree.SubElement(c, f"{{{NS}}}v")
                val = ws[c.get("r")].value
                if val is None or val == "":
                    c.set("t", "str")
                    vnode.text = ""
                elif isinstance(val, bool):
                    c.set("t", "b")
                    vnode.text = "1" if val else "0"
                elif isinstance(val, (int, float)):
                    c.attrib.pop("t", None)
                    vnode.text = repr(float(val)) if isinstance(val, float) else str(val)
                elif hasattr(val, "toordinal"):
                    from openpyxl.utils.datetime import to_excel
                    c.attrib.pop("t", None)
                    vnode.text = repr(float(to_excel(val)))
                elif isinstance(val, str) and re.fullmatch(r"#(N/A|VALUE!|REF!|DIV/0!|NUM!|NAME\?|NULL!)", val):
                    c.set("t", "e")
                    vnode.text = val
                else:
                    c.set("t", "str")
                    vnode.text = str(val)
                filled += 1
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zout.writestr(item, data)
    zout.close()
    return filled


if __name__ == "__main__":
    print(main(*sys.argv[1:4]), "formulecellen van cache voorzien")
