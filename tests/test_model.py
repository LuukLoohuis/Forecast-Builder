"""Regressietests voor het gebouwde werkboek.

Draaien: python3 -m pytest tests/ -q   (of: python3 tests/test_model.py)
Heeft LibreOffice (soffice) nodig om formules door te rekenen; zonder LibreOffice worden de rekentests overgeslagen.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile

import pytest
from openpyxl import load_workbook

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HIER))

import build  # noqa: E402
from forecast_builder import data as D, layout as LY  # noqa: E402

SOFFICE = shutil.which("soffice") or shutil.which("libreoffice")


def herbereken(pad):
    """Rekent het werkboek door met LibreOffice (in-place) en geeft het pad terug."""
    map_ = os.path.dirname(pad)
    subprocess.run([SOFFICE, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                    os.path.join(map_, "calc"), pad], check=True, capture_output=True, timeout=300)
    return os.path.join(map_, "calc", os.path.basename(pad))


def bouw_en_bereken(project, map_, naam):
    paden = build.bouw(project, os.path.join(map_, naam), None)
    return load_workbook(herbereken(paden[0]), data_only=True)


@pytest.fixture(scope="module")
def tmp():
    with tempfile.TemporaryDirectory() as d:
        yield d


@pytest.fixture(scope="module")
def basis(tmp):
    if not SOFFICE:
        pytest.skip("LibreOffice niet gevonden")
    return bouw_en_bereken(D.voorbeeld(), tmp, "basis")


def foutcellen(wb):
    """Foutwaarden buiten de plekken waar NA() hoort (grafiekkolommen voorbij de laatste periode, vorige prognose)."""
    n = wb["Model"]["BY8"].value
    chart = {LY.M[k] for k in LY.M if k.startswith("g_")}
    na_ok = {LY.FIX["vorig"], LY.M["g_vorige"], LY.M["g_punt_nu"], LY.M["g_punt_dal"], LY.M["g_eind_upside"], LY.M["g_eind_downside"], LY.M["g_eind_basis"],
             LY.M["g_realisatie"], LY.M["g_scenario"], LY.M["g_eind_scenario"], LY.M["g_punt_uitverkocht"], LY.M["g_punt_alles_transport"],
             LY.M["g_verkocht_pct_lbl"], LY.M["g_getransporteerd_pct_lbl"]}        # %-label alleen in kwartalen met een stapel
    na_ok |= {LY.m_col(blok, k) for blok in ("gv", "gt") for k in range(1, LY.N_TYPES + 1)}   # lege typeblokken: NA()
    na_ok |= set(LY.BT.values())                                                            # bouwtermijnentabel: NA() buiten de gebruikte rijen
    na_ok |= set(LY.PT.values())                                                            # grafiek per type: NA() voorbij de laatste periode
    na_ok |= {LY.vs_col(s, k) for s in range(1, LY.N_VS + 1) for k in ("g", "gdal", "gcf")}  # scenariolijnen: NA() waar gelijk aan de Basis
    fouten = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith("#"):
                    col = "".join(ch for ch in c.coordinate if ch.isalpha())
                    # PowerPoint: NA() buiten de periodes/rijen en in de puntkolommen, de scenariolijn, lege typeblokken en de realisatie
                    ok = v == "#N/A" and ((ws.title == "Model" and (col in na_ok or (col in chart and c.row > 7 + n)))
                                          or ws.title == "PowerPoint")
                    if not ok:
                        fouten.append(f"{ws.title}!{c.coordinate}={v}")
    return fouten


def test_geen_onverwachte_fouten(basis):
    assert foutcellen(basis) == []
    assert basis["Model"]["BY8"].value == len(D.voorbeeld().periodes)


def test_controles_ok(basis):
    d = basis["Dashboard"]
    assert [d[f"W{r}"].value for r in range(LY.D_ROW_CONTROLES + 1, LY.D_ROW_CONTROLES + 7)] == ["OK"] * 6


def test_rente_nagerekend(basis):
    m, d = basis["Model"], basis["Dashboard"]
    n = m["BY8"].value
    rente = d["W6"].value
    start_idx = m[f"BY{LY.H['start_idx']}"].value
    eind = {"basis": start_idx, "up": start_idx + d["X18"].value, "down": start_idx + d["W18"].value}
    kol = {"basis": ("I", "H", LY.M["stand_basis_rente"]), "up": (LY.M["opbr_up"], LY.M["kosten_up"], LY.M["stand_up_raw"]),
           "down": (LY.M["opbr_down"], LY.M["kosten_down"], LY.M["stand_down_raw"])}
    for naam, (opb, kos, stand_col) in kol.items():
        stand = 0.0
        for r in range(8, 8 + n):
            ri = max(0.0, -stand) * rente / 4 if (m[f"G{r}"].value == "Prognose" and m[f"F{r}"].value < eind[naam]) else 0.0
            stand += (m[f"{opb}{r}"].value or 0) - (m[f"{kos}{r}"].value or 0) - ri
            assert abs(m[f"{stand_col}{r}"].value - stand) < 1e-6, (naam, r)


def test_knoppen_op_nul_geeft_basis(tmp):
    if not SOFFICE:
        pytest.skip("LibreOffice niet gevonden")
    p = D.voorbeeld()
    for k in ("shift_down", "shift_up", "uitstel_down", "uitstel_up", "opbr_down", "opbr_up", "kosten_down", "kosten_up"):
        p.params[k] = 0
    wb = bouw_en_bereken(p, tmp, "knoppen0")
    m = wb["Model"]
    n = m["BY8"].value
    for r in range(8, 8 + n):
        b = m[f"{LY.M['stand_basis']}{r}"].value
        assert abs(m[f"{LY.M['stand_up']}{r}"].value - b) < 1e-6
        assert abs(m[f"{LY.M['stand_down']}{r}"].value - b) < 1e-6


def test_lege_rijen_tellen_niet_mee(tmp, basis):
    if not SOFFICE:
        pytest.skip("LibreOffice niet gevonden")
    p = D.voorbeeld()
    leeg = {k: None for k in p.periodes[0]}
    p.periodes = [dict(leeg)] * 5 + p.periodes[:3] + [dict(leeg)] + p.periodes[3:]
    p.verkocht = [[None] * 10] * 5 + p.verkocht[:3] + [[None] * 10] + p.verkocht[3:]
    p.transport = [[None] * 10] * 5 + p.transport[:3] + [[None] * 10] + p.transport[3:]
    wb = bouw_en_bereken(p, tmp, "legerijen")
    for r in range(8, 70):
        a, b = basis["Model"][f"BY{r}"].value, wb["Model"][f"BY{r}"].value
        assert a == b or (isinstance(a, float) and abs(a - b) < 1e-6), f"BY{r}: {a!r} != {b!r}"
    assert [wb["Model"][f"B{r}"].value for r in range(8, 12)] == [7, 8, 9, 11]   # positie binnen Invoer!$B$7:$B$500 (koprij = 1)


if __name__ == "__main__":
    sys.exit(pytest.main([__file__, "-q"]))


def test_bouwtermijnen_als_pct_van_aanneemsom(tmp, basis):
    """Bouwtermijnen die samen 100% zijn (van de aanneemsom) geven hetzelfde model als termijnen die samen 100% − grond zijn."""
    p = D.voorbeeld()
    for t in p.types:
        som = sum(float(pct or 0) for _, pct, _ in t.termijnen)
        if som:
            t.termijnen = [(naam, (float(pct or 0) / som) if pct is not None else None, kw) for naam, pct, kw in t.termijnen]
    wb = bouw_en_bereken(p, tmp, "aanneemsom")
    m, mb = wb["Model"], basis["Model"]
    n = basis["Model"]["BY8"].value
    for naam in ("model_basis", "model_up", "stand_basis", "stand_up", "stand_down"):
        for r in range(LY.ROW1, LY.ROW1 + n):
            a, b = m[f"{LY.M[naam]}{r}"].value, mb[f"{LY.M[naam]}{r}"].value
            assert abs(float(a) - float(b)) < 1e-3, (naam, r, a, b)
    assert wb["Dashboard"][f"W{LY.D_ROW_CONTROLES + 9}"].value == "OK"
    assert all(abs(float(wb["Woningtypes"].cell(LY.WT_R_TOTAAL, LY.wt_col(k, 1)).value) - 1) < 1e-9
               for k in range(1, len(p.types) + 1) if p.types[k - 1].termijnen)


def test_extra_opbrengsten_per_woning(tmp, basis):
    """Extra's (€ per woning): bij transport en per bouwkwartaal tellen ze op bij de modelopbrengst, de koopsomknop niet."""
    p = D.voorbeeld()
    t = p.types[0]
    t.extras = [("Kadastrale kosten", 1000, None), ("Kopersmeerwerk", 2000, 3)]
    wb = bouw_en_bereken(p, tmp, "extras")
    m, mb = wb["Model"], basis["Model"]
    n = basis["Model"]["BY8"].value
    laatste = LY.ROW1 + n - 1
    extra = float(t.aantal) * 3000
    assert abs(float(m[f"{LY.M['model_basis']}{laatste}"].value) - float(mb[f"{LY.M['model_basis']}{laatste}"].value) - extra) < 1e-3
    assert abs(float(m[f"BY{LY.H['model_opbr']}"].value) - float(mb[f"BY{LY.H['model_opbr']}"].value) - extra) < 1e-3
    # de koopsomknop upside (+2 % opbrengsten staat al aan in het voorbeeld) schaalt de extra's niet mee
    assert abs(float(m[f"{LY.M['model_up']}{laatste}"].value) - float(mb[f"{LY.M['model_up']}{laatste}"].value) - extra) < 1e-3
    # tussentijds: in elke rij ligt de extra tussen 0 en aantal × 3000 en loopt hij op
    vorige = 0.0
    for r in range(LY.ROW1, laatste + 1):
        d = float(m[f"{LY.M['model_basis']}{r}"].value) - float(mb[f"{LY.M['model_basis']}{r}"].value)
        assert -1e-6 <= d <= extra + 1e-6 and d >= vorige - 1e-6, (r, d)
        vorige = d
    assert wb["Woningtypes"].cell(LY.WT_R_X_TOTAAL, LY.wt_col(1, 1)).value == 3000


def test_daeb_fees(tmp, basis):
    """DAEB-type (voorbeeldtype 5): fees tellen als bedragen voor het hele type (bedrag of % van het componenttotaal), met het kwartaal
    uit 'jaar Qk', een bouwkwartaal of 'actuals'; koopsom en bouwtermijnen tellen niet; de knoppen raken de fees niet."""
    mb = basis["Model"]
    n = mb["BY8"].value
    laatste = LY.ROW1 + n - 1
    daeb = D.voorbeeld().types[4]
    assert daeb.soort == "DAEB"
    fee = LY.m_col("fee", 5)
    assert mb[f"{fee}5"].value == 1
    totaal = 1192252 * (0.30 + 0.30 + 0.38 + 0.02) + 1636540 * (0.40 + 0.40 + 0.20)
    assert abs(float(mb[f"{fee}6"].value) - totaal) < 1e-3
    assert abs(float(mb[f"{fee}{laatste}"].value) - totaal) < 1e-3
    # 'actuals' valt in het laatste gerealiseerde kwartaal (Q4 '26 = rij van idx 8108), '2027 Q1' een kwartaal later,
    # bouwkwartaal 1 = start bouw (2027 Q3), 5 = vier kwartalen later
    idx = {r: mb[f"{LY.FIX['idx']}{r}"].value for r in range(LY.ROW1, laatste + 1)}
    rij = {v: r for r, v in idx.items()}
    ak = 1192252
    assert abs(float(mb[f"{fee}{rij[8108]}"].value) - ak * 0.98) < 1e-3
    assert abs(float(mb[f"{fee}{rij[8109]}"].value) - (ak + 1636540 * 0.40)) < 1e-3
    assert abs(float(mb[f"{fee}{rij[8111]}"].value) - (ak + 1636540 * 0.80)) < 1e-3
    assert abs(float(mb[f"{fee}{rij[8115]}"].value) - totaal) < 1e-3
    # de fees zitten in model_basis en in alle scenario-kolommen: scenario's verschillen niet door de fees
    p = D.voorbeeld()
    p.types[4].fee_termijnen = []
    p.types[4].fee_comp = []
    wb = bouw_en_bereken(p, tmp, "daeb_zonder_fees")
    m = wb["Model"]
    for r in range(LY.ROW1, laatste + 1):
        d = float(mb[f"{LY.M['model_basis']}{r}"].value) - float(m[f"{LY.M['model_basis']}{r}"].value)
        assert abs(d - float(mb[f"{fee}{r}"].value)) < 1e-3, (r, d)
        for naam in ("model_up", "model_down", "model_scn"):
            a = float(mb[f"{LY.M[naam]}{r}"].value) - float(mb[f"{LY.M['model_basis']}{r}"].value)
            b = float(m[f"{LY.M[naam]}{r}"].value) - float(m[f"{LY.M['model_basis']}{r}"].value)
            assert abs(a - b) < 1e-3, (naam, r, a, b)
        for naam in ("stand_up", "stand_down", "stand_scn"):
            assert abs(float(mb[f"{LY.M[naam]}{r}"].value) - float(m[f"{LY.M[naam]}{r}"].value)) < 1e-3, (naam, r)
    # zonder geplande fees geen opbrengsten DAEB: dan één marge (totaal)
    assert m[f"BY{LY.H['opbr_daeb']}"].value == 0 and m[f"BY{LY.H['marge_split']}"].value == 0
    assert m[f"BY{LY.H['marge_tekst']}"].value.startswith("marge ") and "niet-DAEB" not in m[f"BY{LY.H['marge_tekst']}"].value
    # vorige prognose van de fees (blok 6): 'na omgevingsvergunning' stond in 2026 Q4 (nu 2027 Q1), 'bij gevelsluiting' in bouwkwartaal 4
    # (nu 5, geen vorige start bouw: t.o.v. de huidige start 2027 Q3); zonder vorig kwartaal ('na akkoord SO', actuals) staat vorig = nu
    BT = LY.BT
    rij_bt = lambda i: LY.BT_ROW1 + 4 * LY.N_TERMIJNEN + i - 1
    assert (mb[f"{BT['idxb']}{rij_bt(4)}"].value, mb[f"{BT['idxr']}{rij_bt(4)}"].value) == (8109, 8108)
    assert (mb[f"{BT['idxb']}{rij_bt(8)}"].value, mb[f"{BT['idxr']}{rij_bt(8)}"].value) == (8115, 8114)
    assert mb[f"{BT['idxr']}{rij_bt(1)}"].value == mb[f"{BT['idxb']}{rij_bt(1)}"].value == 8108
    assert mb[f"{BT['fout']}{rij_bt(4)}"].value == 0
    # controles en labels
    d = basis["Dashboard"]
    assert d[f"W{LY.D_ROW_CONTROLES + 12}"].value == "OK" and d[f"W{LY.D_ROW_CONTROLES + 13}"].value == "OK"
    assert mb[f"BY{LY.H['daeb_types']}"].value == 1
    assert mb[f"BY{LY.H['verk_bij_start']}"].value == sum(float(mb[f"{LY.m_col('vcum', k)}5"].value) for k in range(1, 5))
    assert mb[f"BY{LY.H['won_met_start']}"].value == 40 + 20 + 8 + 24
    # fee-component zonder volledige planning: controle slaat aan
    p = D.voorbeeld()
    p.types[4].fee_termijnen[3] = ("na omgevingsvergunning", None, None)         # 98% gepland
    wb = bouw_en_bereken(p, tmp, "daeb_98pct")
    assert str(wb["Dashboard"][f"W{LY.D_ROW_CONTROLES + 12}"].value).startswith("LET OP")
    assert abs(float(wb["Model"][f"{fee}6"].value) - (totaal - ak * 0.02)) < 1e-3
    # ongeldig vorig kwartaal (blok 6): controle 'geldig kwartaal' slaat aan, de huidige planning blijft staan
    p = D.voorbeeld()
    p.types[4].fee_vorig[3] = "ooit"
    wb = bouw_en_bereken(p, tmp, "daeb_vorig_fout")
    assert str(wb["Dashboard"][f"W{LY.D_ROW_CONTROLES + 13}"].value).startswith("LET OP")
    assert wb["Model"][f"{LY.BT['idxr']}{rij_bt(4)}"].value == 99999 and wb["Model"][f"{LY.BT['idxb']}{rij_bt(4)}"].value == 8109
    assert abs(float(wb["Model"][f"{fee}6"].value) - totaal) < 1e-3


def test_grafiek_per_type(basis):
    """Selectieblok van de grafiek per woningtype: type 1 (keuzecel), som van de onderdelen = koopsom x aantal."""
    m = basis["Model"]
    n = m["BY8"].value
    laatste = LY.ROW1 + n - 1
    assert m[f"BY{LY.H['pt_keuze']}"].value == 1 and m[f"BY{LY.H['pt_naam']}"].value == "Rijwoning"
    assert abs(float(m[f"{LY.PT['cum']}{laatste}"].value) - 40 * 385000 / 1e6) < 1e-6
    som = sum(float(m[f"{LY.PT[x]}{r}"].value) for x in ("grond", "bouw", "extra", "fee1", "fee2") for r in range(LY.ROW1, laatste + 1))
    assert abs(som - 40 * 385000 / 1e6) < 1e-6
    assert m[f"BY{LY.H['pt_lbl_grond']}"].value == "Grondtermijn" and (m[f"BY{LY.H['pt_lbl_fee1']}"].value in (None, ""))


def test_geen_kringverwijzingen(tmp):
    """Excel meldt een kringverwijzing zodra een formule (ook via INDEX over een bereik) afhangt van een cel die van haar afhangt."""
    sys.path.insert(0, os.path.join(os.path.dirname(HIER), "tools"))
    import kringverwijzing
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "kring"), None, caches=False)
    uit = kringverwijzing.kringen(paden[0])
    assert uit == [], [[kringverwijzing.adres(n) for n in scc[:8]] for scc, _ in uit]


def test_tijdas_twee_jaar_voor_eerste_termijn(basis):
    """De tijd-as van de bouwtermijnengrafiek begint op 1 januari van het jaar twee jaar vóór de eerste termijn (huidig of vorige
    prognose), maar niet vóór de eerste periode; de Excel-grafiek krijgt dezelfde ondergrens als het Model (BY97)."""
    m = basis["Model"]
    BT = LY.BT
    eerste = min(min(m[f"{BT['idxb']}{r}"].value, m[f"{BT['idxr']}{r}"].value)
                 for r in range(LY.BT_ROW1, LY.BT_ROWN + 1) if m[f"{BT['gebruikt']}{r}"].value == 1)
    assert m[f"BY{LY.H['bt_min']}"].value == eerste
    jaar_periodes = int((m[f"{LY.FIX['idx']}{LY.ROW1}"].value - 1) // 4)
    verwacht = max(jaar_periodes, (eerste - 1) // 4 - 2)
    assert m[f"BY{LY.H['jaar_first']}"].value == verwacht and m[f"BY{LY.H['t_first']}"].value == verwacht * 4 + 1
    assert eerste == 2026 * 4 + 3 and verwacht == 2024          # voorbeeld: Rijwoning volgens de vorige prognose vanaf 2026 Q3


def test_grafiek_per_type_tijdvak(basis):
    """Het tijdvak van de grafiek per type begint twee jaar vóór de eerste opbrengst van het gekozen type (niet vóór de eerste
    periode); pt_start/pt_n sturen de benoemde bereiken en het PowerPoint-blok; totaal per kwartaal en eindpunt staan in het blok."""
    m = basis["Model"]
    n = m["BY8"].value
    tot = {r: m[f"{LY.PT['totaal']}{r}"].value for r in range(LY.ROW1, LY.ROW1 + n)}
    eerste = min(m[f"{LY.FIX['idx']}{r}"].value for r, v in tot.items() if isinstance(v, (int, float)) and v > 1e-6)
    assert m[f"BY{LY.H['pt_first']}"].value == eerste
    jaar_periodes = (m[f"{LY.FIX['idx']}{LY.ROW1}"].value - 1) // 4
    verwacht_jaar = max(jaar_periodes, (eerste - 1) // 4 - 2)
    assert m[f"BY{LY.H['pt_jaar_first']}"].value == verwacht_jaar
    start = m[f"BY{LY.H['pt_start']}"].value
    assert 1 <= start <= n and m[f"{LY.FIX['idx']}{LY.ROW1 + start - 1}"].value >= verwacht_jaar * 4 + 1
    assert start == 1 or m[f"{LY.FIX['idx']}{LY.ROW1 + start - 2}"].value < verwacht_jaar * 4 + 1
    assert m[f"BY{LY.H['pt_n']}"].value == n - start + 1
    laatste = LY.ROW1 + n - 1
    assert abs(float(m[f"{LY.PT['eind']}{laatste}"].value) - float(m[f"{LY.PT['cum']}{laatste}"].value)) < 1e-9
    assert all(abs(float(tot[r]) - sum(float(m[f"{LY.PT[k]}{r}"].value) for k in ("grond", "bouw", "extra", "fee1", "fee2"))) < 1e-9
               for r in range(LY.ROW1, LY.ROW1 + n))
    pp = basis["PowerPoint"]
    c0 = pp[LY.PP_BLOK["pt"] + "1"].column
    assert pp.cell(7, c0 + LY.PP_PT_KOL - 1).value == "Realisatie" and pp.cell(7, c0 + 7).value == "Totaal per kwartaal"
    assert pp.cell(LY.PP_KPI_ROW1, c0).value == m[f"{LY.FIX['kwartaal']}{LY.ROW1 + start - 1}"].value


def test_verkoopgrafiek_venster_en_percentage(basis):
    """De verkoop- en transportgrafiek toont alleen de periodes vanaf 1 januari van het actuals-jaar (vp_start/vp_n);
    blok 'vp' = verkocht cumulatief als deel van het aantal per type, g_verkocht_pct = totaal."""
    m = basis["Model"]
    d = basis["Dashboard"]
    n = m["BY8"].value
    jaar_act = d[LY.D["actuals_jaar"]].value
    start = m[f"BY{LY.H['vp_start']}"].value
    assert 1 <= start <= n and m[f"{LY.FIX['idx']}{LY.ROW1 + start - 1}"].value >= jaar_act * 4 + 1
    assert start == 1 or m[f"{LY.FIX['idx']}{LY.ROW1 + start - 2}"].value < jaar_act * 4 + 1
    assert m[f"BY{LY.H['vp_n']}"].value == n - start + 1
    laatste = LY.ROW1 + n - 1
    assert abs(float(m[f"{LY.M['g_verkocht_pct']}{laatste}"].value) - float(m[f"{LY.M['g_verkocht_cum']}{laatste}"].value) / m[f"BY{LY.H['woningen']}"].value) < 1e-9
    # label alleen in kwartalen met een stapel; tekst per type met het verkochte deel t/m actuals
    for r in range(LY.ROW1, LY.ROW1 + n):
        v = m[f"{LY.M['g_verkocht']}{r}"].value
        lbl = m[f"{LY.M['g_verkocht_pct_lbl']}{r}"].value
        if isinstance(v, (int, float)) and v > 0:
            assert abs(float(lbl) - float(m[f"{LY.M['g_verkocht_pct']}{r}"].value)) < 1e-9
        else:
            assert lbl in (None, "#N/A") or str(lbl).startswith("#")
    tekst = m[f"BY{LY.H['vp_tekst']}"].value
    assert tekst.startswith("Rijwoning ") and "%" in tekst and ", Tweekapper " in tekst
    chip = d[f"{LY.D_LEGENDA_CELLEN[0]}{LY.D_ROW_LEGENDA}"].value
    assert chip.startswith("Rijwoning ") and chip.endswith("%")
    pp = basis["PowerPoint"]
    c0 = pp[LY.PP_BLOK["vp"] + "1"].column
    assert pp.cell(LY.PP_KPI_ROW1, c0).value == m[f"{LY.FIX['kwartaal']}{LY.ROW1 + start - 1}"].value
    # blok van dia 6: 30 kolommen, de laatste vier het cumulatieve percentage (lijn) en het label (alleen bij een stapel)
    koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_VP_KOL)]
    assert koppen[-4:] == ["Verkocht cumulatief %", "Label verkocht %", "Getransporteerd cumulatief %", "Label transport %"]
    assert pp.cell(7, c0 + LY.PP_VP_KOL).value is None
    r0 = LY.ROW1 + start - 1
    pct = pp.cell(LY.PP_KPI_ROW1, c0 + LY.PP_VP_KOL - 4)
    assert abs(float(pct.value) - float(m[f"{LY.M['g_verkocht_pct']}{r0}"].value)) < 1e-9 and pct.number_format == "0%"
    tpct = pp.cell(LY.PP_KPI_ROW1 + n - start, c0 + LY.PP_VP_KOL - 2)
    assert abs(float(tpct.value) - float(m[f"{LY.M['g_getransporteerd_pct']}{laatste}"].value)) < 1e-9
    assert pp.cell(LY.PP_KPI_ROW1 + n - start + 1, c0 + LY.PP_VP_KOL - 4).value == "#N/A"


def test_kerncijfers_dia2(basis):
    """Dia 2: verkocht bij start bouw telt per type t/m het kwartaal van start bouw (DAEB niet); de margekaart (KPI8, was rente)
    toont niet-DAEB (DAEB heeft geen marge: projectresultaat / (opbrengsten - fees DAEB)) en totaal; bandbreedte (KPI5)
    heeft geen regel meer op tab PowerPoint (de macro haalt die kaart van de dia, net als de voetnoot)."""
    m = basis["Model"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    rij = {m[f"{LY.FIX['idx']}{r}"].value: r for r in range(LY.ROW1, LY.ROW1 + hb("n"))}
    for k in range(1, 5):
        vc, start = LY.m_col("vcum", k), m[f"{LY.m_col('verv', k)}6"].value
        assert m[f"{vc}5"].value == m[f"{vc}{rij[start]}"].value, k              # cumulatief t/m het startkwartaal
    vc1 = LY.m_col("vcum", 1)
    assert m[f"{vc1}5"].value > m[f"{vc1}{rij[m[LY.m_col('verv', 1) + '6'].value] - 1}"].value   # voorbeeld: ook in dat kwartaal verkocht
    assert m[f"{LY.m_col('vcum', 5)}5"].value == 0                              # DAEB telt niet mee
    assert hb("verk_bij_start") == sum(m[f"{LY.m_col('vcum', k)}5"].value for k in range(1, 5))
    assert abs(hb("verk_bij_start_pct") - hb("verk_bij_start") / hb("won_met_start")) < 1e-12
    fees = 1192252 + 1636540
    assert abs(hb("opbr_daeb") - fees) < 1e-6 and hb("woningen_daeb") == 30
    assert abs(hb("marge") - hb("eind_basis") / hb("opbr_totaal")) < 1e-12
    assert abs(hb("marge_nd") - hb("eind_basis") / (hb("opbr_totaal") - fees)) < 1e-12 and hb("marge_nd") > hb("marge")
    assert hb("marge_split") == 1
    assert hb("marge_tekst").startswith("marge niet-DAEB ") and " · totaal " in hb("marge_tekst")
    pp = basis["PowerPoint"]
    regels, r = {}, LY.PP_KPI_ROW1
    while pp[f"C{r}"].value:
        if pp[f"B{r}"].value == 2:
            regels[pp[f"F{r}"].value] = pp[f"D{r}"].value
        r += 1
    assert {v.split("_")[0] for v in regels if v[3].isdigit()} == {"KPI1", "KPI2", "KPI3", "KPI4", "KPI6", "KPI7", "KPI8", "KPI9"}
    assert regels["KPI3_SUB"] == f'eindsaldo in {hb("laatste_kw")}'
    assert regels["KPI8_LABEL"] == "MARGE NIET-DAEB / TOTAAL" and regels["KPI8_SUB"] == "DAEB heeft geen marge"
    assert regels["KPI8_WAARDE"].replace(",", ".") == f'{hb("marge_nd") * 100:.1f}% / {hb("marge") * 100:.1f}%'
    assert "VOETNOOT" not in regels
    assert regels["KPI6_LABEL"] == "VERKOCHT BIJ START BOUW" and regels["KPI6_WAARDE"].endswith("%")
    assert regels["KPI7_WAARDE"] == f'{int(hb("verkocht_actuals"))} / {int(hb("transport_actuals"))}'
    assert regels["KPI7_SUB"] == f'van {int(hb("woningen"))} woningen · t/m {hb("kw_actuals")}'
    assert regels["KPI_SUBTITEL"].startswith("Opbrengsten €")
    assert regels["KPI_SUBTITEL"].endswith(f'{int(hb("woningen"))} woningen, waarvan {int(hb("woningen_daeb"))} DAEB')
    assert basis["Dashboard"]["B3"].value.startswith("Projectresultaat ") and hb("marge_tekst") in basis["Dashboard"]["B3"].value


def test_norm_met_oud_label_wordt_gelezen(tmp):
    """Een werkboek van vóór de hernoeming ('Norm verkocht vóór start bouw') houdt zijn norm bij inlezen (build.py --from)."""
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "normlabel"), None, caches=False)
    wb = load_workbook(paden[0])
    d = wb["Dashboard"]
    rij = next(r for r in range(1, LY.D_ROW_CONTROLES + 1) if d[f"V{r}"].value == LY.D_LABELS["norm"])
    d[f"V{rij}"].value = "Norm verkocht vóór start bouw"
    d[LY.D["norm"]].value = 0.65
    pad = os.path.join(tmp, "normlabel_oud.xlsx")
    wb.save(pad)
    assert D.lees_werkboek(pad).params["norm"] == 0.65


def test_dia2_grafiekblok_balken_en_lasten(tmp, basis):
    """Sjabloon v17: het blok van de kleine cashflowgrafiek (jaarlabel bij het eerste kwartaal van elk jaar, één-punts-reeksen
    nu/dieptepunt/break-even/eindsaldo/€0, '(uit)' zonder vorige prognose), de as-grenzen, de balkregels (aandeel 0..1) en de
    kaart 'grootste lasten' uit de Dashboard-sectie; de sectie wordt bij inlezen teruggelezen."""
    m, pp, d = basis["Model"], basis["PowerPoint"], basis["Dashboard"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    n, pos_nu, pos_be = hb("n"), hb("pos_actuals"), hb("be_basis_pos")
    c0 = pp[LY.PP_BLOK["kc"] + "1"].column
    koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_KC_KOL)]
    assert koppen == ["Kwartaal", "cum. cashflow", "vorige prognose (uit)", "nu", "dieptepunt", "break-even", "eindsaldo", "€0"]
    assert pp.cell(7, c0 + LY.PP_KC_KOL).value is None
    cel = lambda i, j: pp.cell(LY.PP_KPI_ROW1 + i - 1, c0 + j).value  # noqa: E731  rij i (1..n), kolom j (0..7)
    for i in range(1, n + 1):
        kw = m[f"{LY.FIX['kw']}{LY.ROW1 + i - 1}"].value
        jaar = (m[f"{LY.FIX['idx']}{LY.ROW1 + i - 1}"].value - 1) // 4
        assert cel(i, 0) == (jaar if (i == 1 or kw == 1) else None), i
        assert abs(float(cel(i, 1)) - float(m[f"{LY.M['g_stand']}{LY.ROW1 + i - 1}"].value)) < 1e-9
        for j, pos in ((3, pos_nu), (5, pos_be), (6, n), (7, 1)):
            v = cel(i, j)
            if i == pos:
                assert abs(float(v) - (0 if j == 7 else float(cel(i, 1)))) < 1e-9, (i, j)
            else:
                assert v == "#N/A", (i, j, v)
    assert cel(n + 1, 0) is None and cel(n + 1, 1) == "#N/A"
    lo, hi = min(0, hb("dal_basis")), max(0, hb("top_basis"))
    assert abs(hb("kc_min") - (lo - 0.24 * (hi - lo)) / 1e6) < 1e-9 and abs(hb("kc_max") - (hi + 0.12 * (hi - lo)) / 1e6) < 1e-9
    regels, r = {}, LY.PP_KPI_ROW1
    while pp[f"C{r}"].value:
        if pp[f"B{r}"].value == 2:
            regels[pp[f"F{r}"].value] = pp[f"D{r}"].value
        r += 1
    assert r - 1 < int(LY.PP_CEL_SJABLOON[1:]) - 1 and pp[f"C{int(LY.PP_CEL_SJABLOON[1:])}"].value == LY.PP_LABEL_SJABLOON
    assert abs(regels["KPI6_BALK"] - hb("verk_bij_start_pct")) < 1e-12 and regels["KPI6_NORM"] == 0.7
    assert abs(regels["KPI7_BALK_V"] - hb("verkocht_actuals") / hb("woningen")) < 1e-12
    assert abs(regels["KPI7_BALK_T"] - hb("transport_actuals") / hb("woningen")) < 1e-12
    assert regels["GRAFIEK_LEGENDA_VORIG"] is None
    lasten = D.voorbeeld().params["lasten"]
    assert regels["KPI9_LABEL"] == "GROOTSTE LASTEN"
    for i, (naam, bedrag) in enumerate(lasten, start=1):
        assert regels[f"KPI9_NAAM{i}"] == naam and regels[f"KPI9_BEDRAG{i}"].replace(",", ".") == f"€{bedrag / 1e6:.1f}M"
        assert abs(regels[f"KPI9_BALK{i}"] - bedrag / lasten[0][1]) < 1e-12
    assert d[f"V{LY.D_ROW_LASTEN}"].value == LY.D_LABEL_LASTEN and d[LY.D["last1_naam"]].value == lasten[0][0]
    # ronde reis: de sectie wordt teruggelezen; zonder lasten blijven de regels leeg (de macro haalt de kaart dan weg)
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "lasten_rond"), None, caches=False)
    assert D.lees_werkboek(paden[0]).params["lasten"] == lasten
    p = D.voorbeeld()
    p.params["lasten"] = []
    p.kosten_posten = []                    # anders komen de lasten automatisch van tab CF kosten (test_cf_kosten_en_schakelaar)
    wb = bouw_en_bereken(p, tmp, "zonder_lasten")
    pp2 = wb["PowerPoint"]
    leeg = {pp2[f"F{r}"].value: pp2[f"D{r}"].value for r in range(LY.PP_KPI_ROW1, LY.PP_KPI_ROW1 + 80) if pp2[f"F{r}"].value}
    assert leeg["KPI9_NAAM1"] is None and leeg["KPI9_BALK1"] == 0 and wb["Model"][f"BY{LY.H['lasten_max']}"].value == 0


def _pp_regels(pp, dia):
    regels, r = {}, LY.PP_KPI_ROW1
    while pp[f"C{r}"].value:
        if pp[f"B{r}"].value == dia:
            regels[pp[f"F{r}"].value] = pp[f"D{r}"].value
        r += 1
    return regels


def test_verkoopscenarios(basis):
    """Tab Scenario's: een leeg scenario is precies de Basis; later starten (start bouw +2 kw) verandert het eindsaldo alleen met de
    extra rente; later verkopen (verkoop +4 kw) verandert de lijn maar niet het eindsaldo. De lijn staat alleen waar het scenario
    afwijkt; dia 4 en de cashflowgrafiek tonen alleen ingevulde scenario's."""
    m, pp = basis["Model"], basis["PowerPoint"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    n = hb("n")
    rij = lambda i: LY.ROW1 + i  # noqa: E731
    stand = lambda col: [m[f"{col}{rij(i)}"].value for i in range(n)]  # noqa: E731
    basislijn = stand(LY.M["stand_basis"])
    # voorbeeld: scenario 1 verkoop Q1 2028, scenario 2 'Later starten' start bouw Q3 2027, scenario 3 leeg
    assert [hb(f"vs{s}_aan") for s in (1, 2, 3)] == [1, 1, 0] and hb("vs_pp_aan") == 1
    assert max(abs(a - b) for a, b in zip(stand(LY.vs_col(3, "stand")), basislijn)) < 1e-6
    assert stand(LY.vs_col(3, "g")) == ["#N/A"] * n
    assert (hb("vs3_dal"), hb("vs3_eind"), hb("vs3_rente")) == (hb("dal_basis"), hb("eind_basis"), 0)
    assert (hb("vs2_d_sb"), hb("vs2_d_v"), hb("vs2_d_t")) == (2, 0, 2)
    assert hb("vs2_rente") > 0 and abs(hb("vs2_eind") - (hb("eind_basis") - hb("vs2_rente") + hb("rente_basis_tot"))) < 1
    assert hb("vs2_dal") < hb("dal_basis") and hb("vs2_d_be") == 1 and hb("vs2_vbsb_pct") >= hb("verk_bij_start_pct")
    assert (hb("vs1_d_sb"), hb("vs1_d_v")) == (0, 4)
    assert abs(hb("vs1_eind") - (hb("eind_basis") - hb("vs1_rente") + hb("rente_basis_tot"))) < 1
    assert hb("vs1_vbsb_pct") < hb("verk_bij_start_pct")
    for s in (1, 2):
        lijn, g = stand(LY.vs_col(s, "stand")), stand(LY.vs_col(s, "g"))
        afwijkend = [i for i in range(n) if abs(lijn[i] - basislijn[i]) > 1]
        assert afwijkend, s
        for i in range(n):                                   # zichtbaar = afwijkend of naast een afwijkend kwartaal
            zichtbaar = any(abs(j - i) <= 1 for j in afwijkend)
            assert (g[i] != "#N/A") == zichtbaar, (s, i)
            if zichtbaar:
                assert abs(g[i] - lijn[i] / 1e6) < 1e-6, (s, i)
        assert lijn[:min(afwijkend) - 1] == basislijn[:min(afwijkend) - 1]      # de gedeelde geschiedenis blijft van de Basis
    # PowerPoint: cashflowgrafiek met de ingevulde scenario's, dia 4 zonder kaart en legenda voor scenario 3
    c0 = pp[LY.PP_BLOK["cf"] + "1"].column
    koppen = [pp.cell(7, c0 + j).value for j in range(LY.PP_CF_KOL)]
    assert koppen[-3:] == [f"Scenario {s} (uit)" for s in (1, 2, 3)]                 # lijnen in de cashflowgrafiek: standaard uit
    assert hb("vs_cf") == 0 and stand(LY.vs_col(1, "gcf")) == ["#N/A"] * n
    assert hb("vs1_leg") == "Scenario 1 · verkoop Q1 ’28" and hb("vs2_leg") == "Scenario 2 · Later starten"
    regels = _pp_regels(pp, 4)
    assert all(regels[f"VS_K3_{veld}"] is None for veld in ("NAAM", "SUB", "WAARDE", "LABEL", "DAL", "BE", "VERKOCHT", "DELTA"))
    assert regels["VS_LEG3"] is None and regels["VS_LEG1"] == hb("vs1_leg") and regels["VS_K2_NAAM"] == "Scenario 2 · Later starten"
    assert regels["VS_TITEL"].startswith("Later starten: financiering tot −€18")
    assert regels["VS_K2_DELTA"].endswith("extra financiering · break-even 1 kwartaal later")
    assert any(str(k).startswith("SC_") for k in _pp_regels(pp, 5))                   # de scenario-dia schoof van 4 naar 5


def test_scenarios_niet_naar_powerpoint(tmp):
    """Naar PowerPoint = nee: de scenario's rekenen op het tabblad, maar dia 4 krijgt geen tekst (de macro haalt hem weg) en de
    scenariolijnen in de cashflowgrafiek staan op '(uit)'. De invoer op tab Scenario's wordt bij inlezen teruggelezen."""
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "vs_rond"), None, caches=False)
    terug = D.lees_werkboek(paden[0]).params
    assert terug["vs"] == D.voorbeeld().params["vs"] and terug["vs_pp"] == "ja"
    p = D.voorbeeld()
    p.params["vs_pp"] = "nee"
    wb = bouw_en_bereken(p, tmp, "vs_nee")
    m, pp = wb["Model"], wb["PowerPoint"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    assert hb("vs_aan") == 2 and hb("vs_pp_aan") == 0 and hb("vs2_dal") < hb("dal_basis")
    c0 = pp[LY.PP_BLOK["cf"] + "1"].column
    assert [pp.cell(7, c0 + j).value for j in range(LY.PP_CF_KOL)][-3:] == [f"Scenario {s} (uit)" for s in (1, 2, 3)]
    regels = _pp_regels(pp, 4)
    assert all(regels[f"VS_K{s}_WAARDE"] is None and regels[f"VS_LEG{s}"] is None for s in (1, 2, 3))


def test_leeg_werkboek_zonder_fouten(tmp):
    """Het lege sjabloon (build.py --leeg: geen periodes, geen woningtypes) rekent zonder foutwaarden buiten de bedoelde #N/A;
    zonder scenario komt er niets van tab Scenario's in PowerPoint."""
    wb = bouw_en_bereken(D.leeg(), tmp, "leeg")
    fout = [f"{ws.title}!{c.coordinate}={c.value}" for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"]
    assert fout == []
    m = wb["Model"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    assert (hb("vs_aan"), hb("vs_pp_aan")) == (0, 0) and hb("vs_max") > hb("vs_min")


def _verwacht_scenariokosten(basis_kosten, posten, n_act, d_sb, d_v):
    """Python-referentie voor de kosten per periode in een scenario met posten op tab CF kosten: de prognose van de posten op
    'start bouw' schuift d_sb periodes, die op 'verkoop' d_v; wat vóór de eerste prognoseperiode zou vallen komt in die periode,
    wat na de laatste zou vallen in de laatste; gerealiseerde periodes blijven gelijk."""
    n = len(basis_kosten)
    uit = list(basis_kosten)
    for post in posten:
        d = {"start bouw": d_sb, "verkoop": d_v}.get(post["mee"], 0)
        if not d:
            continue
        for i, b in enumerate(post["bedragen"][:n]):
            if not b or i < n_act:
                continue
            j = min(n - 1, max(n_act, i + d))
            uit[i] -= b
            uit[j] += b
    return uit


def test_cf_kosten(basis, tmp):
    """Tab CF kosten: in een scenario schuiven alleen de posten op 'start bouw' (met de start bouw) en 'verkoop' (met de verkoop)
    mee; de rest blijft staan. De tab sluit per prognosekwartaal aan op tab Invoer, telt de kosten t/m het dieptepunt en wordt bij
    inlezen teruggelezen."""
    m, kf = basis["Model"], basis[LY.KF_TAB]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    n, pa = hb("n"), hb("pos_actuals")
    posten = D.voorbeeld().kosten_posten
    basis_kosten = [m[f"{LY.FIX['kosten']}{LY.ROW1 + i}"].value for i in range(n)]
    assert hb("kf_aan") == 1 and hb("kf_verschil") == 0 and hb("kf_controle").startswith("✓")
    som = lambda mee: sum(b or 0 for p_ in posten if p_["mee"] == mee for b in p_["bedragen"][pa:])  # noqa: E731
    assert abs(hb("kf_sb_tot") - som("start bouw")) < 1 and abs(hb("kf_v_tot") - som("verkoop")) < 1
    for s in (1, 2):
        kosten_s = [m[f"{LY.vs_col(s, 'kosten')}{LY.ROW1 + i}"].value for i in range(n)]
        verwacht = _verwacht_scenariokosten(basis_kosten, posten, pa, hb(f"vs{s}_d_sb"), hb(f"vs{s}_d_v"))
        assert max(abs(a - b) for a, b in zip(kosten_s, verwacht)) < 1, (s, kosten_s, verwacht)
        assert abs(sum(kosten_s) - sum(basis_kosten)) < 1                     # alleen het moment verandert
    # scenario 2 (start bouw +2): honoraria en AK blijven staan, dus het eerste prognosekwartaal houdt alleen die posten
    assert abs(m[f"{LY.vs_col(2, 'kosten')}{LY.ROW1 + pa}"].value
               - sum((p_["bedragen"][pa] or 0) for p_ in posten if p_["mee"] != "start bouw")) < 1
    assert "schuift mee met de start bouw" in hb("kf_tekst") and hb("vs_subtitel").endswith(hb("kf_tekst_kort"))
    # kosten t/m het dieptepunt per post (gerealiseerd + prognose t/m de dieptepuntpositie) en de top 5 op de tab
    dal = hb("vs0_dal_pos")
    tm_dal = {p_["naam"]: (p_["gerealiseerd"] or 0) + sum(b or 0 for b in p_["bedragen"][pa:dal]) for p_ in posten}
    rijen = {kf[f"B{r}"].value: r for r in range(LY.KF_R1, LY.KF_RN + 1) if kf[f"B{r}"].value}
    for naam, r in rijen.items():
        assert abs(kf[f"{LY.KF_C['dal']}{r}"].value - tm_dal[naam]) < 1, naam
    assert abs(hb("kf_dal_tot") - sum(tm_dal.values())) < 1
    r_top = LY.KF_R_CONTROLE + 5
    top = [kf[f"B{r_top + k}"].value for k in range(3)]
    assert top == sorted(tm_dal, key=lambda k: -tm_dal[k])[:3]
    assert hb("kl_bron") == 0 and hb("kl1_naam") == "Grondkosten"           # Dashboard-lasten gaan voor
    # ronde reis: de posten (naam, keuze, gerealiseerd, bedragen) komen terug bij inlezen
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "kosten_rond"), None, caches=False)
    terug = D.lees_werkboek(paden[0]).kosten_posten
    assert [(k["naam"], k["mee"], k["gerealiseerd"]) for k in terug] == [(k["naam"], k["mee"], k["gerealiseerd"]) for k in posten]
    for a, b in zip(posten, terug):
        assert b["bedragen"][:n] == a["bedragen"] and not any(b["bedragen"][n:])


def test_cf_kosten_en_schakelaar(tmp):
    """Lijnen in de cashflowgrafiek = ja: de ingevulde scenario's staan als lijn in de cashflowgrafiek (Dashboard en dia 3, met de
    blokjes op het Dashboard). Zonder lasten op het Dashboard komen de drie grootste posten t/m het dieptepunt van tab CF kosten
    op dia 2. Zonder posten op tab CF kosten schuiven, zoals voorheen, alle kosten vanaf de geplande start bouw mee."""
    p = D.voorbeeld()
    p.params["vs_cf"] = "ja"
    p.params["lasten"] = []
    p.params["vs_sel"] = "Scenario 2"                        # tabel per termijn: 'Later starten' (start bouw twee kwartalen later)
    wb = bouw_en_bereken(p, tmp, "vs_cf_ja")
    m, pp, d = wb["Model"], wb["PowerPoint"], wb["Dashboard"]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    n = hb("n")
    assert hb("vs_cf") == 1 and [hb(f"vs{s}_cf") for s in (1, 2, 3)] == [1, 1, 0]
    c0 = pp[LY.PP_BLOK["cf"] + "1"].column
    assert [pp.cell(7, c0 + j).value for j in range(LY.PP_CF_KOL)][-3:] == [hb("vs1_leg"), hb("vs2_leg"), "Scenario 3 (uit)"]
    g, gcf = (([m[f"{LY.vs_col(2, k)}{LY.ROW1 + i}"].value for i in range(n)]) for k in ("g", "gcf"))
    assert g == gcf and any(v != "#N/A" for v in g)
    assert [d[f"{c}43"].value for c in "QRST"] == ["scenario's:", "Scenario 1", "Scenario 2", None]
    # grootste lasten automatisch uit tab CF kosten
    regels = _pp_regels(pp, 2)
    assert hb("kl_bron") == 1 and regels["KPI9_LABEL"] == "LASTEN T/M DIEPTEPUNT"
    namen = [hb(f"kl{i}_naam") for i in (1, 2, 3)]
    assert namen == ["Grondkosten", "Direct doorbelaste projectkosten", "Aanneemsom"]
    assert regels["KPI9_NAAM2"] == "Direct doorbelaste proj…" and regels["KPI9_NAAM1"] == "Grondkosten"
    assert abs(regels["KPI9_BALK1"] - 1) < 1e-12 and 0 < regels["KPI9_BALK3"] < regels["KPI9_BALK2"] < 1
    assert d["Z77"].value.startswith("automatisch: Grondkosten")
    # bouwkosten per termijn in de tabel van scenario 2: twee kwartalen later dan de Basis (zelfde bedragen)
    sc, kf = wb[LY.SC_TAB], wb[LY.KF_TAB]
    r_pos, r_tot = _scenario_rij(sc, "positie in het Model"), _scenario_rij(sc, "Totaal bouwkosten volgens termijnen")
    basis_tot = [kf.cell(LY.KF_R_BK_TOT, LY.KF_COL1 + i).value or 0 for i in range(n)]
    kolommen = [(sc.cell(r_pos, 3 + j).value, sc.cell(r_tot, 3 + j).value) for j in range(LY.SC_N_KW) if sc.cell(r_pos, 3 + j).value not in (None, "")]
    for (pos, v) in kolommen:
        verwacht = kf.cell(LY.KF_R_BK_TOT, LY.KF_COL1 + pos - 3).value if pos - 3 >= 0 else 0
        assert abs(v - (verwacht or 0) / 1000) < 1e-6, (pos, v, verwacht)
    assert sum(v for _, v in kolommen) * 1000 == pytest.approx(sum(basis_tot[:n - 2]))   # de laatste twee kwartalen schuiven buiten de tabel
    # zonder posten: de oude regel (alles vanaf de geplande start bouw schuift mee) en de oude tekst
    p2 = D.voorbeeld()
    p2.kosten_posten = []
    wb2 = bouw_en_bereken(p2, tmp, "zonder_posten")
    m2 = wb2["Model"]
    hb2 = lambda naam: m2[f"BY{LY.H[naam]}"].value  # noqa: E731
    assert hb2("kf_aan") == 0 and hb2("kf_tekst_kort") == "kosten vanaf start bouw schuiven mee met de start bouw"
    assert abs(hb2("vs2_dal") - (-17424665.2328)) < 1 and abs(hb2("vs1_eind") - hb2("eind_basis")) < 1


def test_cf_kosten_controle_en_dieptepunt_in_realisatie(tmp):
    """Actuals een half jaar verder zonder de posten bij te werken: het dieptepunt (Q1 '27) ligt dan in de realisatie, dus de tab en
    dia 2 noemen het 'gerealiseerd t/m' (het FO geeft de gerealiseerde kosten als één bedrag per post); de controle meldt dat kolom E
    niet meer aansluit, en bedragen na de laatste periode gaan voor in de melding en tellen niet mee in het totaal."""
    p = D.voorbeeld()
    p.params["actuals_jaar"], p.params["actuals_kw"] = 2027, 2
    p.params["lasten"] = []
    p.kosten_posten = [dict(k) for k in p.kosten_posten]
    p.kosten_posten.append({"naam": "Na de looptijd", "mee": "vast", "gerealiseerd": None,
                            "bedragen": [None] * len(p.periodes) + [5000]})
    wb = bouw_en_bereken(p, tmp, "dal_in_realisatie")
    m, kf = wb["Model"], wb[LY.KF_TAB]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    assert hb("vs0_dal_pos") <= hb("pos_actuals") and hb("kf_dal_real") == 1
    assert hb("kf_dal_lbl") == f'gerealiseerd t/m {hb("kw_actuals")}' and kf[f"{LY.KF_C['dal']}{LY.KF_R_KOP}"].value == hb("kf_dal_lbl")
    assert _pp_regels(wb["PowerPoint"], 2)["KPI9_LABEL"] == "GEREALISEERDE LASTEN"
    rij = next(r for r in range(LY.KF_R1, LY.KF_RN + 1) if kf[f"B{r}"].value == "Na de looptijd")
    assert kf[f"{LY.KF_C['totaal']}{rij}"].value == 0 and hb("kf_buiten") == 5000
    assert hb("kf_real_gap") > 1000 and hb("kf_controle").startswith("Let op: er staan bedragen")
    # zonder de post na de looptijd: de melding over kolom E
    p.kosten_posten.pop()
    wb2 = bouw_en_bereken(p, tmp, "dal_in_realisatie2")
    hb2 = lambda naam: wb2["Model"][f"BY{LY.H[naam]}"].value  # noqa: E731
    assert hb2("kf_buiten") == 0 and hb2("kf_controle").startswith("Let op: gerealiseerd t/m")


def test_cf_kosten_lege_keuze_is_vast(tmp):
    """Een lege 'schuift mee met' rekent het Model als vast; teruglezen (build.py --from) houdt hem vast in plaats van de standaard
    op de naam."""
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "lege_keuze"), None, caches=False)
    wb = load_workbook(paden[0])
    kf = wb[LY.KF_TAB]
    rij = next(r for r in range(LY.KF_R1, LY.KF_RN + 1) if kf[f"B{r}"].value == "Aanneemsom")
    kf[f"{LY.KF_C['mee']}{rij}"].value = None
    pad = os.path.join(tmp, "lege_keuze_aangepast.xlsx")
    wb.save(pad)
    terug = {k["naam"]: k["mee"] for k in D.lees_werkboek(pad).kosten_posten}
    assert terug["Aanneemsom"] == "vast" and terug["Aansluitkosten"] == "start bouw"


def _scenario_rij(sc, label):
    return next(r for r in range(1, 200) if isinstance(sc[f"B{r}"].value, str) and sc[f"B{r}"].value.startswith(label))


def test_bouwkosten_per_termijn(basis, tmp):
    """Woningtypes blok 7 (bouwkosten per termijn): tab CF kosten blok 3 telt ze per type en per kwartaal op (Basis) en sluit aan op
    de post Aanneemsom; tab Scenario's zet ze per termijn naast de bouwtermijnen van de koper; bij inlezen komen ze terug."""
    m, kf, sc = basis["Model"], basis[LY.KF_TAB], basis[LY.SC_TAB]
    hb = lambda naam: m[f"BY{LY.H[naam]}"].value  # noqa: E731
    n = hb("n")
    p = D.voorbeeld()
    idx = [per["jaar"] * 4 + int(per["kw"][1]) for per in p.periodes]
    per_kw = [0.0] * n
    for k, t in enumerate(p.types, start=1):
        rij = LY.KF_R_BK1 + k - 1
        tot = sum(b for _, b, _ in t.bouwkosten)
        assert (kf[f"{LY.KF_C['totaal']}{rij}"].value or 0) == tot, t.naam
        for _, b, kw in t.bouwkosten:
            i = idx.index(t.start_jaar * 4 + t.start_kw + kw - 1)
            per_kw[i] += b
            assert abs(kf.cell(rij, LY.KF_COL1 + i).value - sum(b2 for _, b2, kw2 in t.bouwkosten if kw2 == kw)) < 1
    aanneemsom = next(k for k in p.kosten_posten if k["naam"] == "Aanneemsom")["bedragen"]
    assert [round(x) for x in per_kw] == [round(a or 0) for a in aanneemsom]
    assert hb("kf_bk_verschil") == 0 and hb("kf_bk_controle").startswith("✓") and hb("kf_bk_tot") == sum(per_kw)
    # tab Scenario's (gekozen: scenario 1, start bouw zoals gepland): totaal per kwartaal = Basis; koper - aannemer
    r_pos = _scenario_rij(sc, "positie in het Model")
    r_tot = _scenario_rij(sc, "Totaal bouwkosten volgens termijnen")
    r_saldo = _scenario_rij(sc, "Bouwtermijnen koper")
    r_k1 = r_pos + 7                                         # eerste bouwtermijn van de koper (na verkocht, getransporteerd, lege rij, grond)
    for j in range(LY.SC_N_KW):
        pos = sc.cell(r_pos, 3 + j).value
        if pos in (None, ""):
            break
        assert abs(sc.cell(r_tot, 3 + j).value - per_kw[pos - 1] / 1000) < 1e-6, j
        koper = sum(sc.cell(r, 3 + j).value or 0 for r in range(r_k1, r_k1 + LY.N_TERMIJNEN))
        assert abs(sc.cell(r_saldo, 3 + j).value - (koper - sc.cell(r_tot, 3 + j).value)) < 1e-6, j
    assert sc[f"B{r_tot - LY.N_TERMIJNEN}"].value == "Aanvang bouw"
    # ronde reis
    paden = build.bouw(D.voorbeeld(), os.path.join(tmp, "bk_rond"), None, caches=False)
    terug = D.lees_werkboek(paden[0])
    assert [t.bouwkosten for t in terug.types] == [t.bouwkosten for t in p.types]
