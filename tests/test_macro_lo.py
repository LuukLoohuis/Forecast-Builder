"""Einde-tot-einde: de macro 'Ophalen uit FO' draait in LibreOffice op het nagebouwde FO en zet hetzelfde neer als fo.lees_fo.
Wordt overgeslagen zonder pyuno en soffice."""
import os
import sys
import tempfile

import pytest

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HIER))
sys.path.insert(0, HIER)

from forecast_builder import fo, data as D  # noqa: E402
import macro_lo  # noqa: E402
from test_fo import maak_fo, maak_fo_bouwkosten, maak_fo_vorig  # noqa: E402


@pytest.mark.skipif(not macro_lo.beschikbaar(), reason="pyuno of soffice ontbreekt")
def test_macro_uit_fo_in_libreoffice():
    with tempfile.TemporaryDirectory() as d:
        fo_pad = os.path.join(d, "fo.xlsx")
        maak_fo(fo_pad)
        uit = os.path.join(d, "uit.xlsx")
        melding = macro_lo.voer_uit(fo_pad, uit, werkmap=d)
        assert melding.startswith("KLAAR:"), melding
        assert "8 periodes" in melding and "2 woningtype(n)" in melding
        p_macro = D.lees_werkboek(uit)
        verschillen = macro_lo.vergelijk(fo.lees_fo(fo_pad), p_macro)
    assert verschillen == [], "\n".join(verschillen)
    # de planning van vóór het ophalen (voorbeeld: Rijwoning start 2027 Q1) is de vorige prognose geworden (blok 4)
    assert (p_macro.types[0].vorig_start_jaar, p_macro.types[0].vorig_start_kw) == (2027, 1)


@pytest.mark.skipif(not macro_lo.beschikbaar(), reason="pyuno of soffice ontbreekt")
def test_vorige_prognose_uit_fo_in_libreoffice():
    """Knop 'Vorige prognose uit FO': na het ophalen uit het FO de vorige prognose uit een ouder FO (start bouw naar blok 4,
    fee-kwartalen op mijlpaal naar blok 6); zelfde uitkomst als fo.voeg_vorige_prognose_toe."""
    with tempfile.TemporaryDirectory() as d:
        fo_pad, pad_vorig = os.path.join(d, "fo.xlsx"), os.path.join(d, "fo_vorig.xlsx")
        maak_fo(fo_pad)
        maak_fo_vorig(pad_vorig)
        uit = os.path.join(d, "uit.xlsx")
        meldingen = macro_lo.voer_uit(fo_pad, uit, werkmap=d, stappen=[("UitFOOphalen", fo_pad), ("VorigePrognoseUitFO", pad_vorig)])
        assert all(m.startswith("KLAAR:") for m in meldingen), meldingen
        p_macro = D.lees_werkboek(uit)
        p = fo.lees_fo(fo_pad)
        fo.voeg_vorige_prognose_toe(p, pad_vorig)
    for t, tm in zip(p.types, p_macro.types):
        assert (t.vorig_start_jaar, t.vorig_start_kw) == (tm.vorig_start_jaar, tm.vorig_start_kw), t.naam
        assert [kw for kw in t.fee_vorig] == [kw for kw in tm.fee_vorig] or (not any(t.fee_vorig) and not any(tm.fee_vorig)), (t.naam, t.fee_vorig, tm.fee_vorig)


@pytest.mark.skipif(not macro_lo.beschikbaar(), reason="pyuno of soffice ontbreekt")
def test_macro_veel_termijnen_samengevoegd():
    """Meer dan tien bouwtermijnen na het splitsen over kwartalen: de macro voegt, net als fo._binnen_limiet, de kleinste delen
    samen bij het grootste deel van dezelfde termijn (in plaats van de laatste termijnen weg te laten)."""
    from openpyxl import load_workbook
    with tempfile.TemporaryDirectory() as d:
        fo_pad = os.path.join(d, "fo.xlsx")
        maak_fo(fo_pad)
        wb = load_workbook(fo_pad)
        wo = wb["CF - opbrengsten"]
        kol = lambda j, q: 11 + (j - 2026) * 4 + (q - 1)  # noqa: E731
        wo.insert_rows(9, 3)                                       # bouwplanning type 1: drie termijnen over meer kwartalen
        for r, (naam, kws) in zip((9, 10, 11), (("Dak", [(2027, 2), (2027, 3), (2027, 4)]), ("Gevel", [(2027, 3), (2027, 4), (2028, 1)]),
                                                 ("Afbouw", [(2028, 1), (2028, 2)]))):
            wo.cell(r, 2).value = naam
            for i, (j, q) in enumerate(kws):
                wo.cell(r, kol(j, q)).value = [5, 3, 2][i] if len(kws) == 3 else [6, 4][i]
        wo.cell(7, kol(2027, 3)).value = 4                         # Casco over twee kwartalen
        wo.cell(7, kol(2027, 2)).value = 6
        wo.insert_rows(36, 3)                                      # omzetblok type 1: de termijnen erbij (percentages samen 100%)
        for r, (naam, pct) in zip((33, 34, 35, 36, 37, 38), (("Start bouw", 0.2), ("Casco", 0.3), ("Oplevering", 0.2), ("Dak", 0.1),
                                                            ("Gevel", 0.1), ("Afbouw", 0.1))):
            wo.cell(r, 2).value, wo.cell(r, 3).value, wo.cell(r, 5).value = naam, pct, pct * 300000
        wb.save(fo_pad)
        p = fo.lees_fo(fo_pad)
        assert len(p.types[0].termijnen) == 10 and [w for w in p.fo["waarschuwingen"] if "samengevoegd" in w]
        uit = os.path.join(d, "uit.xlsx")
        melding = macro_lo.voer_uit(fo_pad, uit, werkmap=d)
        assert melding.startswith("KLAAR:") and "samengevoegd" in melding, melding
        verschillen = macro_lo.vergelijk(p, D.lees_werkboek(uit))
    assert verschillen == [], "\n".join(verschillen)


@pytest.mark.skipif(not macro_lo.beschikbaar(), reason="pyuno of soffice ontbreekt")
def test_macro_bouwkosten_randgevallen():
    """Bouwkosten per termijn: dubbele bloktitels, samenvoegen en afkappen boven tien regels en het gerealiseerde deel: de macro
    zet hetzelfde neer als fo.lees_fo (zie test_fo.test_bouwkosten_randgevallen)."""
    with tempfile.TemporaryDirectory() as d:
        fo_pad = os.path.join(d, "fo.xlsx")
        maak_fo_bouwkosten(fo_pad)
        p = fo.lees_fo(fo_pad)
        uit = os.path.join(d, "uit.xlsx")
        melding = macro_lo.voer_uit(fo_pad, uit, werkmap=d)
        assert melding.startswith("KLAAR:") and "samengevoegd" in melding and "niet overgenomen" in melding, melding
        p_macro = D.lees_werkboek(uit)
        verschillen = macro_lo.vergelijk(p, p_macro)
    assert verschillen == [], "\n".join(verschillen)
    assert p_macro.types[1].bouwkosten[0] == ("Bouw DAEB (gerealiseerd)", 600000, -2)
