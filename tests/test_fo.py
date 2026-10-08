"""FO-lezer (forecast_builder/fo.py): een klein nagebouwd FO-werkboek met dezelfde labels als het echte FO (zonder LibreOffice)."""
import os
import sys
import tempfile

from openpyxl import Workbook

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HIER))

from forecast_builder import fo  # noqa: E402

KW = ["Q1", "Q2", "Q3", "Q4"]


def _kopregels(ws, r, jaar0=2026, n_jaar=3, extra=None):
    """rij r-1: jaartallen (elke vier kolommen), rij r: 'actuals' in J en Q1..Q4 vanaf K; geeft {(jaar, kw): kolom}."""
    ws.cell(r, 10).value = "actuals"
    kaart = {}
    c = 11
    for j in range(n_jaar):
        ws.cell(r - 1, c).value = jaar0 + j
        for q in range(4):
            ws.cell(r, c).value = KW[q]
            kaart[(jaar0 + j, q + 1)] = c
            c += 1
    return kaart


def maak_fo(pad):
    wb = Workbook()
    ws = wb.active
    ws.title = "1. Cashflow"
    ws["B3"], ws["C3"] = "Jaar", "Q"
    rijen = [(2025, None, 100.0, 0.0), (2026, "Q1", 50.0, 0.0), (2026, "Q2", 60.0, 0.0), (2026, "Q3", 70.0, 0.0),
             (2026, "Q4", 500.0, 900.0), (2027, "Q1", 400.0, 800.0), (2027, "Q2", 300.0, 700.0), (2027, "Q3", 200.0, 600.0)]
    for i, (j, q, k, o) in enumerate(rijen):
        r = 4 + i
        ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 6).value = j, q, k, o
        ws.cell(r, 8).value = o - k
    ws.cell(13, 3).value = "Totaal:"
    ws.cell(13, 4).value = sum(k for _, _, k, _ in rijen)
    ws.cell(13, 6).value = sum(o for _, _, _, o in rijen)
    wa = wb.create_sheet("FO - actuals")
    for c, kop in enumerate(["Waarde/TrVal", "Boekjaar", "KOSTEN/OMZET", "Q"], start=11):
        wa.cell(1, c).value = kop
    for i, (j, q) in enumerate([(2026, "Q1"), (2026, "Q2"), (2026, "Q3"), (2025, "Q4")], start=2):
        wa.cell(i, 11).value, wa.cell(i, 12).value, wa.cell(i, 14).value = 10, str(j), q
    wo = wb.create_sheet("CF - opbrengsten")
    # type 1 (koop): bouwplanning
    wo["F4"], wo["G4"] = "# won", 10
    wo["B5"] = "Termijnen 1 - Rijwoning"
    k1 = _kopregels(wo, 5)
    wo["B6"] = "Start bouw"; wo.cell(6, k1[(2027, 1)]).value = 10
    wo["B7"] = "Casco"; wo.cell(7, k1[(2027, 2)]).value = 10
    wo["B8"] = "Oplevering"; wo.cell(8, k1[(2027, 3)]).value = 4; wo.cell(8, k1[(2027, 4)]).value = 6
    wo["B9"] = "Kopersmeerwerk"; wo.cell(9, k1[(2027, 2)]).value = 10
    # type 2 (geen woningen) en type 3 (DAEB)
    wo["F12"], wo["G12"] = "# won", 0
    wo["B13"] = "Termijnen 2 - Leeg"
    _kopregels(wo, 13)
    wo["B14"] = "Start bouw"
    wo["F17"], wo["G17"] = "# won", 20
    wo["B18"] = "Termijnen 3 - DAEB"
    k3 = _kopregels(wo, 18)
    wo["B19"] = "Start bouw"; wo.cell(19, k3[(2027, 2)]).value = 20
    # verkooptempo
    wo["B22"] = "Verkooptempo"
    kv = _kopregels(wo, 22)
    wo["B23"] = "Rijwoning"; wo.cell(23, 10).value = 3; wo.cell(23, kv[(2026, 4)]).value = 7
    wo["B24"] = "Leeg"
    wo["B25"] = "DAEB"; wo.cell(25, kv[(2027, 1)]).value = 20
    # omzetblok type 1
    wo["B28"] = "Omzet HVG Termijnen"
    ko = _kopregels(wo, 28)
    wo["B29"], wo["D29"], wo["E29"] = "Koopsom", 1000000, 100000
    wo.cell(29, ko[(2026, 4)]).value = 7; wo.cell(29, ko[(2027, 1)]).value = 3
    wo["B30"], wo["C30"], wo["D30"], wo["E30"] = "Start bouw", 0.30, 600000, 60000
    wo["B31"], wo["C31"], wo["D31"], wo["E31"] = "Casco", 0.50, 1000000, 100000
    wo["B32"], wo["C32"], wo["D32"], wo["E32"] = "Oplevering", 0.20, 400000, 40000
    wo["B33"], wo["D33"], wo["E33"] = "Kopersmeerwerk", 50000, 5000
    wo["B34"], wo["D34"], wo["E34"] = "Kadastrale kosten", 10000, 1000
    wo.cell(34, 10).value = 10
    # omzetblok type 2 (leeg) en type 3 (DAEB met fees)
    wo["B37"] = "Omzet HVG Termijnen"
    _kopregels(wo, 37)
    wo["B38"] = "Koopsom"
    wo["B41"] = "Omzet HVG Termijnen"
    kd = _kopregels(wo, 41)
    wo["B42"], wo["D42"] = "AK fee", 1000000
    wo["B43"], wo["C43"] = "AK fee na akkoord SO", 0.30; wo.cell(43, 10).value = 300000
    wo["B44"], wo["C44"] = "AK fee na omgevingsvergunning", 0.70; wo.cell(44, kd[(2026, 4)]).value = 700000
    wo["B45"], wo["D45"] = "Bijkomende kosten (excl. AK)", 500000
    wo["B46"], wo["C46"] = "Bijkomende kosten bij start bouw", 1.0; wo.cell(46, kd[(2027, 2)]).value = 500000
    wb.save(pad)


def test_lees_fo():
    with tempfile.TemporaryDirectory() as d:
        pad = os.path.join(d, "fo.xlsx")
        maak_fo(pad)
        p = fo.lees_fo(pad)
    assert len(p.periodes) == 8 and p.periodes[0]["kw"] is None and p.periodes[1]["kw"] == "Q1"
    assert p.fo["opbrengsten"] == 3000.0 and p.fo["kosten"] == 1680.0
    assert (p.params["actuals_jaar"], p.params["actuals_kw"]) == (2026, 3)
    assert [t.naam for t in p.types] == ["Rijwoning", "DAEB"]          # blok zonder woningen overgeslagen
    t = p.types[0]
    assert t.aantal == 10 and t.koopsom == 300000 and abs(t.grond_pct - 1 / 3) < 1e-6 and (t.start_jaar, t.start_kw) == (2027, 1)
    assert t.termijnen == [("Start bouw", 0.3, 1), ("Casco", 0.5, 2), ("Oplevering (1/2)", 0.08, 3), ("Oplevering (2/2)", 0.12, 4)]
    assert t.extras == [("Kopersmeerwerk", 5000.0, 2), ("Kadastrale kosten", 1000.0, None)]
    daeb = p.types[1]
    assert daeb.soort == "DAEB" and daeb.koopsom is None and daeb.termijnen == []
    assert daeb.fee_comp == [("AK fee", 1000000), ("Bijkomende kosten (excl. AK)", 500000)]
    assert daeb.fee_termijnen[0] == ("AK fee na akkoord SO", 0.3, "actuals")
    assert daeb.fee_termijnen[1] == ("AK fee na omgevingsvergunning", 0.7, "2026 Q4")
    assert daeb.fee_termijnen[5] == ("Bijkomende kosten bij start bouw", 1.0, "2027 Q2")
    # verkocht: actuals (3) in het actuals-kwartaal Q3 '26 (periode-index 3), 7 in Q4 '26; DAEB 20 in Q1 '27 (tweede kolom)
    assert p.verkocht[3] == [3, None] and p.verkocht[4] == [7, None] and p.verkocht[5] == [None, 20]
    # transport van de rij Koopsom: 7 in Q4 '26, 3 in Q1 '27
    assert p.transport[4] == [7, None] and p.transport[5] == [3, None]
